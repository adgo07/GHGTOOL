"""仅本地：从交付EXE的PYZ读取代码，阻止业务模块回退源码。不是原生EXE界面操作。"""
from pathlib import Path
import importlib.abc, importlib.machinery, importlib.util
import json, sys, os
from datetime import date
from decimal import Decimal, localcontext
from hashlib import sha256
from PyInstaller.archive.readers import CArchiveReader
from unittest.mock import patch  # host verification tool, loaded before bundle import hook


root=Path(sys.argv[1]).resolve();output=Path(sys.argv[2]).resolve();output.mkdir(parents=True,exist_ok=True)
archive=CArchiveReader(str(root/'QingzhouCarbonAccounting.exe')).open_embedded_archive('PYZ.pyz')
class BundleLoader(importlib.machinery.SourceFileLoader):
    def get_code(self, fullname):return archive.extract(fullname)
    def get_source(self, fullname):return None
class BundleFinder(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname in archive.toc:
            package=archive.toc[fullname][0] == 1
            filename=root/Path(*fullname.split('.'))/ '__init__.py' if package else root/Path(*fullname.split('.')).with_suffix('.py')
            loader=BundleLoader(fullname,str(filename))
            return importlib.util.spec_from_loader(fullname,loader,is_package=package)
        if fullname.split('.')[0] in {'apps','packages','resources','CoolProp','numpy','lxml','PySide6','shiboken6'}:
            spec=importlib.machinery.PathFinder.find_spec(fullname,path or [str(root)])
            if spec is not None and spec.origin and Path(spec.origin).resolve().is_relative_to(root):return spec
            raise ModuleNotFoundError('包内业务模块缺失，禁止回退源码：'+fullname)
for name in list(sys.modules):
    if name.split('.')[0] in {'numpy','CoolProp','lxml'}: del sys.modules[name]
sys.path.insert(0,str(root))
sys.meta_path.insert(0,BundleFinder())
sys._MEIPASS=str(root)
sys.frozen=True
handles=[]
if hasattr(os,'add_dll_directory'):
    for folder in (root,root/'lxml',root/'PySide6',root/'shiboken6',root/'numpy.libs',root/'CoolProp.libs'):
        if folder.is_dir():handles.append(os.add_dll_directory(str(folder)))

from dataclasses import replace
from openpyxl import load_workbook
from packages.core.errors import DomainValidationError
from io import BytesIO
from zipfile import ZipFile
import shutil,tempfile,math
from packages.core.models import AccountingPeriod,PeriodType
from packages.core.decimal_policy import DecimalPolicy
from packages.standards.carbon_material import *
from packages.standards.carbon_material import SOURCE_PURCHASED_HEAT,SOURCE_EXPORTED_HEAT,SOURCE_FUEL,SOURCE_FUME
from packages.standards.carbon_material_input_policy import InputOrigin
from packages.application.carbon_accounting import CarbonAccountingUseCase,create_g06_parameter_resolver
from packages.persistence import SQLiteRecordRepository,SQLiteCatalogRepository,SQLiteProjectWorkspaceRepository
from packages.application.reporting.model import build_report_model
from packages.infrastructure.reporting import render_report_docx
from packages.application import ProjectWorkspaceService
from packages.application.canonical_input_codec import encode_canonical_input,decode_canonical_input
from packages.excel.templates import ExcelTemplateService
import packages.standards.carbon_material as domain
import packages.standards.carbon_material_steam_if97 as physics
manifest=json.loads((root/'build-manifest.json').read_text())
assert manifest['source_commit']==manifest['pr_head_sha']==sys.argv[3]
assert ALGORITHM_VERSION=='CAR-SM01-2026-10-20-EXTERNAL.1'
assert Path(domain.__file__).resolve().is_relative_to(root)
resolver=create_g06_parameter_resolver(SQLiteCatalogRepository(root/'databases/catalog.sqlite'))
calc=CarbonMaterialCalculator(parameter_resolver=resolver)
records=SQLiteRecordRepository(output/'records.sqlite');formal=CarbonAccountingUseCase(calc,records)
period=AccountingPeriod(PeriodType.ANNUAL,date(2026,1,1),date(2026,12,31))
def value(line,*,export=(),report_period=period):
 return CarbonMaterialInput('actual-heat','synthetic','实际新包热力合成',report_period,boundary_confirmed=True,purchased_heat=(line,),exported_heat=export,
 source_states=(EmissionSourceState(SOURCE_PURCHASED_HEAT,EmissionSourceStatus.INVOLVED),EmissionSourceState(SOURCE_EXPORTED_HEAT,EmissionSourceStatus.INVOLVED if export else EmissionSourceStatus.NOT_INVOLVED)))
def line(**kw):
 fields=dict(line_id='actual-heat',amount=InputValue('1','t'),enthalpy='2800',manual_enthalpy=True,source_boundary='OUTSIDE_ENTERPRISE');fields.update(kw);return HeatInput(**fields)
def frozen_word(store,rid,save=None,*,today=None):
 with patch.object(CarbonMaterialCalculator,'calculate',side_effect=AssertionError('frozen')),patch.object(resolver,'resolve',side_effect=AssertionError('frozen')),patch.object(physics,'_runtime',side_effect=AssertionError('frozen')):
  model=build_report_model(store.get(rid),raw_input_snapshot=store.get_raw_input_snapshot(rid),trace_snapshot=store.get_trace_snapshot(rid),provenance_snapshot=store.get_provenance_snapshot(rid),reporting_snapshot=store.get_reporting_snapshot(rid),report_qualification=store.get_report_qualification(rid),today=today)
  word=render_report_docx(model)
  with ZipFile(BytesIO(word)) as z:xml=z.read('word/document.xml')
 if save:save.write_bytes(word)
 return model,xml
known= formal.calculate(value(line(pressure_mpa='999')))
assert known.successful and known.result.total_amount==Decimal('.308')
# Official printed point is checked with explicit physical tolerance, emissions use full adopted h.
auto_line=line(enthalpy=None,manual_enthalpy=False,steam_kind=SteamKind.SUPERHEATED,pressure_mpa='.0035',temperature_c='26.85')
auto=formal.calculate(value(auto_line));assert auto.successful,auto.problems
trace=records.get_trace_snapshot(auto.record.record_id);text=json.dumps(trace,ensure_ascii=False)
for key in ('property_library_output_j_per_kg_lexeme','property_working_pressure_pa_lexeme','property_working_temperature_k_lexeme','actual_phase','GHGTOOL_COOLPROP_IF97_PROPERTY_V1','ae81610e7d23efc57f9d051c8e70a4d66e87537f'):assert key in text,key
result=physics.evaluate_if97_steam(steam_kind=SteamKind.SUPERHEATED,pressure_mpa=Decimal('.0035'),temperature_c=Decimal('26.85'),policy=calc.policy,numeric_profile_id=calc.numeric_profile_id)
assert abs(result.enthalpy_kj_per_kg-Decimal('2549.91145'))<=Decimal('0.000005')
assert auto.result.total_amount==calc.policy.multiply(result.enthalpy_kj_per_kg,Decimal('.00011'))
for ambient in (6,17,60):
 with localcontext() as ctx:
  ctx.prec=ambient;ctx.rounding='ROUND_DOWN';again=calc.calculate(value(auto_line));assert again.result.total_amount==auto.result.total_amount
model,xml=frozen_word(records,auto.record.record_id,output/'auto-frozen.docx');assert b'IF97' in xml and b'8.0.0' in xml
import CoolProp.CoolProp as cp,numpy,numpy._core._multiarray_umath as np_native
assert Path(cp.__file__).resolve().is_relative_to(root),cp.__file__
assert Path(np_native.__file__).resolve().is_relative_to(root),np_native.__file__
assert cp.get_global_param_string('gitrevision')==physics.IMPLEMENTATION_REVISION
sat_line=line(enthalpy=None,manual_enthalpy=False,steam_kind=SteamKind.SATURATED,pressure_mpa='.1')
sat=formal.calculate(value(sat_line));assert sat.successful,sat.problems
sat_h=physics.evaluate_if97_steam(steam_kind=SteamKind.SATURATED,pressure_mpa=Decimal('.1'),temperature_c=None,policy=calc.policy,numeric_profile_id=calc.numeric_profile_id)
assert sat_h.quality_lexeme=='1' and sat_h.working_temperature_k_lexeme is None
assert sat.result.total_amount==calc.policy.multiply(sat_h.enthalpy_kj_per_kg,Decimal('.00011'))
zero_factor=ParameterValue('actual-measured-heat','0','tCO2/GJ',source_kind=ParameterSourceKind.MEASURED)
with patch.object(resolver,'resolve',side_effect=AssertionError('measured no default')):
 zero=formal.calculate(value(line(factor_mode=HeatFactorMode.MEASURED,factor=zero_factor)))
 assert zero.successful and zero.result.total_amount==0
 negative=formal.calculate(value(line(factor_mode=HeatFactorMode.MEASURED,factor=replace(zero_factor,value=Decimal('-1')))))
 assert not negative.successful
before=len(records.list_all())
for bad in (line(source_boundary='UNCONFIRMED'),line(source_boundary='INSIDE_ENTERPRISE',fuel_emissions_accounted=False),line(enthalpy=None,manual_enthalpy=True,pressure_mpa='1'),replace(auto_line,pressure_mpa='.1',temperature_c='90'),replace(auto_line,pressure_mpa='1000'),line(enthalpy=None,manual_enthalpy=False,pressure_mpa='22.06399999999999999999999999999999999999'),replace(auto_line,pressure_mpa='22.06399999999999999999999999999999999999',temperature_c='426.85'),line(factor_mode=HeatFactorMode.MEASURED)):
 outcome=formal.calculate(value(bad));assert not outcome.successful and len(records.list_all())==before
sat_state=cp.AbstractState('IF97','Water');sat_state.update(cp.PQ_INPUTS,100000.,1.)
tsat=Decimal(repr(float(sat_state.T())))
del sat_state  # verifier-only probe; release before module shutdown
folded_c=calc.policy.add(calc.policy.subtract(tsat,Decimal('273.15')),Decimal('1e-30'))
folded=formal.calculate(value(replace(auto_line,pressure_mpa='.1',temperature_c=folded_c)))
assert not folded.successful and len(records.list_all())==before
assert any('工作温度' in p.message for p in folded.problems)
adjacent_pa=math.nextafter(22064000.,0.)
adjacent=physics.evaluate_if97_steam(steam_kind=SteamKind.SATURATED,pressure_mpa=calc.policy.divide(Decimal(repr(adjacent_pa)),Decimal('1000000')),temperature_c=None,policy=calc.policy,numeric_profile_id=calc.numeric_profile_id)
assert adjacent.working_pressure_pa_lexeme==repr(adjacent_pa) and adjacent.quality_lexeme=='1'
legal_k=float(str(calc.policy.add(tsat,Decimal('.01'))))
legal_temperatures=[]
for working_k in (legal_k,math.nextafter(legal_k,math.inf)):
 current=physics.evaluate_if97_steam(steam_kind=SteamKind.SUPERHEATED,pressure_mpa=Decimal('.1'),temperature_c=calc.policy.subtract(Decimal(repr(working_k)),Decimal('273.15')),policy=calc.policy,numeric_profile_id=calc.numeric_profile_id)
 assert current.working_temperature_k_lexeme==repr(working_k)
 legal_temperatures.append(current.working_temperature_k_lexeme)
with patch.object(physics,'_runtime',side_effect=DomainValidationError('unavailable')):
 assert not formal.calculate(value(auto_line)).successful
 assert formal.calculate(value(line())).successful
for raw in ('0','-1'):
 outcome=formal.calculate(value(line(enthalpy=raw)));assert outcome.successful,outcome.problems
 assert any(p.code=='CAR-VAL-STEAM-MANUAL-RANGE' and p.level.value=='WARNING' for p in outcome.problems)
 assert records.get_raw_input_snapshot(outcome.record.record_id)['purchased_heat'][0]['enthalpy']['value']==raw
 internal=replace(line(),amount=None,enthalpy=None,manual_enthalpy=False,source_boundary='INSIDE_ENTERPRISE',fuel_emissions_accounted=True)
with patch.object(calc,'_enthalpy',side_effect=AssertionError('internal not applicable')),patch.object(calc,'_heat_factor',side_effect=AssertionError('internal not applicable')):
 inside=formal.calculate(value(internal));assert inside.successful and inside.result.total_amount==0
m,x=frozen_word(records,inside.record.record_id);assert '不重复计购热' in x.decode()
def measured(name,val,unit):return ParameterValue(name,val,unit,source_kind=ParameterSourceKind.MEASURED)
fuel=FuelInput('same-boundary-fuel',FuelPath.MASS,'10',measured('carbon','.03','tC/GJ'),measured('oxidation','.98','ratio'),measured('lhv','8','GJ/t'),fuel_type=FuelType.OTHER,fuel_label='同边界燃料',origin=InputOrigin.CUSTOM)
fuel_only=CarbonMaterialInput('actual-fuel','synthetic','原燃料完整保留',period,boundary_confirmed=True,fuel_inputs=(fuel,),source_states=(EmissionSourceState(SOURCE_FUEL,EmissionSourceStatus.INVOLVED),))
baseline=calc.calculate(fuel_only);assert baseline.successful
internal_with_fuel=replace(fuel_only,purchased_heat=(internal,),source_states=(*fuel_only.source_states,EmissionSourceState(SOURCE_PURCHASED_HEAT,EmissionSourceStatus.INVOLVED)))
inside_fuel=formal.calculate(internal_with_fuel);assert inside_fuel.successful and inside_fuel.result.total_amount==baseline.result.total_amount
assert not any(item.emission_source_id==SOURCE_PURCHASED_HEAT for item in inside_fuel.result.lines)

purchased=line();exported=replace(line(),line_id='exported',amount=InputValue('.5','t'))
mixed=formal.calculate(value(purchased,export=(exported,)));assert mixed.successful and mixed.result.total_amount==Decimal('.154')
# PR52-F01: actual formal Record -> shared B.9 -> frozen Word, both directions.
report_unit_checks=[]
for p,pu,t,tu in (('.0035','MPa','26.85','C'),('3500','Pa','300','K'),('3500','Pa','26.85','C'),('.0035','MPa','300','K'),('.1','MPa',None,None),('100000','Pa',None,None)):
 state_line=line(enthalpy=None,manual_enthalpy=False,steam_kind=SteamKind.SUPERHEATED if t else SteamKind.SATURATED,
     pressure_mpa=InputValue(p,pu),temperature_c=InputValue(t,tu) if t else None)
 state_export=replace(state_line,line_id='unit-export',amount=InputValue('.5','t'))
 unit_result=formal.calculate(value(state_line,export=(state_export,)));assert unit_result.successful,unit_result.problems
 unit_id=unit_result.record.record_id
 unit_raw=records.get_raw_input_snapshot(unit_id);unit_trace=records.get_trace_snapshot(unit_id)
 unit_model,unit_xml=frozen_word(records,unit_id)
 unit_table=next(section for section in unit_model.sections if section.section_id=='b9').tables[0]
 unit_headers=''.join(cell.value for header in unit_table.header_rows for cell in header.cells)
 assert 'MPa' in unit_headers and '℃' in unit_headers
 expected=('3.50E-3','26.85') if t else ('0.10','不适用')
 from xml.etree import ElementTree as ET
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 word_b9=next(table for table in ET.fromstring(unit_xml).findall('.//w:tbl',ns) if '蒸汽压力' in ''.join(table.itertext()))
 word_b9_rows=[[''.join(cell.itertext()) for cell in row.findall('w:tc',ns)] for row in word_b9.findall('w:tr',ns)]
 for direction,unit_row in zip(('购入','输出'),unit_table.rows):
  assert unit_row.cells[0].value==direction
  assert (unit_row.cells[3].value,unit_row.cells[4].value)==expected,(pu,tu,direction,unit_row)
  assert tuple(next(row for row in word_b9_rows if row[0]==direction)[3:5])==expected
 assert records.get_raw_input_snapshot(unit_id)==unit_raw and records.get_trace_snapshot(unit_id)==unit_trace
 report_unit_checks.append({'pressure_unit':pu,'temperature_unit':tu,'display_pressure_mpa':expected[0],'display_temperature_c':expected[1],'directions':['购入','输出'],'frozen_only_and_original_evidence_unchanged':True})
periods=[]
for a,b,expected in ((date(2024,1,1),date(2024,12,31),'heat_default_gbt32151_34_c3'),(date(2025,1,1),date(2025,12,31),'heat_default_gbt32151_34_c3'),(date(2026,1,1),date(2026,12,31),'heat_default_2025')):
 outcome=formal.calculate(value(line(),report_period=AccountingPeriod(PeriodType.ANNUAL,a,b)));assert outcome.successful,outcome.problems
 snapshot=next(s for s in outcome.parameter_snapshots if s.parameter_id=='heat_emission_factor_default');assert snapshot.factor_id==expected
 periods.append({'period':str(a.year),'factor':snapshot.factor_id,'source':snapshot.source_id})
# Preserve historical actual Record by copying its database, never reconstructing old behavior.
legacy=Path(sys.argv[4]);old=json.loads((legacy/'created-proof.json').read_text());history=[]
with tempfile.TemporaryDirectory() as folder:
 db=Path(folder)/'records.sqlite';shutil.copy2(legacy/'records.sqlite',db);store=SQLiteRecordRepository(db)
 for item in old['records']:
  model,xml=frozen_word(store,item['record_id'],today=date(2026,10,18))
  assert sha256(repr(model).encode()).hexdigest()==item['model_sha256']
  assert sha256(xml).hexdigest()==item['document_xml_sha256'];history.append(item)
# Actual packaged Qt/Application/Excel/Canonical/SQLite/Word, programmatically offscreen.
import PySide6.QtWidgets as widgets
from PySide6.QtCore import QCoreApplication
from PySide6.QtWidgets import QApplication,QMessageBox
assert Path(widgets.__file__).resolve().is_relative_to(root)
os.environ['QT_QPA_PLATFORM']='offscreen';os.environ['QT_PLUGIN_PATH']=str(root/'PySide6/plugins')
QCoreApplication.setLibraryPaths([str(root/'PySide6/plugins')]);app=QApplication.instance() or QApplication([])
from apps.carbon_accounting_desktop.app import create_main_window
from apps.carbon_accounting_desktop.config import AppConfig
from packages.ui.view_models import AppRoute
from contextlib import contextmanager
qt_root=output/'qt';qt_root.mkdir(exist_ok=True);os.environ['LOCALAPPDATA']=str(qt_root/'app-data')
@contextmanager
def fresh_qt():
 projects=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(qt_root/'projects.sqlite'));store=SQLiteRecordRepository(qt_root/'records.sqlite')
 window=create_main_window(AppConfig(catalog_database=root/'databases/catalog.sqlite',records_database=qt_root/'records.sqlite',projects_database=qt_root/'projects.sqlite',user_database=qt_root/'user.sqlite'),record_repository=store,project_service=projects)
 shell=window.centralWidget();shell.navigate(AppRoute.NEW_ACCOUNTING);app.processEvents()
 try:yield shell.pages[AppRoute.NEW_ACCOUNTING],projects,store,shell
 finally:
  with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):window.close()
  window.deleteLater();app.processEvents()
with fresh_qt() as (gui,projects,store,shell):
 gui.enterprise_name.setText('新包热力整链');gui.period_year.setValue(2026);gui.boundary_confirmed.setChecked(True)
 for status in gui._source_statuses.values():status.setCurrentIndex(status.findData('NOT_INVOLVED'))
 status=gui._source_statuses[SOURCE_PURCHASED_HEAT];status.setCurrentIndex(status.findData('INVOLVED'))
 row=gui._energy_family_rows['heat'][0];row.heat_boundary.setCurrentIndex(row.heat_boundary.findData('OUTSIDE_ENTERPRISE'))
 row.amount.setText('1');row.enthalpy_mode.setCurrentIndex(row.enthalpy_mode.findData('AUTO'));row.heat_kind.setCurrentIndex(row.heat_kind.findData(SteamKind.SUPERHEATED));row.pressure.setText('.0035');row.temperature.setText('26.85');row.enthalpy.setText('9999')
 gui_value=gui._input(increment=False);assert gui_value.purchased_heat[0].enthalpy is None
 expected=gui.calculator.calculate(gui_value);assert expected.successful
 book=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
 for cell,val in {'B3':'过热蒸汽','C3':1,'D3':.0035,'E3':26.85,'F3':9999,'H3':'自动物性','I3':'缺省值','J3':'边界外'}.items():book['B.9'][cell]=val
 path=output/'auto-input.xlsx';book.save(path);book.close()
 page=shell.pages[AppRoute.EXCEL_IMPORT];context=page.excel_adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name='新包热力整链')
 preview=page.excel_adapter.import_preview(path,context);assert preview.units[0].can_calculate,preview.units[0].errors
 canonical=preview.units[0].input_value;assert canonical.purchased_heat[0].enthalpy is None
 assert decode_canonical_input(encode_canonical_input(canonical))==canonical
 workspace=projects.new_workspace('实际包IF97');projects.save(replace(workspace,units=(replace(workspace.units[0],canonical_input=canonical,form_state=gui._capture_form_state()),)))
 project_id=workspace.project_id;assert not store.list_all()
with fresh_qt() as (gui,projects,store,shell):
 recovered=projects.get(project_id);assert recovered.units[0].canonical_input==canonical
 gui._restore_form_state(recovered.units[0].form_state);row=gui._energy_family_rows['heat'][0]
 assert row.enthalpy.text()=='9999' and row.heat_boundary.currentData()=='OUTSIDE_ENTERPRISE'
 outcome=CarbonAccountingUseCase(gui.calculator,store).calculate(recovered.units[0].canonical_input);assert outcome.successful and outcome.result.total_amount==expected.result.total_amount
 rid=outcome.record.record_id;model,xml=frozen_word(store,rid,output/'qt-auto-frozen.docx');assert b'IF97' in xml and b'8.0.0' in xml
 row.pressure.setText('.1');row.temperature.setText('90');bad=CarbonAccountingUseCase(gui.calculator,store).calculate(gui._input(increment=False));assert not bad.successful and len(store.list_all())==1
proof={'source_commit':manifest['source_commit'],'tested_merge_sha':manifest['tested_merge_sha'],'algorithm':ALGORITHM_VERSION,'execution':'ACTUAL_EXE_PYZ_AND_PACKAGED_NATIVE_COOLPROP_NUMPY_QT_PROGRAMMATIC_OFFSCREEN_NOT_NATIVE_UI','source_fallback':False,'native_coolprop':str(Path(cp.__file__).relative_to(root)),'native_numpy':str(Path(np_native.__file__).relative_to(root)),'if97':{'enthalpy':str(result.enthalpy_kj_per_kg),'emission':str(auto.result.total_amount),'phase':result.actual_phase,'working_pa':result.working_pressure_pa_lexeme,'working_k':result.working_temperature_k_lexeme,'library_jkg':result.library_output_j_per_kg_lexeme},'ambient_independence':[6,17,60],'report_unit_checks':report_unit_checks,'working_boundary_checks':{'critical_pressure_collapse_rejected_both_paths':True,'temperature_fold_to_saturation_rejected':True,'adjacent_valid_pressure_pa':repr(adjacent_pa),'adjacent_valid_gas_temperatures_k':legal_temperatures,'no_epsilon_or_clamping':True},'invalid_cases_zero_added_records':True,'known_runtime_independent':True,'nonpositive_known_warning_frozen':True,'internal_no_unrelated_values':True,'internal_original_fuel_preserved':str(inside_fuel.result.total_amount),'saturated_quality_one_enthalpy':str(sat_h.enthalpy_kj_per_kg),'measured_zero_no_default_and_negative_blocked':True,'purchase_export_once_total':str(mixed.result.total_amount),'real_c3_periods':periods,'actual_old_ci273_frozen_records':history,'qt_project_reopen_excel_formal_frozen_word':True,'qt_invalid_state_zero_added_records':True,'qt_records':len(store.list_all()),'commercial_release':False}
(output/'proof.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(json.dumps(proof,ensure_ascii=False))

# B-EXTERNAL actual packaged producer, input state machine and frozen-only consumers.
from packages.standards.carbon_material_external_co2 import ExternalCO2Input,evaluate_external_co2_supplement
import packages.standards.carbon_material_external_co2 as external_module
assert Path(external_module.__file__).resolve().is_relative_to(root)
from packages.application import carbon_material_input_codec as industry_codec
for bad_fields in (None,False,[],"invalid"):
 payload=json.loads(industry_codec.encode_canonical_input(value(line())));payload['value']['fields']=bad_fields
 try:industry_codec.decode_canonical_input(json.dumps(payload))
 except industry_codec.CanonicalInputCodecError:pass
 else:raise AssertionError('corrupted payload fields accepted')
from packages.application.reporting.software_supplement import software_supplement_display
for raw,expected_display in [('1e100','1.00E+100'),('1e-100','1.00E-100')]:
 node=evaluate_external_co2_supplement(ExternalCO2Input(True,raw),carbon_total_tco2=Decimal(100),carbon_complete=True,policy=calc.policy,numeric_profile_id=calc.numeric_profile_id)
 assert dict(software_supplement_display({'software_supplement':node})[0])['外部 CO₂（tCO₂）']==expected_display
external_store=SQLiteRecordRepository(output/'external-records.sqlite')
external_service=CarbonAccountingUseCase(calc,external_store)
external_states=[]
for ext,status,total in [(None,'HISTORICAL_NOT_RECORDED',None),(ExternalCO2Input(False,'NaN'),'NOT_INCLUDED',None),(ExternalCO2Input(True,None),'PENDING',None),(ExternalCO2Input(True,'0'),'COMPLETED_ZERO',Decimal('.308')),(ExternalCO2Input(True,'20'),'COMPLETED_POSITIVE',Decimal('20.308'))]:
 outcome=external_service.calculate(replace(value(line()),external_co2=ext,other_activity_present=True,transport_present=True))
 assert outcome.successful,outcome.problems
 rid=outcome.record.record_id;trace=external_store.get_trace_snapshot(rid);node=trace['software_supplement']
 assert outcome.result.total_amount==Decimal('.308') and node['status']==status
 assert (None if node['supplementary_total_tco2'] is None else Decimal(node['supplementary_total_tco2']))==total
 assert not node['enterprise_completeness_verified'] and not node['external_calculation_verified']
 raw=external_store.get_raw_input_snapshot(rid)
 with patch.object(external_module,'evaluate_external_co2_supplement',side_effect=AssertionError('frozen')):
  model,xml=frozen_word(external_store,rid)
 assert model.schema_version=='3.1.0' and external_store.get_raw_input_snapshot(rid)==raw and external_store.get_trace_snapshot(rid)==trace
 assert any(section.section_id=='software-supplement' for section in model.sections)
 assert not any(item.emission_source_id=='CAR-SUP-OTHER-ACTIVITY-CO2' for item in outcome.result.lines)
 external_states.append({'status':status,'et':str(outcome.result.total_amount),'supplementary_total':node['supplementary_total_tco2']})
for ext in [ExternalCO2Input(True,'-1'),ExternalCO2Input(True,'NaN'),ExternalCO2Input(True,'20',unit='tCO2e'),ExternalCO2Input(True,None,known_previously_counted=True)]:
 before=len(external_store.list_all());outcome=external_service.calculate(replace(value(line()),external_co2=ext))
 assert not outcome.successful and len(external_store.list_all())==before
full_et=Decimal('100.1234567890123456789012345678901234567')
for ambient in (6,17,60):
 with localcontext() as ctx:
  ctx.prec=ambient
  node=evaluate_external_co2_supplement(ExternalCO2Input(True,'1e-37'),carbon_total_tco2=full_et,carbon_complete=False,policy=calc.policy,numeric_profile_id=calc.numeric_profile_id)
 assert Decimal(node['supplementary_total_tco2'])==Decimal('100.1234567890123456789012345678901234568')
 assert not node['core_coverage_complete']
qt_root=output/'external-qt';qt_root.mkdir(parents=True,exist_ok=True)
with fresh_qt() as (gui,projects,store,shell):
 gui.enterprise_name.setText('外部补充合成验证');gui.period_year.setValue(2026);gui.boundary_confirmed.setChecked(True)
 for status in gui._source_statuses.values():status.setCurrentIndex(status.findData('NOT_INVOLVED'))
 status=gui._source_statuses[SOURCE_PURCHASED_HEAT];status.setCurrentIndex(status.findData('INVOLVED'))
 row=gui._energy_family_rows['heat'][0];row.heat_boundary.setCurrentIndex(row.heat_boundary.findData('OUTSIDE_ENTERPRISE'))
 row.amount.setText('1');row.enthalpy_mode.setCurrentIndex(row.enthalpy_mode.findData('MANUAL'));row.enthalpy.setText('2800')
 gui.external_included.setChecked(True);gui.external_amount.setText('NaN');gui.external_label.setText('未纳入草稿');gui.external_included.setChecked(False)
 state=gui._capture_form_state();draft=projects.new_workspace('外部草稿')
 projects.save(replace(draft,units=(replace(draft.units[0],form_state=state),)))
 draft_id=draft.project_id
with fresh_qt() as (gui,projects,store,shell):
 gui._restore_form_state(projects.get(draft_id).units[0].form_state)
 assert gui.external_amount.text()=='NaN' and not gui.external_amount.isEnabled()
 gui.external_included.setChecked(True);gui.external_amount.setText('20')
 start=len(store.list_all())
 for count in (1,2):
  gui._run_calculation();assert len(store.list_all())==start+count
  assert '20.31' in gui.external_result_summary.text()
 gui.external_amount.setText('-1');gui._run_calculation();assert len(store.list_all())==start+2
 page=shell.pages[AppRoute.EXCEL_IMPORT];context=page.excel_adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name='新包软件补充')
 for choice,amount,expected_state,valid in [('不纳入','NaN','NOT_INCLUDED',True),('纳入',None,'PENDING',True),('纳入',0,'COMPLETED_ZERO',True),('纳入',20,'COMPLETED_POSITIVE',True),('纳入',-1,None,False)]:
  wb=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()));assert tuple(wb.sheetnames)==tuple([f'B.{i}' for i in range(1,10)]+['软件补充'])
  for cell,val in {'B3':'饱和蒸汽','C3':1,'F3':2800,'H3':'已知焓','I3':'缺省值','J3':'边界外'}.items():wb['B.9'][cell]=val
  wb['软件补充']['B4']=choice;wb['软件补充']['B5']=amount;wb['软件补充']['B10']='有';wb['软件补充']['B11']='有'
  path=output/'external-input.xlsx';wb.save(path);wb.close();before=len(store.list_all())
  preview=page.excel_adapter.import_preview(path,context);assert preview.units[0].can_calculate==valid,preview.units[0].errors
  assert len(store.list_all())==before
  if not valid:continue
  canonical=preview.units[0].input_value;assert canonical.external_co2.amount_lexeme==(None if amount is None else str(amount))
  assert canonical.other_activity_present and canonical.transport_present
  assert decode_canonical_input(encode_canonical_input(canonical))==canonical
  workspace=projects.new_workspace('外部Canonical');projects.save(replace(workspace,units=(replace(workspace.units[0],canonical_input=canonical),)))
  reopened=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(qt_root/'projects.sqlite')).get(workspace.project_id)
  outcome=CarbonAccountingUseCase(gui.calculator,store).calculate(reopened.units[0].canonical_input);assert outcome.successful,outcome.problems
  assert store.get_trace_snapshot(outcome.record.record_id)['software_supplement']['status']==expected_state
  frozen_word(store,outcome.record.record_id)
 # v11 must retain the v7+ material conversion path in the separate parser.
 wb=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()));sheet=wb['B.3']
 for r,mass,percentage in [(4,'80','90'),(5,'80','10'),(6,'20','80'),(7,'30','40'),(8,'10','25'),(9,'20','5')]:
  sheet[f'C{r}']='合成材料';sheet[f'D{r}']=Decimal(mass);sheet[f'E{r}']=Decimal(percentage);sheet[f'F{r}']='计算值' if r in (4,5) else '实测值'
  if r in (4,5):
   for col,v in {'I':'干基','J':'干基','K':20,'L':'按收到基水分归一'}.items():sheet[f'{col}{r}']=v
 path=output/'v11-material.xlsx';wb.save(path);wb.close();before=len(store.list_all())
 preview=page.excel_adapter.import_preview(path,context);assert preview.units[0].can_calculate,preview.units[0].errors
 assert len(store.list_all())==before
 canonical=preview.units[0].input_value;assert canonical.calcinations[0].material_rows[0].mass_basis=='DRY'
 outcome=CarbonAccountingUseCase(gui.calculator,store).calculate(canonical);assert outcome.successful,outcome.problems
 frozen_word(store,outcome.record.record_id)
# Real packaged partial monitoring and missing supplement stay separate.
from datetime import datetime,timedelta
from packages.standards.carbon_material_fume_segments import FumeMonitoringSegment
with localcontext() as reference:
 reference.prec=40
 expected_partial=Decimal(1000)*1000*40*Decimal('.02')*1*1*24*(Decimal(44)/12)*Decimal('1e-9')
 expected_partial_sum=expected_partial+Decimal(20)
for raw,state,total in [('20','COMPLETED_POSITIVE',expected_partial_sum),(None,'PENDING',None)]:
 start=datetime(2025,1,1)
 segment=FumeMonitoringSegment('source-A','stack-A','segment-A',start,start+timedelta(days=1),FumeIncinerationInput(q='1000',qvar='1000',hm='40',fch='.02',fox='1',duration='1',instance_id='source-A'))
 partial=CarbonMaterialInput('actual-partial','synthetic','实际包部分覆盖合成',AccountingPeriod(PeriodType.ANNUAL,date(2025,1,1),date(2025,12,31)),boundary_confirmed=True,source_states=(EmissionSourceState(SOURCE_FUME,EmissionSourceStatus.INVOLVED),),fume_segments=(segment,),fume_execution_version='FUME_V2',external_co2=ExternalCO2Input(True,raw))
 outcome=external_service.calculate(partial);assert outcome.successful,outcome.problems
 assert outcome.result.total_amount==expected_partial
 trace=external_store.get_trace_snapshot(outcome.record.record_id);node=trace['software_supplement']
 assert not trace['fume_monitoring']['complete'] and not node['core_coverage_complete'] and node['status']==state
 assert (None if node['supplementary_total_tco2'] is None else Decimal(node['supplementary_total_tco2']))==total
 assert not node['enterprise_completeness_verified'];frozen_word(external_store,outcome.record.record_id)
for unit in ('tCO2','tCO₂','tC','tCO₂e','kgCO2'):
 before=len(external_store.list_all());outcome=external_service.calculate(replace(value(line()),external_co2=ExternalCO2Input(True,'20',unit=unit)))
 assert outcome.successful==(unit in {'tCO2','tCO₂'})
 assert len(external_store.list_all())==before+(1 if outcome.successful else 0)
proof.update(stage='B-EXTERNAL',v11_material_conversion_preserved=True,external_actual_partial_fume=True,external_unit_matrix=['tCO2','tCO₂','tC','tCO₂e','kgCO2'],external_states=external_states,external_invalid_zero_added_records=True,external_full_precision_ambient_independence=True,external_gui_repeat_append=True,external_excel_preview_zero_records=True,external_sqlite_project_draft_and_canonical_reopen=True,external_frozen_consumers_no_recompute=True,template_version=ExcelTemplateService.default().get().template_version,template_sha256=ExcelTemplateService.default().get().sha256)
(output/'proof.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(json.dumps(proof,ensure_ascii=False))
