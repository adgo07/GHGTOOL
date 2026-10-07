"""Runtime Excel template and auditable ingress for GB/T 32151.34—2024.

The adapter reads the workbook's saved OOXML numeric lexemes into Decimal,
normalizes standard table rows into the existing Domain input, and delegates
calculation to the shared Application preview workflow.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
from pathlib import Path
import re
from typing import Iterable, Mapping, Sequence
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

from packages.application.carbon_accounting import CarbonAccountingPreviewUseCase
from packages.application.project_workspaces import AccountingUnitType
from packages.core.decimal_policy import DecimalPolicy
from packages.core.errors import DomainValidationError
from packages.core.models import (
    ActivityDataSource,
    AccountingPeriod,
    ElectricityAcquisitionMode,
    ElectricityAttribute,
    ElectricityProofStatus,
    ElectricityProofType,
    PeriodType,
    ReviewStatus,
    ValueType,
)
from packages.core.parameter_resolution import ElectricityConsumptionDetail, ParameterResolver
from packages.standards.carbon_material import (
    STANDARD_ID,
    STANDARD_VERSION,
    ActivityDataEvidence,
    BakingInput,
    CalcinationInput,
    CarbonMaterialCalculationOutcome,
    CarbonMaterialCalculator,
    CarbonMaterialInput,
    CarbonReportingData,
    CarbonateComponent,
    EmissionSourceState,
    EmissionSourceStatus,
    ElectricityOutputLine,
    FGDInput,
    FuelInput,
    FuelPath,
    FuelType,
    FumeIncinerationInput,
    GraphitizationInput,
    HeatInput,
    InputValue,
    MaterialBasis,
    MaterialComponentKind,
    MeasuredFactorEvidence,
    ParameterSourceKind,
    ParameterValue,
    SteamKind,
    MaterialInputLine,
    SOURCE_BAKING,
    SOURCE_CALCINATION,
    SOURCE_EXPORTED_ELECTRICITY,
    SOURCE_EXPORTED_HEAT,
    SOURCE_FGD,
    SOURCE_FUME,
    SOURCE_FUEL,
    SOURCE_GRAPHITIZATION,
    SOURCE_PURCHASED_ELECTRICITY,
    SOURCE_PURCHASED_HEAT,
)
from packages.standards.carbon_material_normalization import (
    MaterialDataSource,
    MaterialRole,
    normalize_material_inputs,
)


TEMPLATE_ID = "GHGTOOL_GBT_32151_34_2024"
TEMPLATE_VERSION = "2.0.0"
INGRESS_POLICY_ID = "GHGTOOL_EXCEL_INGRESS_R2"
METADATA_SHEET = "__metadata__"
VISIBLE_SHEETS = (
    "填写说明", "基本信息", "B.2 化石燃料", "B.3 原料煅烧", "B.4 焙烧／炭化",
    "B.5 石墨化", "B.6 烟气焚烧", "B.7 烟气脱硫", "B.8 电力", "B.9 热力",
)
REQUIRED_SHEETS = frozenset((*VISIBLE_SHEETS, METADATA_SHEET))
SHEET_NAME_ALIASES: Mapping[str, tuple[str, ...]] = {
    "B.4 焙烧／炭化": ("B.4 焙烧_炭化",),
}
INITIAL_ROWS = 20

SOURCE_LABELS: Mapping[str, str] = {
    SOURCE_FUEL: "化石燃料",
    SOURCE_CALCINATION: "煅烧",
    SOURCE_BAKING: "焙烧/炭化",
    SOURCE_GRAPHITIZATION: "石墨化",
    SOURCE_FUME: "烟气焚烧",
    SOURCE_FGD: "烟气脱硫",
    SOURCE_PURCHASED_ELECTRICITY: "购入电力",
    SOURCE_EXPORTED_ELECTRICITY: "输出电力",
    SOURCE_PURCHASED_HEAT: "购入热力",
    SOURCE_EXPORTED_HEAT: "输出热力",
}
SOURCE_IDS = tuple(SOURCE_LABELS)

UNIT_TYPE_LABELS = {
    "全厂": AccountingUnitType.WHOLE_SITE,
    "工序": AccountingUnitType.PROCESS,
    "其他": AccountingUnitType.OTHER,
}
PERIOD_LABELS = {"年度": PeriodType.ANNUAL, "月度": PeriodType.MONTHLY, "自定义": PeriodType.CUSTOM}
FUEL_TYPE_LABELS: Mapping[str, FuelType] = {
    "无烟煤": FuelType.ANTHRACITE,
    "烟煤": FuelType.BITUMINOUS_COAL,
    "褐煤": FuelType.LIGNITE,
    "洗精煤": FuelType.CLEANED_COAL,
    "其他洗煤": FuelType.OTHER_CLEANED_COAL,
    "型煤": FuelType.BRIQUETTE,
    "其他煤制品": FuelType.OTHER_COAL_PRODUCTS,
    "焦炭": FuelType.COKE,
    "石油焦": FuelType.PETROLEUM_COKE,
    "原油": FuelType.CRUDE_OIL,
    "燃料油": FuelType.FUEL_OIL,
    "汽油": FuelType.GASOLINE,
    "柴油": FuelType.DIESEL,
    "一般煤油": FuelType.KEROSENE,
    "液化天然气": FuelType.LIQUEFIED_NATURAL_GAS,
    "液化石油气": FuelType.LIQUEFIED_PETROLEUM_GAS,
    "石脑油": FuelType.NAPHTHA,
    "焦油": FuelType.TAR,
    "粗苯": FuelType.CRUDE_BENZENE,
    "天然气": FuelType.NATURAL_GAS,
    "高炉煤气": FuelType.BLAST_FURNACE_GAS,
    "转炉煤气": FuelType.CONVERTER_GAS,
    "焦炉煤气": FuelType.COKE_OVEN_GAS,
    "炼厂干气": FuelType.REFINERY_DRY_GAS,
    "其他煤气": FuelType.OTHER_GAS,
    "煤（请提供适用参数）": FuelType.COAL,
    "其他燃料（请提供适用参数）": FuelType.OTHER,
}
FUEL_LABEL_BY_TYPE = {value: key for key, value in FUEL_TYPE_LABELS.items()}
FUEL_C1_SUBJECT_IDS = {
    fuel_type: fuel_type.value.lower()
    for fuel_type in FuelType
    if fuel_type not in {FuelType.COAL, FuelType.OTHER}
}
FUEL_PATH_LABELS = {"体积": FuelPath.VOLUME, "质量": FuelPath.MASS, "热量": FuelPath.HEAT}
CARBONATE_LABELS = {
    "CaCO₃": "car-par-c2-caco3",
    "MgCO₃": "car-par-c2-mgco3",
    "Na₂CO₃": "car-par-c2-na2co3",
    "NaHCO₃": "car-par-c2-nahco3",
    "FeCO₃": "car-par-c2-feco3",
    "MnCO₃": "car-par-c2-mnco3",
    "BaCO₃": "car-par-c2-baco3",
    "Li₂CO₃": "car-par-c2-li2co3",
    "K₂CO₃": "car-par-c2-k2co3",
    "SrCO₃": "car-par-c2-srco3",
    "CaMg(CO₃)₂": "car-par-c2-camgco3-2",
}


UNIT_HEADERS = ("核算单元名称（必填）", "类型（必填）", "是否启用（必填）", "核算边界说明（选填）")
FUEL_HEADERS = (
    "核算单元（必填）", "燃料品种（必填）", "计量方式（必填）", "活动量（按计量方式）",
    "活动数据来源（选填）", "活动来源说明/编号（选填）",
    "低位发热量", "低位发热量来源类型", "低位发热量来源说明/编号",
    "单位热值含碳量", "单位热值含碳量来源类型", "单位热值含碳量来源说明/编号",
    "碳氧化率", "碳氧化率来源类型", "碳氧化率来源说明/编号", "自定义燃料名称（选填）",
)
ELECTRICITY_HEADERS = (
    "核算单元（必填）", "方向（必填）", "电量（MWh，必填）", "输出电力因子（软件自动；实测时填写）",
    "活动数据来源（选填）", "活动来源说明/编号（选填）", "参数来源类型（实测/计算/指定时填写）", "参数来源说明/编号（实测/计算/指定时填写）",
    "取得方式（普通购入留空）", "电力属性（普通购入留空）", "证明类型（需要时填写）", "证明状态（需要时填写）",
)
HEAT_HEADERS = (
    "核算单元（必填）", "方向（必填）", "蒸汽量（t，必填）", "蒸汽类型（必填）",
    "焓值（kJ/kg；可留空查表）", "压力（MPa；查表时填写）", "温度（℃；过热蒸汽查表时填写）",
    "排放因子（软件自动；实测时填写）", "活动数据来源（选填）", "活动来源说明/编号（选填）",
    "参数来源类型（实测/计算/指定时填写）", "参数来源说明/编号（实测/计算/指定时填写）",
)
MATERIAL_HEADERS = (
    "核算单元（必填）", "过程实例（必填）", "物料类别（必填）", "物料名称（必填）", "活动量（t，必填）",
    "固定碳含量（%）", "固定碳数据来源（选填）", "挥发分（%）", "挥发分数据来源（选填）",
)
FUME_HEADERS = (
    "核算单元（必填）", "实例名称（必填）", "烟气流量（Nm³/h，必填）", "焦油含量（mg/Nm³，必填）",
    "低位发热量（GJ/t，必填）", "单位热值含碳量（tC/GJ，必填）", "碳氧化率（必填）", "运行时间（d，必填）",
    "活动数据来源（必填）", "活动来源说明/编号（选填）", "参数来源类型（实测/计算/指定时填写）", "参数来源说明/编号（参数必填）",
)
FGD_HEADERS = (
    "核算单元（必填）", "设施/批次（必填）", "碳酸盐种类（必填）", "脱硫剂消耗量（t，必填）",
    "组分含量（默认90%；覆盖时填写）", "C.2排放因子（自动；覆盖时填写）", "转化率（默认100%；覆盖时填写）",
    "数据来源（选填）", "参数来源类型（实测/计算/指定时填写）", "来源说明/编号（覆盖时填写）",
)

SECTION_HEADERS: Mapping[str, tuple[str, ...]] = {
    "核算单元清单": UNIT_HEADERS,
    "B.2 化石燃料": FUEL_HEADERS,
    "B.8 电力": ELECTRICITY_HEADERS,
    "B.9 热力": HEAT_HEADERS,
    "B.3 煅烧": MATERIAL_HEADERS,
    "B.4 焙烧/炭化": MATERIAL_HEADERS,
    "B.5 石墨化": MATERIAL_HEADERS,
    "B.6 烟气焚烧": FUME_HEADERS,
    "B.7 烟气脱硫": FGD_HEADERS,
}

SHEET_FOR_MARKER: Mapping[str, str] = {
    "核算单元清单": "基本信息",
    "B.2 化石燃料": "B.2 化石燃料",
    "B.3 煅烧": "B.3 原料煅烧",
    "B.4 焙烧/炭化": "B.4 焙烧／炭化",
    "B.5 石墨化": "B.5 石墨化",
    "B.6 烟气焚烧": "B.6 烟气焚烧",
    "B.7 烟气脱硫": "B.7 烟气脱硫",
    "B.8 电力": "B.8 电力",
    "B.9 热力": "B.9 热力",
}

SECTION_DESCRIPTIONS = {
    "核算单元清单": "每行一个范围；停用单元的数据会提示并忽略。",
    "B.2 化石燃料": "填写具体燃料和活动量。C.1 缺省参数由软件按燃料品种自动选取；低位发热量、含碳量和碳氧化率分别填写适用来源。",
    "B.8 电力": "不同来源逐行填写。购入电力因子按来源属性和期间自动选取；输出因子可留空按标准取值。自定义输出因子需展开右侧来源栏。",
    "B.9 热力": "不同来源逐行填写。右侧折叠栏可录入焓值或压力/温度并选择实测因子；焓值留空时按现有蒸汽表查值，标准热力因子留空自动采用。",
    "B.3 煅烧": "按实际物料逐行填写质量、固定碳含量和适用的挥发分；同一过程可填写多行。空白表示没有该类物流，0表示明确为零。",
    "B.4 焙烧/炭化": "按实际投入和输出物料逐行填写；粉尘、碎屑和副产品可留空。",
    "B.5 石墨化": "按实际投入和输出物料逐行填写；粉尘、碎屑和副产品可留空。",
    "B.6 烟气焚烧": "填写标准计算所需的监测值。单位热值含碳量为企业参数，需填写可追溯的来源编号。",
    "B.7 烟气脱硫": "每种碳酸盐组分逐行填写。必须选择种类；C.2 因子自动匹配，不能把未知组分当作 CaCO₃。",
}

PROCESS_SPECS = {
    "B.3 煅烧": {
        "source": SOURCE_CALCINATION,
        "categories": {"待煅烧原料": "raw", "煅后料": "calcined", "欠烧煅料": "underburned", "炭粉尘": "dust"},
        "required": {"raw", "calcined"},
    },
    "B.4 焙烧/炭化": {
        "source": SOURCE_BAKING,
        "categories": {"填充料": "filler", "待焙烧/炭化品": "green", "粉尘/碎屑/副产品": "byproducts", "焙烧/炭化品": "product"},
        "required": {"green", "product"},
    },
    "B.5 石墨化": {
        "source": SOURCE_GRAPHITIZATION,
        "categories": {"保温料/电阻料": "packing", "待石墨化品": "green", "粉尘/碎屑/残块/副产品": "byproducts", "石墨化产品": "product"},
        "required": {"green", "product"},
    },
}

MATERIAL_ROLES = {
    "B.3 煅烧": {
        "待煅烧原料": MaterialRole.CALCINATION_FEED,
        "煅后料": MaterialRole.CALCINED_PRODUCT,
        "欠烧煅料": MaterialRole.UNDERBURN_RECOVERED,
        "碳粉尘": MaterialRole.CARBON_DUST,
    },
    "B.4 焙烧/炭化": {
        "填充料": MaterialRole.BAKING_FILLER,
        "待焙烧/炭化品": MaterialRole.GREEN_BAKING_PRODUCT,
        "焙烧/炭化产品": MaterialRole.BAKED_PRODUCT,
        "粉尘/碎屑/副产品": MaterialRole.BAKING_BYPRODUCT,
    },
    "B.5 石墨化": {
        "保温料/电阻料": MaterialRole.GRAPHITIZATION_PACKING,
        "待石墨化品": MaterialRole.GREEN_GRAPHITIZATION_PRODUCT,
        "石墨化产品": MaterialRole.GRAPHITIZED_PRODUCT,
        "粉尘/碎屑/残块/副产品": MaterialRole.GRAPHITIZATION_BYPRODUCT,
    },
}


class WorkbookFatalError(ValueError):
    """The workbook cannot safely be interpreted as this supported template."""


@dataclass(frozen=True, slots=True)
class WorkbookProvenance:
    workbook_sha256: str
    template_id: str
    template_version: str
    standard_id: str
    standard_version: str
    ingress_policy_id: str
    imported_at: datetime


@dataclass(frozen=True, slots=True)
class NumericCellEvidence:
    sheet: str
    cell: str
    raw_cell_type: str
    workbook_value_repr: str
    serialized_numeric_text: str
    normalized_decimal: Decimal


@dataclass(frozen=True, slots=True)
class ImportMessage:
    code: str
    message: str
    location: str | None = None


@dataclass(frozen=True, slots=True)
class UnitCalculationPreview:
    unit_id: str
    name: str
    unit_type: AccountingUnitType
    input_value: CarbonMaterialInput | None
    calculation: CarbonMaterialCalculationOutcome | None
    errors: tuple[ImportMessage, ...] = ()
    warnings: tuple[ImportMessage, ...] = ()

    @property
    def can_calculate(self) -> bool:
        return self.input_value is not None and self.calculation is not None and self.calculation.successful and not self.errors

    @property
    def result(self):
        return self.calculation.result if self.can_calculate and self.calculation else None

    @property
    def source_breakdown(self) -> Mapping[str, Decimal]:
        if not self.can_calculate or self.calculation is None or self.calculation.result is None:
            return {}
        totals = {source_id: Decimal(0) for source_id in SOURCE_IDS}
        policy = DecimalPolicy()
        for line in self.calculation.result.lines:
            source_id = getattr(line, "emission_source_id", None)
            if source_id in totals:
                totals[source_id] = policy.add(totals[source_id], line.amount)
        return totals


@dataclass(frozen=True, slots=True)
class WorkbookImportPreview:
    provenance: WorkbookProvenance
    units: tuple[UnitCalculationPreview, ...]
    warnings: tuple[ImportMessage, ...]
    numeric_evidence: tuple[NumericCellEvidence, ...]


@dataclass(slots=True)
class _UnitContext:
    name: str
    unit_id: str
    unit_type: AccountingUnitType
    enterprise_id: str
    enterprise_name: str | None
    period: AccountingPeriod | None
    boundary_confirmed: bool
    boundary_description: str | None
    source_states: dict[str, EmissionSourceStatus] = field(default_factory=dict)
    errors: list[ImportMessage] = field(default_factory=list)
    warnings: list[ImportMessage] = field(default_factory=list)
    payloads: dict[str, list[object]] = field(default_factory=dict)
    activity_evidence: list[ActivityDataEvidence] = field(default_factory=list)
    factor_evidence: list[MeasuredFactorEvidence] = field(default_factory=list)


@dataclass(frozen=True, slots=True)
class _MaterialRow:
    category: str
    name: str
    mass: Decimal
    fixed_carbon: Decimal | None
    volatile_matter: Decimal | None
    evidence_ids: tuple[str, ...]

    def amount(self) -> MaterialAmount:
        return MaterialAmount(self.mass, self.fixed_carbon, self.volatile_matter)


def _table_name(prefix: str) -> str:
    return prefix.replace(".", "").replace("/", "").replace(" ", "")


def _set_list_validation(sheet, values: str, cells: str) -> None:
    formula = values if values.startswith("=") else f'"{values}"'
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.errorTitle = "请选择列表中的选项"
    validation.error = "请使用下拉选项；导入时会再次校验。"
    sheet.add_data_validation(validation)
    validation.add(cells)


def _set_numeric_validation(sheet, column: int, start: int, end: int, *, ratio: bool = False) -> None:
    if ratio:
        validation = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
        validation.errorTitle = "比例范围为 0 到 1"
        validation.error = "请输入 0 到 1 之间的实际比例值。"
    else:
        validation = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        validation.errorTitle = "数值不能为负"
        validation.error = "请输入非负数值；导入时会再次校验。"
    sheet.add_data_validation(validation)
    from openpyxl.utils import get_column_letter
    letter = get_column_letter(column)
    validation.add(f"{letter}{start}:{letter}{end}")


def _style_title(sheet, title: str, columns: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    cell = sheet.cell(1, 1, title)
    cell.font = Font(name="Microsoft YaHei", size=16, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="176B64")
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet.sheet_view.showGridLines = False


def _write_section(
    sheet,
    marker: str,
    description: str,
    headers: Sequence[str],
    row: int,
    table_name: str,
    *,
    required_columns: set[int],
    automatic_columns: set[int] | None = None,
) -> tuple[int, int, int]:
    width = len(headers)
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    title = sheet.cell(row, 1, marker)
    title.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=12)
    title.fill = PatternFill("solid", fgColor="176B64")
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 24
    sheet.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=width)
    note = sheet.cell(row + 1, 1, description)
    note.font = Font(name="Microsoft YaHei", size=9, color="404040")
    note.fill = PatternFill("solid", fgColor="F2F2F2")
    note.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row + 1].height = 30
    header_row = row + 2
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        if column in required_columns:
            color = "176B64"
        elif column in (automatic_columns or set()):
            color = "7F7F7F"
        else:
            color = "4472C4"
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = min(max(len(header) * 1.55, 14), 34)
    sheet.row_dimensions[header_row].height = 44
    first_data = header_row + 1
    last_data = first_data + INITIAL_ROWS - 1
    for data_row in range(first_data, last_data + 1):
        for column in range(1, width + 1):
            cell = sheet.cell(data_row, column)
            cell.fill = PatternFill("solid", fgColor="FFF2CC" if column in required_columns else "E7E6E6" if column in (automatic_columns or set()) else "DDEBF7")
            cell.alignment = Alignment(vertical="center", wrap_text=(column in required_columns))
    sheet.freeze_panes = f"A{first_data}"
    ref = f"A{header_row}:{sheet.cell(last_data, width).coordinate}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    return header_row, first_data, last_data


def _metadata_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet(METADATA_SHEET)
    sheet.append(("key", "value"))
    for key, value in (
        ("template_id", TEMPLATE_ID),
        ("template_version", TEMPLATE_VERSION),
        ("standard_id", STANDARD_ID),
        ("standard_version", STANDARD_VERSION),
        ("ingress_policy_id", INGRESS_POLICY_ID),
    ):
        sheet.append((key, value))
    sheet.sheet_state = "hidden"


def create_template_bytes() -> bytes:
    """Build the ten-sheet R2 input workbook in memory."""

    workbook = Workbook()
    workbook.remove(workbook.active)
    notes = workbook.create_sheet("填写说明")
    _style_title(notes, "GB/T 32151.34—2024 炭素材料生产企业数据录入模板 R2", 2)
    notes.column_dimensions["A"].width = 22
    notes.column_dimensions["B"].width = 92
    guidance = (
        ("用途", "本模板用于录入核算数据并在软件中逐核算单元预览；不会在 Excel 内计算排放量。"),
        ("核算单元", "先在“基本信息”登记期间、边界和核算单元；B.2～B.9 每条数据必须关联一个已启用单元。"),
        ("多行填写", "燃料、物料、电力和热力均可按实际数量继续增加行；表内预置空行不是业务数量上限。"),
        ("数值", "请使用数值单元格，最多15位有效数字；不接受公式、文本数字、日期、布尔值或非有限数。空白表示未提供，0表示明确为零。"),
        ("标准参数", "缺省参数由软件按标准和核算期间匹配。需要填写来源的字段可填写来源类型和简短说明；说明缺失时软件按业务规则提醒。"),
        ("蒸汽", "蒸汽量使用吨，压力使用绝压MPa；可留空焓值按附录C.4/C.5自动确定，也可填写手动焓值。"),
        ("错误隔离", "每个启用核算单元单独校验；一个单元的问题不会阻止其他有效单元预览。"),
        ("工作表", "请保留10个可见工作表及隐藏模板信息页。Excel不允许工作表名称含半角斜线，因此B.4页使用全角斜线。"),
        ("范围", "Excel导入预览不保存项目、工作区或正式核算记录；正式写入闭环和Excel结果导出尚未开放。"),
    )
    for row, (label, explanation) in enumerate(guidance, start=3):
        notes.cell(row, 1, label).font = Font(name="Microsoft YaHei", bold=True, color="176B64")
        notes.cell(row, 1).fill = PatternFill("solid", fgColor="E2F0D9")
        notes.cell(row, 2, explanation).font = Font(name="Microsoft YaHei", size=10)
        notes.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="center")
        notes.cell(row, 2).fill = PatternFill("solid", fgColor="F8F9FA")
        notes.row_dimensions[row].height = 34
    notes.freeze_panes = "A3"

    basic = workbook.create_sheet("基本信息")
    _style_title(basic, "核算基本信息与核算单元", 4)
    values = (("企业名称", ""), ("核算期间类型", "年度"), ("开始日期", None), ("结束日期", None), ("已核对核算边界", "否"), ("公共边界说明", ""))
    for row, (label, value) in enumerate(values, start=3):
        basic.cell(row, 1, label).font = Font(name="Microsoft YaHei", bold=True)
        basic.cell(row, 1).fill = PatternFill("solid", fgColor="F2F2F2")
        basic.cell(row, 2, value).fill = PatternFill("solid", fgColor="FFF2CC" if label not in {"企业名称", "公共边界说明"} else "DDEBF7")
        basic.cell(row, 2).alignment = Alignment(vertical="center", wrap_text=True)
        basic.row_dimensions[row].height = 24
    basic.column_dimensions["A"].width = 25
    basic.column_dimensions["B"].width = 36
    basic.column_dimensions["C"].width = 16
    basic.column_dimensions["D"].width = 46
    basic["B5"].number_format = "yyyy-mm-dd"
    basic["B6"].number_format = "yyyy-mm-dd"
    basic.merge_cells("B8:D8")
    basic.merge_cells("A9:D9")
    basic["A9"] = "企业名称可以留空；每行登记一个核算范围。"
    basic["A9"].fill = PatternFill("solid", fgColor="F2F2F2")
    basic["A9"].alignment = Alignment(wrap_text=True)
    unit_area = _write_section(basic, "核算单元清单", "每个核算单元独立校验和预览；业务表中的单元名称需与此处一致。", UNIT_HEADERS, 10, "AccountingUnits", required_columns={1, 2, 3})
    unit_header, unit_first, unit_last = unit_area
    _set_list_validation(basic, "全厂,工序,其他", f"B{unit_first}:B1000")
    _set_list_validation(basic, "是,否", f"C{unit_first}:C1000")
    _set_list_validation(basic, "年度,月度,自定义", "B4")
    _set_list_validation(basic, "是,否", "B7")
    workbook.defined_names.add(DefinedName("UnitNames", attr_text="'基本信息'!$A$" + str(unit_first) + ":$A$1000"))

    sheet_specs = (
        ("B.2 化石燃料", "B.2 化石燃料", FUEL_HEADERS, "FuelRows", {1, 2, 3}, set()),
        ("B.3 原料煅烧", "B.3 煅烧", MATERIAL_HEADERS, "CalcinationRows", {1, 2, 3, 4, 5}, set()),
        ("B.4 焙烧／炭化", "B.4 焙烧/炭化", MATERIAL_HEADERS, "BakingRows", {1, 2, 3, 4, 5}, set()),
        ("B.5 石墨化", "B.5 石墨化", MATERIAL_HEADERS, "GraphitizationRows", {1, 2, 3, 4, 5}, set()),
        ("B.6 烟气焚烧", "B.6 烟气焚烧", FUME_HEADERS, "FumeRows", {1, 2, 3, 4, 5, 6, 7, 8}, set()),
        ("B.7 烟气脱硫", "B.7 烟气脱硫", FGD_HEADERS, "FGDRows", {1, 2, 3, 4}, {6}),
        ("B.8 电力", "B.8 电力", ELECTRICITY_HEADERS, "ElectricityRows", {1, 2, 3}, {4}),
        ("B.9 热力", "B.9 热力", HEAT_HEADERS, "HeatRows", {1, 2, 3, 4}, {8}),
    )
    areas = {}
    for sheet_name, marker, headers, table_name, required, automatic in sheet_specs:
        sheet = workbook.create_sheet(sheet_name)
        _style_title(sheet, sheet_name, len(headers))
        area = _write_section(sheet, marker, SECTION_DESCRIPTIONS[marker], headers, 3, table_name, required_columns=required, automatic_columns=automatic)
        areas[marker] = (sheet, area)
        _set_list_validation(sheet, "=UnitNames", f"A{area[1]}:A1000")

    fuel_sheet, fuel_area = areas["B.2 化石燃料"]
    f1 = fuel_area[1]
    _set_list_validation(fuel_sheet, ",".join(FUEL_TYPE_LABELS), f"B{f1}:B1000")
    _set_list_validation(fuel_sheet, "体积,质量,热量", f"C{f1}:C1000")
    _set_list_validation(fuel_sheet, "计量/仪表记录,生产/能源台账,结算记录,检测报告,用户指定", f"E{f1}:E1000")
    for col in (8, 11, 14):
        letter = get_column_letter(col)
        _set_list_validation(fuel_sheet, "缺省值,实测值,计算值,用户指定", f"{letter}{f1}:{letter}1000")
    for col in (4, 7, 10, 13):
        _set_numeric_validation(fuel_sheet, col, f1, 1000, ratio=col == 13)

    categories = {
        "B.3 煅烧": "待煅烧原料,煅后料,欠烧煅料,碳粉尘",
        "B.4 焙烧/炭化": "填充料,待焙烧/炭化品,焙烧/炭化产品,粉尘/碎屑/副产品",
        "B.5 石墨化": "保温料/电阻料,待石墨化品,石墨化产品,粉尘/碎屑/残块/副产品",
    }
    for marker, options in categories.items():
        sheet, area = areas[marker]
        first = area[1]
        _set_list_validation(sheet, options, f"C{first}:C1000")
        for column in (5, 6, 8):
            _set_numeric_validation(sheet, column, first, 1000, ratio=column in {6, 8})
        for column in (7, 9):
            letter = get_column_letter(column)
            _set_list_validation(sheet, "缺省值,实测值,化学计算", f"{letter}{first}:{letter}1000")

    fume_sheet, fume_area = areas["B.6 烟气焚烧"]
    ff = fume_area[1]
    for col in range(3, 9):
        _set_numeric_validation(fume_sheet, col, ff, 1000, ratio=col == 7)
    _set_list_validation(fume_sheet, "计量/仪表记录,生产/能源台账,检测报告,用户指定", f"I{ff}:I1000")
    _set_list_validation(fume_sheet, "实测值,计算值,用户指定", f"K{ff}:K1000")

    fgd_sheet, fgd_area = areas["B.7 烟气脱硫"]
    fg = fgd_area[1]
    _set_list_validation(fgd_sheet, ",".join(CARBONATE_LABELS), f"C{fg}:C1000")
    _set_list_validation(fgd_sheet, "计量/仪表记录,生产/能源台账,检测报告,用户指定", f"H{fg}:H1000")
    _set_list_validation(fgd_sheet, "实测值,计算值,用户指定", f"I{fg}:I1000")
    for col in (4, 5, 6, 7):
        _set_numeric_validation(fgd_sheet, col, fg, 1000, ratio=col in {5, 7})

    electricity_sheet, electricity_area = areas["B.8 电力"]
    ef = electricity_area[1]
    for options, col in (("购入,输出", 2), ("购入,自发自用", 9), ("常规,非化石,化石", 10), ("无,合同和结算,GEC,月度原始记录", 11), ("未提供,有效,无效", 12), ("计量/仪表记录,生产/能源台账,结算记录,检测报告,用户指定", 5), ("实测值,计算值,用户指定", 7)):
        letter = get_column_letter(col)
        _set_list_validation(electricity_sheet, options, f"{letter}{ef}:{letter}1000")
    _set_numeric_validation(electricity_sheet, 3, ef, 1000)
    _set_numeric_validation(electricity_sheet, 4, ef, 1000)

    heat_sheet, heat_area = areas["B.9 热力"]
    hf = heat_area[1]
    for options, col in (("购入,输出", 2), ("饱和蒸汽,过热蒸汽", 4), ("计量/仪表记录,生产/能源台账,结算记录,检测报告,用户指定", 9), ("实测值,计算值,用户指定", 11)):
        letter = get_column_letter(col)
        _set_list_validation(heat_sheet, options, f"{letter}{hf}:{letter}1000")
    for col in (3, 5, 6, 7, 8):
        _set_numeric_validation(heat_sheet, col, hf, 1000)
    _metadata_sheet(workbook)
    import io
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()

def write_template(destination: str | Path) -> Path:
    path = Path(destination)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("模板文件扩展名必须为 .xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(create_template_bytes())
    return path


def significant_digit_count(value: Decimal) -> int:
    if not value.is_finite():
        raise ValueError("numeric value must be finite")
    digits = list(value.as_tuple().digits)
    while len(digits) > 1 and digits[0] == 0:
        digits.pop(0)
    return len(digits) if any(digits) else 1


_NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def _worksheet_xml_paths(archive: ZipFile) -> dict[str, str]:
    workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    rels_root = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_targets = {node.attrib["Id"]: node.attrib["Target"] for node in rels_root.findall(f"{{{_NS_PKG_REL}}}Relationship")}
    paths: dict[str, str] = {}
    for node in workbook_root.findall(f"{{{_NS_MAIN}}}sheets/{{{_NS_MAIN}}}sheet"):
        name = node.attrib["name"]
        target = rel_targets[node.attrib[f"{{{_NS_REL}}}id"]]
        paths[name] = target.lstrip("/") if target.startswith("/") else target if target.startswith("xl/") else f"xl/{target}"
    return paths


def _serialized_numeric_values(source: bytes) -> dict[tuple[str, str], str]:
    values: dict[tuple[str, str], str] = {}
    import io
    with ZipFile(io.BytesIO(source)) as archive:
        for sheet_name, member in _worksheet_xml_paths(archive).items():
            root = ElementTree.fromstring(archive.read(member))
            for cell in root.findall(f".//{{{_NS_MAIN}}}c"):
                reference = cell.attrib.get("r")
                if not reference or cell.attrib.get("t") not in (None, "n"):
                    continue
                value = cell.find(f"{{{_NS_MAIN}}}v")
                if value is not None and value.text is not None:
                    values[(sheet_name, reference)] = value.text
    return values


class _CellReader:
    def __init__(self, numeric_xml: Mapping[tuple[str, str], str]):
        self.numeric_xml = numeric_xml
        self.evidence: list[NumericCellEvidence] = []

    def error(self, ctx: _UnitContext, code: str, message: str, location: str) -> None:
        ctx.errors.append(ImportMessage(code, message, location))

    def number(self, ctx: _UnitContext, sheet: str, row: int, cell, *, required: bool = False) -> Decimal | None:
        location = f"{sheet}!{cell.coordinate}"
        if cell.data_type == "f":
            self.error(ctx, "EXCEL-FORMULA-REJECTED", "数值不接受 Excel 公式；请将来源数据粘贴为数值。", location)
            return None
        serialized = self.numeric_xml.get((sheet, cell.coordinate))
        if cell.value is None and serialized is None:
            if required:
                self.error(ctx, "EXCEL-NUMBER-REQUIRED", "请填写此数值；空白表示未提供，0 表示明确为零。", location)
            return None
        if cell.data_type != "n" or isinstance(cell.value, (bool, date, datetime)):
            self.error(ctx, "EXCEL-NUMBER-CELL-REQUIRED", "此字段必须是 Excel 数值单元格；文本数字、公式和日期均不接受。", location)
            return None
        if serialized is None:
            self.error(ctx, "EXCEL-NUMERIC-SERIALIZATION-MISSING", "无法读取此数值在工作簿中的保存内容。", location)
            return None
        try:
            value = Decimal(serialized)
        except InvalidOperation:
            self.error(ctx, "EXCEL-NUMBER-INVALID", "此数值不是有效的十进制数。", location)
            return None
        self.evidence.append(NumericCellEvidence(sheet, cell.coordinate, cell.data_type, repr(cell.value), serialized, value))
        if not value.is_finite():
            self.error(ctx, "EXCEL-NUMBER-NONFINITE", "NaN 和 Infinity 不能作为核算输入。", location)
            return None
        digits = significant_digit_count(value)
        if digits > 15:
            self.error(ctx, "EXCEL-NUMBER-SIGNIFICANT-DIGITS", f"此数值有 {digits} 位有效数字；Excel 入口最多接受 15 位。", location)
            return None
        return value

    def text(self, ctx: _UnitContext, sheet: str, cell, *, required: bool = False) -> str | None:
        location = f"{sheet}!{cell.coordinate}"
        if cell.data_type == "f":
            self.error(ctx, "EXCEL-FORMULA-REJECTED", "此字段不接受公式，请填写固定文字或数值。", location)
            return None
        if cell.value is None:
            if required:
                self.error(ctx, "EXCEL-TEXT-REQUIRED", "请填写此字段。", location)
            return None
        if not isinstance(cell.value, str):
            self.error(ctx, "EXCEL-TEXT-CELL-REQUIRED", "此字段必须以文字填写。", location)
            return None
        value = cell.value.strip()
        if required and not value:
            self.error(ctx, "EXCEL-TEXT-REQUIRED", "请填写此字段。", location)
        return value or None


def _stable_token(prefix: str, *parts: object) -> str:
    payload = "|".join(str(part) for part in parts)
    return f"{prefix}-{hashlib.sha256(payload.encode('utf-8')).hexdigest()[:20]}"


def _table_area(sheet, marker: str, headers: Sequence[str]) -> tuple[int, int, int]:
    # Row 1 is the visible page title and can equal the section marker; the
    # actual tabular section marker begins below that title row.
    markers = [row for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value == marker]
    if len(markers) != 1:
        raise WorkbookFatalError(f"工作表“{sheet.title}”中“{marker}”区域缺失或重复。")
    marker_row = markers[0]
    header_row = marker_row + 2
    actual = tuple(sheet.cell(header_row, column).value for column in range(1, len(headers) + 1))
    if actual != tuple(headers):
        raise WorkbookFatalError(f"工作表“{sheet.title}”中“{marker}”表头与当前模板不匹配。")
    later_markers = [
        row for row in range(header_row + 1, sheet.max_row + 1)
        if isinstance(sheet.cell(row, 1).value, str) and sheet.cell(row, 1).value in SECTION_HEADERS
    ]
    first_data = header_row + 1
    last_data = (min(later_markers) - 1) if later_markers else sheet.max_row
    return header_row, first_data, max(first_data - 1, last_data)


def _cell(sheet, row: int, header_row: int, headers: Sequence[str], header: str):
    try:
        column = headers.index(header) + 1
    except ValueError as exc:
        raise WorkbookFatalError(f"工作表“{sheet.title}”缺少列“{header}”。") from exc
    return sheet.cell(row, column)


def _location(sheet, row: int, header_row: int, headers: Sequence[str], header: str) -> str:
    return f"{sheet.title}!{_cell(sheet, row, header_row, headers, header).coordinate}"


def _has_payload(sheet, row: int, headers: Sequence[str]) -> bool:
    return any(sheet.cell(row, column).value is not None for column in range(2, len(headers) + 1))


def _iter_payload_rows(sheet, area: tuple[int, int, int], headers: Sequence[str]) -> Iterable[int]:
    _header_row, first, last = area
    for row in range(first, last + 1):
        if _has_payload(sheet, row, headers):
            yield row


def _metadata(workbook) -> dict[str, str]:
    sheet = workbook[METADATA_SHEET]
    values: dict[str, str] = {}
    for row in range(2, sheet.max_row + 1):
        key, value = sheet.cell(row, 1).value, sheet.cell(row, 2).value
        if isinstance(key, str) and isinstance(value, str):
            values[key] = value
    return values


def _date_cell(reader: _CellReader, ctx: _UnitContext, sheet, row: int, header: str) -> date | None:
    cell_value = sheet.cell(row, 2)
    label = sheet.cell(row, 1).value
    if label != header:
        raise WorkbookFatalError(f"核算单元页缺少项目字段“{header}”。")
    if cell_value.data_type == "f":
        reader.error(ctx, "EXCEL-FORMULA-REJECTED", "日期不接受公式。", f"{sheet.title}!{cell_value.coordinate}")
        return None
    if isinstance(cell_value.value, datetime):
        return cell_value.value.date()
    if isinstance(cell_value.value, date):
        return cell_value.value
    if isinstance(cell_value.value, str):
        try:
            return date.fromisoformat(cell_value.value.strip())
        except ValueError:
            pass
    reader.error(ctx, "EXCEL-DATE-INVALID", "日期请填写 Excel 日期或 YYYY-MM-DD。", f"{sheet.title}!{cell_value.coordinate}")
    return None


def _project_value(reader: _CellReader, ctx: _UnitContext, sheet, label: str, *, required: bool = False) -> str | None:
    for row in range(3, 9):
        if sheet.cell(row, 1).value == label:
            return reader.text(ctx, sheet.title, sheet.cell(row, 2), required=required)
    raise WorkbookFatalError(f"核算单元页缺少项目字段“{label}”。")


def _read_unit_contexts(workbook, digest: str, reader: _CellReader, warnings: list[ImportMessage]) -> tuple[dict[str, _UnitContext], set[str]]:
    sheet = workbook["基本信息"]
    project_probe = _UnitContext("项目", "project", AccountingUnitType.WHOLE_SITE, "enterprise.invalid", "", None, False, None)
    enterprise = _project_value(reader, project_probe, sheet, "企业名称") or None
    period_label = _project_value(reader, project_probe, sheet, "核算期间类型", required=True)
    boundary_label = _project_value(reader, project_probe, sheet, "已核对核算边界", required=True)
    boundary_description = _project_value(reader, project_probe, sheet, "公共边界说明")
    start = _date_cell(reader, project_probe, sheet, 5, "开始日期")
    end = _date_cell(reader, project_probe, sheet, 6, "结束日期")
    period = None
    period_type = PERIOD_LABELS.get(period_label or "")
    if period_type is None:
        reader.error(project_probe, "EXCEL-PERIOD-TYPE-INVALID", "请选择年度、月度或自定义期间。", "基本信息!B4")
    elif start is not None and end is not None:
        try:
            period = AccountingPeriod(period_type, start, end)
        except DomainValidationError as exc:
            reader.error(project_probe, "EXCEL-PERIOD-INVALID", f"核算期间无效：{exc}", "基本信息!B5:B6")
    elif not project_probe.errors:
        reader.error(project_probe, "EXCEL-PERIOD-REQUIRED", "请填写开始日期和结束日期。", "基本信息!B5:B6")
    boundary_confirmed = boundary_label == "是"
    if boundary_label not in {"是", "否"}:
        reader.error(project_probe, "EXCEL-BOUNDARY-REQUIRED", "请选择是否已核对核算边界。", "基本信息!B7")
    enterprise_id = _stable_token("enterprise", enterprise)
    contexts: dict[str, _UnitContext] = {}
    disabled: set[str] = set()
    unit_area = _table_area(sheet, "核算单元清单", UNIT_HEADERS)
    header_row, first, last = unit_area
    for row in range(first, last + 1):
        if not _has_payload(sheet, row, UNIT_HEADERS):
            continue
        name = reader.text(project_probe, sheet.title, _cell(sheet, row, header_row, UNIT_HEADERS, UNIT_HEADERS[0]), required=True)
        type_label = reader.text(project_probe, sheet.title, _cell(sheet, row, header_row, UNIT_HEADERS, UNIT_HEADERS[1]), required=True)
        enabled_label = reader.text(project_probe, sheet.title, _cell(sheet, row, header_row, UNIT_HEADERS, UNIT_HEADERS[2]), required=True)
        unit_boundary = reader.text(project_probe, sheet.title, _cell(sheet, row, header_row, UNIT_HEADERS, UNIT_HEADERS[3]))
        if not name:
            warnings.append(ImportMessage("EXCEL-UNIT-NAME-MISSING", "核算单元名称缺失，此行不能关联业务数据。", f"基本信息!A{row}"))
            continue
        unit_type = UNIT_TYPE_LABELS.get(type_label or "")
        if unit_type is None:
            unit_type = AccountingUnitType.OTHER
        enabled = enabled_label == "是"
        if enabled_label not in {"是", "否"}:
            enabled = True
        if not enabled:
            disabled.add(name)
            continue
        unit_id = _stable_token("unit", enterprise_id, name)
        ctx = _UnitContext(name, unit_id, unit_type, enterprise_id, enterprise, period, boundary_confirmed, unit_boundary or boundary_description)
        ctx.errors.extend(project_probe.errors)
        if type_label not in UNIT_TYPE_LABELS:
            ctx.errors.append(ImportMessage("EXCEL-UNIT-TYPE-INVALID", "请选择全厂、工序或其他。", f"基本信息!B{row}"))
        if enabled_label not in {"是", "否"}:
            ctx.errors.append(ImportMessage("EXCEL-UNIT-ENABLED-INVALID", "请选择是否启用此核算单元。", f"基本信息!C{row}"))
        if not boundary_confirmed:
            ctx.errors.append(ImportMessage("EXCEL-BOUNDARY-UNCONFIRMED", "请先确认企业核算边界后再预览。", "基本信息!B7"))
        if name in contexts:
            contexts[name].errors.append(ImportMessage("EXCEL-UNIT-DUPLICATE", f"核算单元“{name}”重复，无法可靠关联数据。", f"基本信息!A{row}"))
            ctx.errors.append(ImportMessage("EXCEL-UNIT-DUPLICATE", f"核算单元“{name}”重复，无法可靠关联数据。", f"基本信息!A{row}"))
            continue
        contexts[name] = ctx
    if not contexts and not disabled:
        warnings.append(ImportMessage("EXCEL-UNIT-NONE", "请至少填写一个已启用的核算单元。", "基本信息"))
    return contexts, disabled


def _context_for_row(
    reader: _CellReader,
    sheet,
    row: int,
    header_row: int,
    headers: Sequence[str],
    contexts: Mapping[str, _UnitContext],
    disabled: set[str],
    warnings: list[ImportMessage],
) -> _UnitContext | None:
    cell_value = _cell(sheet, row, header_row, headers, headers[0])
    if not isinstance(cell_value.value, str) or not cell_value.value.strip():
        warnings.append(ImportMessage("EXCEL-ROW-UNASSOCIATED", "数据行没有核算单元，已忽略。", f"{sheet.title}!{cell_value.coordinate}"))
        return None
    name = cell_value.value.strip()
    if name in disabled:
        warnings.append(ImportMessage("EXCEL-DISABLED-UNIT-DATA-IGNORED", f"“{name}”未启用，此行数据已忽略。", f"{sheet.title}!{cell_value.coordinate}"))
        return None
    ctx = contexts.get(name)
    if ctx is None:
        warnings.append(ImportMessage("EXCEL-UNKNOWN-UNIT-DATA-IGNORED", f"无法关联到已启用核算单元“{name}”，此行数据已忽略。", f"{sheet.title}!{cell_value.coordinate}"))
    return ctx


def _mark(ctx: _UnitContext, source_id: str) -> None:
    ctx.source_states[source_id] = EmissionSourceStatus.INVOLVED


def _activity_source(label: str | None) -> ActivityDataSource:
    return {
        "计量/仪表记录": ActivityDataSource.METER,
        "生产/能源台账": ActivityDataSource.PRODUCTION_LEDGER,
        "结算记录": ActivityDataSource.ENERGY_BILL,
        "检测报告": ActivityDataSource.TEST_REPORT,
        "实测值": ActivityDataSource.TEST_REPORT,
        "化学计算": ActivityDataSource.OTHER,
        "用户指定": ActivityDataSource.MANUAL,
    }.get(label or "", ActivityDataSource.MANUAL)


PARAMETER_SOURCE_KINDS = {
    "实测值": ParameterSourceKind.MEASURED,
    "化学计算": ParameterSourceKind.CALCULATED,
    "计算值": ParameterSourceKind.CALCULATED,
    "用户指定": ParameterSourceKind.USER_DEFINED,
}


def _parse_optional_text(reader: _CellReader, ctx: _UnitContext, sheet, row: int, header_row: int, headers: Sequence[str], header: str) -> str | None:
    return reader.text(ctx, sheet.title, _cell(sheet, row, header_row, headers, header))


def _parse_number(reader: _CellReader, ctx: _UnitContext, sheet, row: int, header_row: int, headers: Sequence[str], header: str, *, required: bool = False) -> Decimal | None:
    return reader.number(ctx, sheet.title, row, _cell(sheet, row, header_row, headers, header), required=required)


def _evidence_for_row(
    ctx: _UnitContext,
    source_id: str,
    sheet,
    row: int,
    kind: str | None,
    reference: str | None,
    *,
    parameter: bool = False,
) -> tuple[str, ...]:
    if not kind and not reference:
        return ()
    evidence_id = _stable_token("excel-parameter" if parameter else "excel-activity", ctx.unit_id, sheet.title, row, reference or "")
    applies_to = f"{SOURCE_LABELS[source_id]}工作表输入"
    note = f"表格数据来源：{kind or '未注明'}；单元格位置：{sheet.title} 第 {row} 行。"
    if parameter:
        ctx.factor_evidence.append(MeasuredFactorEvidence(
            evidence_id=evidence_id,
            applies_to=applies_to,
            source_ids=(source_id,),
            source_reference=reference,
            referenced_standard=STANDARD_ID,
            reason=note,
        ))
    else:
        ctx.activity_evidence.append(ActivityDataEvidence(
            evidence_id=evidence_id,
            applies_to=applies_to,
            source_ids=(source_id,),
            source_reference=reference,
            note=note,
        ))
    return (evidence_id,)


def _activity(value: Decimal | None, unit: str, source_type: ActivityDataSource, reference: str | None, evidence_ids: tuple[str, ...] = ()) -> InputValue | None:
    if value is None:
        return None
    return InputValue(value, unit, source_type=source_type, source_reference=reference, evidence_ref_ids=evidence_ids)


def _factor_for(resolver: ParameterResolver | None, parameter_id: str, period: AccountingPeriod | None):
    if resolver is None:
        return None
    try:
        factors = resolver.repository.list_factors(parameter_id)
    except (AttributeError, TypeError):
        return None
    matches = [
        factor for factor in factors
        if factor.parameter_id == parameter_id
        and STANDARD_ID in factor.applicable_standard_ids
        and factor.review_status is ReviewStatus.VERIFIED
        and factor.value_type in {ValueType.STANDARD_DEFAULT, ValueType.STANDARD_SPECIFIED}
        and not (period is not None and (
            (factor.valid_from is not None and factor.valid_from > period.end)
            or (factor.valid_to is not None and factor.valid_to < period.start)
            or (period.period_type is PeriodType.CUSTOM and (
                (factor.valid_from is not None and period.start < factor.valid_from)
                or (factor.valid_to is not None and period.end > factor.valid_to)
            ))
        ))
    ]
    return sorted(matches, key=lambda factor: factor.factor_id)[0] if matches else None


_PARAMETER_SOURCE_KINDS_FROM_VALUE = {
    ValueType.STANDARD_DEFAULT: ParameterSourceKind.STANDARD_DEFAULT,
    ValueType.STANDARD_SPECIFIED: ParameterSourceKind.STANDARD_SPECIFIED,
    ValueType.GOVERNMENT_PUBLISHED: ParameterSourceKind.OFFICIAL_PUBLISHED,
    ValueType.MEASURED: ParameterSourceKind.MEASURED,
    ValueType.DERIVED: ParameterSourceKind.CALCULATED,
    ValueType.SYSTEM_CONSTANT: ParameterSourceKind.CALCULATED,
    ValueType.SCIENTIFIC_REFERENCE: ParameterSourceKind.USER_DEFINED,
    ValueType.HISTORICAL: ParameterSourceKind.USER_DEFINED,
}


def _parameter_from_factor(factor, reason: str) -> ParameterValue | None:
    if factor is None:
        return None
    return ParameterValue(
        parameter_id=factor.parameter_id,
        value=factor.value,
        unit=factor.unit,
        source_kind=_PARAMETER_SOURCE_KINDS_FROM_VALUE[factor.value_type],
        source_id=factor.source_id,
        source_version=str(factor.version),
        source_location=factor.source_location,
        selection_reason=reason,
        factor_id=factor.factor_id,
        factor_year=factor.factor_year,
    )


def _make_parameter(
    value: Decimal | None,
    *,
    parameter_id: str,
    unit: str,
    source_kind_label: str | None,
    source_reference: str | None,
    evidence_ids: tuple[str, ...],
    default: ParameterValue | None,
    ctx: _UnitContext,
    location: str,
) -> ParameterValue | None:
    if value is None:
        if source_kind_label and source_kind_label not in {"缺省值"}:
            ctx.errors.append(ImportMessage("EXCEL-PARAMETER-VALUE-REQUIRED", "已选择自定义参数来源，请填写对应数值。", location))
        return default
    if source_kind_label == "缺省值":
        if default is None:
            ctx.errors.append(ImportMessage("EXCEL-PARAMETER-DEFAULT-UNAVAILABLE", "此参数没有可用的标准缺省值，请填写实际数值及来源。", location))
            return None
        if value != default.value:
            ctx.errors.append(ImportMessage("EXCEL-PARAMETER-DEFAULT-MISMATCH", "选择标准缺省值时请留空数值；如需使用其他数值，请选择相应来源类型。", location))
            return None
        return default
    if default is not None and value == default.value and not source_kind_label and not source_reference:
        return default
    kind = PARAMETER_SOURCE_KINDS.get(source_kind_label or "实测值", ParameterSourceKind.MEASURED)
    if not source_kind_label:
        ctx.warnings.append(ImportMessage("EXCEL-PARAMETER-SOURCE-UNSPECIFIED", "自定义参数未选择来源类型，已按实测值记录；请核对数据来源。", location))
    if not source_reference:
        ctx.warnings.append(ImportMessage("EXCEL-PARAMETER-REFERENCE-MISSING", "参数来源说明未填写；当前仍可继续预览，建议补充追溯信息。", location))
    return ParameterValue(
        parameter_id=parameter_id,
        value=value,
        unit=unit,
        source_kind=kind,
        source_id=f"USER-EXCEL-{hashlib.sha256((source_reference or 'user-input').encode('utf-8')).hexdigest()[:12]}",
        source_version="user-input",
        source_location=f"来源说明/编号：{source_reference}" if source_reference else None,
        selection_reason=f"工作簿中按“{source_kind_label or '未注明'}”提供；来源：{source_reference or '未提供'}。",
        evidence_ref_ids=evidence_ids,
    )


def _parse_fuels(workbook, contexts, disabled, reader, warnings, resolver) -> None:
    sheet = workbook[SHEET_FOR_MARKER["B.2 化石燃料"]]
    area = _table_area(sheet, "B.2 化石燃料", FUEL_HEADERS)
    header_row = area[0]
    for row in _iter_payload_rows(sheet, area, FUEL_HEADERS):
        ctx = _context_for_row(reader, sheet, row, header_row, FUEL_HEADERS, contexts, disabled, warnings)
        if ctx is None:
            continue
        _mark(ctx, SOURCE_FUEL)
        fuel_label = _parse_optional_text(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[1])
        path_label = _parse_optional_text(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[2])
        fuel_type = FUEL_TYPE_LABELS.get(fuel_label or "")
        fuel_path = FUEL_PATH_LABELS.get(path_label or "")
        if fuel_type is None:
            reader.error(ctx, "EXCEL-FUEL-TYPE-INVALID", "请选择明确的燃料品种；不得用“煤”代替具体煤种。", _location(sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[1]))
        if fuel_path is None:
            reader.error(ctx, "EXCEL-FUEL-PATH-INVALID", "请选择体积、质量或热量计量方式。", _location(sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[2]))
        if fuel_type is None or fuel_path is None:
            continue
        activity = _parse_number(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[3], required=True)
        activity_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[4])
        activity_ref = _parse_optional_text(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[5])
        lhv_value = _parse_number(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[6])
        lhv_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[7])
        lhv_ref = _parse_optional_text(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[8])
        carbon_value = _parse_number(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[9])
        carbon_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[10])
        carbon_ref = _parse_optional_text(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[11])
        oxidation_value = _parse_number(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[12])
        oxidation_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[13])
        oxidation_ref = _parse_optional_text(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[14])
        custom_name = _parse_optional_text(reader, ctx, sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[15])
        subject = FUEL_C1_SUBJECT_IDS.get(fuel_type)
        if fuel_type in {FuelType.OTHER, FuelType.COAL} and not custom_name:
            reader.error(ctx, "EXCEL-FUEL-CUSTOM-NAME-REQUIRED", "请填写实际燃料名称，软件不会为未明确燃料套用标准缺省值。", _location(sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[15]))
        default_carbon = _parameter_from_factor(_factor_for(resolver, f"{subject}_carbon_content", ctx.period) if subject else None, "采用附录 C.1 对应燃料的单位热值含碳量。")
        default_oxidation = _parameter_from_factor(_factor_for(resolver, f"{subject}_oxidation_rate", ctx.period) if subject else None, "采用附录 C.1 对应燃料的碳氧化率。")
        default_lhv = _parameter_from_factor(_factor_for(resolver, f"{subject}_lhv", ctx.period) if subject and fuel_path is not FuelPath.HEAT else None, "采用附录 C.1 对应燃料的低位发热量。")
        if default_lhv is not None and fuel_path is not FuelPath.HEAT:
            expected = FuelPath.VOLUME if default_lhv.unit == "GJ/10⁴Nm³" else FuelPath.MASS if default_lhv.unit == "GJ/t" else None
            if expected is not None and fuel_path is not expected and lhv_value is None:
                reader.error(ctx, "EXCEL-FUEL-PATH-MISMATCH", "计量方式与附录 C.1 参数单位不一致；请选择适用方式或填写有依据的实测参数。", _location(sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[2]))
        carbon_unit = "tC/GJ" if fuel_path is FuelPath.HEAT or lhv_value is not None or default_lhv is not None else "tC/10^4Nm3" if fuel_path is FuelPath.VOLUME else "tC/t"
        lhv_unit = {FuelPath.VOLUME: "GJ/10⁴Nm³", FuelPath.MASS: "GJ/t", FuelPath.HEAT: "GJ/GJ"}[fuel_path]
        lhv_ids = _evidence_for_row(ctx, SOURCE_FUEL, sheet, row, lhv_kind, lhv_ref, parameter=True) if lhv_value is not None else ()
        carbon_ids = _evidence_for_row(ctx, SOURCE_FUEL, sheet, row, carbon_kind, carbon_ref, parameter=True) if carbon_value is not None else ()
        oxidation_ids = _evidence_for_row(ctx, SOURCE_FUEL, sheet, row, oxidation_kind, oxidation_ref, parameter=True) if oxidation_value is not None else ()
        lhv = _make_parameter(lhv_value, parameter_id=f"{subject or 'fuel'}_lhv", unit=lhv_unit, source_kind_label=lhv_kind, source_reference=lhv_ref, evidence_ids=lhv_ids, default=default_lhv, ctx=ctx, location=_location(sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[6]))
        carbon = _make_parameter(carbon_value, parameter_id=f"{subject or 'fuel'}_carbon_content", unit=carbon_unit, source_kind_label=carbon_kind, source_reference=carbon_ref, evidence_ids=carbon_ids, default=default_carbon, ctx=ctx, location=_location(sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[9]))
        oxidation = _make_parameter(oxidation_value, parameter_id=f"{subject or 'fuel'}_oxidation_rate", unit="ratio", source_kind_label=oxidation_kind, source_reference=oxidation_ref, evidence_ids=oxidation_ids, default=default_oxidation, ctx=ctx, location=_location(sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[12]))
        if fuel_path is FuelPath.HEAT and lhv_value is not None:
            reader.error(ctx, "EXCEL-FUEL-LHV-NOT-APPLICABLE", "热量路径的活动量已是热量，不应重复填写低位发热量。", _location(sheet, row, header_row, FUEL_HEADERS, FUEL_HEADERS[6]))
            lhv = None
        activity_ids = _evidence_for_row(ctx, SOURCE_FUEL, sheet, row, activity_kind, activity_ref)
        activity_unit = {FuelPath.VOLUME: "ten_thousand_Nm3", FuelPath.MASS: "t", FuelPath.HEAT: "GJ"}[fuel_path]
        ctx.payloads.setdefault("fuel_inputs", []).append(FuelInput(
            fuel_id=_stable_token("fuel", ctx.unit_id, sheet.title, row),
            fuel_type=fuel_type,
            fuel_label=custom_name or (fuel_label if fuel_type in {FuelType.OTHER, FuelType.COAL} else None),
            path=fuel_path,
            activity=_activity(activity, activity_unit, _activity_source(activity_kind), activity_ref, activity_ids),
            carbon_content=carbon,
            oxidation_rate=oxidation,
            lower_heating_value=lhv,
        ))

def _material_source(label: str | None) -> MaterialDataSource:
    return {
        "缺省值": MaterialDataSource.STANDARD_DEFAULT,
        "实测值": MaterialDataSource.MEASURED,
        "化学计算": MaterialDataSource.CHEMICAL_CALCULATION,
    }.get(label or "实测值", MaterialDataSource.MEASURED)


def _parse_process_materials(workbook, contexts, disabled, reader, warnings, resolver) -> None:
    process_for_marker = {
        "B.3 煅烧": ("calcination", SOURCE_CALCINATION, "calcinations", CalcinationInput, "car-par-k1", "k1"),
        "B.4 焙烧/炭化": ("baking", SOURCE_BAKING, "bakings", BakingInput, "car-par-k2", "k2"),
        "B.5 石墨化": ("graphitization", SOURCE_GRAPHITIZATION, "graphitizations", GraphitizationInput, "car-par-k3", "k3"),
    }
    groups: dict[tuple[str, str, str], list[MaterialInputLine]] = {}
    seen_instances: dict[tuple[str, str], tuple[object, ...]] = {}
    for marker, spec in PROCESS_SPECS.items():
        sheet = workbook[SHEET_FOR_MARKER[marker]]
        area = _table_area(sheet, marker, MATERIAL_HEADERS)
        header_row = area[0]
        for row in _iter_payload_rows(sheet, area, MATERIAL_HEADERS):
            ctx = _context_for_row(reader, sheet, row, header_row, MATERIAL_HEADERS, contexts, disabled, warnings)
            if ctx is None:
                continue
            prefix, source_id, payload_key, payload_type, parameter_id, parameter_name = process_for_marker[marker]
            _mark(ctx, source_id)
            instance = _parse_optional_text(reader, ctx, sheet, row, header_row, MATERIAL_HEADERS, MATERIAL_HEADERS[1])
            category = _parse_optional_text(reader, ctx, sheet, row, header_row, MATERIAL_HEADERS, MATERIAL_HEADERS[2])
            name = _parse_optional_text(reader, ctx, sheet, row, header_row, MATERIAL_HEADERS, MATERIAL_HEADERS[3]) or ""
            role = MATERIAL_ROLES[marker].get(category or "")
            if role is None:
                reader.error(ctx, "EXCEL-MATERIAL-ROLE-INVALID", "请选择本表对应的物料类别。", _location(sheet, row, header_row, MATERIAL_HEADERS, MATERIAL_HEADERS[2]))
                continue
            instance = instance or f"过程{row}"
            if not _parse_optional_text(reader, ctx, sheet, row, header_row, MATERIAL_HEADERS, MATERIAL_HEADERS[1]):
                reader.error(ctx, "EXCEL-PROCESS-NAME-REQUIRED", "请填写过程实例名称。", _location(sheet, row, header_row, MATERIAL_HEADERS, MATERIAL_HEADERS[1]))
            mass = _parse_number(reader, ctx, sheet, row, header_row, MATERIAL_HEADERS, MATERIAL_HEADERS[4])
            fixed = _parse_number(reader, ctx, sheet, row, header_row, MATERIAL_HEADERS, MATERIAL_HEADERS[5])
            fixed_source = _parse_optional_text(reader, ctx, sheet, row, header_row, MATERIAL_HEADERS, MATERIAL_HEADERS[6])
            volatile = _parse_number(reader, ctx, sheet, row, header_row, MATERIAL_HEADERS, MATERIAL_HEADERS[7])
            volatile_source = _parse_optional_text(reader, ctx, sheet, row, header_row, MATERIAL_HEADERS, MATERIAL_HEADERS[8])
            line = MaterialInputLine(
                line_id=_stable_token("material", ctx.unit_id, marker, row, instance),
                role=role,
                name=name,
                mass_t=mass,
                fixed_carbon_percent=fixed,
                fixed_carbon_source=_material_source(fixed_source),
                volatile_matter_percent=volatile,
                volatile_matter_source=_material_source(volatile_source),
            )
            key = (ctx.name, marker, instance)
            groups.setdefault(key, []).append(line)
            seen_instances[(ctx.name, marker)] = (prefix, source_id, payload_key, payload_type, parameter_id, parameter_name)
    for (unit_name, marker, instance), lines in groups.items():
        ctx = contexts[unit_name]
        prefix, source_id, payload_key, payload_type, parameter_id, parameter_name = process_for_marker[marker]
        k_value = _parameter_from_factor(
            _factor_for(resolver, parameter_id, ctx.period),
            f"采用 GB/T 32151.34—2024 第5.2.{ {'calcination': '2', 'baking': '3', 'graphitization': '4'}[prefix] }条一般取值。",
        )
        payload = payload_type(
            instance_id=_stable_token("process", ctx.unit_id, marker, instance),
            material_rows=tuple(lines),
            **{parameter_name: k_value},
        )
        ctx.payloads.setdefault(payload_key, []).append(payload)

def _parse_fume(workbook, contexts, disabled, reader, warnings) -> None:
    sheet = workbook[SHEET_FOR_MARKER["B.6 烟气焚烧"]]
    area = _table_area(sheet, "B.6 烟气焚烧", FUME_HEADERS)
    header_row = area[0]
    labels = ("烟气流量（Nm³/h，必填）", "焦油含量（mg/Nm³，必填）", "低位发热量（GJ/t，必填）", "单位热值含碳量（tC/GJ，必填）", "碳氧化率（必填）", "运行时间（d，必填）")
    attrs = ("q", "qvar", "hm", "fch", "fox", "duration")
    units = {"q": "Nm3/h", "qvar": "mg/Nm3", "hm": "GJ/t", "fox": "ratio", "duration": "d"}
    for row in _iter_payload_rows(sheet, area, FUME_HEADERS):
        ctx = _context_for_row(reader, sheet, row, header_row, FUME_HEADERS, contexts, disabled, warnings)
        if ctx is None:
            continue
        _mark(ctx, SOURCE_FUME)
        instance = _parse_optional_text(reader, ctx, sheet, row, header_row, FUME_HEADERS, FUME_HEADERS[1])
        if not instance:
            reader.error(ctx, "EXCEL-FUME-NAME-REQUIRED", "请填写烟气焚烧实例名称。", _location(sheet, row, header_row, FUME_HEADERS, FUME_HEADERS[1]))
            instance = f"烟气焚烧{row}"
        numbers = {attr: _parse_number(reader, ctx, sheet, row, header_row, FUME_HEADERS, label, required=True) for attr, label in zip(attrs, labels, strict=True)}
        activity_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, FUME_HEADERS, FUME_HEADERS[8])
        activity_ref = _parse_optional_text(reader, ctx, sheet, row, header_row, FUME_HEADERS, FUME_HEADERS[9])
        parameter_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, FUME_HEADERS, FUME_HEADERS[10])
        parameter_ref = _parse_optional_text(reader, ctx, sheet, row, header_row, FUME_HEADERS, FUME_HEADERS[11])
        activity_ids = _evidence_for_row(ctx, SOURCE_FUME, sheet, row, activity_kind, activity_ref)
        factor_ids = _evidence_for_row(ctx, SOURCE_FUME, sheet, row, parameter_kind, parameter_ref, parameter=True)
        fch_value = _make_parameter(numbers["fch"], parameter_id="CAR-PAR-P04A-FCH", unit="tC/GJ", source_kind_label=parameter_kind, source_reference=parameter_ref, evidence_ids=factor_ids, default=None, ctx=ctx, location=_location(sheet, row, header_row, FUME_HEADERS, FUME_HEADERS[5]))
        for key, value in tuple(numbers.items()):
            if key == "fch":
                numbers[key] = fch_value
            else:
                numbers[key] = _activity(value, units[key], _activity_source(activity_kind), activity_ref, activity_ids)
        ctx.payloads.setdefault("fume_incinerations", []).append(FumeIncinerationInput(**numbers, instance_id=_stable_token("fume", ctx.unit_id, instance)))


def _parse_fgd(workbook, contexts, disabled, reader, warnings, resolver) -> None:
    sheet = workbook[SHEET_FOR_MARKER["B.7 烟气脱硫"]]
    area = _table_area(sheet, "B.7 烟气脱硫", FGD_HEADERS)
    header_row = area[0]
    components: dict[tuple[str, str], list[CarbonateComponent]] = {}
    for row in _iter_payload_rows(sheet, area, FGD_HEADERS):
        ctx = _context_for_row(reader, sheet, row, header_row, FGD_HEADERS, contexts, disabled, warnings)
        if ctx is None:
            continue
        _mark(ctx, SOURCE_FGD)
        facility = _parse_optional_text(reader, ctx, sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[1])
        carbonate = _parse_optional_text(reader, ctx, sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[2])
        amount = _parse_number(reader, ctx, sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[3], required=True)
        fraction_value = _parse_number(reader, ctx, sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[4])
        factor_value = _parse_number(reader, ctx, sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[5])
        conversion_value = _parse_number(reader, ctx, sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[6])
        activity_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[7])
        parameter_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[8])
        source_ref = _parse_optional_text(reader, ctx, sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[9])
        if not facility:
            reader.error(ctx, "EXCEL-FGD-FACILITY-REQUIRED", "请填写脱硫设施或批次名称。", _location(sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[1]))
        parameter_id = CARBONATE_LABELS.get(carbonate or "")
        if not parameter_id:
            reader.error(ctx, "EXCEL-CARBONATE-REQUIRED", "请选择碳酸盐种类；未知种类不能默认按 CaCO₃ 计算。", _location(sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[2]))
        if amount is None or not facility or not parameter_id:
            continue
        if not any(value is not None for value in (fraction_value, factor_value, conversion_value)) and (parameter_kind or source_ref):
            reader.error(ctx, "EXCEL-PARAMETER-SOURCE-UNUSED", "填写了参数来源，但没有填写任何自定义参数值。", _location(sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[8]))
        factor_ids = _evidence_for_row(ctx, SOURCE_FGD, sheet, row, parameter_kind, source_ref, parameter=True)
        activity_ids = _evidence_for_row(ctx, SOURCE_FGD, sheet, row, activity_kind, source_ref)
        default_fraction = _parameter_from_factor(_factor_for(resolver, "car-par-p04b-i", ctx.period), "採用 GB/T 32151.34—2024 脱硫碳酸盐含量标准缺省值。")
        default_conversion = _parameter_from_factor(_factor_for(resolver, "car-par-p04b-tr", ctx.period), "采用 GB/T 32151.34—2024 脱硫转化率标准缺省值。")
        default_factor = _parameter_from_factor(_factor_for(resolver, parameter_id, ctx.period), "按所选碳酸盐种类采用附录 C.2 对应因子。")
        fraction = _make_parameter(fraction_value, parameter_id="car-par-p04b-i", unit="ratio", source_kind_label=parameter_kind, source_reference=source_ref, evidence_ids=factor_ids, default=default_fraction, ctx=ctx, location=_location(sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[4]))
        factor = _make_parameter(factor_value, parameter_id=parameter_id, unit="tCO2/t", source_kind_label=parameter_kind, source_reference=source_ref, evidence_ids=factor_ids, default=default_factor, ctx=ctx, location=_location(sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[5]))
        conversion = _make_parameter(conversion_value, parameter_id="car-par-p04b-tr", unit="ratio", source_kind_label=parameter_kind, source_reference=source_ref, evidence_ids=factor_ids, default=default_conversion, ctx=ctx, location=_location(sheet, row, header_row, FGD_HEADERS, FGD_HEADERS[6]))
        components.setdefault((ctx.name, facility), []).append(CarbonateComponent(
            amount=_activity(amount, "t", _activity_source(activity_kind), source_ref, activity_ids),
            carbonate_fraction=fraction,
            emission_factor=factor,
            conversion_rate=conversion,
            carbonate_type=carbonate,
        ))
    for (unit_name, facility), values in components.items():
        contexts[unit_name].payloads.setdefault("fgd_units", []).append(FGDInput(
            components=tuple(values),
            instance_id=_stable_token("fgd", contexts[unit_name].unit_id, facility),
        ))


def _parse_electricity(workbook, contexts, disabled, reader, warnings) -> None:
    sheet = workbook[SHEET_FOR_MARKER["B.8 电力"]]
    area = _table_area(sheet, "B.8 电力", ELECTRICITY_HEADERS)
    header_row = area[0]
    directions = {"购入": "purchased", "输出": "exported"}
    acquisitions = {"购入": ElectricityAcquisitionMode.PURCHASED, "自发自用": ElectricityAcquisitionMode.SELF_CONSUMED}
    attributes = {"常规": ElectricityAttribute.ORDINARY, "非化石": ElectricityAttribute.NONFOSSIL, "化石": ElectricityAttribute.FOSSIL}
    proofs = {"无": ElectricityProofType.NONE, "合同和结算": ElectricityProofType.CONTRACT_AND_SETTLEMENT, "GEC": ElectricityProofType.GEC, "月度原始记录": ElectricityProofType.MONTHLY_ORIGINAL_RECORD}
    statuses = {"未提供": ElectricityProofStatus.NOT_PROVIDED, "有效": ElectricityProofStatus.VALID, "无效": ElectricityProofStatus.INVALID}
    for row in _iter_payload_rows(sheet, area, ELECTRICITY_HEADERS):
        ctx = _context_for_row(reader, sheet, row, header_row, ELECTRICITY_HEADERS, contexts, disabled, warnings)
        if ctx is None:
            continue
        direction_label = _parse_optional_text(reader, ctx, sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[1])
        direction = directions.get(direction_label or "")
        if direction is None:
            _mark(ctx, SOURCE_PURCHASED_ELECTRICITY)
            _mark(ctx, SOURCE_EXPORTED_ELECTRICITY)
            reader.error(ctx, "EXCEL-ENERGY-DIRECTION-REQUIRED", "请选择购入或输出。", _location(sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[1]))
            continue
        source_id = SOURCE_PURCHASED_ELECTRICITY if direction == "purchased" else SOURCE_EXPORTED_ELECTRICITY
        _mark(ctx, source_id)
        amount = _parse_number(reader, ctx, sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[2], required=True)
        factor_value = _parse_number(reader, ctx, sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[3])
        activity_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[4])
        activity_ref = _parse_optional_text(reader, ctx, sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[5])
        parameter_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[6])
        parameter_ref = _parse_optional_text(reader, ctx, sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[7])
        activity_ids = _evidence_for_row(ctx, source_id, sheet, row, activity_kind, activity_ref)
        if amount is None:
            continue
        if direction == "purchased":
            if factor_value is not None or parameter_kind or parameter_ref:
                reader.error(ctx, "EXCEL-PURCHASED-ELECTRICITY-FACTOR-AUTO", "购入电力因子由软件按核算期间、属性和证明信息自动匹配；请勿填写输出电力因子栏。", _location(sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[3]))
            acquisition_label = _parse_optional_text(reader, ctx, sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[8]) or "购入"
            attribute_label = _parse_optional_text(reader, ctx, sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[9]) or "常规"
            proof_label = _parse_optional_text(reader, ctx, sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[10]) or "无"
            status_label = _parse_optional_text(reader, ctx, sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[11]) or "未提供"
            acquisition, attribute, proof, status = acquisitions.get(acquisition_label), attributes.get(attribute_label), proofs.get(proof_label), statuses.get(status_label)
            if None in (acquisition, attribute, proof, status):
                reader.error(ctx, "EXCEL-ELECTRICITY-OPTION-INVALID", "请检查电力取得方式、属性和证明选项。", f"{sheet.title}!A{row}")
                continue
            if attribute is ElectricityAttribute.NONFOSSIL and proof is ElectricityProofType.NONE:
                reader.error(ctx, "EXCEL-NONFOSSIL-PROOF-REQUIRED", "非化石电力需提供适用证明类型和有效状态。", _location(sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[10]))
            if ctx.period is None:
                ctx.errors.append(ImportMessage("EXCEL-PERIOD-REQUIRED", "购入电力需要有效核算期间。", f"{sheet.title}!A{row}"))
                continue
            ctx.payloads.setdefault("electricity_details", []).append(ElectricityConsumptionDetail(
                detail_id=_stable_token("electricity", ctx.unit_id, row),
                enterprise_id=ctx.enterprise_id,
                standard_id=STANDARD_ID,
                accounting_period=ctx.period,
                electricity_amount=amount,
                electricity_unit="MWh",
                acquisition_mode=acquisition,
                attribute=attribute,
                proof_type=proof,
                proof_status=status,
            ))
        else:
            if factor_value is None and (parameter_kind or parameter_ref):
                reader.error(ctx, "EXCEL-PARAMETER-SOURCE-UNUSED", "填写了参数来源，但没有填写自定义输出电力因子。", _location(sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[6]))
            factor_ids = _evidence_for_row(ctx, source_id, sheet, row, parameter_kind, parameter_ref, parameter=True) if factor_value is not None else ()
            factor = _make_parameter(factor_value, parameter_id="electricity_emission_factor_measured", unit="tCO2/MWh", source_kind_label=parameter_kind, source_reference=parameter_ref, evidence_ids=factor_ids, default=None, ctx=ctx, location=_location(sheet, row, header_row, ELECTRICITY_HEADERS, ELECTRICITY_HEADERS[3]))
            ctx.payloads.setdefault("exported_electricity", []).append(ElectricityOutputLine(
                line_id=_stable_token("power-output", ctx.unit_id, row),
                amount=_activity(amount, "MWh", _activity_source(activity_kind), activity_ref, activity_ids),
                factor=factor,
                unit="MWh",
            ))


def _parse_heat(workbook, contexts, disabled, reader, warnings) -> None:
    sheet = workbook[SHEET_FOR_MARKER["B.9 热力"]]
    area = _table_area(sheet, "B.9 热力", HEAT_HEADERS)
    header_row = area[0]
    directions = {"购入": "purchased", "输出": "exported"}
    steam_kinds = {"饱和蒸汽": SteamKind.SATURATED, "过热蒸汽": SteamKind.SUPERHEATED}
    for row in _iter_payload_rows(sheet, area, HEAT_HEADERS):
        ctx = _context_for_row(reader, sheet, row, header_row, HEAT_HEADERS, contexts, disabled, warnings)
        if ctx is None:
            continue
        direction_label = _parse_optional_text(reader, ctx, sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[1])
        direction = directions.get(direction_label or "")
        if direction is None:
            _mark(ctx, SOURCE_PURCHASED_HEAT)
            _mark(ctx, SOURCE_EXPORTED_HEAT)
            reader.error(ctx, "EXCEL-HEAT-DIRECTION-REQUIRED", "请选择购入或输出。", _location(sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[1]))
            continue
        source_id = SOURCE_PURCHASED_HEAT if direction == "purchased" else SOURCE_EXPORTED_HEAT
        _mark(ctx, source_id)
        amount = _parse_number(reader, ctx, sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[2], required=True)
        steam_label = _parse_optional_text(reader, ctx, sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[3],)
        if not steam_label:
            reader.error(ctx, "EXCEL-STEAM-KIND-REQUIRED", "请选择饱和蒸汽或过热蒸汽。", _location(sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[3]))
        enthalpy = _parse_number(reader, ctx, sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[4])
        pressure = _parse_number(reader, ctx, sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[5])
        temperature = _parse_number(reader, ctx, sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[6])
        factor_value = _parse_number(reader, ctx, sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[7])
        activity_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[8])
        activity_ref = _parse_optional_text(reader, ctx, sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[9])
        parameter_kind = _parse_optional_text(reader, ctx, sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[10])
        parameter_ref = _parse_optional_text(reader, ctx, sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[11])
        steam = steam_kinds.get(steam_label or "")
        if steam is None:
            steam = SteamKind.SATURATED
        if enthalpy is None and pressure is None:
            reader.error(ctx, "EXCEL-STEAM-STATE-REQUIRED", "请填写焓值，或填写压力和温度以使用附录C.4/C.5蒸汽表。", f"{sheet.title}!A{row}")
        if steam is SteamKind.SUPERHEATED and enthalpy is None and temperature is None:
            reader.error(ctx, "EXCEL-STEAM-TEMPERATURE-REQUIRED", "过热蒸汽查表需要填写温度。", _location(sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[6]))
        if amount is None:
            continue
        if factor_value is None and (parameter_kind or parameter_ref):
            reader.error(ctx, "EXCEL-PARAMETER-SOURCE-UNUSED", "填写了参数来源，但没有填写自定义热力因子。", _location(sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[10]))
        activity_ids = _evidence_for_row(ctx, source_id, sheet, row, activity_kind, activity_ref)
        factor_ids = _evidence_for_row(ctx, source_id, sheet, row, parameter_kind, parameter_ref, parameter=True) if factor_value is not None else ()
        factor = _make_parameter(factor_value, parameter_id="heat_emission_factor_measured", unit="tCO2/GJ", source_kind_label=parameter_kind, source_reference=parameter_ref, evidence_ids=factor_ids, default=None, ctx=ctx, location=_location(sheet, row, header_row, HEAT_HEADERS, HEAT_HEADERS[7]))
        value = HeatInput(
            line_id=_stable_token("heat", ctx.unit_id, direction, row),
              amount=_activity(amount, "t", _activity_source(activity_kind), activity_ref, activity_ids),
            enthalpy=_activity(enthalpy, "kJ/kg", _activity_source(activity_kind), activity_ref, activity_ids),
            factor=factor,
              unit="t",
            steam_kind=steam,
            pressure_mpa=_activity(pressure, "MPa", _activity_source(activity_kind), activity_ref, activity_ids),
              temperature_c=_activity(temperature, "C", _activity_source(activity_kind), activity_ref, activity_ids),
              steam_amount_t=_activity(amount, "t", _activity_source(activity_kind), activity_ref, activity_ids),
        )
        ctx.payloads.setdefault("purchased_heat" if direction == "purchased" else "exported_heat", []).append(value)


def _make_input(ctx: _UnitContext, digest: str) -> CarbonMaterialInput:
    states = tuple(
        EmissionSourceState(source_id, ctx.source_states.get(source_id, EmissionSourceStatus.NOT_INVOLVED))
        for source_id in SOURCE_IDS
    )
    report = CarbonReportingData(
        boundary_description=ctx.boundary_description,
        activity_evidence=tuple(ctx.activity_evidence),
        measured_factor_evidence=tuple(ctx.factor_evidence),
    )
    return CarbonMaterialInput(
        input_id=_stable_token("input", ctx.unit_id, digest),
        enterprise_id=ctx.enterprise_id,
        enterprise_name=ctx.enterprise_name,
        period=ctx.period,
        boundary_confirmed=ctx.boundary_confirmed,
        boundary_component_ids=(_stable_token("boundary", ctx.unit_id),),
        source_states=states,
        fuel_inputs=tuple(ctx.payloads.get("fuel_inputs", ())),
        calcinations=tuple(ctx.payloads.get("calcinations", ())),
        bakings=tuple(ctx.payloads.get("bakings", ())),
        graphitizations=tuple(ctx.payloads.get("graphitizations", ())),
        fume_incinerations=tuple(ctx.payloads.get("fume_incinerations", ())),
        fgd_units=tuple(ctx.payloads.get("fgd_units", ())),
        electricity_details=tuple(ctx.payloads.get("electricity_details", ())),
        exported_electricity=tuple(ctx.payloads.get("exported_electricity", ())),
        purchased_heat=tuple(ctx.payloads.get("purchased_heat", ())),
        exported_heat=tuple(ctx.payloads.get("exported_heat", ())),
        reporting_data=report,
    )


class ExcelWorkbookImporter:
    """Read supported workbooks into independent, non-persistent unit previews."""

    def __init__(self, parameter_resolver: ParameterResolver | None = None):
        self.parameter_resolver = parameter_resolver

    def import_preview(self, file_path: str | Path) -> WorkbookImportPreview:
        path = Path(file_path)
        if path.suffix.lower() != ".xlsx":
            raise WorkbookFatalError("仅支持 .xlsx 工作簿。")
        try:
            content = path.read_bytes()
        except OSError as exc:
            raise WorkbookFatalError(f"无法读取工作簿：{exc}") from exc
        digest = hashlib.sha256(content).hexdigest().upper()
        try:
            import io
            workbook = load_workbook(io.BytesIO(content), data_only=False, read_only=False)
            numeric_xml = _serialized_numeric_values(content)
        except (BadZipFile, KeyError, OSError, ValueError, ElementTree.ParseError) as exc:
            raise WorkbookFatalError(f"工作簿损坏或结构无法解析：{exc}") from exc
        try:
            for canonical, aliases in SHEET_NAME_ALIASES.items():
                matches = [name for name in (canonical, *aliases) if name in workbook.sheetnames]
                if len(matches) > 1:
                    raise WorkbookFatalError(f"工作簿中存在重复的“{canonical}”工作表，无法确定应读取哪一张。")
                if matches and matches[0] != canonical:
                    # WPS rewrites the full-width slash to an underscore when it saves.
                    # Normalize this known alias in memory only; never rewrite the user file.
                    workbook[matches[0]].title = canonical
            missing = sorted(REQUIRED_SHEETS - set(workbook.sheetnames))
            if missing:
                raise WorkbookFatalError(f"缺少模板工作表：{', '.join(missing)}")
            for sheet in VISIBLE_SHEETS:
                if workbook[sheet].sheet_state != "visible":
                    raise WorkbookFatalError(f"用户工作表“{sheet}”必须可见。")
            if workbook[METADATA_SHEET].sheet_state == "visible":
                raise WorkbookFatalError("模板信息页必须保持隐藏。")
            for marker, headers in SECTION_HEADERS.items():
                _table_area(workbook[SHEET_FOR_MARKER[marker]], marker, headers)
            metadata = _metadata(workbook)
            expected = {
                "template_id": TEMPLATE_ID,
                "template_version": TEMPLATE_VERSION,
                "standard_id": STANDARD_ID,
                "standard_version": STANDARD_VERSION,
                "ingress_policy_id": INGRESS_POLICY_ID,
            }
            for key, value in expected.items():
                if metadata.get(key) != value:
                    raise WorkbookFatalError(f"模板 {key} 不匹配或版本无法识别。")

            warnings: list[ImportMessage] = []
            for name in sorted(set(workbook.sheetnames) - REQUIRED_SHEETS):
                warnings.append(ImportMessage("EXCEL-UNKNOWN-SHEET", f"自定义工作表“{name}”不参与核算，已忽略。", name))
            reader = _CellReader(numeric_xml)
            contexts, disabled = _read_unit_contexts(workbook, digest, reader, warnings)
            _parse_fuels(workbook, contexts, disabled, reader, warnings, self.parameter_resolver)
            _parse_process_materials(workbook, contexts, disabled, reader, warnings, self.parameter_resolver)
            _parse_fume(workbook, contexts, disabled, reader, warnings)
            _parse_fgd(workbook, contexts, disabled, reader, warnings, self.parameter_resolver)
            _parse_electricity(workbook, contexts, disabled, reader, warnings)
            _parse_heat(workbook, contexts, disabled, reader, warnings)

            imported_at = datetime.now(timezone.utc)
            provenance = WorkbookProvenance(digest, metadata["template_id"], metadata["template_version"], metadata["standard_id"], metadata["standard_version"], metadata["ingress_policy_id"], imported_at)
            previews: list[UnitCalculationPreview] = []
            preview_use_case = CarbonAccountingPreviewUseCase(
                CarbonMaterialCalculator(parameter_resolver=self.parameter_resolver)
            )
            for ctx in contexts.values():
                input_value = None
                try:
                    input_value = _make_input(ctx, digest)
                except (DomainValidationError, TypeError, ValueError) as exc:
                    ctx.errors.append(ImportMessage("EXCEL-DOMAIN-INPUT-INVALID", f"无法构造该核算单元的数据：{exc}"))
                calculation = None
                if input_value is not None and not ctx.errors:
                    outcome = preview_use_case.calculate(input_value, calculated_at=imported_at)
                    for problem in outcome.problems:
                        message = ImportMessage(problem.code, problem.message, problem.field_id)
                        if getattr(problem.level, "value", "ERROR") == "ERROR":
                            ctx.errors.append(message)
                        else:
                            ctx.warnings.append(message)
                    calculation = outcome
                previews.append(UnitCalculationPreview(ctx.unit_id, ctx.name, ctx.unit_type, input_value, calculation, tuple(ctx.errors), tuple(ctx.warnings)))
            return WorkbookImportPreview(provenance, tuple(previews), tuple(warnings), tuple(reader.evidence))
        finally:
            workbook.close()


__all__ = [
    "TEMPLATE_ID", "TEMPLATE_VERSION", "INGRESS_POLICY_ID", "METADATA_SHEET", "VISIBLE_SHEETS", "SHEET_NAME_ALIASES", "SECTION_HEADERS", "SHEET_FOR_MARKER",
    "FUEL_HEADERS", "ELECTRICITY_HEADERS", "HEAT_HEADERS", "MATERIAL_HEADERS", "FUME_HEADERS", "FGD_HEADERS",
    "WorkbookFatalError", "WorkbookProvenance", "NumericCellEvidence", "ImportMessage", "UnitCalculationPreview",
    "WorkbookImportPreview", "ExcelWorkbookImporter", "create_template_bytes", "write_template", "significant_digit_count",
]
