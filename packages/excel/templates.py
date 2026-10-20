"""Approved per-standard Excel resources; downloads always copy verified bytes."""
from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from importlib.resources import files
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from packages.standards.carbon_material import STANDARD_ID, STANDARD_VERSION
from .appendix_b_legacy_structure import LEGACY_HEADER_CELLS, LEGACY_HEADER_MERGES

APPROVED_SOURCE_SHA256 = "c805e446994e0863221063109f2b425a545d0c7db88544ba5583102219c1d9e4"
APPROVED_TEMPLATE_V2_SHA256 = "58b38c15d487bdc173f573a50e5b8f3446b5bb4df4693dd696a83742da5235b4"
APPROVED_TEMPLATE_V2_RESOURCE = "excel_templates/gb_t_32151_34_2024_appendix_b_v2.xlsx"
APPROVED_TEMPLATE_V3_SHA256 = "f303ad02ca03e0464f736dafc39967abd06366a902f1b502dd17470656eb1718"
APPROVED_TEMPLATE_V3_RESOURCE = "excel_templates/gb_t_32151_34_2024_appendix_b_v3.xlsx"
APPROVED_TEMPLATE_V4_SHA256 = "e9b83462eabd97d5c7bd9aae1b4fa97a27d80006a205907c6671e626e8e57f6f"
APPROVED_TEMPLATE_V4_RESOURCE = "excel_templates/gb_t_32151_34_2024_appendix_b_v4.xlsx"
APPROVED_TEMPLATE_V5_SHA256 = "7438082b59e6ba294bb953ccc44db0c7a3a16eacbf4f8ddd83f063ce492bb890"
APPROVED_TEMPLATE_V5_RESOURCE = "excel_templates/gb_t_32151_34_2024_appendix_b_v5.xlsx"
APPROVED_TEMPLATE_V6_RESOURCE = "excel_templates/gb_t_32151_34_2024_appendix_b_v6.xlsx"
APPROVED_TEMPLATE_V6_SHA256 = "ca61cf3133c2663c78ca7409d23f7dd20726e43a8c79adca2aa78faecb3c031a"
APPROVED_TEMPLATE_V7_RESOURCE = "excel_templates/gb_t_32151_34_2024_appendix_b_v7.xlsx"
APPROVED_TEMPLATE_V7_SHA256 = "e4203f0d5bf9133df5b6abde3bfede8b82905d11d795cd29ba590ea60f7ebb5a"
APPROVED_TEMPLATE_V8_RESOURCE = "excel_templates/gb_t_32151_34_2024_appendix_b_v8.xlsx"
APPROVED_TEMPLATE_V8_SHA256 = "a8ee7272ee4ac3665997f7ce93f78e8fd78c58cb4215ee477c01b76b6f87dff4"
APPROVED_TEMPLATE_V9_RESOURCE = "excel_templates/gb_t_32151_34_2024_appendix_b_v9.xlsx"
APPROVED_TEMPLATE_V9_SHA256 = "0fa20cbc126fadd2807fe02bea9b26b4a8855285e6b90fc171985557a895fae9"
APPROVED_TEMPLATE_V10_RESOURCE = "excel_templates/gb_t_32151_34_2024_appendix_b_v10.xlsx"
APPROVED_TEMPLATE_V10_SHA256 = "705a9326f6910bceaad34ce955975a6150731fa8ed2b50ddfe8c2c5188deb729"
APPROVED_TEMPLATE_RESOURCE = "excel_templates/gb_t_32151_34_2024_appendix_b_v11.xlsx"
APPROVED_TEMPLATE_SHA256 = "b1e843a1eb8980c5a9bf4ee11d5b2369abcb89ebf4ca5b8659d8cccaefd166c4"
APPENDIX_B_SHEETS = tuple(f"B.{number}" for number in range(1, 10))
CURRENT_TEMPLATE_SHEETS = (*APPENDIX_B_SHEETS, "软件补充")


@dataclass(frozen=True, slots=True)
class ExcelTemplate:
    standard_id: str
    standard_version: str
    template_id: str
    template_version: str
    resource_path: str
    sha256: str

    @property
    def approved_master_path(self):
        return files("resources").joinpath(self.resource_path)


APPENDIX_B_TEMPLATE = ExcelTemplate(
    STANDARD_ID, STANDARD_VERSION,
    "GHGTOOL_GBT_32151_34_2024_APPENDIX_B", "2.9.0",
    APPROVED_TEMPLATE_RESOURCE, APPROVED_TEMPLATE_SHA256,
)


class ExcelTemplateService:
    """Small explicit registry, with no template generation or business formulas."""

    def __init__(self, templates: tuple[ExcelTemplate, ...] = (APPENDIX_B_TEMPLATE,)):
        self._templates = {item.standard_id: item for item in templates}
        if len(self._templates) != len(templates):
            raise ValueError("同一标准不能登记多个当前正式模板。")

    @classmethod
    def default(cls) -> "ExcelTemplateService":
        return cls()

    def get(self, standard_id: str = STANDARD_ID) -> ExcelTemplate:
        try:
            return self._templates[standard_id]
        except KeyError as exc:
            raise ValueError("此标准尚未登记经批准的 Excel 模板。") from exc

    def read_bytes(self, standard_id: str = STANDARD_ID) -> bytes:
        template = self.get(standard_id)
        content = template.approved_master_path.read_bytes()
        if sha256(content).hexdigest() != template.sha256:
            raise ValueError("正式 Excel 模板校验失败，请重新安装完整软件。")
        return content

    def copy_to(self, target: str | Path, standard_id: str = STANDARD_ID) -> Path:
        # Read and verify before touching the destination. Never re-save with openpyxl.
        content = self.read_bytes(standard_id)
        destination = Path(target)
        if destination.suffix.lower() != ".xlsx":
            raise ValueError("模板文件必须使用 .xlsx 扩展名。")
        master = self.get(standard_id).approved_master_path
        if isinstance(master, Path) and destination.resolve() == master.resolve():
            raise ValueError("不能覆盖软件内的正式模板资源。")
        destination.write_bytes(content)
        return destination

    def validate_structure(self, workbook, standard_id: str = STANDARD_ID) -> str:
        """Recognize approved headers while allowing inserted business data rows.

        Filled workbooks have their own hash. Formatting changes from Excel/WPS
        are allowed; fixed headers and their column grouping must remain intact.
        """
        if standard_id != STANDARD_ID:
            raise ValueError("此标准尚未实现正式 Excel 输入映射。")
        if tuple(workbook.sheetnames) not in (APPENDIX_B_SHEETS, CURRENT_TEMPLATE_SHEETS):
            raise ValueError("工作簿必须保留原顺序的 B.1 至 B.9；新版另有软件补充页。")
        has_supplement = tuple(workbook.sheetnames) == CURRENT_TEMPLATE_SHEETS
        is_v21 = (
            workbook["B.2"]["M2"].value == "名称类型"
            and workbook["B.2"]["N2"].value == "FCh来源"
            and workbook["B.7"]["I2"].value == "名称类型"
            and workbook["B.7"]["J2"].value == "因子来源"
        )
        # Inserting approved activity rows moves the protected footer. Version
        # identity follows its text, not the original absolute row number.
        notes = tuple(value for (value,) in workbook["B.2"].iter_rows(
            min_row=20, min_col=1, max_col=1, values_only=True) if isinstance(value, str))
        is_v23 = is_v21 and any(note.startswith("燃料输入规则 V5：") for note in notes)
        is_v22 = is_v21 and any(note.startswith(
            "来源栏用于确认参数来源；缺省值须有当前期间适用的 Resolver 绑定。") for note in notes)
        if is_v23 and is_v22:
            raise ValueError("B.2包含冲突的模板版本说明，请使用一个完整的正式模板版本。")
        is_v20 = (
            workbook["B.2"]["M2"].value == "名称类型"
            and workbook["B.7"]["I2"].value == "名称类型"
        )
        if not is_v21 and not is_v20:
            for name in APPENDIX_B_SHEETS:
                sheet = workbook[name]
                if sheet.sheet_state != "visible":
                    raise ValueError(f"{name} 工作表必须保持可见。")
                for coordinate,value in LEGACY_HEADER_CELLS[name].items():
                    if sheet[coordinate].value != value:
                        raise ValueError(f"{name}!{coordinate} 表头与已登记模板版本不一致。")
                end = 3 if name in {"B.2","B.3","B.4","B.5"} else 2
                actual = sorted(str(a) for a in sheet.merged_cells.ranges if a.min_row <= end)
                if actual != LEGACY_HEADER_MERGES[name]:
                    raise ValueError(f"{name} 表头合并区域与已登记模板版本不一致。")
                width = max(sheet[coordinate].column for coordinate in LEGACY_HEADER_CELLS[name])
                if any(c.value is not None for row in sheet.iter_rows(max_row=end,min_col=width+1) for c in row):
                    raise ValueError(f"{name} 存在未登记的新增表头，请使用软件当前模板。")
            return "1.0.0"
        is_v24 = is_v23 and workbook["B.3"]["I3"].value == "质量基准"
        v25_markers = [any(isinstance(row[0].value, str) and row[0].value.startswith("材料输入规则 V7（2.5.0）：")
                           for row in workbook[name].iter_rows(min_col=1, max_col=1)) for name in ("B.3", "B.4", "B.5")]
        if any(v25_markers) and (not all(v25_markers) or not is_v24):
            raise ValueError("材料说明包含冲突模板版本，请使用完整受控模板。")
        v26 = workbook["B.7"]["K2"].value == "含量来源" and workbook["B.7"]["L2"].value == "转化率来源"
        has_v26_note = any(isinstance(row[0].value, str) and row[0].value.startswith("数据提示：脱硫输入规则 V8（2.6.0）")
                           for row in workbook["B.7"].iter_rows(min_col=1, max_col=1))
        partial_v26_headers = any(workbook["B.7"][cell].value is not None for cell in ("K2", "L2", "M2"))
        if v26 != has_v26_note or (partial_v26_headers and not v26) or (v26 and not all(v25_markers)):
            raise ValueError("脱硫来源栏与模板说明版本冲突。")
        v27 = workbook["B.6"]["H2"].value == "排放源编号"
        has_v27_note = any(isinstance(row[0].value,str) and row[0].value.startswith("数据提示：烟气分段输入规则 V9（2.7.0）") for row in workbook["B.6"].iter_rows(min_col=1,max_col=1))
        partial_v27 = any(workbook["B.6"][f"{col}2"].value is not None for col in "HIJKLM")
        if v27 != has_v27_note or (partial_v27 and not v27) or (v27 and not v26):
            raise ValueError("烟气分段表头与模板说明版本冲突。")
        v28 = workbook["B.9"]["H2"].value == "焓值方式"
        has_v28_note = any(isinstance(row[0].value, str) and row[0].value.startswith("数据提示：热力输入规则 V10（2.8.0）")
                           for row in workbook["B.9"].iter_rows(min_col=1,max_col=1))
        partial_v28 = any(workbook["B.9"][f"{col}2"].value is not None for col in "HIJKL")
        if v28 != has_v28_note or (partial_v28 and not v28) or (v28 and not v27):
            raise ValueError("热力来源/边界表头与模板说明版本冲突。")
        if has_supplement and not v28:
            raise ValueError("软件补充页必须与完整新版标准九表配套。")
        if has_supplement:
            template_version = "2.9.0"
            reference_bytes = self.read_bytes(standard_id)
        elif v28:
            template_version = "2.8.0"
            reference_bytes = files("resources").joinpath(APPROVED_TEMPLATE_V10_RESOURCE).read_bytes()
            if sha256(reference_bytes).hexdigest() != APPROVED_TEMPLATE_V10_SHA256:
                raise ValueError("旧版2.8模板校验失败。")
        elif v27:
            template_version = "2.7.0"
            reference_bytes = files("resources").joinpath(APPROVED_TEMPLATE_V9_RESOURCE).read_bytes()
            if sha256(reference_bytes).hexdigest() != APPROVED_TEMPLATE_V9_SHA256:
                raise ValueError("旧版2.7模板校验失败。")
        elif v26:
            template_version = "2.6.0"
            reference_bytes = files("resources").joinpath(APPROVED_TEMPLATE_V8_RESOURCE).read_bytes()
            if sha256(reference_bytes).hexdigest() != APPROVED_TEMPLATE_V8_SHA256:
                raise ValueError("旧版2.6模板校验失败。")
        elif all(v25_markers):
            template_version = "2.5.0"
            reference_bytes = files("resources").joinpath(APPROVED_TEMPLATE_V7_RESOURCE).read_bytes()
            if sha256(reference_bytes).hexdigest() != APPROVED_TEMPLATE_V7_SHA256:
                raise ValueError("旧版2.5模板校验失败。")
        elif is_v24:
            template_version = "2.4.0"
            reference_bytes = files("resources").joinpath(APPROVED_TEMPLATE_V6_RESOURCE).read_bytes()
            if sha256(reference_bytes).hexdigest() != APPROVED_TEMPLATE_V6_SHA256:
                raise ValueError("旧版2.4材料模板校验失败。")
        elif is_v23:
            template_version = "2.3.0"
            reference_bytes = files("resources").joinpath(APPROVED_TEMPLATE_V5_RESOURCE).read_bytes()
            if sha256(reference_bytes).hexdigest() != APPROVED_TEMPLATE_V5_SHA256:
                raise ValueError("旧版2.3模板校验失败。")
        elif is_v22:
            template_version = "2.2.0"
            reference_bytes = files("resources").joinpath(APPROVED_TEMPLATE_V4_RESOURCE).read_bytes()
            if sha256(reference_bytes).hexdigest() != APPROVED_TEMPLATE_V4_SHA256:
                raise ValueError("旧版2.2 Excel模板校验失败，无法安全读取该版本工作簿。")
        elif is_v21:
            template_version = "2.1.0"
            legacy_master = files("resources").joinpath(APPROVED_TEMPLATE_V3_RESOURCE)
            reference_bytes = legacy_master.read_bytes()
            if sha256(reference_bytes).hexdigest() != APPROVED_TEMPLATE_V3_SHA256:
                raise ValueError("旧版2.1 Excel模板校验失败，无法安全读取该版本工作簿。")
        else:
            template_version = "2.0.0"
            legacy_master = files("resources").joinpath(APPROVED_TEMPLATE_V2_RESOURCE)
            reference_bytes = legacy_master.read_bytes()
            if sha256(reference_bytes).hexdigest() != APPROVED_TEMPLATE_V2_SHA256:
                raise ValueError("旧版2.0 Excel模板校验失败，无法安全读取该版本工作簿。")
        reference = load_workbook(BytesIO(reference_bytes), data_only=False)
        try:
            for name in APPENDIX_B_SHEETS:
                sheet, approved = workbook[name], reference[name]
                if sheet.sheet_state != "visible":
                    raise ValueError(f"{name} 工作表必须保持可见。")
                header_end = 3 if name in {"B.2", "B.3", "B.4", "B.5"} else 2
                for row in approved.iter_rows(max_row=header_end):
                    for expected in row:
                        actual = sheet[expected.coordinate]
                        if actual.value != expected.value:
                            raise ValueError(f"{name}!{expected.coordinate} 表头与正式模板不一致。")
                expected_merges = {
                    str(area) for area in approved.merged_cells.ranges
                    if area.min_row <= header_end
                }
                actual_merges = {
                    str(area) for area in sheet.merged_cells.ranges
                    if area.min_row <= header_end
                }
                if actual_merges != expected_merges:
                    raise ValueError(f"{name} 表头合并区域与正式模板不一致。")
            if has_supplement:
                sheet, approved = workbook["软件补充"], reference["软件补充"]
                editable = {"B4", "B5", "B6", "B10", "B11"}
                if sheet.sheet_state != "visible":
                    raise ValueError("软件补充页必须保持可见。")
                for row in approved.iter_rows():
                    for expected in row:
                        if expected.coordinate not in editable and sheet[expected.coordinate].value != expected.value:
                            raise ValueError(f"软件补充!{expected.coordinate} 固定说明与正式模板不一致。")
                if set(map(str,sheet.merged_cells.ranges)) != set(map(str,approved.merged_cells.ranges)):
                    raise ValueError("软件补充页合并区域与正式模板不一致。")
                if sheet.max_row > approved.max_row or sheet.max_column > approved.max_column:
                    raise ValueError("软件补充页存在未登记的表外单元格。")
        finally:
            reference.close()
        return template_version
