"""GB/T 32151.34—2024 Appendix B Excel ingress adapter."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
from types import MappingProxyType
from typing import Mapping
import re
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from packages.application.project_workspaces import AccountingUnitType
from packages.application.uat03_parameter_queries import UAT03ParameterQueries
from packages.core.decimal_policy import DecimalPolicy
from packages.core.models import (
    AccountingPeriod, ActivityDataSource, ElectricityAcquisitionMode,
    ElectricityAttribute, ElectricityProofStatus, ElectricityProofType,
    ParameterType, ValueType,
)
from packages.core.parameter_resolution import (
    ElectricityConsumptionDetail, ParameterResolutionContext,
    NATIONAL_ELECTRICITY_PARAMETER_ID,
    PROVINCIAL_ELECTRICITY_PARAMETER_ID,
    REGIONAL_ELECTRICITY_PARAMETER_ID,
    UserProvidedParameterValue,
)
from packages.standards.carbon_material import (
    STANDARD_ID, STANDARD_VERSION, BakingInput, CalcinationInput,
    CarbonMaterialInput, CarbonReportingData, CarbonateComponent,
    EmissionSourceState, EmissionSourceStatus, ElectricityOutputLine,
    FGDInput, FuelInput, FuelPath, FuelType, FumeIncinerationInput,
    GraphitizationInput, HeatFactorMode, HeatInput, InputValue,
    ParameterSourceKind, ParameterValue, SteamKind,
    SOURCE_BAKING, SOURCE_CALCINATION, SOURCE_EXPORTED_ELECTRICITY,
    SOURCE_EXPORTED_HEAT, SOURCE_FGD, SOURCE_FUME, SOURCE_FUEL,
    SOURCE_GRAPHITIZATION, SOURCE_PURCHASED_ELECTRICITY,
    SOURCE_PURCHASED_HEAT,
)
from packages.standards.carbon_material_normalization import MaterialDataSource, MaterialRole, MaterialInputLine
from packages.excel.ingress import (
    ImportMessage, UnitCalculationPreview, WorkbookCellReader,
    WorkbookFatalError, WorkbookImportPreview, WorkbookProvenance,
    read_worksheet_numeric_lexemes,
)
from packages.excel.templates import ExcelTemplateService
from packages.standards.carbon_material_input_policy import (
    InputOrigin, fuel_input_policy, catalog_defaults_allowed, STEAM_KIND_LABELS,
    STANDARD_DEFAULT, CALCULATED, MEASURED, parameter_source_policy,
    validate_parameter_source_policy,
)


EXCEL_APPENDIX_B = "EXCEL_APPENDIX_B"
INGRESS_POLICY_ID = "GHGTOOL_EXCEL_APPENDIX_B_V2"
_SOURCE_IDS = (
    SOURCE_FUEL, SOURCE_CALCINATION, SOURCE_BAKING, SOURCE_GRAPHITIZATION,
    SOURCE_FUME, SOURCE_FGD, SOURCE_PURCHASED_ELECTRICITY,
    SOURCE_EXPORTED_ELECTRICITY, SOURCE_PURCHASED_HEAT, SOURCE_EXPORTED_HEAT,
)
_PROCESS_META = {
    "B.3": (SOURCE_CALCINATION, CalcinationInput, "car-par-k1", "k1", 0),
    "B.4": (SOURCE_BAKING, BakingInput, "car-par-k2", "k2", 1),
    "B.5": (SOURCE_GRAPHITIZATION, GraphitizationInput, "car-par-k3", "k3", 2),
}
_APPENDIX_TABLE_WIDTHS = {
    "B.2": 10,
    "B.3": 5,
    "B.4": 5,
    "B.5": 5,
    "B.6": 7,
    "B.7": 7,
    "B.8": 5,
    "B.9": 7,
}


@dataclass(frozen=True, slots=True)
class AppendixBImportContext:
    period: AccountingPeriod
    boundary_confirmed: bool
    enterprise_name: str | None = None
    enterprise_id: str | None = None
    region: str | None = None
    fuel_path_overrides: Mapping[str, FuelPath] = field(default_factory=dict)

    def __post_init__(self) -> None:
        if not isinstance(self.period, AccountingPeriod):
            raise ValueError("必须明确选择有效核算期间。")
        if not isinstance(self.boundary_confirmed, bool):
            raise ValueError("必须明确确认核算边界状态。")
        for name in ("enterprise_name", "enterprise_id", "region"):
            value = getattr(self, name)
            if value is not None and (not isinstance(value, str) or not value.strip()):
                raise ValueError(f"{name} 不能是空白文字。")
            if value is not None:
                object.__setattr__(self, name, value.strip())
        overrides = dict(self.fuel_path_overrides)
        for location, path in overrides.items():
            if not isinstance(location, str) or not re.fullmatch(r"B\.2!B\d+", location):
                raise ValueError("燃料单位选择必须按 B.2!B<行号> 精确指定。")
            if not isinstance(path, FuelPath):
                raise ValueError("燃料单位选择必须使用 FuelPath。")
        object.__setattr__(self, "fuel_path_overrides", MappingProxyType(overrides))


def _stable_id(prefix: str, *parts: object) -> str:
    raw = "|".join(str(part) for part in parts).encode("utf-8")
    return f"{prefix}-{sha256(raw).hexdigest()[:20]}"


def _has_value(cell) -> bool:
    return cell.value is not None and cell.value != ""


def _is_appendix_footer_row(worksheet, row: int) -> bool:
    """Recognize only marked footer rows whose merges span the whole table.

    Activity labels are user-editable in several Appendix B tables. Prefix-only
    checks can therefore mistake valid names such as "A Sample Batch" or a
    fuel beginning with "使用提示" for a footer and silently discard activity.
    The approved master marks footer rows with a full-width merge; preserve that
    structure and require one of its known footnote/prompt markers as evidence.
    """
    value = worksheet[f"A{row}"].value
    if not isinstance(value, str):
        return False
    table_width = _APPENDIX_TABLE_WIDTHS.get(worksheet.title)
    if table_width is None:
        return False

    spans = sorted(
        (area.min_col, area.max_col)
        for area in worksheet.merged_cells.ranges
        if area.min_row == row and area.max_row == row
    )
    covered_through = 0
    for first, last in spans:
        if first > covered_through + 1:
            break
        covered_through = max(covered_through, last)
        if covered_through >= table_width:
            break
    if covered_through < table_width:
        return False

    text = value.strip()
    if text.startswith(("使用提示：", "数据提示：", "备注：", "数据来源：")):
        return True
    if text == "排放量（tCO₂）" or text.startswith("排放量自动计算并只读"):
        return True
    return re.match(r"^[ab]\s{2,}(?:若|对于|填写|报告主体|脱硫剂名称)", text) is not None


def _factor_value(factor):
    if factor is None:
        return None
    normalized = getattr(factor, "normalized_value", None)
    return normalized if normalized is not None else getattr(factor, "value", None)


def _domain_factor_unit(factor) -> str:
    unit = getattr(factor, "normalized_unit", None) or factor.unit
    return {
        "tCO₂/MWh": "tCO2/MWh",
        "tCO₂/GJ": "tCO2/GJ",
        "tCO₂/t": "tCO2/t",
    }.get(unit, unit)


def _parameter_value(factor, reason: str, *, unit: str | None = None,
                     source_version: str | None = None) -> ParameterValue:
    kinds = {
        ValueType.STANDARD_DEFAULT: ParameterSourceKind.STANDARD_DEFAULT,
        ValueType.STANDARD_SPECIFIED: ParameterSourceKind.STANDARD_SPECIFIED,
        ValueType.GOVERNMENT_PUBLISHED: ParameterSourceKind.OFFICIAL_PUBLISHED,
        ValueType.MEASURED: ParameterSourceKind.MEASURED,
        ValueType.DERIVED: ParameterSourceKind.CALCULATED,
        ValueType.SYSTEM_CONSTANT: ParameterSourceKind.CALCULATED,
        ValueType.SCIENTIFIC_REFERENCE: ParameterSourceKind.USER_DEFINED,
        ValueType.HISTORICAL: ParameterSourceKind.USER_DEFINED,
    }
    factor_year = getattr(factor, "factor_year", None)
    if source_version is None:
        source_version = getattr(factor, "source_version", None)
        if source_version is None:
            candidate = getattr(factor, "version", None)
            # The Catalog-to-Resolver bridge currently places factor_year in
            # Factor.version. Preserve it as factor_year, not as a source edition.
            if candidate is not None and str(candidate) != str(factor_year):
                source_version = candidate
    parameter_unit = unit or _domain_factor_unit(factor)
    return ParameterValue(
        factor.parameter_id, _factor_value(factor), parameter_unit, kinds[factor.value_type],
        factor.source_id, source_version, factor.source_location, reason,
        factor.factor_id, factor_year,
    )


def _user_parameter(value: Decimal, parameter_id: str, unit: str, location: str,
                    source_kind: ParameterSourceKind = ParameterSourceKind.USER_DEFINED,
                    *, selection_reason: str | None = None) -> ParameterValue:
    return ParameterValue(
        parameter_id=parameter_id, value=value, unit=unit, source_kind=source_kind,
        source_id=None, source_version=None, source_location=location,
        selection_reason=selection_reason or "工作簿提供未核验来源；本次按用户提供值进入现有计算链。",
    )


from packages.standards.carbon_material_business_policy import fuel_source_context, resolve_fgd_field

_SOURCE_LABEL_KINDS = {"缺省值": STANDARD_DEFAULT, "计算值": CALCULATED, "实测值": MEASURED}


def _validate_source_choice(field_id, source_label, *, origin, has_resolver_binding,
                            errors, location, field_label, carbonate_component_count=1, resolved_context=None):
    source_kind = _SOURCE_LABEL_KINDS.get(source_label)
    if source_kind is None:
        return False
    issue = validate_parameter_source_policy(
        field_id, source_kind, origin=origin,
        has_resolver_binding=has_resolver_binding,
        carbonate_component_count=carbonate_component_count,
        resolved_context=resolved_context,
    )
    if issue is None:
        return True
    _, code, message = issue
    errors.append(ImportMessage(code, message, location, field_label))
    return False


class AppendixBWorkbookImporter:
    """Convert one approved 9-sheet workbook into one canonical preview unit."""

    def __init__(self, *, preview_use_case, catalog_service,
                 template_service: ExcelTemplateService | None = None) -> None:
        self.policy = DecimalPolicy()
        self.preview_use_case = preview_use_case
        self.catalog_service = catalog_service
        self.template_service = template_service or ExcelTemplateService.default()
        try:
            self._template_numeric_lexemes = read_worksheet_numeric_lexemes(
                self.template_service.read_bytes(STANDARD_ID)
            )
        except (AttributeError, TypeError, ValueError):
            self._template_numeric_lexemes = {}
        self.resolver = getattr(getattr(preview_use_case, "calculator", None), "parameter_resolver", None)
        self.queries = UAT03ParameterQueries(catalog_service, self.resolver) if self.resolver else None
        try:
            repository = catalog_service.repository
            self.parameters = {p.parameter_id.lower(): p for p in repository.list_parameters()}
            self.sources = {source.source_id: source for source in repository.list_sources()}
        except (AttributeError, TypeError):
            self.parameters = {}
            self.sources = {}

    def import_preview(self, file_path: str | Path, *,
                       context: AppendixBImportContext) -> WorkbookImportPreview:
        if not isinstance(context, AppendixBImportContext):
            raise ValueError("必须提供 AppendixBImportContext，并明确核算期间和边界确认状态。")
        path = Path(file_path)
        try:
            content = path.read_bytes()
            canonical_path = str(path.resolve())
        except OSError as exc:
            raise WorkbookFatalError("无法读取所选 Excel 文件。") from exc
        digest = sha256(content).hexdigest()
        reader = WorkbookCellReader(content, workbook_sha256=digest, canonical_path=canonical_path)
        try:
            workbook = load_workbook(BytesIO(content), data_only=False)
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
            raise WorkbookFatalError("文件不是可读取的 Excel 工作簿，或内部结构已损坏。") from exc

        errors: list[ImportMessage] = []
        warnings: list[ImportMessage] = []
        involved: set[str] = set()
        fuels: list[FuelInput] = []
        processes: tuple[list, list, list] = ([], [], [])
        fumes: list[FumeIncinerationInput] = []
        fgd: list[FGDInput] = []
        purchased_power: list[ElectricityConsumptionDetail] = []
        exported_power: list[ElectricityOutputLine] = []
        purchased_heat: list[HeatInput] = []
        exported_heat: list[HeatInput] = []
        try:
            try:
                template_version = self.template_service.validate_structure(workbook, STANDARD_ID)
            except ValueError as exc:
                raise WorkbookFatalError(str(exc)) from exc
            unit_id = _stable_id("appendix-b-unit", digest)
            enterprise_id = context.enterprise_id or _stable_id("appendix-b-enterprise", digest)

            self._check_trailing_activity(workbook, reader, errors, warnings)
            self._parse_fuels(workbook, context, reader, errors, warnings, fuels, involved, unit_id, template_version=template_version)
            self._parse_processes(workbook, context, reader, errors, warnings, processes, involved, unit_id, template_version=template_version)
            self._parse_fumes(workbook, context, reader, errors, warnings, fumes, involved, unit_id)
            self._parse_fgd(workbook, context, reader, errors, warnings, fgd, involved, unit_id, template_version=template_version)
            self._parse_power(workbook, context, reader, errors, warnings, purchased_power,
                              exported_power, involved, unit_id, enterprise_id)
            self._parse_heat(workbook, context, reader, errors, warnings, purchased_heat,
                             exported_heat, involved, unit_id)

            states = tuple(EmissionSourceState(
                source_id,
                EmissionSourceStatus.INVOLVED if source_id in involved else EmissionSourceStatus.NOT_INVOLVED,
            ) for source_id in _SOURCE_IDS)
            canonical = CarbonMaterialInput(
                input_id=_stable_id("appendix-b-input", digest),
                enterprise_id=enterprise_id,
                enterprise_name=context.enterprise_name,
                period=context.period,
                boundary_confirmed=context.boundary_confirmed,
                boundary_component_ids=(_stable_id("appendix-b-boundary", digest),),
                source_states=states,
                fuel_inputs=tuple(fuels),
                calcinations=tuple(processes[0]),
                bakings=tuple(processes[1]),
                graphitizations=tuple(processes[2]),
                fume_incinerations=tuple(fumes),
                fgd_units=tuple(fgd),
                electricity_details=tuple(purchased_power),
                exported_electricity=tuple(exported_power),
                purchased_heat=tuple(purchased_heat),
                exported_heat=tuple(exported_heat),
            )
            imported_at = datetime.now(timezone.utc)
            calculation = (
                self.preview_use_case.calculate(canonical, calculated_at=imported_at)
                if not errors else None
            )
            if calculation is not None:
                for problem in getattr(calculation, "problems", ()):
                    level = getattr(getattr(problem, "level", None), "value", getattr(problem, "level", None))
                    message = ImportMessage(
                        code=getattr(problem, "code", "EXB01_CALCULATION_ISSUE"),
                        message=getattr(problem, "message", "共享计算器返回了未说明的问题。"),
                    )
                    if str(level).upper() == "ERROR":
                        errors.append(message)
                    else:
                        warnings.append(message)
            template = self.template_service.get(STANDARD_ID)
            unit = UnitCalculationPreview(
                unit_id=unit_id,
                name=context.enterprise_name or path.stem or "附录B工作簿",
                unit_type=AccountingUnitType.WHOLE_SITE,
                input_value=canonical,
                calculation=calculation,
                errors=tuple(errors),
                warnings=tuple(warnings),
            )
            return WorkbookImportPreview(
                provenance=WorkbookProvenance(
                    workbook_sha256=digest,
                    template_id=template.template_id,
                    template_version=template_version,
                    standard_id=template.standard_id,
                    standard_version=template.standard_version,
                    ingress_policy_id=INGRESS_POLICY_ID,
                    imported_at=imported_at,
                    canonical_path=canonical_path,
                    source=EXCEL_APPENDIX_B,
                    # Freeze the verified master at ingress, never infer it at export.
                    template_master_sha256=template.sha256 if template_version == template.template_version else None,
                ),
                units=(unit,),
                warnings=tuple(warnings),
                numeric_evidence=tuple(reader.evidence),
            )
        finally:
            workbook.close()

    def _parameter_value(self, factor, reason: str, *, unit: str | None = None) -> ParameterValue:
        source = self.sources.get(getattr(factor, "source_id", None))
        source_version = getattr(source, "version", None)
        return _parameter_value(factor, reason, unit=unit, source_version=source_version)

    def _default_factor(self, parameter_id: str, period: AccountingPeriod, *,
                        subject_id: str | None = None, direction: str | None = None,
                        region: str | None = None,
                        electricity_attribute: ElectricityAttribute | None = None,
                        electricity_acquisition_mode: ElectricityAcquisitionMode | None = None):
        if self.resolver is None:
            return None
        parameter = self.parameters.get(parameter_id.lower())
        if parameter is None:
            return None
        ctx = ParameterResolutionContext(
            parameter_id=parameter.parameter_id,
            standard_id=STANDARD_ID,
            accounting_period=period,
            subject_id=subject_id or parameter.subject_id,
            region=region,
            parameter_type=parameter.parameter_type,
            electricity_attribute=electricity_attribute,
            electricity_acquisition_mode=electricity_acquisition_mode,
            emission_source_type=direction,
            extra_context=(("energy_direction", direction),) if direction else (),
        )
        try:
            resolved = self.resolver.resolve(ctx)
        except (AttributeError, TypeError, ValueError):
            return None
        return resolved.recommended.factor if not resolved.blocked and resolved.recommended else None

    @staticmethod
    def _check_trailing_activity(workbook, reader, errors, warnings):
        numeric_columns = {
            "B.2":("BDFHIL" if workbook["B.2"]["M2"].value == "名称类型" else "BCEGH"), "B.3":("DE" if workbook["B.3"]["C3"].value == "实际名称" else "CD"),
            "B.4":("DE" if workbook["B.4"]["C3"].value == "实际名称" else "CD"),
            "B.5":("DE" if workbook["B.5"]["C3"].value == "实际名称" else "CD"),
            "B.6":"ABCDEF", "B.7":"BDEFH", "B.8":"CD", "B.9":"CDEFG",
        }
        for name in ("B.3", "B.4", "B.5"):
            if workbook[name]["I3"].value == "质量基准":
                numeric_columns[name] += "K"
        for name,columns in numeric_columns.items():
            sheet=workbook[name]
            footer=next((row for row in range(3,sheet.max_row+1)
                         if _is_appendix_footer_row(sheet,row)),None)
            if footer is None:
                continue
            for row in range(footer+1,sheet.max_row+1):
                if _is_appendix_footer_row(sheet,row):
                    continue
                for column in columns:
                    cell=sheet[f"{column}{row}"]
                    value=cell.value
                    if value is None or value=="":
                        continue
                    numeric = cell.data_type in {"n","f"}
                    if isinstance(value,str):
                        try:
                            Decimal(value.strip())
                            numeric=True
                        except (InvalidOperation,ValueError):
                            pass
                    if not numeric:
                        continue
                    location=f"{name}!{column}{row}"
                    reader.observe_unused(workbook,name,f"{column}{row}",warnings,
                                          field_label="表后活动数据",reason="活动数据位于表后说明区域，未纳入核算。")
                    errors.append(ImportMessage(
                        "EXB01_TRAILING_ACTIVITY","活动数据位于表后说明区域；请移到表内有效填写行，不能静默漏算。",
                        location,"活动数据位置",
                    ))

    def _fuel_option(self, label):
        if self.queries is None:
            return None
        from packages.standards.carbon_material_input_policy import match_fuel_option
        return match_fuel_option(self.queries.fuel_options(), label)

    def _fuel_defaults(self, option, period):
        if self.queries is None or option is None:
            return None, None, None
        try:
            defaults = self.queries.fuel_defaults(option.subject_id, period)
        except (AttributeError, TypeError, ValueError):
            return None, None, None
        return tuple(
            item.factor if item else None for item in (
                defaults.lower_heating_value,
                defaults.carbon_content_per_heat,
                defaults.oxidation_rate,
            )
        )

    @staticmethod
    def _fuel_path(label, row, default_lhv, context, errors):
        location = f"B.2!B{row}"
        override = context.fuel_path_overrides.get(location)
        if override is FuelPath.HEAT:
            errors.append(ImportMessage(
                "EXB01_FUEL_PATH_UNSUPPORTED",
                f"“{label}”的消耗量单位栏只支持吨或万立方米，不支持热量路径。",
                location, f"燃料“{label}”的消耗量单位",
            ))
            return None
        if override is not None:
            return override
        unit = getattr(default_lhv, "unit", None)
        if unit == "GJ/t":
            return FuelPath.MASS
        if unit == "GJ/10⁴Nm³":
            return FuelPath.VOLUME
        errors.append(ImportMessage(
            "EXB01_FUEL_PATH_REQUIRED",
            f"“{label}”第 {row} 行的消耗量单位无法从工作簿唯一确定；请明确选择吨或万立方米。",
            location, f"燃料“{label}”的消耗量单位",
        ))
        return None


    def _parse_fuels(self, workbook, context, reader, errors, warnings, output, involved, unit_id, *, template_version="1.0.0"):
        sheet = workbook["B.2"]
        from packages.standards.carbon_material_input_policy import FUEL_QUANTITIES
        current = template_version in {"2.0.0", "2.1.0", "2.2.0", "2.3.0", "2.4.0", "2.5.0", "2.6.0"}
        explicit_source_schema = template_version in {"2.1.0", "2.2.0", "2.3.0", "2.4.0", "2.5.0", "2.6.0"}
        policy_source_schema = template_version in {"2.2.0", "2.3.0", "2.4.0", "2.5.0", "2.6.0"}
        def column(legacy_column):
            return chr(ord(legacy_column)+1) if current and legacy_column >= "C" else legacy_column

        stop = next(
            (r for r in range(4, sheet.max_row + 1) if _is_appendix_footer_row(sheet, r)),
            sheet.max_row + 1,
        )
        for row in range(4, stop):
            if not any(_has_value(sheet[f"{column(col)}{row}"]) for col in ("BCEGHK" if current else "BCEGH")):
                continue
            label = reader.text(workbook, "B.2", f"{column('A')}{row}", errors, required=True, field_label="燃料品种")
            if not label:
                continue
            name_mode = reader.text(workbook,"B.2",f"{column('L')}{row}",errors,required=current,field_label="燃料名称类型") if current else None
            if current and name_mode not in {"目录","自定义"}:
                errors.append(ImportMessage("EXB01_FUEL_IDENTITY_REQUIRED","请选择目录燃料或自定义燃料。",f"B.2!{column('L')}{row}","名称类型"))
                continue
            option = None if name_mode == "自定义" else self._fuel_option(label)
            if current and name_mode == "目录" and option is None:
                errors.append(ImportMessage("EXB01_FUEL_CATALOG_UNRESOLVED","目录中未找到该燃料；自行填写参数请选择自定义。",f"B.2!{column('L')}{row}","名称类型"))
                continue
            fuel_type = (option.fuel_type or FuelType.OTHER) if option else FuelType.OTHER
            origin = InputOrigin.CATALOG if option else InputOrigin.CUSTOM
            default_lhv, default_carbon, default_oxidation = self._fuel_defaults(option, context.period)
            if current:
                unit_label = reader.text(workbook,"B.2",f"C{row}",errors,required=True,field_label="消耗量单位")
                from packages.standards.carbon_material_input_policy import fuel_path_for_unit
                unit_path = fuel_path_for_unit(unit_label)
                path = FuelPath(unit_path) if unit_path is not None else None
                if unit_label is not None and path is None:
                    errors.append(ImportMessage("EXB01_FUEL_UNIT_INVALID","消耗量单位请选择t或10⁴Nm³。",f"B.2!C{row}","消耗量单位"))
            else:
                path = self._fuel_path(label, row, default_lhv, context, errors)
            activity = reader.number(workbook, "B.2", f"{column('B')}{row}", errors, required=True, field_label=f"燃料“{label}”消耗量")
            if path is None:
                continue
            route = reader.text(workbook, "B.2", f"{column('D')}{row}", errors, required=True, field_label=f"燃料“{label}”碳含量路径")
            if route not in ({"计算值", "实测值", "直接计算值"} if template_version in {"2.3.0", "2.4.0", "2.5.0", "2.6.0"} else {"计算值", "实测值"}):
                if route is not None:
                    errors.append(ImportMessage("EXB01_FUEL_CARBON_ROUTE_INVALID", "碳含量路径只能是“计算值”或“实测值”。", f"B.2!{column('D')}{row}", "碳含量路径"))
                route = None
            lhv_cell = f"B.2!{column('E')}{row}"
            direct_cell = f"B.2!{column('C')}{row}"
            lhv_source = reader.text(workbook, "B.2", f"{column('F')}{row}", errors, field_label="低位发热量来源")
            oxidation_source = reader.text(workbook, "B.2", f"{column('I')}{row}", errors, field_label="碳氧化率来源")
            fch_source_label = (
                reader.text(workbook, "B.2", f"N{row}", errors, field_label="单位热值含碳量来源")
                if explicit_source_schema else None
            )
            lhv_source_valid = True
            fch_source_valid = True
            oxidation_source_valid = True
            if policy_source_schema:
                if route == "计算值":
                    if lhv_source is None:
                        errors.append(ImportMessage("EXB01_PARAMETER_SOURCE_REQUIRED", "请确认低位发热量来源。", f"B.2!{column('F')}{row}", "低位发热量来源"))
                        lhv_source_valid = False
                    elif lhv_source not in _SOURCE_LABEL_KINDS:
                        errors.append(ImportMessage("EXB01_PARAMETER_SOURCE_INVALID", "低位发热量来源选择无效。", f"B.2!{column('F')}{row}", "低位发热量来源"))
                        lhv_source_valid = False
                    else:
                        lhv_source_valid = _validate_source_choice(
                            "fuel.lhv", lhv_source, origin=origin,
                            has_resolver_binding=default_lhv is not None, errors=errors,
                            resolved_context=fuel_source_context("fuel.lhv", resolver=getattr(self.queries, "resolver", None), origin=origin, subject_id=option.subject_id if option else None, accounting_period=context.period),
                            location=f"B.2!{column('F')}{row}", field_label="低位发热量来源",
                        )
                    if fch_source_label is None:
                        errors.append(ImportMessage("EXB01_FCH_SOURCE_REQUIRED", "请确认单位热值含碳量来源。", f"B.2!N{row}", "单位热值含碳量来源"))
                        fch_source_valid = False
                    elif fch_source_label not in _SOURCE_LABEL_KINDS:
                        errors.append(ImportMessage("EXB01_FCH_SOURCE_INVALID", "单位热值含碳量来源选择无效。", f"B.2!N{row}", "单位热值含碳量来源"))
                        fch_source_valid = False
                    else:
                        fch_source_valid = _validate_source_choice(
                            "fuel.fch", fch_source_label, origin=origin,
                            has_resolver_binding=default_carbon is not None, errors=errors,
                            resolved_context=fuel_source_context("fuel.fch", resolver=getattr(self.queries, "resolver", None), origin=origin, subject_id=option.subject_id if option else None, accounting_period=context.period),
                            location=f"B.2!N{row}", field_label="单位热值含碳量来源",
                        )
                if oxidation_source is None:
                    errors.append(ImportMessage("EXB01_PARAMETER_SOURCE_REQUIRED", "请确认碳氧化率来源。", f"B.2!{column('I')}{row}", "碳氧化率来源"))
                    oxidation_source_valid = False
                elif oxidation_source not in _SOURCE_LABEL_KINDS:
                    errors.append(ImportMessage("EXB01_PARAMETER_SOURCE_INVALID", "碳氧化率来源选择无效。", f"B.2!{column('I')}{row}", "碳氧化率来源"))
                    oxidation_source_valid = False
                else:
                    oxidation_source_valid = _validate_source_choice(
                        "fuel.oxidation", oxidation_source, origin=origin,
                        has_resolver_binding=default_oxidation is not None, errors=errors,
                        resolved_context=fuel_source_context("fuel.oxidation", resolver=getattr(self.queries, "resolver", None), origin=origin, subject_id=option.subject_id if option else None, accounting_period=context.period),
                        location=f"B.2!{column('I')}{row}", field_label="碳氧化率来源",
                    )
            if explicit_source_schema and not policy_source_schema and route == "计算值":
                source_location = f"B.2!N{row}"
                if fch_source_label is None:
                    errors.append(ImportMessage(
                        "EXB01_FCH_SOURCE_REQUIRED", "请确认单位热值含碳量来源并选择缺省值、计算值或实测值。",
                        source_location, "单位热值含碳量来源",
                    ))
                elif fch_source_label not in {"缺省值", "计算值", "实测值"}:
                    errors.append(ImportMessage(
                        "EXB01_FCH_SOURCE_INVALID", "来源只能选择“缺省值”“计算值”或“实测值”。",
                        source_location, "单位热值含碳量来源",
                    ))
                if route == "计算值" and lhv_source is None:
                    errors.append(ImportMessage(
                        "EXB01_PARAMETER_SOURCE_REQUIRED", "请确认低位发热量来源并选择缺省值、计算值或实测值。",
                        f"B.2!{column('G')}{row}", "低位发热量来源",
                    ))
            if explicit_source_schema and not policy_source_schema and oxidation_source is None:
                errors.append(ImportMessage(
                    "EXB01_PARAMETER_SOURCE_REQUIRED", "请确认碳氧化率来源并选择缺省值、计算值或实测值。",
                    f"B.2!{column('J')}{row}", "碳氧化率来源",
                ))
            carbon_column = column(
                "K" if (
                    explicit_source_schema and fch_source_label in {"计算值", "实测值"}
                    or not explicit_source_schema and current and origin is InputOrigin.CUSTOM
                ) else "G"
            )
            carbon_cell = f"B.2!{carbon_column}{row}"
            if current and not explicit_source_schema and origin is InputOrigin.CATALOG and _has_value(sheet[f"{column('K')}{row}"]):
                errors.append(ImportMessage("EXB01_CATALOG_FCH_CUSTOM_CELL","目录燃料不能使用自定义FCh填写格。",f"B.2!{column('K')}{row}","自定义单位热值含碳量"))
            if current and not explicit_source_schema and origin is InputOrigin.CUSTOM and _has_value(sheet[f"{column('G')}{row}"]):
                errors.append(ImportMessage("EXB01_CUSTOM_FCH_ADOPTED_CELL","自定义燃料不采用目录FCh，请清空H列并填写L列。",f"B.2!{column('G')}{row}","目录采用值"))

            if route in {"实测值", "直接计算值"}:
                direct = reader.number(workbook, "B.2", f"{column('C')}{row}", errors, required=True, field_label="直接含碳量")
                reader.observe_unused(workbook, "B.2", f"{column('E')}{row}", warnings, field_label="低位发热量", reason="实测直接含碳路径不采用低位发热量；原始数值保留为证据。")
                reader.observe_unused(workbook, "B.2", f"{column('G')}{row}", warnings, field_label="单位热值含碳量", reason="实测直接含碳路径不采用单位热值含碳量；原始数值保留为证据。")
                if current:
                    reader.observe_unused(workbook,"B.2",f"{column('K')}{row}",warnings,field_label="其他来源FCh填写值",reason="直接含碳量路径不采用FCh。")
                if explicit_source_schema:
                    reader.observe_unused(workbook,"B.2",f"N{row}",warnings,field_label="单位热值含碳量来源",reason="直接含碳量路径不采用FCh来源选择。")
                if template_version in {"2.3.0", "2.4.0", "2.5.0", "2.6.0"} and route == "直接计算值":
                    for coordinate, name in ((f"{column('E')}{row}", "低位发热量"), (f"{column('G')}{row}", "FCh目录值"), (f"{column('K')}{row}", "FCh填写值")):
                        if _has_value(sheet[coordinate]):
                            errors.append(ImportMessage("EXB01_DIRECT_PATH_VALUE_CONFLICT", "直接计算含碳量路径不采用热值或FCh，请清空冲突数值。", f"B.2!{coordinate}", name))
                lhv_parameter = None
                direct_source_valid = True
                if policy_source_schema and direct is not None:
                    issue = validate_parameter_source_policy(
                        "fuel.direct_carbon", CALCULATED if route == "直接计算值" else MEASURED, origin=origin,
                    )
                    if issue is not None:
                        errors.append(ImportMessage(issue[1], issue[2], direct_cell, "直接含碳量"))
                        direct_source_valid = False
                carbon_parameter = (
                    _user_parameter(direct, f"enterprise_fuel_{unit_id}_{row}_carbon", FUEL_QUANTITIES[path.value].direct_carbon_unit, direct_cell, ParameterSourceKind.CALCULATED if route == "直接计算值" else ParameterSourceKind.MEASURED)
                    if direct is not None and direct_source_valid else None
                )
                if direct is not None:
                    self._warn_missing_source(warnings, direct_cell, f"燃料“{label}”直接含碳量")
            elif route == "计算值":
                reader.observe_unused(workbook, "B.2", f"{column('C')}{row}", warnings, field_label="直接含碳量", reason="计算路径不采用直接含碳量；原始数值保留为证据。")
                lhv_value = reader.number(workbook, "B.2", f"{column('E')}{row}", errors, field_label="低位发热量")
                carbon_value = reader.number(workbook, "B.2", f"{carbon_column}{row}", errors, field_label="单位热值含碳量")
                lhv_unit = FUEL_QUANTITIES[path.value].lhv_unit if path else "GJ/t"
                lhv_parameter = None if explicit_source_schema and (lhv_source is None or (policy_source_schema and not lhv_source_valid)) else self._parameter_with_source(
                    lhv_value, lhv_source, default_lhv,
                    option.subject_id + "_lhv" if option else f"enterprise_fuel_{unit_id}_{row}_lhv",
                    lhv_unit, lhv_cell, f"B.2!{column('F')}{row}", errors, warnings, "低位发热量",
                    allow_user_defined=not explicit_source_schema,
                )
                source_kind_by_label = {
                    "缺省值": "STANDARD_DEFAULT",
                    "计算值": "CALCULATED",
                    "实测值": "MEASURED",
                }
                policy = fuel_input_policy(
                    origin, direct_carbon=False,
                    carbon_source_kind=source_kind_by_label.get(fch_source_label),
                )
                if explicit_source_schema and not policy_source_schema:
                    adopted_cell = f"B.2!{column('G')}{row}"
                    alternate_cell = f"B.2!{column('K')}{row}"
                    if fch_source_label == "缺省值":
                        if origin is InputOrigin.CUSTOM:
                            errors.append(ImportMessage(
                                "EXB01_CUSTOM_FUEL_DEFAULT_FORBIDDEN",
                                "自定义燃料不适用缺省FCh；请选择计算值或实测值并填写L列。",
                                f"B.2!N{row}", "单位热值含碳量来源",
                            ))
                            carbon_parameter = None
                        elif _has_value(sheet[f"{column('K')}{row}"]):
                            errors.append(ImportMessage(
                                "EXB01_FCH_SOURCE_VALUE_CONFLICT", "选择缺省值时请清空L列，填写的显式FCh不能按缺省值处理。",
                                alternate_cell, "其他来源FCh填写值",
                            ))
                            carbon_parameter = None
                        elif default_carbon is None:
                            errors.append(ImportMessage(
                                "EXB01_CATALOG_FCH_UNRESOLVED", "当前期间Resolver没有可用的目录FCh缺省值。",
                                adopted_cell, "单位热值含碳量",
                            ))
                            carbon_parameter = None
                        elif carbon_value is not None and carbon_value != _factor_value(default_carbon):
                            errors.append(ImportMessage(
                                "EXB01_DEFAULT_PARAMETER_MISMATCH",
                                "选择缺省值时，H列必须留空或与当前期间Resolver核验值完全一致。",
                                adopted_cell, "单位热值含碳量",
                            ))
                            carbon_parameter = None
                        else:
                            carbon_parameter = self._parameter_value(
                                default_carbon, "采用当前期间Resolver核验的附录C.1单位热值含碳量。"
                            )
                    elif fch_source_label in {"计算值", "实测值"}:
                        if _has_value(sheet[f"{column('G')}{row}"]):
                            errors.append(ImportMessage(
                                "EXB01_FCH_SOURCE_VALUE_CONFLICT", "选择计算值或实测值时请清空H列，并在L列填写该来源的FCh。",
                                adopted_cell, "目录采用值",
                            ))
                        if carbon_value is None:
                            errors.append(ImportMessage(
                                "EXB01_FCH_EXPLICIT_VALUE_REQUIRED", "已选择FCh来源，必须在L列填写参数值。",
                                alternate_cell, "其他来源FCh填写值",
                            ))
                            carbon_parameter = None
                        else:
                            source_kind = (
                                ParameterSourceKind.CALCULATED
                                if fch_source_label == "计算值"
                                else ParameterSourceKind.MEASURED
                            )
                            carbon_parameter = _user_parameter(
                                carbon_value, f"enterprise_fuel_{unit_id}_{row}_carbon", "tC/GJ",
                                alternate_cell, source_kind,
                                selection_reason=f"按B.2明确选择的{fch_source_label}来源采用单位热值含碳量。",
                            )
                            self._warn_missing_source(warnings, alternate_cell, "单位热值含碳量")
                    else:
                        carbon_parameter = None
                elif policy_source_schema:
                    adopted_cell = f"B.2!{column('G')}{row}"
                    alternate_cell = f"B.2!{column('K')}{row}"
                    if not fch_source_valid:
                        carbon_parameter = None
                    elif fch_source_label == "缺省值":
                        if _has_value(sheet[f"{column('K')}{row}"]):
                            errors.append(ImportMessage("EXB01_FCH_SOURCE_VALUE_CONFLICT", "选择缺省值时请清空L列填写值。", alternate_cell, "单位热值含碳量"))
                            carbon_parameter = None
                        elif default_carbon is None:
                            errors.append(ImportMessage("EXB01_CATALOG_FCH_UNRESOLVED", "当前期间 Resolver 没有可用的单位热值含碳量缺省值。", adopted_cell, "单位热值含碳量"))
                            carbon_parameter = None
                        elif carbon_value is not None and carbon_value != _factor_value(default_carbon):
                            errors.append(ImportMessage("EXB01_DEFAULT_PARAMETER_MISMATCH", "选择缺省值时，H列必须留空或与当前期间 Resolver 核验值完全一致。", adopted_cell, "单位热值含碳量"))
                            carbon_parameter = None
                        else:
                            carbon_parameter = self._parameter_value(default_carbon, "采用当前期间 Resolver 核验的单位热值含碳量。")
                    elif fch_source_label in {"计算值", "实测值"}:
                        if _has_value(sheet[f"{column('G')}{row}"]):
                            errors.append(ImportMessage("EXB01_FCH_SOURCE_VALUE_CONFLICT", "选择计算值或实测值时请清空H列，并在L列填写FCh。", adopted_cell, "单位热值含碳量"))
                        if carbon_value is None:
                            errors.append(ImportMessage("EXB01_FCH_EXPLICIT_VALUE_REQUIRED", "已选择FCh来源，必须在L列填写参数值。", alternate_cell, "单位热值含碳量"))
                            carbon_parameter = None
                        else:
                            source_kind = ParameterSourceKind.CALCULATED if fch_source_label == "计算值" else ParameterSourceKind.MEASURED
                            carbon_parameter = _user_parameter(
                                carbon_value, f"enterprise_fuel_{unit_id}_{row}_carbon", "tC/GJ", alternate_cell, source_kind,
                                selection_reason=f"按B.2明确选择的{fch_source_label}来源采用单位热值含碳量。",
                            )
                            self._warn_missing_source(warnings, alternate_cell, "单位热值含碳量")
                    else:
                        carbon_parameter = None
                elif policy.fch_editable:
                    carbon_parameter = _user_parameter(
                        carbon_value, f"enterprise_fuel_{unit_id}_{row}_carbon", "tC/GJ", carbon_cell,
                    ) if carbon_value is not None else None
                    if carbon_value is not None:
                        self._warn_missing_source(warnings, carbon_cell, "自定义燃料单位热值含碳量")
                elif carbon_value is not None and default_carbon is not None and carbon_value != _factor_value(default_carbon):
                    errors.append(ImportMessage(
                        "EXB01_CATALOG_FCH_READONLY",
                        "目录燃料单位热值含碳量只采用当前期间核验缺省值；2.0及更早模板的显式数值仅可作为旧证据读取。",
                        carbon_cell, "单位热值含碳量",
                    ))
                    carbon_parameter = None
                else:
                    carbon_parameter = self._parameter_value(
                        default_carbon, "采用当前期间Resolver核验的附录C.1单位热值含碳量。"
                    ) if default_carbon else None
            else:
                continue

            if path is None or activity is None:
                continue
            oxidation_value = reader.number(workbook, "B.2", f"{column('H')}{row}", errors, field_label="碳氧化率")
            oxidation_parameter = None if explicit_source_schema and (oxidation_source is None or (policy_source_schema and not oxidation_source_valid)) else self._parameter_with_source(
                oxidation_value, oxidation_source, default_oxidation,
                option.subject_id + "_oxidation_rate" if option else f"enterprise_fuel_{unit_id}_{row}_oxidation",
                "ratio", f"B.2!{column('H')}{row}", f"B.2!{column('I')}{row}", errors, warnings, "碳氧化率",
                percent=True, allow_user_defined=not explicit_source_schema,
            )
            if route == "计算值" and lhv_parameter is not None:
                expected_unit = {FuelPath.MASS: "GJ/t", FuelPath.VOLUME: "GJ/10⁴Nm³", FuelPath.HEAT: "GJ/GJ"}[path]
                if lhv_parameter.unit != expected_unit:
                    errors.append(ImportMessage("EXB01_FUEL_PATH_PARAMETER_UNIT_MISMATCH", "低位发热量目录单位与消耗量单位不一致；请更正单位，或选择实测参数路径。", f"B.2!{column('B')}{row}", "消耗量单位"))
            if route == "计算值" and lhv_parameter is None and (not policy_source_schema or lhv_source_valid):
                errors.append(ImportMessage("EXB01_FUEL_LHV_REQUIRED", "计算路径缺少低位发热量；请填写实测值或使用可核验的目录缺省值。", lhv_cell, "低位发热量"))
            if carbon_parameter is None:
                errors.append(ImportMessage("EXB01_FUEL_CARBON_REQUIRED", "当前路径缺少所需含碳量。", carbon_cell if route == "计算值" else direct_cell, "含碳量"))
            if oxidation_parameter is None and (not policy_source_schema or oxidation_source_valid):
                errors.append(ImportMessage("EXB01_FUEL_OXIDATION_REQUIRED", "缺少碳氧化率；请填写数值或确认当前期间的目录缺省值。", f"B.2!{column('H')}{row}", "碳氧化率"))
            output.append(FuelInput(
                fuel_id=_stable_id("fuel", unit_id, row),
                path=path,
                activity=InputValue(activity, FUEL_QUANTITIES[path.value].activity_unit, source_reference=f"B.2!{column('B')}{row}"),
                carbon_content=carbon_parameter,
                oxidation_rate=oxidation_parameter,
                lower_heating_value=lhv_parameter,
                fuel_type=fuel_type,
                fuel_label=label if fuel_type is FuelType.OTHER else None,
                origin=origin,
                catalog_subject_id=option.subject_id if option else None,
            ))
            involved.add(SOURCE_FUEL)

    def _parameter_with_source(self, value, source_label, default_factor, parameter_id, unit,
                               value_location, source_location, errors, warnings, label, percent=False,
                               allow_user_defined=True):
        allowed_sources = {"缺省值", "计算值", "实测值"}
        if allow_user_defined:
            allowed_sources.add("填写值")
        if source_label is not None and source_label not in allowed_sources:
            errors.append(ImportMessage("EXB01_PARAMETER_SOURCE_INVALID", "来源只能选择“缺省值”“计算值”或“实测值”。", source_location, label))
            return None
        if source_label == "缺省值":
            if value is None:
                return self._parameter_value(default_factor, "采用当前期间 Resolver 核验的标准缺省参数。") if default_factor else None
            normalized = self.policy.divide(value, Decimal("100")) if percent else value
            if default_factor is None or normalized != _factor_value(default_factor):
                errors.append(ImportMessage(
                    "EXB01_DEFAULT_PARAMETER_MISMATCH",
                    "选择“缺省值”时请清空数值，或确保其与当前期间核验值完全一致；其他数值请改选“实测值”。",
                    value_location, label,
                ))
                return self._parameter_value(default_factor, "采用当前期间 Resolver 核验的标准缺省参数。") if default_factor else None
            return self._parameter_value(default_factor, "数值与当前期间 Resolver 核验值完全一致。")
        if source_label in {"计算值", "实测值", "填写值"}:
            if value is None:
                errors.append(ImportMessage("EXB01_EXPLICIT_PARAMETER_REQUIRED", f"选择“{source_label}”后必须填写参数。", value_location, label))
                return None
            normalized = self.policy.divide(value, Decimal("100")) if percent else value
            self._warn_missing_source(warnings, value_location, label)
            source_kind = {
                "计算值": ParameterSourceKind.CALCULATED,
                "实测值": ParameterSourceKind.MEASURED,
                "填写值": ParameterSourceKind.USER_DEFINED,
            }[source_label]
            return _user_parameter(normalized, parameter_id, unit, value_location, source_kind)
        if value is None:
            warnings.append(ImportMessage("EXB01_PARAMETER_SOURCE_UNSPECIFIED", "来源选择为空；此提醒不阻断计算。", source_location, label))
            return self._parameter_value(default_factor, "来源栏为空，采用当前期间 Resolver 核验的标准缺省参数。") if default_factor else None
        normalized = self.policy.divide(value, Decimal("100")) if percent else value
        if default_factor is not None and normalized == _factor_value(default_factor):
            warnings.append(ImportMessage("EXB01_PARAMETER_SOURCE_UNSPECIFIED", "来源选择为空；数值与已核验缺省值一致，已绑定该目录参数。", source_location, label))
            return self._parameter_value(default_factor, "数值与当前期间 Resolver 核验值完全一致。")
        self._warn_missing_source(warnings, value_location, label)
        return _user_parameter(normalized, parameter_id, unit, value_location)

    def _factor_cell(self, value, default_factor, parameter_id, unit, location, warnings, label):
        if value is None:
            return self._parameter_value(default_factor, "采用当前期间 Resolver 核验的标准缺省参数。") if default_factor else None
        if default_factor is not None and value == _factor_value(default_factor):
            return self._parameter_value(default_factor, "数值与当前期间 Resolver 核验值完全一致。")
        warnings.append(ImportMessage("EXB01_SOURCE_REFERENCE_MISSING", "此列没有来源字段；数值作为用户提供参数，来源待补充。", location, label))
        return _user_parameter(value, parameter_id, unit, location)

    @staticmethod
    def _warn_missing_source(warnings, location, label):
        warnings.append(ImportMessage("EXB01_SOURCE_REFERENCE_MISSING", "工作簿未提供外部来源说明；此提醒不阻断预览或计算。", location, label))

    @staticmethod
    def _warn_power_applicability(warnings, location):
        warnings.append(ImportMessage(
            "EXB01_ELECTRICITY_APPLICABILITY_REVIEW",
            "请核对所用电力因子的适用范围与年份；来源说明可补充，此提醒不阻断核算。",
            location, "电力排放因子",
        ))

    def _parse_power(self, workbook, context, reader, errors, warnings, purchased, exported, involved, unit_id, enterprise_id):
        ws = workbook["B.8"]
        stop = next((r for r in range(3, ws.max_row + 1)
                     if _is_appendix_footer_row(ws, r)), ws.max_row + 1)
        for row in range(3, stop):
            if not any(_has_value(ws[f"{col}{row}"]) for col in "BCD"):
                continue
            direction = reader.text(workbook, "B.8", f"A{row}", errors, required=True, field_label="电力方向")
            label = reader.text(
                workbook, "B.8", f"B{row}", errors,
                required=direction == "购入", field_label="电力属性",
            )
            valid_direction = direction in {"购入", "输出"}
            valid_label = (
                label in {"电网电力", "非化石电力"}
                if direction == "购入"
                else label is None or label in {"电网电力", "非化石电力"}
            )
            if not valid_direction and direction is not None:
                errors.append(ImportMessage("EXB01_ELECTRICITY_DIRECTION_INVALID", "电力方向只能是“购入”或“输出”。", f"B.8!A{row}", "电力方向"))
            if valid_direction and not valid_label and label is not None:
                errors.append(ImportMessage("EXB01_ELECTRICITY_ATTRIBUTE_INVALID", "电力属性只能选择“电网电力”或“非化石电力”；输出行可留空。", f"B.8!B{row}", "电力属性"))
            if direction == "输出" and label is not None and valid_label:
                warnings.append(ImportMessage(
                    "EXB01_OUTPUT_ATTRIBUTE_NOT_APPLICABLE",
                    f"输出电力不适用购入电力属性；B列原文“{label}”仅作为填报证据保留，不用于选择输出因子。",
                    f"B.8!B{row}", "电力属性",
                ))
            if not valid_direction or not valid_label:
                reader.observe_unused(workbook, "B.8", f"C{row}", warnings, field_label="电量")
                reader.observe_unused(workbook, "B.8", f"D{row}", warnings, field_label="电力排放因子")
                continue
            amount = reader.number(workbook, "B.8", f"C{row}", errors, required=True, field_label="电量")
            if amount is None:
                reader.observe_unused(workbook, "B.8", f"D{row}", warnings, field_label="电力排放因子")
                continue
            factor_value = reader.number(workbook, "B.8", f"D{row}", errors, field_label="电力排放因子")

            nonfossil_purchase = direction == "购入" and label == "非化石电力"
            region = (
                None if context.region == "全国" else context.region
            ) if direction == "输出" or label == "电网电力" else None
            direction_token = "purchased" if direction == "购入" else "exported"
            geography_level = self._electricity_geography_level(region, context.period)
            parameter_id = "electricity_emission_factor_nonfossil" if nonfossil_purchase else {
                "全国": NATIONAL_ELECTRICITY_PARAMETER_ID,
                "区域": REGIONAL_ELECTRICITY_PARAMETER_ID,
                "省级": PROVINCIAL_ELECTRICITY_PARAMETER_ID,
            }[geography_level]
            if not nonfossil_purchase and self.queries is not None:
                candidates = self.queries.electricity_options(region,context.period)
                default = candidates[0].factor if candidates else None
            else:
                default = self._default_factor(
                    parameter_id,
                    context.period,
                    region=region,
                    direction=direction_token,
                    electricity_attribute=(
                        ElectricityAttribute.NONFOSSIL if nonfossil_purchase
                        else ElectricityAttribute.ORDINARY
                    ),
                    electricity_acquisition_mode=ElectricityAcquisitionMode.PURCHASED,
                )
            resolver_unit = self._parameter_unit(parameter_id, "tCO₂/MWh")

            if nonfossil_purchase:
                if default is None:
                    errors.append(ImportMessage(
                        "EXB01_NONFOSSIL_FACTOR_UNRESOLVED",
                        "当前 Resolver 没有可用的非化石电力零因子；请核验目录，不回退到普通电网因子。",
                        f"B.8!D{row}", "电力排放因子",
                    ))
                elif factor_value is not None and factor_value != _factor_value(default):
                    errors.append(ImportMessage(
                        "EXB01_NONFOSSIL_FACTOR_MISMATCH",
                        "非化石电力排放因子只能采用当前 Resolver 核验值；此单元格数值与核验值不一致。",
                        f"B.8!D{row}", "电力排放因子",
                    ))
                purchased.append(ElectricityConsumptionDetail(
                    detail_id=_stable_id("power", unit_id, row), enterprise_id=enterprise_id,
                    standard_id=STANDARD_ID, accounting_period=context.period,
                    electricity_amount=amount, electricity_unit="MWh",
                    acquisition_mode=ElectricityAcquisitionMode.PURCHASED,
                    attribute=ElectricityAttribute.NONFOSSIL, region=None,
                ))
                involved.add(SOURCE_PURCHASED_ELECTRICITY)
                continue

            if direction == "购入":
                matched_options = self._electricity_catalog_matches(
                    context.region, context.period, factor_value
                ) if factor_value is not None else ()
                selected_option = matched_options[0] if len(matched_options) == 1 else None
                if len(matched_options) > 1:
                    choices = "、".join(dict.fromkeys(
                        f"{item.table_number}（{item.source_document_name or item.factor.source_id}）"
                        for item in matched_options
                    ))
                    warnings.append(ImportMessage(
                        "EXB01_ELECTRICITY_FACTOR_BASIS_AMBIGUOUS",
                        f"该数值与多个已核验电力因子口径一致：{choices}；未自动绑定目录因子，请核对来源。此值暂按用户提供参数处理。",
                        f"B.8!D{row}", "电力排放因子",
                    ))
                selected_id = (
                    selected_option.factor.factor_id if selected_option is not None
                    else default.factor_id if factor_value is None and default is not None
                    else None
                )
                if selected_option is not None:
                    region = selected_option.region
                    parameter_id = selected_option.factor.parameter_id
                    basis = f"{selected_option.table_number}（{selected_option.label}）"
                    source_name = selected_option.source_document_name or selected_option.factor.source_id
                    source_location = selected_option.factor.source_location or "来源位置未记录"
                    selection_reason = (
                        f"B.8 数值唯一匹配{basis}；来源{source_name}，{source_location}。"
                        "用户以该数值明确选择此目录因子。"
                    )
                elif selected_id is not None:
                    selection_reason = "采用当前 Resolver 推荐的核验电力因子。"
                else:
                    selection_reason = None
                override = None
                if factor_value is not None and selected_id is None:
                    override = UserProvidedParameterValue(
                        factor_value, resolver_unit, source_reference=None,
                        selection_reason=(
                            "附录B数值匹配多个目录因子口径，未自动绑定；请核对来源，暂作为用户提供参数。"
                            if len(matched_options) > 1
                            else "附录B提供未核验来源的电力排放因子。"
                        ),
                    )
                    if len(matched_options) <= 1:
                        self._warn_power_applicability(warnings, f"B.8!D{row}")
                elif factor_value is None and default is None:
                    errors.append(ImportMessage("EXB01_ELECTRICITY_FACTOR_UNRESOLVED", "当前期间和地区没有可用的核验电力因子；请补充地区或更正核算期间。", f"B.8!D{row}", "电力排放因子"))
                purchased.append(ElectricityConsumptionDetail(
                    detail_id=_stable_id("power", unit_id, row), enterprise_id=enterprise_id,
                    standard_id=STANDARD_ID, accounting_period=context.period,
                    electricity_amount=amount, electricity_unit="MWh",
                    acquisition_mode=ElectricityAcquisitionMode.PURCHASED,
                    attribute=ElectricityAttribute.ORDINARY, region=region,
                    selected_factor_id=selected_id,
                    factor_selection_reason=selection_reason,
                    factor_override=override,
                ))
                involved.add(SOURCE_PURCHASED_ELECTRICITY)
                continue

            output_attribute_note = (
                f"；B列原文“{label}”按 PR37 作为不适用购入属性的填报证据保留"
                if label is not None else ""
            )
            output_matches = self._electricity_catalog_matches(
                context.region, context.period, factor_value
            ) if factor_value is not None else ()
            if len(output_matches) == 1:
                selected_option = output_matches[0]
                region = selected_option.region
                parameter_id = selected_option.factor.parameter_id
                basis = selected_option.label
                source_name = selected_option.source_document_name or selected_option.factor.source_id
                source_location = selected_option.factor.source_location or "来源位置未记录"
                parameter = self._parameter_value(
                    selected_option.factor,
                    f"B.8 数值唯一匹配{basis}；来源{source_name}，{source_location}。"
                    f"用户以该数值明确选择此目录因子{output_attribute_note}。",
                    unit="tCO2/MWh",
                )
            elif len(output_matches) > 1:
                choices = "、".join(dict.fromkeys(
                    f"{item.table_number}（{item.source_document_name or item.factor.source_id}）"
                    for item in output_matches
                ))
                warnings.append(ImportMessage(
                    "EXB01_ELECTRICITY_FACTOR_BASIS_AMBIGUOUS",
                    f"该数值与多个已核验电力因子口径一致：{choices}；未自动绑定目录因子，请核对来源。此值暂按用户提供参数处理。",
                    f"B.8!D{row}", "输出电力因子",
                ))
                parameter = _user_parameter(
                    factor_value, parameter_id, "tCO2/MWh", f"B.8!D{row}",
                    selection_reason=(
                        "附录B数值匹配多个目录因子口径，未自动绑定；请核对来源，暂作为用户提供参数。"
                        + output_attribute_note
                    ),
                )
            elif factor_value is None:
                parameter = self._parameter_value(
                    default, f"采用当前 Resolver 核验的输出电力适用因子{output_attribute_note}。"
                ) if default is not None else None
                if parameter is None:
                    errors.append(ImportMessage("EXB01_ELECTRICITY_FACTOR_UNRESOLVED", "当前期间和地区没有可用的核验输出电力因子；请填写可追溯的因子值。", f"B.8!D{row}", "电力排放因子"))
            elif default is not None and factor_value == _factor_value(default):
                parameter = self._parameter_value(
                    default, f"数值与当前 Resolver 核验的输出电力因子完全一致{output_attribute_note}。"
                )
            else:
                self._warn_power_applicability(warnings, f"B.8!D{row}")
                parameter = _user_parameter(
                    factor_value, parameter_id, "tCO2/MWh", f"B.8!D{row}",
                    selection_reason=(
                        "附录B提供未核验来源的输出电力排放因子" + output_attribute_note + "。"
                    ),
                )
            exported.append(ElectricityOutputLine(
                line_id=_stable_id("power-export", unit_id, row),
                amount=InputValue(amount, "MWh", source_reference=f"B.8!C{row}"),
                factor=parameter, unit="MWh", region=region,
            ))
            involved.add(SOURCE_EXPORTED_ELECTRICITY)

    def _electricity_geography_level(self, region: str | None, period: AccountingPeriod) -> str:
        if region in (None, "", "全国") or self.queries is None:
            return "全国" if region in (None, "", "全国") else "省级"
        choices = self.queries.electricity_regions_by_level(period)
        if region in choices.get("区域", ()):
            return "区域"
        if region in choices.get("省级", ()):
            return "省级"
        return "省级"

    def _electricity_catalog_matches(self, region: str | None, period: AccountingPeriod, value: Decimal):
        """Match an entered B.8 factor to one unique, Resolver-approved source basis."""
        if self.queries is None:
            return ()
        options = list(self.queries.electricity_catalog_options("全国", None, period))
        choices = self.queries.electricity_regions_by_level(period)
        if region in choices.get("区域", ()):
            options.extend(self.queries.electricity_catalog_options("区域", region, period))
        elif region in choices.get("省级", ()):
            options.extend(self.queries.electricity_catalog_options("省级", region, period))
        matches = tuple(item for item in options if value == _factor_value(item.factor))
        return matches

    def _parameter_unit(self, parameter_id, fallback):
        # Resolver overrides are validated in the Catalog's declared canonical unit.
        parameter = self.parameters.get(parameter_id.lower())
        return str(getattr(parameter, "canonical_unit", getattr(parameter, "default_unit", fallback))) if parameter else fallback

    def _is_template_numeric_seed(self, reader, sheet: str, cell_ref: str) -> bool:
        expected = self._template_numeric_lexemes.get((sheet, cell_ref))
        actual = reader.numeric_lexemes.get((sheet, cell_ref))
        return expected is not None and actual == expected

    def _parse_heat(self, workbook, context, reader, errors, warnings, purchased, exported, involved, unit_id):
        ws = workbook["B.9"]
        stop = next((r for r in range(3, ws.max_row + 1)
                     if _is_appendix_footer_row(ws, r)), ws.max_row + 1)
        for row in range(3, stop):
            if _has_value(ws[f"F{row}"]):
                reader.observe_unused(workbook, "B.9", f"F{row}", warnings, field_label="蒸汽焓值", reason="焓值为模板只读/自动路径；工作簿内容不作为正式焓值。")
            factor_is_template_seed = self._is_template_numeric_seed(reader, "B.9", f"G{row}")
            active = any(_has_value(ws[f"{col}{row}"]) for col in "BCDE") or (
                _has_value(ws[f"G{row}"]) and not factor_is_template_seed
            )
            if not active:
                continue

            direction = reader.text(workbook, "B.9", f"A{row}", errors, required=True, field_label="热力方向")
            kind_label = reader.text(workbook, "B.9", f"B{row}", errors, required=True, field_label="热力类别")
            if direction not in {"购入", "输出"}:
                if direction is not None:
                    errors.append(ImportMessage("EXB01_HEAT_DIRECTION_INVALID", "热力方向只能是“购入”或“输出”。", f"B.9!A{row}", "热力方向"))
                for col, label in (("C", "动力总量"), ("D", "蒸汽压力"), ("E", "蒸汽温度"), ("G", "热力排放因子")):
                    reader.observe_unused(workbook, "B.9", f"{col}{row}", warnings, field_label=label)
                continue
            if kind_label not in STEAM_KIND_LABELS:
                if kind_label is not None:
                    errors.append(ImportMessage("EXB01_HEAT_KIND_UNSUPPORTED", "当前正式模型只支持饱和蒸汽或过热蒸汽；热水及其他热力需要另行核验口径。", f"B.9!B{row}", "热力类别"))
                for col, label in (("C", "动力总量"), ("D", "蒸汽压力"), ("E", "蒸汽温度"), ("G", "热力排放因子")):
                    reader.observe_unused(workbook, "B.9", f"{col}{row}", warnings, field_label=label)
                continue

            amount = reader.number(workbook, "B.9", f"C{row}", errors, required=True, field_label="动力总量")
            pressure = reader.number(workbook, "B.9", f"D{row}", errors, required=True, field_label="蒸汽压力")
            if kind_label == "饱和蒸汽":
                reader.observe_unused(workbook, "B.9", f"E{row}", warnings, field_label="蒸汽温度", reason="饱和蒸汽自动焓值路径不采用温度；原始内容仅保留为证据。")
                temperature = None
                steam_kind = SteamKind.SATURATED
            else:
                temperature = reader.number(workbook, "B.9", f"E{row}", errors, required=True, field_label="蒸汽温度")
                steam_kind = SteamKind.SUPERHEATED
            if factor_is_template_seed:
                reader.observe_unused(workbook, "B.9", f"G{row}", warnings, field_label="热力排放因子", reason="模板预置展示值不是来源凭证；正式缺省因子由 Resolver 按核算期间和方向核验。")
                factor_value = None
            else:
                factor_value = reader.number(workbook, "B.9", f"G{row}", errors, field_label="热力排放因子")
            if amount is None:
                continue

            direction_token = "purchased_heat" if direction == "购入" else "exported_heat"
            default = self._default_factor("heat_emission_factor_default", context.period, direction=direction_token)
            if factor_value is not None and default is not None and factor_value == _factor_value(default):
                # Keep the line on the Calculator's Resolver path so it can bind
                # the factor for this exact heat direction.
                factor = None
                factor_mode = HeatFactorMode.STANDARD_DEFAULT
            elif factor_value is not None:
                self._warn_missing_source(warnings, f"B.9!G{row}", "热力排放因子")
                factor = _user_parameter(factor_value, "heat_emission_factor_measured", "tCO2/GJ", f"B.9!G{row}")
                factor_mode = HeatFactorMode.MEASURED
            else:
                factor = None
                factor_mode = HeatFactorMode.STANDARD_DEFAULT

            heat = HeatInput(
                line_id=_stable_id("heat", unit_id, row),
                amount=InputValue(amount, "t", source_reference=f"B.9!C{row}"),
                enthalpy=None, factor=factor, unit="t", steam_kind=steam_kind,
                pressure_mpa=pressure, temperature_c=temperature, manual_enthalpy=False,
                factor_mode=factor_mode,
            )
            if direction == "购入":
                purchased.append(heat)
                involved.add(SOURCE_PURCHASED_HEAT)
            else:
                exported.append(heat)
                involved.add(SOURCE_EXPORTED_HEAT)

    def _parse_processes(self, workbook, context, reader, errors, warnings, outputs, involved, unit_id, *, template_version="1.0.0"):
        if template_version in {"2.0.0", "2.1.0", "2.2.0", "2.3.0", "2.4.0", "2.5.0", "2.6.0"}:
            from .appendix_b_processes import parse_process_sheets
            return parse_process_sheets(
                self, workbook, context, reader, errors, warnings, outputs,
                involved, unit_id, template_version=template_version,
            )
        return self._parse_legacy_processes(workbook,context,reader,errors,warnings,outputs,involved,unit_id)

    def _parse_legacy_processes(self, workbook, context, reader, errors, warnings, outputs, involved, unit_id):
        for sheet_name, meta in _PROCESS_META.items():
            source_id, model_type, default_parameter_id, parameter_name, index = meta
            ws = workbook[sheet_name]
            stop = next((r for r in range(4, ws.max_row + 1)
                         if _is_appendix_footer_row(ws, r)), ws.max_row + 1)
            active_by_group = {}
            current_group = None
            group_positions = {}
            for row in range(4, stop):
                a = ws[f"A{row}"].value
                if isinstance(a, str) and a.strip():
                    current_group = " ".join(a.split())
                if current_group is None:
                    if any(_has_value(ws[f"{col}{row}"]) for col in "CD"):
                        errors.append(ImportMessage(
                            "EXB01_PROCESS_GROUP_REQUIRED",
                            "该活动行缺少可识别的附录B过程分组；请恢复 A 列组标签。",
                            f"{sheet_name}!A{row}", "过程分组",
                        ))
                    continue
                group_positions[current_group] = group_positions.get(current_group, 0) + 1
                if any(_has_value(ws[f"{col}{row}"]) for col in "CD"):
                    name = reader.text(workbook, sheet_name, f"B{row}", errors, field_label="物料名称")
                    active_by_group.setdefault(current_group, []).append((row, group_positions[current_group], name))
            material_records = {}
            for group, rows in active_by_group.items():
                for row, position, name in rows:
                    role_kind = self._dynamic_process_role(sheet_name, group, position)
                    if role_kind is None:
                        errors.append(ImportMessage("EXB01_PROCESS_ROW_UNMAPPED", "该活动行无法依据附录B的过程分组和固定角色槽唯一映射；请恢复明确的过程分组结构。", f"{sheet_name}!B{row}", "物料名称"))
                        continue
                    role, component = role_kind
                    if not name:
                        errors.append(ImportMessage("EXB01_PROCESS_MATERIAL_NAME_REQUIRED", "活动行缺少物料名称。", f"{sheet_name}!B{row}", "物料名称"))
                        continue
                    mass = reader.number(workbook, sheet_name, f"C{row}", errors, required=_has_value(ws[f"D{row}"]), field_label="活动数据")
                    percentage = reader.number(workbook, sheet_name, f"D{row}", errors, field_label="碳/挥发分含量")
                    source = reader.text(workbook, sheet_name, f"E{row}", errors, field_label="碳/挥发分来源")
                    if source not in (None, "实测值"):
                        errors.append(ImportMessage("EXB01_PROCESS_SOURCE_UNSUPPORTED", "该表的碳/挥发分来源仅支持“实测值”。", f"{sheet_name}!E{row}", "碳/挥发分来源"))
                    key = (role, name)
                    record = material_records.setdefault(key, {"masses": [], "fixed": [], "volatile": []})
                    if mass is not None:
                        record["masses"].append((row, mass))
                    if percentage is not None:
                        record[component].append((row, percentage))
            material_rows = []
            for ordinal, ((role, name), values) in enumerate(material_records.items(), 1):
                masses = values["masses"]
                if len(masses) > 1:
                    if len({mass for _, mass in masses}) > 1:
                        errors.append(ImportMessage("EXB01_PROCESS_PAIRED_MASS_MISMATCH", "同名物料的碳与挥发分行活动数据必须完全一致；不能相加或猜测。", f"{sheet_name}!C{masses[-1][0]}", "活动数据"))
                    mass = masses[0][1]
                else:
                    mass = masses[0][1] if masses else None
                for comp in ("fixed", "volatile"):
                    if len(values[comp]) > 1:
                        errors.append(ImportMessage("EXB01_PROCESS_DUPLICATE_COMPONENT", "同一过程、物料名称和组分出现多个活动行，无法判断是否为新增过程实例。", f"{sheet_name}!D{values[comp][1][0]}", "碳/挥发分含量"))
                fixed = values["fixed"][0][1] if values["fixed"] else None
                volatile = values["volatile"][0][1] if values["volatile"] else None
                material_rows.append(MaterialInputLine(
                    line_id=_stable_id("material", unit_id, sheet_name, ordinal), role=role, name=name, mass_t=mass,
                    fixed_carbon_percent=fixed, fixed_carbon_source=MaterialDataSource.MEASURED,
                    volatile_matter_percent=volatile, volatile_matter_source=MaterialDataSource.MEASURED,
                ))
            if not material_rows:
                continue
            factor = self._default_factor(default_parameter_id, context.period)
            parameter = self._parameter_value(factor, "采用当前期间 Resolver 核验的标准一般参数。") if factor else None
            instance_id = {"B.3": "calcination-1", "B.4": "baking-1", "B.5": "graphitization-1"}[sheet_name]
            process = model_type(instance_id=instance_id, material_rows=tuple(material_rows), **{parameter_name: parameter})
            outputs[index].append(process)
            involved.add(source_id)

    @staticmethod
    def _dynamic_process_role(sheet, group, position):
        normalized_group = "".join((group or "").split())
        if sheet == "B.3":
            slots = {
                "进入煅烧炉的碳": ((MaterialRole.CALCINATION_FEED, "fixed"),),
                "进入煅烧炉的挥发分": ((MaterialRole.CALCINATION_FEED, "volatile"),),
                "输出煅烧炉的碳": (
                    (MaterialRole.CALCINED_PRODUCT, "fixed"),
                    (MaterialRole.UNDERBURN_RECOVERED, "fixed"),
                    (MaterialRole.CARBON_DUST, "fixed"),
                ),
                "输出煅烧炉的挥发分": ((MaterialRole.CALCINED_PRODUCT, "volatile"),),
            }
        elif sheet == "B.4":
            slots = {
                "进入焙烧(炭化)的碳a": (
                    (MaterialRole.BAKING_FILLER, "fixed"),
                    (MaterialRole.GREEN_BAKING_PRODUCT, "fixed"),
                ),
                "进入焙烧(炭化)的挥发分": (
                    (MaterialRole.BAKING_FILLER, "volatile"),
                    (MaterialRole.GREEN_BAKING_PRODUCT, "volatile"),
                ),
                "输出焙烧(炭化)炉的碳": (
                    (MaterialRole.BAKING_BYPRODUCT, "fixed"),
                    (MaterialRole.BAKED_PRODUCT, "fixed"),
                ),
            }
        elif sheet == "B.5":
            slots = {
                "进入石墨化炉的碳": (
                    (MaterialRole.GRAPHITIZATION_PACKING, "fixed"),
                    (MaterialRole.GREEN_GRAPHITIZATION_PRODUCT, "fixed"),
                ),
                "进入石墨化炉的挥发分": ((MaterialRole.GRAPHITIZATION_PACKING, "volatile"),),
                "输出石墨化炉的碳": (
                    (MaterialRole.GRAPHITIZATION_BYPRODUCT, "fixed"),
                    (MaterialRole.GRAPHITIZED_PRODUCT, "fixed"),
                ),
            }
        else:
            return None
        group_slots = slots.get(normalized_group)
        if not group_slots:
            return None
        if len(group_slots) == 1:
            return group_slots[0]
        return group_slots[position - 1] if 1 <= position <= len(group_slots) else None

    def _parse_fumes(self, workbook, context, reader, errors, warnings, output, involved, unit_id):
        ws = workbook["B.6"]
        stop = next(
            (row for row in range(3, ws.max_row + 1) if _is_appendix_footer_row(ws, row)),
            ws.max_row + 1,
        )
        labels = ("烟气流量", "焦油含量", "低位发热量", "单位热值含碳量", "碳氧化率", "报告期时长")
        fch_default = self._default_factor("CAR-PAR-P04A-FCH", context.period)
        for row in range(3, stop):
            if not any(_has_value(ws[f"{col}{row}"]) for col in "ABCDEF"):
                continue
            values = [
                reader.number(workbook, "B.6", f"{col}{row}", errors, required=True, field_label=label)
                for col, label in zip("ABCDEF", labels)
            ]
            if any(value is None for value in values):
                continue
            q, qvar, hm, fch_value, fox_pct, duration = values
            fch = self._factor_cell(
                fch_value, fch_default, "CAR-PAR-P04A-FCH", "tC/GJ",
                f"B.6!D{row}", warnings, "单位热值含碳量",
            )
            fox = self.policy.divide(fox_pct, Decimal("100"))
            output.append(FumeIncinerationInput(
                q=InputValue(q, "Nm3/h", source_reference=f"B.6!A{row}"),
                qvar=InputValue(qvar, "mg/Nm3", source_reference=f"B.6!B{row}"),
                hm=InputValue(hm, "GJ/t", source_reference=f"B.6!C{row}"),
                fch=fch, fox=fox,
                duration=InputValue(duration, "d", source_reference=f"B.6!F{row}"),
                instance_id=f"fume-{row - 2}",
            ))
            involved.add(SOURCE_FUME)

    def _carbonate_option(self, label, period):
        if self.queries is None:
            return None
        import unicodedata
        normalize = lambda value: "".join(unicodedata.normalize("NFKC", str(value)).split()).casefold()
        target = normalize(label)
        matches = [item for item in self.queries.carbonate_options(period)
                   if target in {normalize(item.label), *(normalize(alias) for alias in item.aliases)}]
        return matches[0] if len(matches) == 1 else None

    def _parse_fgd(self, workbook, context, reader, errors, warnings, output, involved, unit_id, *, template_version="1.0.0"):
        ws = workbook["B.7"]
        current = template_version in {"2.0.0", "2.1.0", "2.2.0", "2.3.0", "2.4.0", "2.5.0", "2.6.0"}
        stop = next((row for row in range(3, ws.max_row + 1)
                     if _is_appendix_footer_row(ws, row)), ws.max_row + 1)
        merge_rows: dict[int, list[object]] = {1: [], 2: []}
        for area in ws.merged_cells.ranges:
            if area.min_col in merge_rows and area.min_row >= 3 and area.max_row < stop:
                merge_rows[area.min_col].append(area)
        a_ranges = sorted(merge_rows[1], key=lambda item: item.min_row)
        b_ranges = sorted(merge_rows[2], key=lambda item: item.min_row)
        a_bounds = {(item.min_row, item.max_row): item for item in a_ranges}
        b_bounds = {(item.min_row, item.max_row): item for item in b_ranges}
        merged_rows = {
            row for item in (*a_ranges, *b_ranges)
            for row in range(item.min_row, item.max_row + 1)
        }
        groups: list[tuple[int, int, bool]] = []
        for bounds, a_range in a_bounds.items():
            if bounds in b_bounds:
                groups.append((bounds[0], bounds[1], True))
            else:
                errors.append(ImportMessage(
                    "EXB01_FGD_BATCH_STRUCTURE_INVALID",
                    "A 列批次名称与 B 列消耗量的合并范围必须一致。",
                    f"B.7!A{a_range.min_row}", "碳酸盐批次名称",
                ))
        for bounds, b_range in b_bounds.items():
            if bounds not in a_bounds:
                errors.append(ImportMessage(
                    "EXB01_FGD_BATCH_STRUCTURE_INVALID",
                    "A 列批次名称与 B 列消耗量的合并范围必须一致。",
                    f"B.7!B{b_range.min_row}", "批次消耗量",
                ))

        explicit_starts: list[int] = []
        for row in range(3, stop):
            if row in merged_rows:
                continue
            has_name = _has_value(ws[f"A{row}"])
            has_amount = _has_value(ws[f"B{row}"])
            if has_name and has_amount:
                explicit_starts.append(row)
            elif has_name:
                errors.append(ImportMessage(
                    "EXB01_FGD_BATCH_BOUNDARY_INVALID",
                    "新增批次边界必须在同一行明确填写批次名称和消耗量。",
                    f"B.7!B{row}", "批次消耗量",
                ))
            elif has_amount:
                errors.append(ImportMessage(
                    "EXB01_FGD_BATCH_BOUNDARY_INVALID",
                    "新增批次边界必须在同一行明确填写批次名称和消耗量。",
                    f"B.7!A{row}", "碳酸盐批次名称",
                ))

        all_starts = sorted([start for start, _, _ in groups] + explicit_starts)
        for index, start in enumerate(explicit_starts):
            next_start = next((item for item in all_starts if item > start), stop)
            groups.append((start, next_start - 1, False))
        groups.sort(key=lambda item: item[0])
        grouped_rows = {row for start, end, _ in groups for row in range(start, end + 1)}
        for row in range(3, stop):
            if row in grouped_rows:
                continue
            if any(_has_value(ws[f"{col}{row}"]) for col in ("DEFHJKL" if template_version == "2.6.0" else "DEFHJ" if current else "DEF")):
                errors.append(ImportMessage(
                    "EXB01_FGD_BATCH_BOUNDARY_REQUIRED",
                    "该碳酸盐组分数据没有落在明确的 A/B 批次合并组或新批次边界内。",
                    f"B.7!C{row}", "碳酸盐组分",
                ))

        explicit_ratios = template_version == "2.6.0"
        facilities = {}
        for ordinal, (start, end, _merged) in enumerate(groups, 1):
            batch_name = reader.text(workbook, "B.7", f"A{start}", errors, field_label="脱硫剂 / 批次名称")
            amount = reader.number(workbook, "B.7", f"B{start}", errors, field_label="本批整剂净耗量")
            active_columns = "DEFHJKL" if explicit_ratios else "DEFHJ"
            active_rows = [row for row in range(start, end + 1)
                           if any(_has_value(ws[f"{col}{row}"]) for col in active_columns)]
            if amount is None and batch_name is None and not active_rows:
                continue
            if not batch_name or amount is None:
                errors.append(ImportMessage("EXB01_FGD_BATCH_AMOUNT_REQUIRED", "请完整填写本批次名称及整剂净耗量。", f"B.7!A{start}", "脱硫批次"))
                continue
            if not active_rows:
                errors.append(ImportMessage("EXB01_FGD_COMPONENT_REQUIRED", "本批次缺少碳酸盐组分输入。", f"B.7!C{start}", "碳酸盐组分"))
                continue
            facility_code = reader.text(workbook, "B.7", f"M{start}", errors, field_label="设施编号") if explicit_ratios else None
            # The immutable workbook unit identity namespaces opaque ingress IDs.
            # Coordinates locate the original group only during first parsing;
            # saved Canonical/Project identities are never regenerated from labels.
            instance_id = _stable_id("fgd", unit_id, ordinal, "batch")
            batch_id = _stable_id("fgd-batch", unit_id, ordinal, "batch")
            facility_id = facilities.setdefault(facility_code, _stable_id("fgd-facility", unit_id, 0, facility_code)) if facility_code else _stable_id("fgd-facility", unit_id, ordinal, "independent")
            components = []
            for row in active_rows:
                label = reader.text(workbook, "B.7", f"C{row}", errors, required=True, field_label="碳酸盐组分")
                name_mode = reader.text(workbook, "B.7", f"I{row}", errors, required=current, field_label="名称类型") if current else None
                if name_mode not in {"目录", "自定义"} and current:
                    errors.append(ImportMessage("EXB01_CARBONATE_IDENTITY_REQUIRED", "请选择目录或自定义碳酸盐身份。", f"B.7!I{row}", "名称类型"))
                origin = InputOrigin.CATALOG if name_mode == "目录" else InputOrigin.CUSTOM
                selected = {field: resolve_fgd_field(field, resolver=self.resolver, origin=origin,
                    carbonate_type=label, accounting_period=context.period, component_count=len(active_rows))
                    for field in ("fgd.ratio", "fgd.factor", "fgd.conversion")}
                adopted = {}
                for field, value_col, source_col, parameter_id, unit in (
                    ("fgd.ratio", "D", "K", "CAR-PAR-P04B-I", "ratio"),
                    ("fgd.factor", "E", "J", "fgd_carbonate_emission_factor_user", "tCO2/t"),
                    ("fgd.conversion", "F", "L", "CAR-PAR-P04B-TR", "ratio"),
                ):
                    source_label = reader.text(workbook, "B.7", f"{source_col}{row}", errors,
                        required=True, field_label="参数来源") if field == "fgd.factor" or explicit_ratios else "实测值"
                    result = selected[field]
                    source = _SOURCE_LABEL_KINDS.get(source_label)
                    if source not in {STANDARD_DEFAULT, MEASURED}:
                        code = ("EXB01_CARBONATE_FACTOR_SOURCE_REQUIRED" if source_label is None else "EXB01_CARBONATE_FACTOR_SOURCE_INVALID") if field == "fgd.factor" else "EXB01_FGD_SOURCE_REQUIRED"
                        errors.append(ImportMessage(code, "请明确选择缺省值或实测值，空白不补默认。", f"B.7!{source_col}{row}", "参数来源"))
                        adopted[field] = None
                        continue
                    issue = validate_parameter_source_policy(field, source, origin=origin,
                        resolved_context=result.context, carbonate_component_count=len(active_rows))
                    if issue:
                        code = issue[1]
                        if field == "fgd.ratio" and len(active_rows) > 1:
                            code = "EXB01_MULTI_CARBONATE_CONTENT_REQUIRED"
                        elif source == STANDARD_DEFAULT and field in {"fgd.ratio", "fgd.conversion"}:
                            code = "EXB01_CARBONATE_CONTENT_UNRESOLVED" if field == "fgd.ratio" else "EXB01_CARBONATE_CONVERSION_UNRESOLVED"
                        errors.append(ImportMessage(code, issue[2], f"B.7!{source_col}{row}", "参数来源"))
                        adopted[field] = None
                        continue
                    col = "H" if field == "fgd.factor" and source == MEASURED else value_col
                    raw = reader.number(workbook, "B.7", f"{col}{row}", errors, required=source == MEASURED, field_label="脱硫参数")
                    if field == "fgd.factor":
                        other = "E" if source == MEASURED else "H"
                        if _has_value(ws[f"{other}{row}"]):
                            errors.append(ImportMessage("EXB01_CARBONATE_FACTOR_SOURCE_CONFLICT", "EF来源与填写列冲突，请清空另一来源列。", f"B.7!{other}{row}", "排放因子"))
                    if field == "fgd.ratio" and len(active_rows) > 1 and source == MEASURED and raw is None:
                        errors.append(ImportMessage("EXB01_MULTI_CARBONATE_CONTENT_REQUIRED", "多组分脱硫剂须逐项填写实测含量。", f"B.7!{col}{row}", "组分含量"))
                    normalized = self.policy.divide(raw, Decimal("100")) if unit == "ratio" and raw is not None else raw
                    if source == STANDARD_DEFAULT:
                        if result.factor is None or (normalized is not None and normalized != _factor_value(result.factor)):
                            errors.append(ImportMessage("EXB01_FGD_DEFAULT_MISMATCH", "填写的缺省值与当前适用Resolver值不一致。", f"B.7!{col}{row}", "脱硫参数"))
                            adopted[field] = None
                        elif _has_value(ws[f"{col}{row}"]) and raw is None:
                            adopted[field] = None
                        else:
                            adopted[field] = self._parameter_value(result.factor, "按明确选择的标准缺省来源，采用同次Resolver核验参数。")
                    else:
                        adopted[field] = _user_parameter(normalized, parameter_id, unit, f"B.7!{col}{row}", ParameterSourceKind.MEASURED) if normalized is not None else None
                        self._warn_missing_source(warnings, f"B.7!{col}{row}", "脱硫参数")
                components.append(CarbonateComponent(InputValue(amount, "t", source_reference=f"B.7!B{start}"),
                    adopted["fgd.ratio"], adopted["fgd.factor"], adopted["fgd.conversion"],
                    label, origin, component_id=_stable_id("fgd-component", batch_id, row-start, "component")))
            output.append(FGDInput(components=tuple(components), instance_id=instance_id,
                reagent_name=batch_name, facility_id=facility_id, facility_name=facility_code, batch_id=batch_id, execution_version="FGD_V2"))
            involved.add(SOURCE_FGD)
