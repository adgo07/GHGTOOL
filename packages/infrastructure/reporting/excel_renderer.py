"""Render frozen report models as a readable, display-only Excel workbook."""

from __future__ import annotations

from io import BytesIO
import math
from pathlib import Path
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from packages.application.reporting.model import ReportCell, ReportModel, ReportSection, ReportTable

REPORT_DISCLAIMER = "已保存核算记录的报告显示值，非输入模板，不含Excel核算公式"
_EXPORT_FORMAT_ID = "GHGTOOL_SAVED_RECORD_REPORT_XLSX"
_EXPORT_FORMAT_VERSION = "1"
_METADATA_SHEET = "__export_metadata__"
_FONT_NAME = "Microsoft YaHei"
_TITLE_FILL = PatternFill(fill_type="solid", fgColor="DCEAF2")
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="EAF1F5")
_NOTE_FILL = PatternFill(fill_type="solid", fgColor="F6F8FA")
_BORDER = Border(bottom=Side(style="thin", color="CFD8DE"))
_TITLE_FONT = Font(name=_FONT_NAME, size=15, bold=True, color="173B52")
_HEADER_FONT = Font(name=_FONT_NAME, size=10, bold=True, color="173B52")
_BODY_FONT = Font(name=_FONT_NAME, size=10, color="202A30")
_NOTE_FONT = Font(name=_FONT_NAME, size=9, color="49545B")


def _excel_text(cell, value: object) -> None:
    """Write literal text even when a value starts with a formula prefix."""
    cell.value = "" if value is None else (value if isinstance(value, str) else str(value))
    cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.font = _BODY_FONT
    cell.border = _BORDER


def _safe_sheet_name(candidate: str, existing: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", candidate).strip().strip("'")
    name = name[:31].rstrip().strip("'") or "报告"
    base = name
    index = 2
    while name.casefold() in {item.casefold() for item in existing}:
        suffix = f" ({index})"
        name = f"{base[:31 - len(suffix)].rstrip()}{suffix}"
        index += 1
    existing.add(name)
    return name


def _ordered_sections(sections: tuple[ReportSection, ...]) -> tuple[ReportSection, ...]:
    by_id = {section.section_id: section for section in sections}
    ordered = [by_id.pop(f"b{i}") for i in range(1, 10) if f"b{i}" in by_id]
    ordered.extend(section for section in sections if section.section_id in by_id and section.section_id != "basic")
    return tuple(ordered)


def _display_width(text: str) -> int:
    """Approximate visible width using full-width East Asian glyphs as two units."""
    return sum(
        4 if character == "\t" else
        2 if unicodedata.east_asian_width(character) in {"F", "W"} else
        1
        for character in text
    )


def _line_count(text: object, width: float, font_size: float = 10) -> int:
    if text is None or str(text) == "":
        return 0
    capacity = max(1.0, (width - 1.5) * 10 / max(10, font_size))
    lines = str(text).splitlines() or [str(text)]
    return sum(max(1, math.ceil(_display_width(line) / capacity)) for line in lines)


def _table_shape(table: ReportTable) -> tuple[int, bool, bool]:
    column_count = max([len(table.columns), *(len(row.cells) for row in table.rows), 0])
    has_details = any(cell.unit or cell.source for row in table.rows for cell in row.cells)
    has_notes = any(row.note is not None for row in table.rows)
    return column_count, has_details, has_notes


def _table_output_columns(table: ReportTable) -> int:
    source_columns, has_details, has_notes = _table_shape(table)
    return max(1, source_columns + int(has_details) + int(has_notes))


def _matches_basic_information(table: ReportTable, basic: tuple[tuple[str, str], ...]) -> bool:
    if tuple(table.columns) != ("项目", "内容") or any(row.note is not None for row in table.rows):
        return False
    rows = tuple((
        row.cells[0].value if len(row.cells) > 0 else "",
        row.cells[1].value if len(row.cells) > 1 else "",
    ) for row in table.rows)
    return rows == basic


def _title_row(sheet, row: int, columns: int, title: str) -> None:
    if columns > 1:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    cell = sheet.cell(row=row, column=1)
    _excel_text(cell, title)
    cell.font = _TITLE_FONT
    cell.fill = _TITLE_FILL
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 28
    for column in range(2, columns + 1):
        sheet.cell(row=row, column=column).fill = _TITLE_FILL


def _header_row(sheet, row: int, columns: int) -> None:
    for column in range(1, columns + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = _BORDER
    sheet.row_dimensions[row].height = 32


def _cell_details(table: ReportTable, row: tuple[ReportCell, ...]) -> str:
    details: list[str] = []
    for index, cell in enumerate(row):
        parts: list[str] = []
        if cell.unit:
            parts.append(f"单位：{cell.unit}")
        if cell.source:
            parts.append(f"来源：{cell.source}")
        if not parts:
            continue
        label = table.columns[index] if index < len(table.columns) else f"未标注列{index + 1}"
        details.append(f"{label}：{'；'.join(parts)}")
    return "\n".join(details)


def _write_table(sheet, table: ReportTable, start_row: int) -> tuple[int, int, int]:
    source_columns, has_details, has_notes = _table_shape(table)
    output_columns = max(1, source_columns + int(has_details) + int(has_notes))
    _title_row(sheet, start_row, output_columns, table.title)

    header_row = start_row + 1
    headers = [
        table.columns[index] if index < len(table.columns) else f"未标注列{index + 1}"
        for index in range(source_columns)
    ]
    if has_details:
        headers.append("单位与来源说明")
    if has_notes:
        headers.append("行内说明")
    if not headers:
        headers.append("内容")
    for column, label in enumerate(headers, start=1):
        _excel_text(sheet.cell(row=header_row, column=column), label)
    _header_row(sheet, header_row, output_columns)

    row_number = header_row + 1
    for row in table.rows:
        for index in range(source_columns):
            item = row.cells[index] if index < len(row.cells) else ReportCell()
            _excel_text(sheet.cell(row=row_number, column=index + 1), item.value)
        if has_details:
            detail_column = source_columns + 1
            _excel_text(sheet.cell(row=row_number, column=detail_column), _cell_details(table, row.cells))
        if has_notes:
            _excel_text(sheet.cell(row=row_number, column=output_columns), row.note)
        row_number += 1
    return row_number, header_row, output_columns


def _write_notes(sheet, notes: tuple[str, ...], row_number: int, columns: int) -> int:
    if not notes:
        return row_number
    _title_row(sheet, row_number, columns, "本节说明")
    row_number += 1
    for note in notes:
        if columns > 1:
            sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=columns)
        cell = sheet.cell(row=row_number, column=1)
        _excel_text(cell, note)
        cell.font = _NOTE_FONT
        cell.fill = _NOTE_FILL
        row_number += 1
    return row_number


def _column_width_limit(sheet, column: int, last_row: int) -> tuple[int, int]:
    labels = {
        str(sheet.cell(row=row, column=column).value)
        for row in range(1, last_row + 1)
        if sheet.cell(row=row, column=column).value is not None
    }
    if "单位与来源说明" in labels:
        return 20, 36
    if "行内说明" in labels:
        return 14, 30
    if sheet.title == "基本信息":
        return (36, 44) if column == 2 else (18, 28)
    return 10, 18


def _set_row_height(sheet, row: int, lines: int, font_size: float = 10) -> None:
    if lines <= 0:
        return
    current = sheet.row_dimensions[row].height or 15
    needed = min(390, lines * max(15, font_size * 1.55) + 8)
    sheet.row_dimensions[row].height = max(current, needed)


def _print_layout(sheet, freeze_row: int, last_row: int, last_column: int, title_rows: str) -> None:
    last_row = max(1, last_row)
    last_column = max(1, last_column)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f"A{max(1, freeze_row)}"
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True, autoPageBreaks=False)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = title_rows
    sheet.page_margins.left = sheet.page_margins.right = 0.25
    sheet.page_margins.top = sheet.page_margins.bottom = 0.45
    sheet.print_area = f"A1:{get_column_letter(last_column)}{last_row}"

    widths: dict[int, float] = {}
    for column in range(1, last_column + 1):
        max_line_width = 0
        for row in range(1, last_row + 1):
            value = sheet.cell(row=row, column=column).value
            if value is not None:
                max_line_width = max(
                    max_line_width,
                    max((_display_width(line) for line in str(value).splitlines()), default=0),
                )
        minimum, maximum = _column_width_limit(sheet, column, last_row)
        width = min(maximum, max(minimum, max_line_width + 2))
        widths[column] = float(width)
        sheet.column_dimensions[get_column_letter(column)].width = width

    merged_ranges_by_row: dict[int, list[tuple[int, int, float]]] = {}
    for merged in sheet.merged_cells.ranges:
        if merged.min_row != merged.max_row:
            continue
        total_width = sum(widths.get(column, 10.0) for column in range(merged.min_col, merged.max_col + 1))
        merged_ranges_by_row.setdefault(merged.min_row, []).append((merged.min_col, merged.max_col, total_width))

    for row in range(1, last_row + 1):
        max_lines = 0
        max_font_size = 10.0
        merged_spans = merged_ranges_by_row.get(row, [])
        merged_origins = {start for start, _, _ in merged_spans}
        for column in range(1, last_column + 1):
            cell = sheet.cell(row=row, column=column)
            if cell.value is None or str(cell.value) == "":
                continue
            if column in merged_origins:
                continue
            max_lines = max(max_lines, _line_count(cell.value, widths[column], float(cell.font.sz or 10)))
            max_font_size = max(max_font_size, float(cell.font.sz or 10))
        for start, _, merged_width in merged_spans:
            cell = sheet.cell(row=row, column=start)
            if cell.value is None or str(cell.value) == "":
                continue
            max_lines = max(max_lines, _line_count(cell.value, merged_width, float(cell.font.sz or 10)))
            max_font_size = max(max_font_size, float(cell.font.sz or 10))
        _set_row_height(sheet, row, max_lines, max_font_size)


def _render_basic(sheet, model: ReportModel, section: ReportSection | None) -> tuple[int, int]:
    _title_row(sheet, 1, 2, model.standard_name)
    for row, label, value in (
        (2, "适用标准版本", model.standard_version),
        (3, "核算期间", model.period_text),
        (4, "编制日期", model.prepared_on),
    ):
        _excel_text(sheet.cell(row=row, column=1), label)
        _excel_text(sheet.cell(row=row, column=2), value)
        sheet.cell(row=row, column=1).font = _HEADER_FONT
    _title_row(sheet, 5, 2, REPORT_DISCLAIMER)
    if section is not None:
        _title_row(sheet, 6, 2, section.title)

    _excel_text(sheet.cell(row=7, column=1), "项目")
    _excel_text(sheet.cell(row=7, column=2), "内容")
    _header_row(sheet, 7, 2)
    row_number = 8
    max_columns = 2
    for label, value in model.basic_information:
        _excel_text(sheet.cell(row=row_number, column=1), label)
        _excel_text(sheet.cell(row=row_number, column=2), value)
        row_number += 1

    if section is not None:
        for table in section.tables:
            if _matches_basic_information(table, model.basic_information):
                continue
            row_number, _, table_columns = _write_table(sheet, table, row_number + 1)
            max_columns = max(max_columns, table_columns)
        row_number = _write_notes(sheet, section.notes, row_number + 1, max_columns)

    if model.notices:
        row_number += 1
        _title_row(sheet, row_number, max_columns, "提示")
        row_number += 1
        for notice in model.notices:
            if max_columns > 1:
                sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=max_columns)
            cell = sheet.cell(row=row_number, column=1)
            _excel_text(cell, notice)
            cell.font = _NOTE_FONT
            cell.fill = _NOTE_FILL
            row_number += 1
    if model.supplementary_note is not None:
        row_number += 1
        _title_row(sheet, row_number, max_columns, "补充说明")
        row_number += 1
        if max_columns > 1:
            sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=max_columns)
        cell = sheet.cell(row=row_number, column=1)
        _excel_text(cell, model.supplementary_note)
        cell.font = _NOTE_FONT
    return max(row_number, 7), max_columns


def _render_section(sheet, section: ReportSection, model: ReportModel) -> None:
    columns = max([2, *(_table_output_columns(table) for table in section.tables)])
    _title_row(sheet, 1, columns, section.title)
    _excel_text(sheet.cell(row=2, column=1), "核算期间")
    _excel_text(sheet.cell(row=2, column=2), model.period_text)
    sheet.cell(row=2, column=1).font = _HEADER_FONT
    _title_row(sheet, 3, columns, REPORT_DISCLAIMER)

    row_number = 5
    max_columns = 2
    freeze_row = 6
    for table in section.tables:
        row_number, header_row, table_columns = _write_table(sheet, table, row_number)
        max_columns = max(max_columns, table_columns)
        if freeze_row == 6:
            freeze_row = header_row + 1
        row_number += 1
    row_number = _write_notes(sheet, section.notes, row_number, max_columns)
    _print_layout(sheet, freeze_row, max(3, row_number - 1), max_columns, "1:3")


def _write_metadata(workbook: Workbook, model: ReportModel) -> None:
    sheet = workbook.create_sheet(_METADATA_SHEET)
    sheet.sheet_state = "hidden"
    for row, (key, value) in enumerate((
        ("format_id", _EXPORT_FORMAT_ID),
        ("format_version", _EXPORT_FORMAT_VERSION),
        ("record_id", model.record_id),
        ("report_schema_version", model.schema_version),
    ), start=1):
        _excel_text(sheet.cell(row=row, column=1), key)
        _excel_text(sheet.cell(row=row, column=2), value)


def render_report_xlsx(model: ReportModel, destination: str | Path | None = None) -> bytes:
    """Render a frozen ReportModel without consulting records or calculation services."""
    workbook = Workbook()
    workbook.properties.creator = "GHGTOOL"
    workbook.properties.title = "已保存核算记录报告"

    used_names: set[str] = set()
    basic = workbook.active
    basic.title = _safe_sheet_name("基本信息", used_names)
    basic_section = next((item for item in model.sections if item.section_id == "basic"), None)
    basic_last_row, basic_last_column = _render_basic(basic, model, basic_section)
    _print_layout(basic, 8, basic_last_row, basic_last_column, "1:7")

    for section in _ordered_sections(model.sections):
        preferred = "数据来源" if section.section_id == "evidence" else (section.title or section.section_id)
        sheet = workbook.create_sheet(_safe_sheet_name(preferred, used_names))
        _render_section(sheet, section, model)

    _write_metadata(workbook, model)
    buffer = BytesIO()
    workbook.save(buffer)
    result = buffer.getvalue()
    if destination is not None:
        Path(destination).write_bytes(result)
    return result


__all__ = ["REPORT_DISCLAIMER", "render_report_xlsx"]
