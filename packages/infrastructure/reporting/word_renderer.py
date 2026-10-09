"""Word renderer for immutable application report models."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from packages.application.reporting.model import ReportCell, ReportModel, ReportRow, ReportTable


def _set_run_font(run, *, size: float = 10.5, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _cell_text(cell_value: ReportCell) -> str:
    value = cell_value.value
    if cell_value.unit:
        value = f"{value} {cell_value.unit}".strip()
    if cell_value.source:
        value = f"{value}\n来源：{cell_value.source}" if value else f"来源：{cell_value.source}"
    return value


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    _set_run_font(run, size=9, color="657080")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    run = paragraph.add_run(" 页")
    _set_run_font(run, size=9, color="657080")


def _format_section(section, landscape: bool) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width, section.page_height = Cm(29.7), Cm(21)
        section.left_margin = section.right_margin = Cm(1.5)
    else:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.1)
    section.top_margin = section.bottom_margin = Cm(1.8)
    if not section.footer.paragraphs[0].text:
        _page_field(section.footer.paragraphs[0])


def _merge_cells(table, merges, offset=0):
    for r0, c0, r1, c1 in merges:
        origin_text = table.cell(r0 + offset, c0).text
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                if (r, c) != (r0, c0):
                    table.cell(r + offset, c).text = ""
        merged = table.cell(r0 + offset, c0).merge(table.cell(r1 + offset, c1))
        merged.text = origin_text


def _add_table(document: Document, table_model: ReportTable, *, page_before: bool = False) -> None:
    if table_model.context_label:
        context = document.add_paragraph(table_model.context_label)
        context.paragraph_format.keep_with_next = True
        context.paragraph_format.page_break_before = page_before
        for run in context.runs:
            _set_run_font(run, size=10, bold=True)
    heading = document.add_paragraph()
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(8)
    run = heading.add_run(table_model.title)
    _set_run_font(run, size=11, bold=True, color="000000")
    headers = table_model.header_rows or (ReportRow(tuple(ReportCell(c) for c in table_model.columns)),)
    table = document.add_table(rows=len(headers) + len(table_model.rows), cols=len(table_model.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    section = document.sections[-1]
    available = section.page_width - section.left_margin - section.right_margin
    weights = table_model.column_weights or (1,) * len(table_model.columns)
    total_weight = sum(weights)
    for i, weight in enumerate(weights):
        width = int(available * weight / total_weight)
        table.columns[i].width = width
        for cell in table.columns[i].cells:
            cell.width = width
    # Explicit black grid matches the approved form, including inside merged groups.
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        for attr, value in (("val", "single"), ("sz", "4"), ("color", "000000")):
            border.set(qn("w:" + attr), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", "40"), ("bottom", "40"), ("left", "75"), ("right", "75")):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    table._tbl.tblPr.append(margins)
    for row_index, report_row in enumerate((*headers, *table_model.rows)):
        row = table.rows[row_index]
        is_header = row_index < len(headers)
        if is_header:
            _repeat_header(row)
        for i, cell_model in enumerate(report_row.cells):
            row.cells[i].text = _cell_text(cell_model)
        # Allow automatic height, but avoid splitting one material/source row across pages.
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
    _merge_cells(table, table_model.header_merges)
    _merge_cells(table, table_model.body_merges, len(headers))
    for row_index, row in enumerate(table.rows):
        is_header = row_index < len(headers)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if is_header:
                _shade(cell, "F2F2F2")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
                if row_index >= len(headers) + len(table_model.rows) - (2 if table_model.context_label else 1) and table_model.footnotes:
                    paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    _set_run_font(run, size=9.5, bold=is_header, color="000000")
    notes = (*table_model.footnotes, *(row.note for row in table_model.rows if row.note))
    for note_index, note in enumerate(notes):
        paragraph = document.add_paragraph(note)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_with_next = note_index < len(notes) - 1
        for run in paragraph.runs:
            _set_run_font(run, size=9, color="000000")


def render_report_docx(model: ReportModel, destination: str | Path | None = None) -> bytes:
    """Render a Word file and return its bytes; no record or calculation is modified."""

    document = Document()
    _format_section(document.sections[0], False)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run(model.standard_name)
    _set_run_font(run, size=18, bold=True, color="000000")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"核算期间：{model.period_text}    编制日期：{model.prepared_on}")
    _set_run_font(run, size=10, color="657080")
    run = document.add_paragraph().add_run(f"适用版本：{model.standard_version}")
    _set_run_font(run, size=9, color="657080")

    document.core_properties.version = model.schema_version
    document.core_properties.comments = f"布局：{model.layout_id}；批准模板 SHA256：{model.template_sha256}"
    current_landscape = False
    for report_section in model.sections:
        renderable_tables = tuple(table for table in report_section.tables if table.rows)
        first_orientation = renderable_tables[0].landscape if renderable_tables else current_landscape
        page_before = report_section.section_id.startswith("b") and report_section.section_id != "basic" and first_orientation == current_landscape
        if renderable_tables and first_orientation != current_landscape:
            section = document.add_section()
            _format_section(section, first_orientation)
            current_landscape = first_orientation
        heading = document.add_paragraph()
        heading.paragraph_format.page_break_before = page_before
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(13)
        run = heading.add_run(report_section.title)
        _set_run_font(run, size=14, bold=True, color="000000")
        for table_index, table_model in enumerate(report_section.tables):
            if not table_model.rows:
                empty = document.add_paragraph("该记录未保存此项明细。")
                for run in empty.runs:
                    _set_run_font(run, size=9, color="657080")
                continue
            if table_model.landscape != current_landscape:
                section = document.add_section()
                _format_section(section, table_model.landscape)
                current_landscape = table_model.landscape
            _add_table(document, table_model, page_before=table_index > 0 and bool(table_model.context_label))
        for note in report_section.notes:
            paragraph = document.add_paragraph(note)
            for run in paragraph.runs:
                _set_run_font(run, size=9, color="657080")

    if model.notices:
        heading = document.add_paragraph()
        run = heading.add_run("核算提醒")
        _set_run_font(run, size=13, bold=True, color="8A5A00")
        for notice in model.notices:
            paragraph = document.add_paragraph(style="List Bullet")
            run = paragraph.add_run(notice)
            _set_run_font(run, size=9)
    if model.supplementary_note:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("补充说明：" + model.supplementary_note)
        _set_run_font(run, size=9)
    footer = document.sections[-1].footer.paragraphs[0]
    footer.text = "本报告由已保存的核算记录生成    "
    _page_field(footer)

    output = BytesIO()
    document.save(output)
    content = output.getvalue()
    if destination is not None:
        Path(destination).write_bytes(content)
    return content


__all__ = ["render_report_docx"]
