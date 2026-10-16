"""Build development checks for the business-rule review workbook and linked evidence."""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import sys
import urllib.parse
from collections.abc import Mapping
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORK = REPOSITORY_ROOT / "build/business-rules-authoring"
DEFAULT_STAGE = DEFAULT_WORK / "r11-review"
GENERATOR = "scripts/build_business_rules_review.mjs"
FIXTURE_SCRIPT = "scripts/test_business_rules_review.mjs"
WORKBOOK_NAME = "业务规则审阅表.xlsx"
EXPECTED_SCOPES = {
    "业务规则": ["整体", "允许来源", "默认来源", "数值默认", "读写锁定"],
    "操作与影响": ["整体", "操作"],
    "本次变化": ["整体"],
}
HIDDEN_HEADERS = {
    "原审阅版本",
    "原审阅投影",
    "稳定审阅键",
    "内容指纹",
    "意见基线指纹",
    "原投影键",
    "当前投影键",
    "规则修订",
}
REQUIRED_HEADERS = {
    "审阅意见",
    "处理状态",
    "审阅编号",
    "意见针对属性",
    "意见针对条件",
    "稳定审阅键",
    "内容指纹",
    "意见基线指纹",
    "原投影键",
    "当前投影键",
}
FEEDBACK_HEADERS = ("审阅意见", "处理状态", "审阅编号")


def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def resolve_path(root: Path, value: Path) -> Path:
    return value if value.is_absolute() else root / value


def semantic_json_sha256(value: Any) -> str:
    """Match builder digest(canonical(value)) / JSON.stringify serialization."""
    canonical_json = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("utf-8")
    return sha256_bytes(canonical_json)


def _workbook_condition_scopes(sheet: Any) -> set[tuple[str, str, str]]:
    headers = [cell.value for cell in sheet[5]]
    required = ("稳定审阅键", "内容指纹")
    if not all(name in headers for name in required):
        raise ValueError("business review sheet lacks scope fingerprint columns")
    stable_index = headers.index("稳定审阅键")
    fingerprints_index = headers.index("内容指纹")
    scopes: set[tuple[str, str, str]] = set()
    for row in sheet.iter_rows(min_row=6, values_only=True):
        stable_key = row[stable_index] if stable_index < len(row) else None
        if not _nonempty(stable_key):
            continue
        raw_fingerprints = row[fingerprints_index] if fingerprints_index < len(row) else None
        try:
            fingerprints = json.loads(raw_fingerprints or "{}")
        except (TypeError, json.JSONDecodeError) as exc:
            raise ValueError(f"invalid review scope fingerprints for {stable_key}") from exc
        if not isinstance(fingerprints, Mapping):
            raise ValueError(f"review scope fingerprints must be an object for {stable_key}")
        for key in fingerprints:
            if isinstance(key, str) and "@@" in key:
                prop, condition = key.split("@@", 1)
                scopes.add((str(stable_key), prop, condition))
    return scopes


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _nonempty(value: Any) -> bool:
    return value is not None and value != ""


def _feedback_rows(workbook_path: Path) -> int:
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    try:
        sheet_aliases = {
            "业务规则": ("业务规则", "字段与条件规则"),
            "操作与影响": ("操作与影响",),
            "本次变化": ("本次变化",),
        }
        count = 0
        for canonical_name, aliases in sheet_aliases.items():
            sheet = next((workbook[name] for name in aliases if name in workbook.sheetnames), None)
            if sheet is None:
                raise ValueError(f"workbook missing review sheet: {canonical_name}")
            matrix = sheet.iter_rows(values_only=True)
            header = None
            for row in matrix:
                values = tuple(row)
                if "稳定审阅键" in values:
                    header = values
                    break
            if header is None:
                raise ValueError(f"review sheet lacks stable keys: {sheet.title}")
            header_index = {name: index for index, name in enumerate(header)}
            if not all(name in header_index for name in (*FEEDBACK_HEADERS, "稳定审阅键")):
                raise ValueError(f"review sheet lacks feedback columns: {sheet.title}")
            feedback_indexes = [header_index[name] for name in FEEDBACK_HEADERS]
            key_index = header_index["稳定审阅键"]
            for row in matrix:
                values = tuple(row)
                key = values[key_index] if key_index < len(values) else None
                feedback = [
                    values[index] if index < len(values) else None
                    for index in feedback_indexes
                ]
                if _nonempty(key) and any(_nonempty(value) for value in feedback):
                    count += 1
        return count
    finally:
        workbook.close()


def _catalog_binding_matches(field: Mapping[str, Any], entry: Mapping[str, Any]) -> bool:
    default = (field.get("policy_branch_values") or {}).get("default") or {}
    tags = ((field.get("condition") or {}).get("tags") or {})
    return (
        default.get("state") == "BINDING"
        and default.get("parameter_id") == entry.get("parameter_id")
        and tags.get("binding_id") == entry.get("binding_id")
    )


def build_checks(
    *,
    root: Path,
    stage: Path,
    review_json: Path,
    technical_json: Path,
    fixture: Path,
    output: Path,
    previous: Path | None = None,
) -> dict[str, Any]:
    root = root.resolve()
    stage = resolve_path(root, stage).resolve()
    review_json = resolve_path(root, review_json).resolve()
    technical_json = resolve_path(root, technical_json).resolve()
    fixture = resolve_path(root, fixture).resolve()
    output = resolve_path(root, output).resolve()
    previous_path = resolve_path(root, previous).resolve() if previous else None

    workbook_path = stage / WORKBOOK_NAME
    manifest_path = stage / "review-view-manifest.json"
    details_path: Path

    for required in (workbook_path, manifest_path, review_json, technical_json, fixture):
        if not required.is_file():
            raise FileNotFoundError(f"required review input is missing: {required}")

    raw_review = review_json.read_bytes()
    raw_review.decode("utf-8")
    json.loads(raw_review.decode("utf-8"))
    projection = sha256_bytes(raw_review.replace(b"\r\n", b"\n"))
    manifest = read_json(manifest_path)
    if manifest.get("projection_sha256") != projection or manifest.get("details_projection_bound") is not True:
        raise ValueError("review-view-manifest must bind the details to this review projection")

    details_name = manifest.get("details")
    if not isinstance(details_name, str) or not details_name:
        raise ValueError("review-view-manifest.details must name the generated details file")
    details_path = (stage / details_name).resolve()
    if not details_path.is_relative_to(stage):
        raise ValueError("details file must remain inside the review stage directory")
    if not details_path.is_file():
        raise FileNotFoundError(f"review details file is missing: {details_path}")

    raw_technical = technical_json.read_bytes()
    raw_technical.decode("utf-8")
    technical = json.loads(raw_technical.decode("utf-8"))
    technical_hash = sha256_bytes(raw_technical.replace(b"\r\n", b"\n"))
    technical_metadata = technical.get("metadata")
    if not isinstance(technical_metadata, Mapping):
        raise ValueError("technical review JSON must contain metadata")
    if technical_metadata.get("source_compact_sha256") != projection:
        raise ValueError("technical review JSON is not bound to this review projection")
    if manifest.get("details_revision") != technical_metadata.get("revision"):
        raise ValueError("review-view-manifest details revision does not match technical review JSON")
    actual_details_counts = {}
    for section in ("fields", "operations", "changes"):
        rows = technical.get(section)
        if not isinstance(rows, list):
            raise ValueError(f"technical review JSON must contain a {section} array")
        actual_details_counts[section] = len(rows)
    if manifest.get("details_counts") != actual_details_counts:
        raise ValueError(
            "review-view-manifest details counts do not match technical review JSON: "
            f"manifest={manifest.get('details_counts')}, actual={actual_details_counts}"
        )
    details_semantic_hash = semantic_json_sha256(technical)
    if manifest.get("details_sha256") != details_semantic_hash:
        raise ValueError("review-view-manifest semantic details SHA-256 does not match technical review JSON")
    technical_fields = technical.get("fields")
    if not isinstance(technical_fields, list):
        raise ValueError("technical review JSON must contain a fields array")
    if any(
        not isinstance(row, dict)
        or not isinstance(row.get("evaluation_status"), str)
        or not row["evaluation_status"]
        for row in technical_fields
    ):
        raise ValueError("technical review fields must carry an evaluation_status")
    technical_combinations = {
        "total": len(technical_fields),
        "READY": sum(row["evaluation_status"] == "READY" for row in technical_fields),
        "BLOCKED": sum(row["evaluation_status"] == "BLOCKED" for row in technical_fields),
    }

    catalog_snapshot = manifest.get("catalog_snapshot")
    if not isinstance(catalog_snapshot, Mapping):
        raise ValueError("review-view-manifest.catalog_snapshot must be an object")
    entries = catalog_snapshot.get("entries")
    if not isinstance(entries, Mapping) or not entries:
        raise ValueError("catalog snapshot must contain bound entries")
    catalog_values: dict[str, Any] = {}
    for field_id, entry in entries.items():
        if not isinstance(entry, Mapping) or entry.get("field_id") != field_id:
            raise ValueError(f"catalog snapshot entry has an invalid field identity: {field_id}")
        matches = [
            row
            for row in technical_fields
            if isinstance(row, dict)
            and row.get("field_id") == field_id
            and _catalog_binding_matches(row, entry)
        ]
        if not matches:
            raise ValueError(f"catalog snapshot entry is not bound in technical review: {field_id}")
        if not isinstance(entry.get("value"), str) or not entry["value"].strip():
            raise ValueError(f"catalog snapshot value is missing: {field_id}")
        catalog_values[field_id] = entry["value"]

    details_html = details_path.read_text(encoding="utf-8")
    scope_section = re.search(
        r'<section\b[^>]*\bid="review-scope"[^>]*>([\s\S]*?)</section>',
        details_html,
    )
    if not scope_section:
        raise ValueError("same-source details HTML lacks the review-scope provenance section")
    digest_markers = re.findall(
        r"正式投影SHA256：([0-9a-f]{64})；详情输入SHA256：([0-9a-f]{64})",
        scope_section.group(1),
    )
    if len(digest_markers) != 1:
        raise ValueError("same-source details HTML must contain one projection/details digest pair")
    html_projection_hash, html_details_hash = digest_markers[0]
    if html_projection_hash != projection:
        raise ValueError("details HTML projection SHA-256 does not match the current review projection")
    if html_details_hash != details_semantic_hash:
        raise ValueError("details HTML semantic SHA-256 does not match the current technical review JSON")
    if html_details_hash != manifest.get("details_sha256"):
        raise ValueError("details HTML and review-view-manifest semantic digests do not match")
    visible_scopes = [
        (
            urllib.parse.unquote(key),
            html.unescape(prop),
            html.unescape(condition),
        )
        for key, prop, condition in re.findall(
            r'<pre data-review-row="([^"]+)" data-review-property="([^"]+)">([\s\S]*?)</pre>',
            details_html,
        )
    ]
    if not visible_scopes:
        raise ValueError("same-source details HTML contains no review condition scopes")
    if len(visible_scopes) != len(set(visible_scopes)):
        raise ValueError("same-source details HTML contains duplicate review condition scopes")

    workbook = load_workbook(workbook_path, data_only=False)
    try:
        if workbook.sheetnames != list(EXPECTED_SCOPES):
            raise ValueError(f"unexpected review sheets: {workbook.sheetnames}")
        workbook_scopes = _workbook_condition_scopes(workbook["业务规则"])
        if set(visible_scopes) != workbook_scopes:
            missing_from_html = sorted(workbook_scopes - set(visible_scopes))[:3]
            missing_from_workbook = sorted(set(visible_scopes) - workbook_scopes)[:3]
            raise ValueError(
                "same-source details HTML scopes do not match business review fingerprints: "
                f"missing_from_html={missing_from_html}, missing_from_workbook={missing_from_workbook}"
            )
        checks: dict[str, Any] = {
            "candidate_only": True,
            "review_mode": "compact_business_v2",
            "review_json_file_sha256": sha256_file(review_json),
            "workbook_sha256": sha256_file(workbook_path),
            "generator": GENERATOR,
            "generator_sha256": sha256_file(root / GENERATOR),
            "fixture_script": FIXTURE_SCRIPT,
            "fixture_script_sha256": sha256_file(root / FIXTURE_SCRIPT),
            "details_file": details_name,
            "details_file_sha256": sha256_file(details_path),
            "details_source_revision": manifest.get("details_revision"),
            "details_source_counts": manifest.get("details_counts"),
            "details_semantic_sha256": manifest.get("details_sha256"),
            "details_projection_bound": True,
            "serialization_aliases": {
                "review_mode": "compact_business_v2",
                "manifest.view_revision": manifest.get("view_revision"),
                "meaning": "既有序列化命名别名；保留各自接口和值，不据此改变业务语义。",
            },
            "sheets": {},
            "table_ranges": {},
            "opinion_and_identity_columns_in_table": True,
            "freeze_panes": "D6",
            "formulas": 0,
            "error_cells": 0,
            "maximum_cell_characters": 0,
            "maximum_row_height_points": 0,
            "view_scope": {
                "power_rows": "按真实输入路径与方向拆分，不作为新增业务对象",
                "visible": "来源/默认/锁定及适用条件、操作初始/触发/确认、必要边界、意见属性与条件",
                "details": "同源HTML对象目录保留渠道/消费者/案例/真实绑定/完整条件和理由；定位列为文本路径与锚点",
                "omitted_from_main": "内部镜像与重复技术组合保留在详情，不在主表重复展开",
                "compatibility": "候选未接生产；codec/模板/DB迁移/正式Record/报告结构未改。映射澄清已独立提交。",
            },
            "definition_assets_hash_policy": manifest.get("definition_assets_hash_policy"),
            "definition_assets_sha256": manifest.get("definition_assets_sha256"),
            "historical_electricity_evidence": {
                "file": "electricity-paths.json",
                "sha256": sha256_file(root / "docs/architecture/business-rules/evidence/electricity-paths.json"),
                "as_of": "2026-10-15",
                "observed_head": "d279f505951b68750db133cf4279b2b76e0323bb",
                "scope": "固定历史观察，不把旧head/状态伪更新成当前验证；现行口径以当前仓库Mapping和有效Owner软件决定为准。",
                "roundtrip_distinction": "test_business_rules_excel_evidence覆盖构造OOXML缓存；原生Excel/WPS保存回导另列，不互相替代",
            },
            "native_office_validation": False,
            "native_office_status": "NOT_RUN",
            "native_office_deferral": "DEFERRED",
            "native_scope": (
                "Owner允许延至下阶段首次正式使用审阅表前完成原生Excel/WPS副本验证；"
                "本次未运行，不记为PASS或豁免，不启动CU。"
            ),
        }

        manifest_tables = manifest.get("tables")
        if not isinstance(manifest_tables, list):
            raise ValueError("review-view-manifest.tables must be an array")
        manifest_rows = {
            item.get("sheet"): item.get("rows")
            for item in manifest_tables
            if isinstance(item, dict)
        }
        feedback_declared = manifest.get("retained_feedback")
        if isinstance(feedback_declared, bool) or not isinstance(feedback_declared, int) or feedback_declared < 0:
            raise ValueError("review-view-manifest.retained_feedback must be a non-negative integer")

        retained_feedback_rows = 0
        for sheet in workbook:
            headers = [cell.value for cell in sheet[5]]
            if not all(isinstance(header, str) and header for header in headers):
                raise ValueError(f"review sheet has an empty or non-text header: {sheet.title}")
            if len(set(headers)) != len(headers):
                raise ValueError(f"review sheet has duplicate headers: {sheet.title}")
            if not REQUIRED_HEADERS.issubset(headers):
                raise ValueError(f"review sheet is missing feedback or identity headers: {sheet.title}")

            hidden_expected = {
                index + 1 for index, header in enumerate(headers) if header in HIDDEN_HEADERS
            }
            hidden_actual = {
                index
                for index in range(1, sheet.max_column + 1)
                if sheet.column_dimensions.get(get_column_letter(index))
                and sheet.column_dimensions[get_column_letter(index)].hidden
            }
            outlined_actual = {
                index
                for index in range(1, sheet.max_column + 1)
                if sheet.column_dimensions.get(get_column_letter(index))
                and sheet.column_dimensions[get_column_letter(index)].outlineLevel > 0
            }
            if hidden_actual != hidden_expected or outlined_actual != hidden_expected:
                raise ValueError(f"hidden review metadata columns changed: {sheet.title}")

            scope_column = headers.index("意见针对属性") + 1
            matching_validations = []
            for validation in sheet.data_validations.dataValidation:
                ranges = [range_boundaries(str(ref)) for ref in validation.sqref.ranges]
                if any(left == scope_column and right == scope_column for left, _, right, _ in ranges):
                    matching_validations.append((validation, ranges))
            if len(matching_validations) != 1:
                raise ValueError(f"review scope dropdown is missing or ambiguous: {sheet.title}")
            validation, ranges = matching_validations[0]
            expected_range = (scope_column, 6, scope_column, sheet.max_row)
            allowed = EXPECTED_SCOPES[sheet.title]
            if (
                validation.type != "list"
                or ranges != [expected_range]
                or validation.formula1 != '"' + ",".join(allowed) + '"'
            ):
                raise ValueError(f"review scope dropdown does not match its sheet: {sheet.title}")

            if sheet["A1"].value != "审阅数据SHA256：" + projection:
                raise ValueError(f"review sheet projection marker is stale: {sheet.title}")
            if sheet.freeze_panes != "D6" or len(sheet.tables) != 1:
                raise ValueError(f"review table or frozen panes changed: {sheet.title}")
            table = next(iter(sheet.tables.values()))
            actual_table_range = range_boundaries(table.ref)
            expected_table_range = (1, 5, len(headers), sheet.max_row)
            if actual_table_range != expected_table_range:
                raise ValueError(f"review table does not cover the complete sheet: {sheet.title}")

            keys: list[Any] = []
            for row in sheet.iter_rows(min_row=6, values_only=True):
                keys.append(row[headers.index("稳定审阅键")])
                feedback_values = [
                    row[headers.index(name)] for name in FEEDBACK_HEADERS
                ]
                if _nonempty(row[headers.index("稳定审阅键")]) and any(
                    _nonempty(value) for value in feedback_values
                ):
                    retained_feedback_rows += 1
                for header_name in ("内容指纹", "意见基线指纹"):
                    json.loads(row[headers.index(header_name)])
            if not all(keys) or len(keys) != len(set(keys)):
                raise ValueError(f"review table has missing or duplicate stable keys: {sheet.title}")

            max_height = max(
                (dimension.height or 0 for dimension in sheet.row_dimensions.values()),
                default=0,
            )
            if max_height > 409:
                raise ValueError(f"review sheet has an unreadable row height: {sheet.title}")
            checks["maximum_row_height_points"] = max(
                checks["maximum_row_height_points"], max_height
            )
            checks["table_ranges"][sheet.title] = {
                "range": table.ref,
                "columns": headers,
                "maximum_row_height_points": max_height,
                "data_validations": len(sheet.data_validations.dataValidation),
            }
            checks["sheets"][sheet.title] = len(keys)

            for row in sheet.iter_rows():
                for cell in row:
                    checks["formulas"] += cell.data_type == "f"
                    checks["error_cells"] += cell.data_type == "e"
                    if isinstance(cell.value, str):
                        checks["maximum_cell_characters"] = max(
                            checks["maximum_cell_characters"], len(cell.value)
                        )
                        if len(cell.value) >= 32767:
                            raise ValueError(f"review cell exceeds Excel text limit: {sheet.title}!{cell.coordinate}")

        if checks["formulas"] or checks["error_cells"]:
            raise ValueError("review workbook contains formula or error cells")
        if checks["sheets"] != manifest_rows:
            raise ValueError("review-view-manifest table row counts do not match the workbook")
        if retained_feedback_rows != feedback_declared:
            raise ValueError(
                "retained feedback row count does not match review-view-manifest.retained_feedback: "
                f"manifest={feedback_declared}, workbook={retained_feedback_rows}"
            )
        previous_feedback_rows = _feedback_rows(previous_path) if previous_path else None
        if previous_feedback_rows is not None and previous_feedback_rows != feedback_declared:
            raise ValueError(
                "previous workbook feedback row count does not match generated manifest: "
                f"previous={previous_feedback_rows}, manifest={feedback_declared}"
            )
        checks["feedback_retention"] = {
            "manifest_retained_feedback": feedback_declared,
            "workbook_feedback_rows": retained_feedback_rows,
            "previous_workbook_feedback_rows": previous_feedback_rows,
            "status": "VERIFIED",
        }
    finally:
        workbook.close()

    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    from scripts.business_rules_review_hashes import verify_definition_assets

    verify_definition_assets(root, checks)
    fixture_report = read_json(fixture)
    if fixture_report.get("status") != "PASS":
        raise ValueError("feedback fixture report status must be PASS")
    if fixture_report.get("builder_sha256") != checks["generator_sha256"]:
        raise ValueError("feedback fixture was not generated by the current workbook builder")
    if fixture_report.get("formal_projection_sha256") != projection:
        raise ValueError("feedback fixture is not bound to this review projection")
    if fixture_report.get("formal_details_sha256") != technical_hash:
        raise ValueError("feedback fixture does not bind the current raw technical input")
    if fixture_report.get("details_projection_bound") is not True:
        raise ValueError("feedback fixture must report details bound to its formal projection")

    checks["condition_scopes"] = {
        "count": len(visible_scopes),
        "unique": len(set(visible_scopes)),
        "source": "HTML与业务规则内容指纹双向核对；详情摘要绑定当前technical JSON",
    }
    checks["catalog_snapshot"] = catalog_snapshot
    checks["catalog_snapshot_values"] = catalog_values
    checks["feedback_test"] = {
        "report": fixture.name,
        "report_sha256": sha256_file(fixture),
        "command": (
            "node scripts/test_business_rules_review.mjs --builder scripts/build_business_rules_review.mjs "
            "--input <review-json> --technical-input <technical-json> "
            "--output <stage>/review-feedback-fixture.json"
        ),
        "method": "真实Table数据范围逆转、ArtifactTool导出/重开/再生成；不冒充原生排序",
        "input_relation": "从本轮真实投影选择对象；意见跨版本、删除、属性及技术详情变更按fixture报告核验",
        "status": "PASS",
        "actual_scenario_count": fixture_report.get("actual_scenario_count"),
        "input_mode": fixture_report.get("input_mode"),
        "formal_projection_sha256": fixture_report["formal_projection_sha256"],
        "formal_details_sha256": fixture_report["formal_details_sha256"],
    }
    checks["technical_combinations"] = technical_combinations

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(checks, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    return {
        "status": "PASS",
        "checks": str(output),
        "sheets": checks["sheets"],
        "technical_combinations": technical_combinations,
        "catalog_snapshot_values": catalog_values,
        "condition_scopes": checks["condition_scopes"],
        "retained_feedback": checks["feedback_retention"],
        "workbook_sha256": checks["workbook_sha256"],
        "checks_sha256": sha256_file(output),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="生成业务规则审阅资产检查摘要（开发工具，不是生产依赖）。"
    )
    parser.add_argument("--root", type=Path, default=REPOSITORY_ROOT)
    parser.add_argument("--stage", type=Path, default=DEFAULT_STAGE)
    parser.add_argument("--review-json", type=Path, default=DEFAULT_WORK / "review-r11.json")
    parser.add_argument("--technical-json", type=Path, default=DEFAULT_WORK / "technical-r11.json")
    parser.add_argument("--fixture", type=Path, default=DEFAULT_STAGE / "review-feedback-fixture.json")
    parser.add_argument("--output", type=Path, default=DEFAULT_STAGE / "review-artifact-checks.json")
    parser.add_argument(
        "--previous",
        type=Path,
        help="可选：上一版工作簿；只核对意见非空行数与manifest及新表一致，意见关联由JS夹具核验。",
    )
    args = parser.parse_args(argv)
    try:
        result = build_checks(
            root=args.root,
            stage=args.stage,
            review_json=args.review_json,
            technical_json=args.technical_json,
            fixture=args.fixture,
            output=args.output,
            previous=args.previous,
        )
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError, KeyError) as exc:
        parser.exit(2, f"business-rule review checks failed: {exc}\n")
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
