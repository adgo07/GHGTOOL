"""Small Appendix B fixture helpers shared by focused and migrated importer tests."""
from __future__ import annotations

from datetime import date
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from io import BytesIO
from xml.etree import ElementTree

from openpyxl import load_workbook

from packages.core.models import AccountingPeriod, PeriodType
from packages.excel.appendix_b import AppendixBImportContext
from packages.excel.templates import ExcelTemplateService
from packages.standards.carbon_material import FuelPath

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def appendix_b_context(*, year: int = 2025, region: str | None = None,
                       enterprise_name: str = "附录B测试企业",
                       fuel_path_overrides=None) -> AppendixBImportContext:
    return AppendixBImportContext(
        period=AccountingPeriod(PeriodType.ANNUAL, date(year, 1, 1), date(year, 12, 31)),
        boundary_confirmed=True, enterprise_name=enterprise_name, region=region,
        fuel_path_overrides=fuel_path_overrides or {"B.2!B4": FuelPath.VOLUME},
    )


def write_valid_appendix_b_workbook(path: str | Path, *, all_sources: bool = False) -> Path:
    """Write a master-based workbook; `all_sources=True` fills one complete source row/group each."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()), data_only=False)
    if all_sources:
        fuel = workbook["B.2"]
        fuel["B4"], fuel["C4"], fuel["D4"], fuel["H4"], fuel["I4"] = 1, 0.5, "实测值", 98, "实测值"
        for sheet, assignments in {
            "B.3": {4:(100,90), 5:(100,5), 6:(80,95), 7:(0,0), 8:(0,0), 9:(80,1)},
            "B.4": {4:(40,90), 5:(50,80), 6:(40,5), 7:(50,3), 8:(1,0), 9:(45,92)},
            "B.5": {4:(40,80), 5:(35,90), 6:(40,4), 7:(1,0), 8:(32,95)},
        }.items():
            ws = workbook[sheet]
            for row, (mass, percent) in assignments.items():
                ws[f"C{row}"], ws[f"D{row}"] = mass, percent
        fume = workbook["B.6"]
        for cell, value in {"A3":100,"B3":10,"C3":30,"D3":0.02,"E3":98,"F3":365}.items():
            fume[cell] = value
        fgd = workbook["B.7"]
        fgd["A3"], fgd["B3"], fgd["D3"], fgd["F3"] = "脱硫批次A", 2, 90, 95
        fgd["D4"], fgd["F4"] = 10, 90
        power = workbook["B.8"]
        power["B3"], power["C3"] = "电网电力", 100
        power["B4"], power["C4"], power["D4"] = "电网电力", 2, 0.5
        heat = workbook["B.9"]
        heat["B3"], heat["C3"], heat["D3"] = "饱和蒸汽", 1000, 1
        heat["B4"], heat["C4"], heat["D4"], heat["G4"] = "饱和蒸汽", 100, 1, 0.12
    workbook.save(target)
    workbook.close()
    return target


def rewrite_numeric_lexeme(path: str | Path, sheet_name: str, cell_reference: str, lexical_value: str) -> None:
    """Rewrite one saved OOXML numeric lexical token without openpyxl normalization."""
    source = Path(path)
    output = BytesIO()
    with ZipFile(source, "r") as original:
        workbook_root = ElementTree.fromstring(original.read("xl/workbook.xml"))
        rels_root = ElementTree.fromstring(original.read("xl/_rels/workbook.xml.rels"))
        targets = {node.attrib["Id"]: node.attrib["Target"] for node in rels_root.findall(f"{{{_PKG_REL_NS}}}Relationship")}
        member = None
        for node in workbook_root.findall(f"{{{_MAIN_NS}}}sheets/{{{_MAIN_NS}}}sheet"):
            if node.attrib["name"] == sheet_name:
                target_member = targets[node.attrib[f"{{{_REL_NS}}}id"]]
                member = target_member.lstrip("/") if target_member.startswith("/") else target_member
                if not member.startswith("xl/"):
                    member = f"xl/{member}"
                break
        if member is None:
            raise AssertionError(f"worksheet not found: {sheet_name}")
        with ZipFile(output, "w", ZIP_DEFLATED) as rewritten:
            for info in original.infolist():
                data = original.read(info.filename)
                if info.filename == member:
                    root = ElementTree.fromstring(data)
                    cell = root.find(f".//{{{_MAIN_NS}}}c[@r='{cell_reference}']")
                    value = cell.find(f"{{{_MAIN_NS}}}v") if cell is not None else None
                    if value is None:
                        raise AssertionError(f"numeric cell not found: {sheet_name}!{cell_reference}")
                    value.text = lexical_value
                    data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
                rewritten.writestr(info, data)
    source.write_bytes(output.getvalue())


__all__ = ["appendix_b_context", "rewrite_numeric_lexeme", "write_valid_appendix_b_workbook"]
