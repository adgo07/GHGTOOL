"""脱硫实际Qt/Excel/SQLite/正式Record/冻结Word整链。"""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QMessageBox
import tests.test_b_material_draft_scope as material_fixtures
from packages.application.project_workspaces import ProjectWorkspaceService
from packages.persistence import SQLiteProjectWorkspaceRepository, SQLiteRecordRepository
from packages.excel.templates import ExcelTemplateService
from packages.ui.view_models import AppRoute
from packages.standards.carbon_material import SOURCE_FGD

class FGDChainTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])
    page = material_fixtures.MaterialDraftScopeTests.page

    @staticmethod
    def prepare(gui):
        gui.enterprise_name.setText("脱硫整链")
        gui.period_year.setValue(2026)
        gui.boundary_confirmed.setChecked(True)
        for status in gui._source_statuses.values():
            status.setCurrentIndex(status.findData("NOT_INVOLVED"))
        gui._source_statuses[SOURCE_FGD].setCurrentIndex(gui._source_statuses[SOURCE_FGD].findData("INVOLVED"))

    @staticmethod
    def select(gui, component, name):
        component["carbonate_type"].setEditText(name)
        component["carbonate_type"].confirm_name(force=True)

    @staticmethod
    def measured(component, field, value):
        combo = component[f"{field}_source"]
        combo.setCurrentIndex(combo.findData("MEASURED"))
        component[field].setText(value)

    def test_gui_blank_rows_source_roundtrip_and_sqlite_scope(self):
        with self.page() as (gui, projects, root, shell, records):
            self.prepare(gui)
            batch = gui._process_rows["fgd"][0]
            batch["reagent_name"].setText("同名批次")
            batch["group_amount"].setText("10")
            first = batch["components"][0]
            self.select(gui, first, "CaCO3")
            blank = gui._add_fgd_component(batch)
            self.assertTrue(gui._fgd_ratio_default_allowed(first))
            self.assertEqual(len(gui._input(increment=False).fgd_units[0].components), 1)
            self.select(gui, blank, "MgCO3")
            self.assertIsNone(first["i_source"].currentData())
            self.assertEqual(first["i"].text(), "")
            self.measured(first, "i", "40")
            self.measured(first, "tr", "80")
            self.measured(blank, "i", "30")
            self.measured(blank, "tr", "90")
            value = gui._input(increment=False)
            outcome = gui.calculator.calculate(value)
            self.assertTrue(outcome.successful, outcome.problems)
            self.assertEqual(outcome.result.total_amount, Decimal("2.8174"))
            first["tr_source"].setCurrentIndex(first["tr_source"].findData("STANDARD_DEFAULT"))
            first["tr_source"].setCurrentIndex(first["tr_source"].findData("MEASURED"))
            self.assertEqual(first["tr"].text(), "80")
            # A second batch has the same local component ID and different data.
            other = gui._add_process_row("fgd", instance_id="other-batch", component_ids=(first["component_id"],))
            other["reagent_name"].setText("同名批次")
            other["group_amount"].setText("20")
            self.select(gui, other["components"][0], "CaCO3")
            self.measured(other["components"][0], "tr", "70")
            state = gui._capture_form_state()
            workspace = projects.new_workspace("脱硫草稿")
            projects.save(replace(workspace, units=(replace(workspace.units[0], form_state=state),)))
            restored = ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/"projects.sqlite")).get(workspace.project_id)
            gui._restore_form_state(restored.units[0].form_state)
            first = gui._process_rows["fgd"][0]["components"][0]
            other = next(row for row in gui._process_rows["fgd"] if row["instance_id"] == "other-batch")["components"][0]
            self.assertEqual((first["tr"].text(), other["tr"].text()), ("80", "70"))
            for component, expected in ((first,"80"),(other,"70")):
                component["tr_source"].setCurrentIndex(component["tr_source"].findData("STANDARD_DEFAULT"))
                component["tr_source"].setCurrentIndex(component["tr_source"].findData("MEASURED"))
                self.assertEqual(component["tr"].text(), expected)
            unit = gui._process_rows["fgd"][0]
            gui._remove_fgd_component(unit, unit["components"][1]["component_id"])
            self.assertTrue(gui._fgd_ratio_default_allowed(first))
            self.assertEqual((first["i_source"].currentData(),first["i"].text()),("MEASURED","40"))
            self.assertEqual(len(records.list_all()), 0)

    def test_excel_gui_parity_project_reopen_formal_record_and_frozen_word(self):
        with self.page() as (gui, projects, root, shell, records):
            self.prepare(gui)
            batch = gui._process_rows["fgd"][0]
            batch["reagent_name"].setText("批次一")
            batch["group_amount"].setText("10")
            first = batch["components"][0]
            self.select(gui, first, "CaCO3")
            self.measured(first,"i","40")
            self.measured(first,"tr","80")
            second=gui._add_fgd_component(batch)
            self.select(gui,second,"MgCO3")
            self.measured(second,"i","30")
            self.measured(second,"tr","90")
            expected=gui.calculator.calculate(gui._input(increment=False))
            self.assertTrue(expected.successful,expected.problems)
            workbook=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
            ws=workbook["B.7"]
            for cell,value in {"A3":"批次一","B3":10,"D3":40,"F3":80,"J3":"缺省值","K3":"实测值","L3":"实测值",
                "D4":30,"F4":90,"J4":"缺省值","K4":"实测值","L4":"实测值","M3":"设备A"}.items():ws[cell]=value
            path=root/"fgd.xlsx";workbook.save(path);workbook.close()
            page=shell.pages[AppRoute.EXCEL_IMPORT]
            context=page.excel_adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name="脱硫整链")
            preview=page.excel_adapter.import_preview(path,context)
            self.assertTrue(preview.units[0].can_calculate,preview.units[0].errors)
            self.assertEqual(len(records.list_all()),0)
            canonical=preview.units[0].input_value
            self.assertEqual(gui.calculator.calculate(canonical).result.total_amount,expected.result.total_amount)
            self.assertEqual(expected.result.total_amount,Decimal("2.8174"))
            # Persist actual Canonical and ingress evidence then reopen through SQLite.
            workspace=projects.new_workspace("脱硫表格")
            from packages.application.carbon_accounting import CarbonAccountingUseCase
            from packages.application.canonical_input_codec import encode_canonical_input,decode_canonical_input
            self.assertEqual(decode_canonical_input(encode_canonical_input(canonical)),canonical)
            provenance={"source":"EXCEL_APPENDIX_B","workbook": {"template_version":preview.provenance.template_version,
                "template_master_sha256":preview.provenance.template_master_sha256,"sha256":preview.provenance.workbook_sha256}}
            projects.save(replace(workspace,units=(replace(workspace.units[0],canonical_input=canonical,ingress_provenance=provenance),)))
            reopened=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/"projects.sqlite")).get(workspace.project_id)
            usecase=CarbonAccountingUseCase(gui.calculator,records)
            outcome=usecase.calculate(reopened.units[0].canonical_input,ingress_provenance=reopened.units[0].ingress_provenance)
            self.assertTrue(outcome.successful,outcome.problems)
            self.assertEqual(len(records.list_all()),1)
            record_id=outcome.record.record_id
            snapshots=SQLiteRecordRepository(root/"records.sqlite")
            from packages.application.reporting.model import build_report_model
            with patch.object(type(gui.calculator),"calculate",side_effect=AssertionError("不得重算冻结记录")), patch.object(gui._parameter_resolver,"resolve",side_effect=AssertionError("不得回查当前因子")):
                model=build_report_model(snapshots.get(record_id),raw_input_snapshot=snapshots.get_raw_input_snapshot(record_id),
                    trace_snapshot=snapshots.get_trace_snapshot(record_id),reporting_snapshot=snapshots.get_reporting_snapshot(record_id),
                    provenance_snapshot=snapshots.get_provenance_snapshot(record_id),report_qualification=snapshots.get_report_qualification(record_id))
                self.assertEqual(model.input_template_version,"2.6.0")
                table=next(table for section in model.sections for table in section.tables if table.table_id=="b7")
                self.assertEqual(len(table.rows),2)
                self.assertIn("实测",table.rows[0].cells[3].source or "")
                from packages.infrastructure.reporting import render_report_docx
                from docx import Document
                document=Document(BytesIO(render_report_docx(model)))
                text=" ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                self.assertIn("批次一",text)
                self.assertIn("2.82",text)

    def test_excel_measured_blank_zero_and_multi_default_are_distinct(self):
        with self.page() as (gui,projects,root,shell,records):
            self.prepare(gui);page=shell.pages[AppRoute.EXCEL_IMPORT]
            context=page.excel_adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name="边界")
            for name, source, amount,valid in (("blank","实测值",None,False),("zero","实测值",0,True),("default","缺省值",None,True)):
                with self.subTest(name=name):
                    workbook=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()));ws=workbook["B.7"]
                    for cell,value in {"A3":"批次","B3":10,"D3":amount,"J3":"缺省值","K3":source,"L3":"缺省值"}.items():ws[cell]=value
                    path=root/(name+".xlsx");workbook.save(path);workbook.close()
                    preview=page.excel_adapter.import_preview(path,context)
                    self.assertEqual(preview.units[0].can_calculate,valid,preview.units[0].errors)
                    if valid:
                        expected=Decimal("0") if name=="zero" else Decimal("3.96")
                        self.assertEqual(gui.calculator.calculate(preview.units[0].input_value).result.total_amount,expected)
            self.assertEqual(len(records.list_all()),0)

    def test_gui_two_batches_same_facility_and_blank_source_never_falls_back(self):
        with self.page() as (gui,projects,root,shell,records):
            self.prepare(gui)
            first=gui._process_rows["fgd"][0]
            first["reagent_name"].setText("同名")
            first["group_amount"].setText("10")
            self.select(gui,first["components"][0],"CaCO3")
            second=gui._add_process_row("fgd")
            second["reagent_name"].setText("同名")
            second["group_amount"].setText("20")
            second["facility_selector"].setCurrentIndex(second["facility_selector"].findData(first["facility_id"]))
            self.select(gui,second["components"][0],"CaCO3")
            from packages.application.carbon_accounting import CarbonAccountingUseCase
            usecase=CarbonAccountingUseCase(gui.calculator,records)
            value=gui._input(increment=False)
            self.assertEqual(value.fgd_units[0].facility_id,value.fgd_units[1].facility_id)
            self.assertNotEqual(value.fgd_units[0].batch_id,value.fgd_units[1].batch_id)
            result=usecase.calculate(value)
            self.assertTrue(result.successful,result.problems)
            self.assertEqual(result.record.calculation_result.total_amount,Decimal("11.88"))
            component=second["components"][0]
            self.measured(component,"tr","")
            rejected=usecase.calculate(gui._input(increment=False))
            self.assertFalse(rejected.successful)
            self.assertEqual(len(records.list_all()),1)
            component["tr"].setText("0")
            zero=usecase.calculate(gui._input(increment=False))
            self.assertTrue(zero.successful,zero.problems)
            self.assertEqual(zero.record.calculation_result.total_amount,Decimal("3.96"))
            self.assertEqual(len(records.list_all()),2)
            self.assertNotEqual(result.record.record_id,zero.record.record_id)

    def test_old_gui_upgrade_cancel_then_confirm_preserves_original_and_values(self):
        with self.page() as (gui,projects,root,shell,records):
            self.prepare(gui)
            batch=gui._process_rows["fgd"][0]
            batch["reagent_name"].setText("历史批次")
            batch["group_amount"].setText("10")
            self.select(gui,batch["components"][0],"CaCO3")
            self.measured(batch["components"][0],"tr","75")
            state=gui._capture_form_state();state.pop("fgd_identities")
            gui._restore_form_state(state)
            before=gui._capture_form_state()
            self.assertEqual(gui._process_rows["fgd"][0]["execution_version"],"FGD_V1")
            self.assertIsNone(gui._process_rows["fgd"][0]["batch_id"])
            with patch.object(QMessageBox,"question",return_value=QMessageBox.StandardButton.No):gui._upgrade_fgd_rules()
            self.assertEqual(gui._capture_form_state(),before)
            with patch.object(QMessageBox,"question",return_value=QMessageBox.StandardButton.Yes):gui._upgrade_fgd_rules()
            after=gui._capture_form_state()
            self.assertEqual(after["fgd_upgrade_original_state"],{key:value for key,value in before.items() if key!="fgd_upgrade_original_state"})
            self.assertEqual(gui._process_rows["fgd"][0]["execution_version"],"FGD_V2")
            self.assertEqual(gui._process_rows["fgd"][0]["components"][0]["tr"].text(),"75")
            self.assertEqual(len(records.list_all()),0)

    def test_controlled_v8_projection_preserves_v7_and_other_parts(self):
        from scripts.project_appendix_b_v8 import generated_bytes,SOURCE,TARGET,SOURCE_SHA
        from hashlib import sha256
        from zipfile import ZipFile
        self.assertEqual(sha256(SOURCE.read_bytes()).hexdigest(),SOURCE_SHA)
        self.assertEqual(TARGET.read_bytes(),generated_bytes())
        with ZipFile(SOURCE) as old,ZipFile(TARGET) as new:
            self.assertEqual(old.namelist(),new.namelist())
            changed=[name for name in old.namelist() if old.read(name)!=new.read(name)]
            self.assertEqual(changed,["xl/worksheets/sheet7.xml"])
        workbook=load_workbook(TARGET)
        try:
            self.assertEqual(ExcelTemplateService.default().validate_structure(workbook),"2.6.0")
            self.assertEqual(workbook["B.7"]["K2"].value,"含量来源")
            self.assertEqual(workbook["B.7"]["L2"].value,"转化率来源")
            workbook["B.7"]["L2"]=None
            with self.assertRaisesRegex(ValueError,"版本冲突"):ExcelTemplateService.default().validate_structure(workbook)
        finally:workbook.close()

    def test_explicit_source_on_unfinished_row_is_not_hidden(self):
        with self.page() as (gui,projects,root,shell,records):
            self.prepare(gui)
            batch=gui._process_rows["fgd"][0]
            batch["reagent_name"].setText("批次")
            batch["group_amount"].setText("10")
            self.select(gui,batch["components"][0],"CaCO3")
            unfinished=gui._add_fgd_component(batch)
            self.assertFalse(gui._fgd_component_active(unfinished))
            self.measured(unfinished,"tr","")
            self.assertTrue(gui._fgd_component_active(unfinished))
            self.assertFalse(gui._fgd_ratio_default_allowed(batch["components"][0]))
            from packages.application.carbon_accounting import CarbonAccountingUseCase
            result=CarbonAccountingUseCase(gui.calculator,records).calculate(gui._input(increment=False))
            self.assertFalse(result.successful)
            self.assertEqual(len(records.list_all()),0)

    def test_ambiguous_legacy_projections_are_all_retained_without_cross_batch_paste(self):
        with self.page() as (gui,projects,root,shell,records):
            self.prepare(gui)
            first=gui._process_rows["fgd"][0]
            first["reagent_name"].setText("第一批")
            first["group_amount"].setText("10")
            a=first["components"][0]
            self.select(gui,a,"CaCO3");self.measured(a,"i","40")
            second=gui._add_process_row("fgd",component_ids=(a["component_id"],))
            second["reagent_name"].setText("第二批");second["group_amount"].setText("20")
            b=second["components"][0]
            self.select(gui,b,"MgCO3");self.measured(b,"i","80")
            state=gui._capture_form_state();key=a["component_id"]
            old={
                "fgd_carbonate_texts":{key:"旧碳酸盐"},
                "fgd_carbonate_confirmations":{key:{"confirmed_name":"旧碳酸盐","legacy_custom_identity":True}},
                "fgd_factor_provenance":{key:{"mode":"MEASURED","factor_id":None}},
                "fgd_ratio_modes":{key:{"i":"MEASURED","tr":"MEASURED"}},
            }
            state.update(old);state.pop("fgd_source_drafts")
            gui._restore_form_state(state)
            restored=gui._capture_form_state()
            retained=restored["fgd_legacy_unresolved"]["projections"]
            for projection in old.values():self.assertIn(projection,retained)
            self.assertEqual([row["components"][0]["i"].text() for row in gui._process_rows["fgd"]],["40","80"])
            for row in gui._process_rows["fgd"]:
                self.assertIn("未自动采用",row["components"][0]["_widget"].toolTip())
            workspace=projects.new_workspace("旧歧义")
            projects.save(replace(workspace,units=(replace(workspace.units[0],form_state=restored),)))
            reopened=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/"projects.sqlite")).get(workspace.project_id)
            self.assertEqual(reopened.units[0].form_state["fgd_legacy_unresolved"],restored["fgd_legacy_unresolved"])
            self.assertEqual(len(records.list_all()),0)

    def test_old_canonical_project_upgrade_copies_sqlite_without_rewriting_original(self):
        with self.page() as (gui,projects,root,shell,records):
            self.prepare(gui)
            batch=gui._process_rows["fgd"][0]
            batch["reagent_name"].setText("旧批次");batch["group_amount"].setText("10")
            first=batch["components"][0]
            self.select(gui,first,"CaCO3");self.measured(first,"i","90");self.measured(first,"tr","75")
            second=gui._add_fgd_component(batch)
            self.select(gui,second,"MgCO3");self.measured(second,"i","10")
            current=gui._input(increment=False)
            from packages.standards.carbon_material import InputValue
            old_batch=replace(current.fgd_units[0],facility_id=None,facility_name=None,batch_id=None,execution_version="FGD_V1",
                components=tuple(replace(component,component_id=None,amount=InputValue("10","t",source_reference=f"旧组分资料{i}")) for i,component in enumerate(current.fgd_units[0].components)))
            old=replace(current,fgd=old_batch,fgd_units=(old_batch,))
            workspace=projects.new_workspace("旧Canonical脱硫")
            workspace=replace(workspace,units=(replace(workspace.units[0],canonical_input=old),))
            projects.save(workspace)
            from packages.application.carbon_material_implementation import CarbonMaterialExcel
            self.assertTrue(CarbonMaterialExcel.can_upgrade_workspace(workspace))
            upgraded=CarbonMaterialExcel.upgrade_workspace(workspace)
            projects.save(upgraded)
            reopened=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/"projects.sqlite"))
            self.assertEqual(reopened.get(workspace.project_id).units[0].canonical_input,old)
            new=reopened.get(upgraded.project_id).units[0].canonical_input
            self.assertNotEqual(upgraded.project_id,workspace.project_id)
            self.assertNotEqual(upgraded.units[0].unit_id,workspace.units[0].unit_id)
            self.assertEqual(new.fgd_units[0].execution_version,"FGD_V2")
            self.assertTrue(all(component.component_id for component in new.fgd_units[0].components))
            self.assertEqual(gui.calculator.calculate(new).result.total_amount,Decimal("3.492"))
            self.assertEqual(len(records.list_all()),0)

    def test_reimport_keeps_ingress_identity_but_saved_projects_have_independent_unit_ids(self):
        with self.page() as (gui,projects,root,shell,records):
            self.prepare(gui)
            workbook=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
            for cell,value in {"A3":"重复导入批次","B3":10,"J3":"缺省值","K3":"缺省值","L3":"缺省值"}.items():workbook["B.7"][cell]=value
            path=root/"同一填表文件.xlsx";workbook.save(path);workbook.close()
            page=shell.pages[AppRoute.EXCEL_IMPORT]
            page.import_enterprise_name.setText("同一企业")
            page.import_period_type.setCurrentIndex(page.import_period_type.findData("ANNUAL"))
            page.import_period_start.setText("2026-01-01")
            page.import_period_end.setText("2026-12-31")
            page.import_boundary_confirmed.setChecked(True)
            saved=[];parsed=[]
            for name in ("独立项目甲","独立项目乙"):
                page._preview_workbook(path)
                preview=page._last_preview
                self.assertIsNotNone(preview)
                self.assertTrue(preview.units[0].can_calculate,preview.units[0].errors)
                parsed.append(preview.units[0])
                page.project_name.setText(name)
                with patch.object(QMessageBox,"warning",side_effect=AssertionError("独立项目保存不应碰撞")):
                    page._save_preview_as_project()
                saved.append(page._workspace)
            self.assertEqual(parsed[0].unit_id,parsed[1].unit_id)
            self.assertEqual(parsed[0].input_value,parsed[1].input_value)
            self.assertNotEqual(saved[0].project_id,saved[1].project_id)
            self.assertNotEqual(saved[0].units[0].unit_id,saved[1].units[0].unit_id)
            for workspace in saved:
                self.assertNotEqual(workspace.units[0].unit_id,parsed[0].unit_id)
                self.assertEqual(workspace.units[0].ingress_provenance["accounting_unit"]["importer_unit_id"],parsed[0].unit_id)
            reopened=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/"projects.sqlite"))
            self.assertEqual(len(reopened.list_all()),2)
            for workspace in saved:
                self.assertEqual(reopened.get(workspace.project_id).units[0].canonical_input,parsed[0].input_value)
            self.assertEqual(len(records.list_all()),0)

    def test_legacy_codec_groups_enforce_amount_without_rewriting_frozen_record(self):
        import json
        from packages.application.carbon_material_input_codec import encode_canonical_input, decode_canonical_input
        from packages.application.carbon_accounting import CarbonAccountingUseCase
        from packages.application.reporting.model import build_report_model
        from packages.infrastructure.reporting import render_report_docx
        from packages.standards.carbon_material import InputValue
        from docx import Document
        with self.page() as (gui,projects,root,shell,records):
            self.prepare(gui)
            batch=gui._process_rows["fgd"][0]
            batch["reagent_name"].setText("旧分组");batch["group_amount"].setText("10")
            first=batch["components"][0]
            self.select(gui,first,"CaCO3");self.measured(first,"i","40")
            second=gui._add_fgd_component(batch)
            self.select(gui,second,"MgCO3");self.measured(second,"i","30")
            current=gui._input(increment=False)
            def legacy(value,version):
                document=json.loads(encode_canonical_input(value));document["version"]=version
                def strip(node):
                    if isinstance(node,dict):
                        if node.get("type")=="fgd_input":
                            for key in ("facility_id","facility_name","batch_id","execution_version"):
                                node["fields"].pop(key,None)
                        if node.get("type")=="carbonate_component":node["fields"].pop("component_id",None)
                        for child in node.values():strip(child)
                    elif isinstance(node,list):
                        for child in node:strip(child)
                strip(document)
                return decode_canonical_input(json.dumps(document,ensure_ascii=False))
            usecase=CarbonAccountingUseCase(gui.calculator,records)
            for version in (1,2):
                with self.subTest(codec=version):
                    restored=legacy(current,version)
                    self.assertEqual(restored.fgd_units[0].execution_version,"FGD_V1")
                    self.assertIsNone(restored.fgd_units[0].batch_id)
                    components=restored.fgd_units[0].components
                    bad_batch=replace(restored.fgd_units[0],components=(components[0],replace(components[1],amount=InputValue("20","t"))))
                    bad=replace(restored,fgd=bad_batch,fgd_units=(bad_batch,))
                    rejected=usecase.calculate(bad)
                    self.assertFalse(rejected.successful)
                    self.assertTrue(any(p.code=="CAR-VAL-FGD-BATCH-AMOUNT" for p in rejected.problems))
                    self.assertEqual(len(records.list_all()),0)
            restored=legacy(current,2)
            components=restored.fgd_units[0].components
            # Different sources and units represent the same whole-reagent amount.
            equal_batch=replace(restored.fgd_units[0],components=(components[0],replace(components[1],amount=InputValue("10000","kg",source_reference="另一资料"))))
            equal=replace(restored,fgd=equal_batch,fgd_units=(equal_batch,))
            outcome=usecase.calculate(equal)
            self.assertTrue(outcome.successful,outcome.problems)
            self.assertEqual(outcome.result.total_amount,Decimal("3.326"))
            stored=SQLiteRecordRepository(root/"records.sqlite")
            record_id=outcome.record.record_id
            before=stored.get_raw_input_snapshot(record_id)
            def model():
                return build_report_model(stored.get(record_id),raw_input_snapshot=stored.get_raw_input_snapshot(record_id),
                    trace_snapshot=stored.get_trace_snapshot(record_id),reporting_snapshot=stored.get_reporting_snapshot(record_id),
                    provenance_snapshot=stored.get_provenance_snapshot(record_id),report_qualification=stored.get_report_qualification(record_id))
            frozen=model()
            self.assertTrue(gui.calculator.calculate(restored).successful)
            # Distinct old groups retain independent quantities; no guessed merge.
            split=(replace(restored.fgd_units[0],instance_id="old-a",components=(components[0],)),
                replace(restored.fgd_units[0],instance_id="old-b",components=(replace(components[1],amount=InputValue("20","t")),)))
            split_outcome=gui.calculator.calculate(replace(restored,fgd=split[0],fgd_units=split))
            self.assertTrue(split_outcome.successful,split_outcome.problems)
            self.assertEqual(split_outcome.result.total_amount,Decimal("4.892"))
            self.assertFalse(usecase.calculate(bad).successful)
            self.assertEqual(len(records.list_all()),1)
            with patch.object(type(gui.calculator),"calculate",side_effect=AssertionError("旧记录不得重算")), patch.object(gui._parameter_resolver,"resolve",side_effect=AssertionError("旧记录不得回查")):
                self.assertEqual(stored.get_raw_input_snapshot(record_id),before)
                self.assertEqual(model(),frozen)
                document=Document(BytesIO(render_report_docx(model())))
                text=" ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                self.assertIn("旧分组",text)
                self.assertIn("3.33",text)
