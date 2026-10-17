"""Independent asset coverage checks for A; no production adoption."""
import ast
import hashlib
import html
import json
import os
import posixpath
import re
import subprocess
import sys
import tempfile
import unittest
import urllib.parse
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from decimal import Decimal


ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/"data-source/carbon_accounting/business-rules"
REVIEW_ARTIFACT_ROOT = Path(os.environ.get(
    "GHG_BUSINESS_RULES_REVIEW_ARTIFACT_DIR",
    ROOT / "docs/architecture/business-rules/evidence",
))
PLAN=ROOT/"docs/architecture/business-rules/APPROVED_IMPLEMENTATION_V1_2.md"

class CandidateAssetCoverageTests(unittest.TestCase):
    def documents(self):
        paths=(SOURCE/"common/gbt_32150_2025.json",SOURCE/"standards/gbt_32151_34_2024.json",SOURCE/"software_policy.json")
        self.assertTrue(all(p.is_file() for p in paths), "三份候选Canonical资产必须齐全")
        return [json.loads(p.read_text(encoding="utf-8"),parse_float=Decimal) for p in paths]

    def test_definition_text_hash_normalizes_crlf_and_rejects_changed_content(self):
        import shutil
        import tempfile

        from scripts.business_rules_review_hashes import (
            DEFINITION_ASSET_HASH_POLICY,
            EXPECTED_DEFINITION_ASSETS,
            definition_text_sha256,
            verify_definition_assets,
        )

        raw = (ROOT / EXPECTED_DEFINITION_ASSETS[0]).read_bytes()
        lf_bytes = raw.replace(b"\r\n", b"\n")
        crlf_bytes = lf_bytes.replace(b"\n", b"\r\n")
        self.assertEqual(definition_text_sha256(lf_bytes), definition_text_sha256(crlf_bytes))
        self.assertNotEqual(definition_text_sha256(lf_bytes), definition_text_sha256(lf_bytes + b" "))

        checks = {
            "definition_assets_hash_policy": DEFINITION_ASSET_HASH_POLICY,
            "definition_assets_sha256": {
                relative: definition_text_sha256((ROOT / relative).read_bytes())
                for relative in EXPECTED_DEFINITION_ASSETS
            },
        }
        with tempfile.TemporaryDirectory() as directory:
            changed_root = Path(directory)
            for relative in EXPECTED_DEFINITION_ASSETS:
                source = ROOT / relative
                destination = changed_root / relative
                destination.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(source, destination)
            self.assertEqual(verify_definition_assets(changed_root, checks), len(EXPECTED_DEFINITION_ASSETS))
            changed = changed_root / EXPECTED_DEFINITION_ASSETS[0]
            changed_definition = json.loads(changed.read_text(encoding="utf-8"))
            changed_definition["revision"] += ".modified"
            changed.write_text(json.dumps(changed_definition, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "SHA-256 mismatch"):
                verify_definition_assets(changed_root, checks)

    def test_registered_definition_asset_hashes_match_checked_in_manifest(self):
        from scripts.business_rules_review_hashes import (
            DEFINITION_ASSET_HASH_POLICY,
            EXPECTED_DEFINITION_ASSETS,
            verify_definition_assets,
        )

        checks_path = REVIEW_ARTIFACT_ROOT / "review-artifact-checks.json"
        checks = json.loads(checks_path.read_text(encoding="utf-8"))
        self.assertEqual(checks["definition_assets_hash_policy"], DEFINITION_ASSET_HASH_POLICY)
        self.assertEqual(set(checks["definition_assets_sha256"]), set(EXPECTED_DEFINITION_ASSETS))
        self.assertEqual(verify_definition_assets(ROOT, checks), len(EXPECTED_DEFINITION_ASSETS))

    def test_review_workbook_hash_uses_raw_xlsx_bytes(self):
        checks = json.loads(
            (REVIEW_ARTIFACT_ROOT / "review-artifact-checks.json")
            .read_text(encoding="utf-8")
        )
        workbook_path = REVIEW_ARTIFACT_ROOT / "业务规则审阅表.xlsx"
        self.assertEqual(checks["workbook_sha256"], hashlib.sha256(workbook_path.read_bytes()).hexdigest())

    def test_approved_plan_bytes_remain_identical(self):
        self.assertEqual(hashlib.sha256(PLAN.read_bytes()).hexdigest(),"98abf84ea7f7af3887d60c3c1e128c272d73ccbf73892d76581ad1771bfad876")

    def test_business_policies_link_existing_mapping_without_internal_mirrors(self):
        documents=self.documents()
        fields=[f for d in documents for f in d['fields']]
        self.assertTrue(fields)
        self.assertEqual(documents[2]['fields'],[], '软件内部属性不是长期业务政策')
        mapping=(ROOT/'specs/carbon_accounting/GB_T_32151_34_2024_MAPPING.md').read_text(encoding='utf-8')
        for field in fields:
            with self.subTest(field=field['id']):
                self.assertFalse(field['id'].startswith('domain.'))
                self.assertTrue(field['mapping_refs'])
                for reference in field['mapping_refs']:self.assertIn('`'+reference+'`',mapping)
                self.assertTrue(all(p['binding']=='business:'+field['id'] for p in field['channel_projections']))
                self.assertTrue(field['basis'])

    def test_default_review_is_business_summary_and_matrix_is_developer_only(self):
        from packages.application.business_rules_candidate import load_bundle
        paths=(SOURCE/'common/gbt_32150_2025.json',SOURCE/'standards/gbt_32151_34_2024.json',SOURCE/'software_policy.json')
        bundle=load_bundle(paths)
        review=bundle.review_rows()
        defined={f['id'] for d in self.documents() for f in d['fields']}
        self.assertEqual({f['field_id'] for f in review['fields']},defined)
        self.assertEqual(len(review['fields']),len(defined))
        self.assertGreater(len(bundle.technical_review_rows()['fields']),len(review['fields']))
        for row in review['fields']:
            self.assertFalse({'channels','consumers','cases','condition_id','evaluation_status'} & row.keys())
            self.assertTrue(row['mapping_refs'])

    def test_candidate_has_no_production_import(self):
        for folder in (ROOT/"apps",ROOT/"packages"):
            for p in folder.rglob("*.py"):
                if p.name=="business_rules_candidate.py":
                    continue
                tree=ast.parse(p.read_text(encoding="utf-8"),filename=str(p))
                for node in ast.walk(tree):
                    if isinstance(node,ast.ImportFrom):
                        self.assertNotIn("business_rules_candidate",node.module or "",str(p))
                    elif isinstance(node,ast.Import):
                        self.assertTrue(all("business_rules_candidate" not in a.name for a in node.names),str(p))

    def test_actual_g05_snapshots_and_existing_resolver_remain_identical(self):
        from packages.application.business_rules_candidate import load_bundle
        from packages.core.parameter_resolution import default_g05_rules
        from packages.core import EffectiveRuleResolver, RuleContext
        paths=(SOURCE/"common/gbt_32150_2025.json",SOURCE/"standards/gbt_32151_34_2024.json",SOURCE/"software_policy.json")
        bundle=load_bundle(paths)
        common, industry=default_g05_rules()
        loaded={key:rule for package in bundle.packages for key,rule in package.rules.items()}
        for original in (*common,*industry):
            with self.subTest(rule=original.rule_id):
                self.assertEqual(loaded[original.rule_id],original)
        context=RuleContext(standard_id="gbt_32151_34_2024")
        expected=EffectiveRuleResolver().resolve(common,industry,context)
        self.assertEqual(bundle.resolve_rules(context),expected,
                         "已由行业覆写解决的共同冲突不能在候选层再次全局阻断")

    def test_default_parameter_and_verified_binding_ids_are_real_catalog_assets(self):
        catalog=json.loads((ROOT/"data-source/carbon_accounting/catalog.json").read_text(encoding="utf-8"))
        parameters={p["parameter_id"] for p in catalog["parameters"]}
        assets={a["asset_id"]:a for a in catalog["reference_data_assets"]}
        bindings={b["binding_id"]:b for b in catalog["reference_data_bindings"]}
        for doc in self.documents():
            states=[f["default"] for f in doc["fields"]]
            for policy in doc["policies"]:
                for branch in policy["branches"]:
                    states.append(branch["value"])
                    tags=branch["when"]["tags"]
                    if tags.get("binding_status")=="VERIFIED":
                        with self.subTest(policy=policy["id"],branch=branch["id"]):
                            self.assertIn(tags.get("binding_id"),bindings,"不能用非空伪绑定代替资格")
                            binding=bindings[tags["binding_id"]]
                            self.assertEqual(binding["review_status"],"VERIFIED")
                            # Existing registry semantics: empty is no standard filter, not proof of applicability.
                            # Formal eligibility is separately checked by adopted rules/Resolver and dependency states.
                            if binding["applicable_standard_ids"]:
                                self.assertIn("gbt_32151_34_2024",binding["applicable_standard_ids"])
                            self.assertEqual(assets[binding["asset_id"]]["parameter_id"],tags["parameter_id"])
            for state in states:
                if state.get("state")=="BINDING":
                    self.assertIn(state["parameter_id"],parameters,"默认值必须引用真实参数身份")

    def test_review_workbook_matches_current_canonical_projection(self):
        from packages.application.business_rules_candidate import load_bundle, _json_safe
        from openpyxl import load_workbook
        paths=(SOURCE/"common/gbt_32150_2025.json",SOURCE/"standards/gbt_32151_34_2024.json",SOURCE/"software_policy.json")
        review=_json_safe(load_bundle(paths).review_rows())
        expected_hash=hashlib.sha256((json.dumps(review,ensure_ascii=False,indent=2)+"\n").encode("utf-8")).hexdigest()
        path=REVIEW_ARTIFACT_ROOT/"业务规则审阅表.xlsx"
        self.assertTrue(path.is_file(),"必须提供实际生成的三页审阅表")
        workbook=load_workbook(path,read_only=True,data_only=False)
        try:
            self.assertEqual(workbook.sheetnames,["业务规则","操作与影响","本次变化"])
            for name,kind in zip(workbook.sheetnames,("fields","operations","changes")):
                self.assertTrue(review[kind],f"{name}不能只建空页")
                sheet=workbook[name]
                self.assertEqual(sheet["A1"].value,"审阅数据SHA256："+expected_hash,"审阅表不得落后于候选定义")
                headers=[c.value for c in sheet[5]]
                self.assertEqual(headers.count("稳定审阅键"),1)
                self.assertTrue({"审阅意见","处理状态","审阅编号"}<=set(headers))
                key_column=headers.index("稳定审阅键")
                data=[row for row in sheet.iter_rows(min_row=6,values_only=True) if row[key_column]]
                keys=[row[key_column] for row in data]
                expected={r["review_key"] for r in review[kind]}
                self.assertTrue(expected<={row[headers.index("当前投影键")] for row in data})
                extra={row[key_column] for row in data if row[headers.index("当前投影键")] not in expected}
                if extra:
                    self.assertEqual(kind,"changes","旧修订意见只应进入明确的保留变化行")
                    category=headers.index("变化类型")
                    feedback=[headers.index(h) for h in ("审阅意见","处理状态","审阅编号")]
                    for row in data:
                        if row[key_column] in extra:
                            self.assertEqual(row[category],"未完成审阅意见保留")
                            self.assertTrue(any(row[i] for i in feedback))
                self.assertEqual(len(keys),len(set(keys)))
        finally:
            workbook.close()

    def test_generated_locking_distinguishes_catalog_and_unbound_measurements(self):
        from openpyxl import load_workbook
        book=load_workbook(REVIEW_ARTIFACT_ROOT/"业务规则审阅表.xlsx",read_only=True)
        try:
            sheet=book["业务规则"]
            rows={row[0]:row for row in sheet.iter_rows(min_row=6,values_only=True)}
            for name in ("燃料碳氧化率FOx","脱硫碳酸盐含量I","脱硫转化率TR"):
                with self.subTest(name=name):
                    blocks=rows[name][6].split("\n\n")
                    selectable=next(b for b in blocks if b.startswith("数值可填写/选择，来源可填写/选择"))
                    locked=next(b for b in blocks if b.startswith("数值可填写/选择，来源只读"))
                    self.assertIn("绑定资格=来源登记已核验",selectable)
                    self.assertIn("当前来源=实测值",selectable)
                    self.assertIn("须匹配精确登记组合",selectable)
                    self.assertIn("绑定资格=无适用绑定",locked)
                    self.assertIn("当前来源=实测值",locked)
                    self.assertNotIn("绑定资格=来源登记已核验",locked)
                    self.assertEqual(rows[name][2],"比值（0～1）；界面/Excel用百分数（%）")
            carbonate=rows["脱硫碳酸盐含量I"][6].split("\n\n")
            self.assertIn("本批有效组分数=1",next(b for b in carbonate if b.startswith("数值可填写/选择，来源可填写/选择")))
            locked=next(b for b in carbonate if b.startswith("数值可填写/选择，来源只读"))
            self.assertIn("本批有效组分数=1",locked)
            self.assertIn("本批有效组分数≥2",locked)
        finally:
            book.close()

    def test_generated_default_and_large_condition_groups_keep_business_conditions(self):
        from openpyxl import load_workbook
        book=load_workbook(REVIEW_ARTIFACT_ROOT/"业务规则审阅表.xlsx",read_only=True)
        try:
            rows={row[0]:row for row in book["业务规则"].iter_rows(min_row=6,values_only=True)}
            no_valid_combo="此路径无有效普通组合；见待处理项及详情。"
            for name,row in rows.items():
                for column in (3,4,5,6):
                    for block in row[column].split("\n\n"):
                        with self.subTest(name=name,column=column):
                            if block == no_valid_combo:
                                self.assertNotIn("\n适用：",block)
                                self.assertNotIn("可填写/选择",block)
                                continue
                            self.assertIn("\n适用：",block, "合法分支必须显示适用条件")
                            self.assertNotIn("技术组合由测试核验",block)
            defaults=rows["燃料碳氧化率FOx"][5].split("\n\n")
            bound=next(b for b in defaults if b.startswith("缺省路径"))
            self.assertIn("绑定资格=来源登记已核验",bound)
            self.assertIn("当前来源=缺省值",bound)
            unbound=next(b for b in defaults if b.startswith("不自动填写"))
            self.assertIn("当前来源=实测值",unbound)
            self.assertIn("目录参考值未核对",unbound)
            self.assertNotIn("绑定资格=",unbound, "未提供目录快照时不应暗示已核验绑定资格")

            details=(REVIEW_ARTIFACT_ROOT/"业务规则审阅详情.html").read_text(encoding="utf-8")
            electricity_conditions=[
                html.unescape(condition)
                for key,_,condition in re.findall(
                    r'<pre data-review-row="([^"]+)" data-review-property="([^"]+)">([\s\S]*?)</pre>',
                    details,
                )
                if "electricity.factor" in urllib.parse.unquote(html.unescape(key))
            ]
            power="\n".join(electricity_conditions)
            for condition in ("目录表1","目录表2","目录表3","目录表4","目录表5","供电单位实测","企业实测","企业计算","非化石购电零排放","方向=购入","方向=输出"):
                self.assertIn(condition,power)
            self.assertNotIn("SUPPLY_UNIT_MEASURED",power)
            self.assertTrue(any("方向=输出" in condition and "绑定资格=无适用绑定" in condition
                                for condition in electricity_conditions), "HTML须保留非法来源组合以支持复核")
            self.assertTrue(any("非化石购电零排放" in condition and "方向=购入" in condition
                                for condition in electricity_conditions), "HTML须保留仅购入方向的非化石零排放条件")
        finally:
            book.close()

    def test_actual_tables_keep_business_feedback_and_identity_together(self):
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
        book=load_workbook(REVIEW_ARTIFACT_ROOT/"业务规则审阅表.xlsx")
        try:
            for sheet in book:
                headers=[c.value for c in sheet[5]]
                self.assertTrue(all(isinstance(h,str) and h for h in headers))
                self.assertEqual(len(sheet.tables),1)
                table=next(iter(sheet.tables.values()))
                left,top,right,bottom=range_boundaries(table.ref)
                self.assertEqual((left,top,right,bottom),(1,5,len(headers),sheet.max_row))
                for h in ("审阅意见","处理状态","审阅编号","稳定审阅键","意见针对属性","意见针对条件","原审阅版本","原审阅投影","意见关联状态","内容指纹","意见基线指纹"):
                    self.assertLessEqual(headers.index(h)+1,right)
                for i,d in sheet.row_dimensions.items():
                    if d.height is not None:self.assertLessEqual(d.height,409,(sheet.title,i,d.height))
                for row in sheet.iter_rows(min_row=6,values_only=True):
                    self.assertLess(len(str(row[headers.index("内容指纹")])),32767)
        finally:book.close()

    def test_critical_operation_conditions_and_same_source_details_are_visible(self):
        from openpyxl import load_workbook
        book=load_workbook(REVIEW_ARTIFACT_ROOT/"业务规则审阅表.xlsx",read_only=True)
        try:
            sheet=book["操作与影响"];headers=[c.value for c in sheet[5]]
            rows=list(sheet.iter_rows(min_row=6,values_only=True))
            for operation,terms in (("capture_external_co2_supplement",("同一","空白","0","tCO₂")),("resolve_steam_enthalpy",("已知焓","自动","不"))):
                row=next(r for r in rows if operation in r[headers.index("稳定审阅键")])
                visible=row[headers.index("触发条件与处理")]
                self.assertTrue(all(term in visible for term in terms),(operation,visible))
                self.assertIn("初始条件",visible);self.assertIn("确认/取消",visible)
            fields=book["业务规则"];field_headers=[c.value for c in fields[5]]
            fch=next(r for r in fields.iter_rows(min_row=6,values_only=True) if r[0]=="单位热值含碳量FCh")
            self.assertEqual(fch[field_headers.index("单位")],"tC/GJ")
            scope=fch[field_headers.index("业务条件")]
            for term in ("质量","体积","热值","不新增HEAT"):
                self.assertIn(term,scope)
            html=(REVIEW_ARTIFACT_ROOT/"业务规则审阅详情.html").read_text(encoding="utf-8")
            self.assertIn("field|qz.carbon_accounting.industry.gbt_32151_34_2024|electricity.factor",html)
            self.assertIn("正式投影SHA256",html)
            checks=json.loads((REVIEW_ARTIFACT_ROOT/"review-artifact-checks.json").read_text(encoding="utf-8"))
            self.assertEqual(checks["generator_sha256"],hashlib.sha256((ROOT/"scripts/build_business_rules_review.mjs").read_bytes()).hexdigest())
        finally:book.close()

    def test_workbook_uses_object_specific_dropdowns_and_real_ooxml_column_hiding(self):
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries

        workbook_path = REVIEW_ARTIFACT_ROOT / "业务规则审阅表.xlsx"
        expected_scopes = {
            "业务规则": ["整体", "允许来源", "默认来源", "数值默认", "读写锁定"],
            "操作与影响": ["整体", "操作"],
            "本次变化": ["整体"],
        }
        hidden_names = {
            "原审阅版本", "原审阅投影", "稳定审阅键", "内容指纹", "意见基线指纹",
            "原投影键", "当前投影键", "规则修订",
        }
        visible_required = {
            "业务规则": {"审阅意见", "处理状态", "审阅编号", "意见针对属性", "意见针对条件", "意见关联状态", "业务条件"},
            "操作与影响": {"审阅意见", "处理状态", "审阅编号", "意见针对属性", "意见针对条件", "意见关联状态", "触发条件与处理"},
            "本次变化": {"审阅意见", "处理状态", "审阅编号", "意见针对属性", "意见针对条件", "意见关联状态", "变化类型"},
        }
        ns = {
            "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
            "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            "pkg": "http://schemas.openxmlformats.org/package/2006/relationships",
        }
        book = load_workbook(workbook_path, read_only=False, data_only=False)
        try:
            self.assertEqual(book.sheetnames, ["业务规则", "操作与影响", "本次变化"])
            with zipfile.ZipFile(workbook_path) as archive:
                workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
                relations_xml = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
                relationships = {item.attrib["Id"]: item.attrib["Target"] for item in relations_xml.findall("pkg:Relationship", ns)}
                part_by_sheet = {}
                for item in workbook_xml.findall("main:sheets/main:sheet", ns):
                    target = relationships[item.attrib[f"{{{ns['rel']}}}id"]]
                    part_by_sheet[item.attrib["name"]] = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))

                for sheet_name, allowed in expected_scopes.items():
                    with self.subTest(sheet=sheet_name):
                        sheet = book[sheet_name]
                        headers = [cell.value for cell in sheet[5]]
                        self.assertTrue(visible_required[sheet_name] <= set(headers))
                        self.assertTrue(hidden_names <= set(headers))
                        self.assertEqual(len(sheet.tables), 1)
                        table = next(iter(sheet.tables.values()))
                        _, top, right, bottom = range_boundaries(table.ref)
                        self.assertEqual((top, right, bottom), (5, len(headers), sheet.max_row))

                        property_column = headers.index("意见针对属性") + 1
                        sheet_xml = ET.fromstring(archive.read(part_by_sheet[sheet_name]))
                        validations = sheet_xml.findall("main:dataValidations/main:dataValidation", ns)
                        matching = []
                        for validation in validations:
                            refs = validation.attrib.get("sqref", "").split()
                            cells = [range_boundaries(ref) for ref in refs]
                            if any(left == property_column and right == property_column for left, _, right, _ in cells):
                                matching.append((validation, cells))
                        self.assertEqual(len(matching), 1, "scope dropdown must be encoded in worksheet OOXML")
                        validation, cells = matching[0]
                        self.assertEqual(validation.attrib.get("type"), "list")
                        self.assertEqual(cells, [(property_column, 6, property_column, sheet.max_row)])
                        formula = validation.findtext("main:formula1", namespaces=ns)
                        self.assertIsNotNone(formula)
                        self.assertTrue(formula.startswith('"') and formula.endswith('"'), formula)
                        self.assertEqual(formula[1:-1].split(","), allowed)

                        hidden_columns = set()
                        outlined_columns = set()
                        for col in sheet_xml.findall("main:cols/main:col", ns):
                            for index in range(int(col.attrib["min"]), int(col.attrib["max"]) + 1):
                                if col.attrib.get("hidden") in {"1", "true"}:
                                    hidden_columns.add(index)
                                if int(col.attrib.get("outlineLevel", "0")) > 0:
                                    outlined_columns.add(index)
                        expected_hidden = {i + 1 for i, header in enumerate(headers) if header in hidden_names}
                        self.assertEqual(hidden_columns, expected_hidden, "only technical identity/fingerprint columns may be hidden")
                        self.assertEqual(outlined_columns, expected_hidden, "hidden technical columns must be expandable groups")
        finally:
            book.close()

    def test_historical_visible_scope_set_matches_workbook_in_both_directions(self):
        from openpyxl import load_workbook

        workbook_path = REVIEW_ARTIFACT_ROOT / "业务规则审阅表.xlsx"
        html_path = REVIEW_ARTIFACT_ROOT / "业务规则审阅详情.html"
        text = html_path.read_text(encoding="utf-8")
        visible = [
            (urllib.parse.unquote(html.unescape(key)), html.unescape(prop), html.unescape(condition))
            for key, prop, condition in re.findall(
                r'<pre data-review-row="([^"]+)" data-review-property="([^"]+)">([\s\S]*?)</pre>', text
            )
        ]
        self.assertEqual(len(visible), 450, "Issue010 table5 conventional branches append eight visible scopes; historical branch IDs are retained")
        self.assertEqual(len(visible), len(set(visible)), "visible condition scopes must remain unique")

        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
        try:
            sheet = workbook["业务规则"]
            headers = [cell.value for cell in sheet[5]]
            stable_index = headers.index("稳定审阅键")
            fingerprint_index = headers.index("内容指纹")
            projected = set()
            for row in sheet.iter_rows(min_row=6, values_only=True):
                stable_key = row[stable_index]
                if not stable_key:
                    continue
                fingerprints = json.loads(row[fingerprint_index])
                for key in fingerprints:
                    if "@@" in key:
                        prop, condition = key.split("@@", 1)
                        projected.add((stable_key, prop, condition))
            self.assertEqual(set(visible), projected, "HTML and workbook scope fingerprints must correspond bidirectionally")
        finally:
            workbook.close()

    def test_catalog_snapshot_checks_mapping_values_and_exact_catalog_bindings(self):
        from packages.application.catalog_queries import CatalogQueryService
        from packages.persistence.catalog_builder import build_catalog_database
        from packages.persistence.catalog_repository import SQLiteCatalogRepository

        catalog_source = ROOT / "data-source/carbon_accounting/catalog.json"
        mapping = (ROOT / "specs/carbon_accounting/GB_T_32151_34_2024_MAPPING.md").read_text(encoding="utf-8")
        lines = mapping.splitlines()

        def line_for(reference):
            return next(line for line in lines if f"`{reference}`" in line)

        ratio_line = line_for("CAR-FLD-P04B-I")
        conversion_line = line_for("CAR-FLD-P04B-TR")
        heat_line = line_for("CAR-PAR-C3-HEAT-011")
        ratio_match = re.search(r"新建初始化\s*([0-9]+(?:\.[0-9]+)?)%", ratio_line)
        conversion_match = re.search(r"新建初始化\s*([0-9]+(?:\.[0-9]+)?)%", conversion_line)
        heat_match = re.search(r"\|\s*([0-9]+(?:\.[0-9]+)?)\s*[；;]", heat_line)
        self.assertIsNotNone(ratio_match, ratio_line)
        self.assertIsNotNone(conversion_match, conversion_line)
        self.assertIsNotNone(heat_match, heat_line)
        expected = {
            "fgd.ratio": (Decimal(ratio_match.group(1)) / Decimal(100), "ratio"),
            "fgd.conversion": (Decimal(conversion_match.group(1)) / Decimal(100), "ratio"),
            "heat.factor": (Decimal(heat_match.group(1)), "tCO₂/GJ"),
        }

        with tempfile.TemporaryDirectory(prefix="r11-catalog-snapshot-test-") as directory:
            root = Path(directory)
            review_path = root / "review.json"
            technical_path = root / "technical.json"
            snapshot_path = root / "catalog-snapshot.json"
            candidate_paths = (
                SOURCE / "common/gbt_32150_2025.json",
                SOURCE / "standards/gbt_32151_34_2024.json",
                SOURCE / "software_policy.json",
            )
            result = subprocess.run(
                [
                    sys.executable,
                    str(ROOT / "scripts/validate_business_rules_candidate.py"),
                    "--review-json", str(review_path),
                    "--technical-review-json", str(technical_path),
                    "--catalog-snapshot", str(snapshot_path),
                    "--catalog-source", str(catalog_source),
                    *(str(path) for path in candidate_paths),
                ],
                cwd=ROOT,
                text=True,
                capture_output=True,
                check=False,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
            technical = json.loads(technical_path.read_text(encoding="utf-8"))
            self.assertEqual(snapshot["schema_version"], 1)
            self.assertEqual(snapshot["status"], "CHECKED")
            self.assertEqual(snapshot["source_compact_sha256"], technical["metadata"]["source_compact_sha256"])
            self.assertEqual(snapshot["catalog"]["canonical_source_lf_sha256"],
                             hashlib.sha256(catalog_source.read_bytes().replace(b"\r\n", b"\n")).hexdigest())
            self.assertTrue(snapshot["catalog"]["data_version"])
            self.assertEqual(set(snapshot["entries"]), set(expected))

            with tempfile.TemporaryDirectory(prefix="r11-catalog-snapshot-db-") as database_directory:
                database_path = Path(database_directory) / "catalog.sqlite"
                build_catalog_database(source_path=catalog_source, output_path=database_path)
                repository = SQLiteCatalogRepository(database_path)
                service = CatalogQueryService(repository)
                assets = service.list_reference_data_assets()
                bindings = service.list_reference_data_bindings()
                source_tables = service.list_source_tables()
                sources = service.list_sources()
                parameters = repository.list_parameters()
                for field_id, (expected_value, expected_unit) in expected.items():
                    with self.subTest(field=field_id):
                        entry = snapshot["entries"][field_id]
                        matching_technical = [
                            row for row in technical["fields"]
                            if row.get("field_id") == field_id
                            and (row.get("policy_branch_values") or {}).get("default", {}).get("state") == "BINDING"
                        ]
                        self.assertTrue(matching_technical, "candidate technical projection must identify the BINDING branch")
                        identities = set()
                        for row in matching_technical:
                            tags = (row.get("conditions_by_property") or {}).get("default", {}).get("tags", {})
                            identities.add((row["policy_branch_values"]["default"]["parameter_id"], tags.get("binding_id")))
                        self.assertEqual(identities, {(entry["parameter_id"], entry["binding_id"])})

                        self.assertEqual(Decimal(entry["value"]), expected_value, "catalog value must match the approved Mapping")
                        self.assertEqual(entry["unit"], expected_unit)
                        binding = next(item for item in bindings if item.binding_id == entry["binding_id"])
                        asset = next(item for item in assets if item.asset_id == entry["asset_id"])
                        parameter = next(item for item in parameters if item.parameter_id == entry["parameter_id"])
                        table = next(item for item in source_tables if item.source_table_id == entry["source_table_id"])
                        source = next(item for item in sources if item.source_id == entry["source_id"])
                        self.assertEqual(binding.asset_id, asset.asset_id)
                        self.assertEqual(binding.source_table_id, table.source_table_id)
                        self.assertEqual(table.source_id, source.source_id)
                        self.assertEqual(asset.parameter_id, parameter.parameter_id)
                        self.assertEqual(asset.value, Decimal(entry["value"]))
                        self.assertEqual(asset.unit, entry["unit"])
                        self.assertEqual(entry["source_version"], source.version)
                        self.assertEqual(entry["source_table_number"], table.display_number)
                        self.assertEqual(entry["review_status"], "VERIFIED")
                        self.assertIn(entry["candidate_standard_id"], binding.applicable_standard_ids)

    def test_all_packages_remain_candidate_and_reference_explicit_revisions(self):
        for doc in self.documents():
            self.assertEqual(doc["status"],"CANDIDATE")
            for adopted in doc["adopts"]:
                for value in adopted.values():
                    if isinstance(value,str):
                        self.assertNotIn(value.lower(),{"latest","*","all"})

    def test_actual_candidate_property_branches_are_structurally_composable(self):
        from packages.application.business_rules_candidate import load_bundle
        paths=(SOURCE/"common/gbt_32150_2025.json",SOURCE/"standards/gbt_32151_34_2024.json",SOURCE/"software_policy.json")
        rows=load_bundle(paths).technical_review_rows()["fields"]
        unexpected=[(row["field_id"],row.get("condition_id"),row["evaluation_status"])
                    for row in rows if row["evaluation_status"] not in {"READY","BLOCKED"}]
        self.assertEqual(unexpected,[],
            "真实候选资产可明确阻断非法/待核来源，但不应留下策略间无法组合的结构缺项")


    def test_review_checks_cli_binds_current_html_and_runs_without_node(self):
        import shutil
        import tempfile

        from openpyxl import load_workbook

        from scripts.business_rules_review_checks import _feedback_rows, semantic_json_sha256
        from scripts.business_rules_review_hashes import (
            DEFINITION_ASSET_HASH_POLICY,
            EXPECTED_DEFINITION_ASSETS,
            definition_text_sha256,
        )

        source_checks = json.loads(
            (REVIEW_ARTIFACT_ROOT / "review-artifact-checks.json").read_text(encoding="utf-8")
        )
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            stage = root / "input" / "review-stage"
            stage.mkdir(parents=True)
            review_path = root / "review.json"
            technical_path = root / "technical.json"
            snapshot_path = root / "catalog-snapshot.json"
            candidate_paths = (
                SOURCE / "common/gbt_32150_2025.json",
                SOURCE / "standards/gbt_32151_34_2024.json",
                SOURCE / "software_policy.json",
            )
            generated = subprocess.run(
                [
                    sys.executable,
                    str(ROOT / "scripts/validate_business_rules_candidate.py"),
                    "--review-json", str(review_path),
                    "--technical-review-json", str(technical_path),
                    "--catalog-snapshot", str(snapshot_path),
                    "--catalog-source", str(ROOT / "data-source/carbon_accounting/catalog.json"),
                    *(str(path) for path in candidate_paths),
                ],
                cwd=ROOT,
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(generated.returncode, 0, generated.stdout + generated.stderr)

            review_bytes = review_path.read_bytes()
            projection = hashlib.sha256(review_bytes.replace(b"\r\n", b"\n")).hexdigest()
            technical = json.loads(technical_path.read_text(encoding="utf-8"))
            snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
            self.assertEqual(technical["metadata"]["source_compact_sha256"], projection)

            workbook_path = stage / "业务规则审阅表.xlsx"
            details_path = stage / "业务规则审阅详情.html"
            fixture_path = stage / "review-feedback-fixture.json"
            for name, destination in (
                ("业务规则审阅表.xlsx", workbook_path),
                ("业务规则审阅详情.html", details_path),
                ("review-feedback-fixture.json", fixture_path),
            ):
                shutil.copyfile(REVIEW_ARTIFACT_ROOT / name, destination)

            workbook = load_workbook(workbook_path, data_only=False)
            try:
                for sheet in workbook.worksheets:
                    self.assertEqual(sheet["A1"].value, "审阅数据SHA256：" + projection)
                business_sheet = workbook["业务规则"]
                headers = [cell.value for cell in business_sheet[5]]
                opinion_column = headers.index("审阅意见") + 1
                original_feedback = _feedback_rows(workbook_path)
                empty_cell = next(
                    (
                        business_sheet.cell(row_number, opinion_column)
                        for row_number in range(6, business_sheet.max_row + 1)
                        if business_sheet.cell(row_number, opinion_column).value in (None, "")
                    ),
                    None,
                )
                self.assertIsNotNone(empty_cell, "the review workbook needs an empty feedback row")
                empty_cell.value = "保留用户意见"
                workbook.save(workbook_path)
            finally:
                workbook.close()
            previous_path = root / "previous.xlsx"
            shutil.copyfile(workbook_path, previous_path)
            retained_feedback = _feedback_rows(workbook_path)
            self.assertEqual(retained_feedback, original_feedback + 1)

            asset_hashes = {
                relative: definition_text_sha256((ROOT / relative).read_bytes())
                for relative in EXPECTED_DEFINITION_ASSETS
            }
            manifest = {
                "view_revision": "compact-business-v2",
                "definition_assets_hash_policy": DEFINITION_ASSET_HASH_POLICY,
                "definition_assets_sha256": asset_hashes,
                "projection_sha256": projection,
                "details": details_path.name,
                "details_revision": technical["metadata"]["revision"],
                "details_counts": {
                    section: len(technical[section])
                    for section in ("fields", "operations", "changes")
                },
                "details_sha256": semantic_json_sha256(technical),
                "details_projection_bound": True,
                "catalog_snapshot": {
                    "catalog": snapshot["catalog"],
                    "entries": snapshot["entries"],
                },
                "tables": [
                    {
                        "sheet": name,
                        "table": source_checks["table_ranges"][name]["range"],
                        "rows": source_checks["sheets"][name],
                        "headers": source_checks["table_ranges"][name]["columns"],
                    }
                    for name in source_checks["sheets"]
                ],
                "retained_feedback": retained_feedback,
            }
            (stage / "review-view-manifest.json").write_text(
                json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )

            output_path = root / "nested" / "checks" / "review-artifact-checks.json"
            python_path = str(Path(sys.executable).parent)
            self.assertIsNone(shutil.which("node", path=python_path))
            command = [
                sys.executable,
                str(ROOT / "scripts/business_rules_review_checks.py"),
                "--root", str(ROOT),
                "--stage", str(stage),
                "--review-json", str(review_path),
                "--technical-json", str(technical_path),
                "--fixture", str(fixture_path),
                "--output", str(output_path),
                "--previous", str(previous_path),
            ]
            check_env = {**os.environ, "PATH": python_path, "PYTHONIOENCODING": "cp1252"}

            def run_checks():
                return subprocess.run(
                    command,
                    cwd=ROOT,
                    env=check_env,
                    capture_output=True,
                    text=True,
                )

            completed = run_checks()
            self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
            checks = json.loads(output_path.read_text(encoding="utf-8"))
            self.assertEqual(checks["review_mode"], "compact_business_v2")
            self.assertEqual(checks["serialization_aliases"]["manifest.view_revision"], "compact-business-v2")
            self.assertEqual(checks["feedback_retention"]["manifest_retained_feedback"], retained_feedback)
            self.assertEqual(checks["feedback_retention"]["workbook_feedback_rows"], retained_feedback)
            self.assertEqual(checks["feedback_retention"]["previous_workbook_feedback_rows"], retained_feedback)
            self.assertEqual(checks["feedback_test"]["formal_projection_sha256"], projection)
            self.assertEqual(checks["feedback_test"]["formal_details_sha256"], hashlib.sha256(
                technical_path.read_bytes().replace(b"\r\n", b"\n")
            ).hexdigest())

            original_fixture = fixture_path.read_bytes()
            invalid_fixture = json.loads(original_fixture)
            invalid_fixture["formal_projection_sha256"] = "0" * 64
            fixture_path.write_text(json.dumps(invalid_fixture), encoding="utf-8")
            try:
                with self.subTest(rejected_binding="fixture projection"):
                    rejected = run_checks()
                    self.assertNotEqual(rejected.returncode, 0)
                    self.assertIn("feedback fixture is not bound", rejected.stderr)
            finally:
                fixture_path.write_bytes(original_fixture)

            original_technical = technical_path.read_bytes()
            invalid_technical = json.loads(original_technical)
            invalid_technical["metadata"]["source_compact_sha256"] = "0" * 64
            technical_path.write_text(
                json.dumps(invalid_technical, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            try:
                with self.subTest(rejected_binding="technical source projection"):
                    rejected = run_checks()
                    self.assertNotEqual(rejected.returncode, 0)
                    self.assertIn("not bound to this review projection", rejected.stderr)
            finally:
                technical_path.write_bytes(original_technical)

            original_html = details_path.read_text(encoding="utf-8")
            for label, pattern, message in (
                ("mixed HTML projection", r"(?<=正式投影SHA256：)[0-9a-f]{64}", "details HTML projection SHA-256"),
                ("mixed HTML semantic details", r"(?<=详情输入SHA256：)[0-9a-f]{64}", "details HTML semantic SHA-256"),
            ):
                mixed_html, substitutions = re.subn(pattern, "0" * 64, original_html, count=1)
                self.assertEqual(substitutions, 1)
                details_path.write_text(mixed_html, encoding="utf-8")
                try:
                    with self.subTest(rejected_mixed_source=label):
                        rejected = run_checks()
                        self.assertNotEqual(rejected.returncode, 0)
                        self.assertIn(message, rejected.stderr)
                finally:
                    details_path.write_text(original_html, encoding="utf-8")

            scope_match = re.search(
                r'(<pre data-review-row="[^"]+" data-review-property="[^"]+">)([\s\S]*?)(</pre>)',
                original_html,
            )
            self.assertIsNotNone(scope_match)
            altered_scope = (
                original_html[:scope_match.start()]
                + scope_match.group(1)
                + scope_match.group(2)
                + " 混源"
                + scope_match.group(3)
                + original_html[scope_match.end():]
            )
            details_path.write_text(altered_scope, encoding="utf-8")
            try:
                with self.subTest(rejected_binding="HTML scope set"):
                    rejected = run_checks()
                    self.assertNotEqual(rejected.returncode, 0)
                    self.assertIn("HTML scopes do not match business review fingerprints", rejected.stderr)
            finally:
                details_path.write_text(original_html, encoding="utf-8")
