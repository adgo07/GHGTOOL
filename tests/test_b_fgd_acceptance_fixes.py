from packages.ui.dialogs import MessageBox as QMessageBox
"""PR50 F01-F03实际offscreen Qt/Excel/Application/SQLite退修回归。"""
import os
os.environ.setdefault("QT_QPA_PLATFORM","offscreen")
from contextlib import contextmanager
from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import json,tempfile,unittest
from unittest.mock import patch
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication
from apps.carbon_accounting_desktop.app import create_main_window
from apps.carbon_accounting_desktop.config import AppConfig
from packages.application import ProjectWorkspaceService
from packages.persistence import SQLiteProjectWorkspaceRepository,SQLiteRecordRepository,build_catalog_database
from packages.reference_data import DEFAULT_SOURCE_PATH
from packages.excel.templates import ExcelTemplateService
from packages.ui.view_models import AppRoute
import tests.test_b_fgd_chain as fgd_fixtures

class FGDAcceptanceFixTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app=QApplication.instance() or QApplication([])
        cls.probes=[]
    @classmethod
    def tearDownClass(cls):
        if os.environ.get("QZ_FGD_REPRO_PATH"):
            path=Path(os.environ["QZ_FGD_REPRO_PATH"]);path.parent.mkdir(parents=True,exist_ok=True)
            path.write_text(json.dumps({'head':os.environ.get('QZ_FGD_REPRO_HEAD'),'evidence':'actual offscreen Qt/Excel/Application/SQLite','probes':cls.probes},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    @contextmanager
    def fresh_page(self,root):
        catalog=root/'catalog.sqlite'
        if not catalog.exists():build_catalog_database(DEFAULT_SOURCE_PATH,catalog)
        projects=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/'projects.sqlite'))
        records=SQLiteRecordRepository(root/'records.sqlite')
        window=create_main_window(AppConfig(catalog_database=catalog,records_database=root/'records.sqlite',projects_database=root/'projects.sqlite'),record_repository=records,project_service=projects)
        shell=window.centralWidget();shell.navigate(AppRoute.NEW_ACCOUNTING);self.app.processEvents()
        try:yield shell.pages[AppRoute.NEW_ACCOUNTING],projects,records,shell
        finally:
            with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):window.close()
            window.deleteLater();self.app.processEvents()
    def prepare_valid(self,gui):
        fgd_fixtures.FGDChainTests.prepare(gui)
        batch=gui._process_rows['fgd'][0]
        batch['reagent_name'].setText('完整第一批');batch['group_amount'].setText('10')
        fgd_fixtures.FGDChainTests.select(gui,batch['components'][0],'CaCO3')
        return batch
    def test_f01_named_amount_batch_without_components_blocks_formal_record(self):
        with tempfile.TemporaryDirectory() as directory,self.fresh_page(Path(directory)) as (gui,projects,records,shell):
            self.prepare_valid(gui)
            second=gui._add_process_row('fgd',instance_id='incomplete-second')
            second['reagent_name'].setText('已录入第二批');second['group_amount'].setText('20')
            value=gui._input(increment=False)
            gui._run_calculation()
            self.probes.append({'issue':'F01','ui_batches':2,'canonical_batches':len(value.fgd_units),'records_after_formal':len(records.list_all()),'gui_complete':gui._calculation_has_result})
            self.assertEqual(len(records.list_all()),0)
            self.assertFalse(gui._calculation_has_result)
            self.assertEqual(len(value.fgd_units),2)
            self.assertEqual(second['group_amount'].text(),'20')
            fgd_fixtures.FGDChainTests.select(gui,second['components'][0],'CaCO3')
            gui._run_calculation()
            self.assertEqual(len(records.list_all()),1)
            self.assertTrue(gui._calculation_has_result)
            self.assertEqual(records.list_all()[0].calculation_result.total_amount,Decimal('11.88'))
    def test_f02_identity_only_excel_component_is_not_ignored(self):
        with tempfile.TemporaryDirectory() as directory,self.fresh_page(Path(directory)) as (gui,projects,records,shell):
            root=Path(directory);self.prepare_valid(gui)
            workbook=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
            ws=workbook['B.7']
            for cell,value in {'A3':'身份待补批次','B3':10,'J3':'缺省值','K3':'缺省值','L3':'缺省值','C4':'待补实测的自定义碳酸盐','I4':'自定义'}.items():ws[cell]=value
            path=root/'identity-only.xlsx';workbook.save(path);workbook.close()
            adapter=shell.pages[AppRoute.EXCEL_IMPORT].excel_adapter
            preview=adapter.import_preview(path,adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name='退修复现'))
            unit=preview.units[0]
            formal=gui.calculation_use_case.calculate(unit.input_value) if unit.input_value is not None else None
            self.probes.append({'issue':'F02','preview_can_calculate':unit.can_calculate,'canonical_components':len(unit.input_value.fgd_units[0].components) if unit.input_value else None,'formal_success':formal.successful if formal else None,'records_after_formal':len(records.list_all()),'errors':[e.code for e in unit.errors]})
            self.assertFalse(unit.can_calculate)
            if formal is not None:self.assertFalse(formal.successful)
            self.assertEqual(len(records.list_all()),0)
            # Same incomplete identity is retained by GUI as well as Excel.
            batch=gui._process_rows['fgd'][0]
            component=gui._add_fgd_component(batch)
            fgd_fixtures.FGDChainTests.select(gui,component,'待补实测的自定义碳酸盐')
            gui._run_calculation()
            self.assertFalse(gui._calculation_has_result)
            self.assertEqual(len(records.list_all()),0)
            # Complete both components with independent approved measured data.
            workbook=load_workbook(path);ws=workbook['B.7']
            for cell,value in {'D3':40,'F3':80,'K3':'实测值','L3':'实测值','D4':30,'F4':90,'H4':.522,'J4':'实测值','K4':'实测值','L4':'实测值'}.items():ws[cell]=value
            workbook.save(path);workbook.close()
            preview=adapter.import_preview(path,adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name='退修复验'))
            self.assertTrue(preview.units[0].can_calculate,preview.units[0].errors)
            formal=gui.calculation_use_case.calculate(preview.units[0].input_value)
            self.assertTrue(formal.successful,formal.problems)
            self.assertEqual(formal.result.total_amount,Decimal('2.8174'))
            self.assertEqual(len(records.list_all()),1)
    def test_f03_upgraded_facility_survives_real_new_window_and_services(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            with self.fresh_page(root) as (gui,projects,records,shell):
                batch=self.prepare_valid(gui)
                fgd_fixtures.FGDChainTests.measured(batch['components'][0],'tr','75')
                state=gui._capture_form_state();state.pop('fgd_identities')
                gui._restore_form_state(state)
                before=gui._capture_form_state()
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):gui._upgrade_fgd_rules()
                self.assertEqual(gui._capture_form_state(),before)
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.Yes):gui._upgrade_fgd_rules()
                batch=gui._process_rows['fgd'][0];identifier=batch['facility_id'];display=batch['facility_selector'].currentText()
                gui.project_name.setText('设施新窗口复现');self.assertTrue(gui._save_project())
                project_id=gui._workspace.project_id
            with self.fresh_page(root) as (gui,projects,records,shell):
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):self.assertTrue(gui.open_project_from_home(project_id))
                batch=gui._process_rows['fgd'][0];selector=batch['facility_selector']
                value=gui._input(increment=False).fgd_units[0]
                self.probes.append({'issue':'F03','saved_id':identifier,'restored_id':batch['facility_id'],'saved_display':display,'restored_display':selector.currentText(),'findData':selector.findData(identifier),'canonical_facility_name':value.facility_name,'records':len(records.list_all())})
                self.assertEqual(batch['facility_id'],identifier)
                self.assertGreaterEqual(selector.findData(identifier),0)
                self.assertEqual(selector.currentData(),identifier)
                self.assertEqual(selector.currentText(),display)
                self.assertEqual(value.facility_name,display)
                self.assertEqual(batch['components'][0]['tr'].text(),'75')
                self.assertEqual(len(records.list_all()),0)

    def test_f01_blank_placeholder_ignored_but_named_blank_and_zero_batch_preserved(self):
        with tempfile.TemporaryDirectory() as directory,self.fresh_page(Path(directory)) as (gui,projects,records,shell):
            self.prepare_valid(gui)
            second=gui._add_process_row('fgd',instance_id='placeholder')
            self.assertEqual(len(gui._input(increment=False).fgd_units),1)
            second['reagent_name'].setText('待补用量')
            gui._run_calculation()
            self.assertFalse(gui._calculation_has_result);self.assertEqual(len(records.list_all()),0)
            second['group_amount'].setText('0')
            gui._run_calculation()
            self.assertFalse(gui._calculation_has_result);self.assertEqual(len(records.list_all()),0)
            fgd_fixtures.FGDChainTests.select(gui,second['components'][0],'CaCO3')
            self.assertEqual(len(gui._input(increment=False).fgd_units),2)
            gui._run_calculation()
            self.assertTrue(gui._calculation_has_result)
            self.assertEqual(len(records.list_all()),1)
            self.assertEqual(records.list_all()[0].calculation_result.total_amount,Decimal('3.96'))
    def test_f02_static_placeholders_identity_source_value_and_zero(self):
        with tempfile.TemporaryDirectory() as directory,self.fresh_page(Path(directory)) as (gui,projects,records,shell):
            self.prepare_valid(gui);root=Path(directory)
            adapter=shell.pages[AppRoute.EXCEL_IMPORT].excel_adapter
            context=adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name='活动行保护')
            def preview(changes):
                workbook=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()));ws=workbook['B.7']
                for cell,value in {'A3':'第一批','B3':10,'J3':'缺省值','K3':'缺省值','L3':'缺省值',**changes}.items():ws[cell]=value
                path=root/'variants.xlsx';workbook.save(path);workbook.close()
                return adapter.import_preview(path,context).units[0]
            untouched=preview({})
            self.assertTrue(untouched.can_calculate,untouched.errors)
            self.assertEqual(len(untouched.input_value.fgd_units[0].components),1)
            for changes in ({'C4':'只改名称'},{'I4':'自定义'},{'J4':'缺省值'},{'D4':0},{'C6':'新批次待补身份'}):
                with self.subTest(changes=changes):
                    result=preview(changes);self.assertFalse(result.can_calculate)
                    self.assertEqual(len(records.list_all()),0)
            zero=preview({'D3':40,'F3':100,'K3':'实测值','L3':'实测值','D4':0,'F4':0,'J4':'缺省值','K4':'实测值','L4':'实测值'})
            self.assertTrue(zero.can_calculate,zero.errors)
            self.assertEqual(len(zero.input_value.fgd_units[0].components),2)
            self.assertEqual(zero.calculation.result.total_amount,Decimal('1.76'))
    def test_f03_new_v2_shared_facility_deleted_owner_and_source_drafts_reopen(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            with self.fresh_page(root) as (gui,projects,records,shell):
                self.prepare_valid(gui)
                owner=gui._add_process_row('fgd',instance_id='facility-owner')
                owner['reagent_name'].setText('前批');owner['group_amount'].setText('10')
                fgd_fixtures.FGDChainTests.select(gui,owner['components'][0],'CaCO3')
                shared=gui._add_process_row('fgd',instance_id='facility-survivor')
                shared['reagent_name'].setText('后批');shared['group_amount'].setText('20')
                fgd_fixtures.FGDChainTests.select(gui,shared['components'][0],'CaCO3')
                identifier=owner['facility_id'];selector=shared['facility_selector']
                selector.setCurrentIndex(selector.findData(identifier));display=selector.currentText()
                fgd_fixtures.FGDChainTests.measured(shared['components'][0],'tr','75')
                gui._remove_process_row('fgd','facility-owner')
                self.assertEqual(shared['facility_id'],identifier)
                self.assertEqual(shared['facility_selector'].currentText(),display)
                gui.project_name.setText('删除前批后保存');self.assertTrue(gui._save_project());project_id=gui._workspace.project_id
            with self.fresh_page(root) as (gui,projects,records,shell):
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):self.assertTrue(gui.open_project_from_home(project_id))
                shared=next(row for row in gui._process_rows['fgd'] if row['instance_id']=='facility-survivor')
                selector=shared['facility_selector'];self.assertEqual(selector.currentData(),identifier);self.assertEqual(selector.currentText(),display)
                for row in gui._process_rows['fgd']:
                    ids=[row['facility_selector'].itemData(i) for i in range(row['facility_selector'].count())]
                    self.assertEqual(len(ids),len(set(ids)));self.assertIn(identifier,ids)
                component=shared['components'][0]
                component['tr_source'].setCurrentIndex(component['tr_source'].findData('STANDARD_DEFAULT'))
                component['tr_source'].setCurrentIndex(component['tr_source'].findData('MEASURED'))
                self.assertEqual(component['tr'].text(),'75')
                self.assertEqual(len(records.list_all()),0)
    def test_f03_unknown_saved_facility_is_kept_with_business_notice(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            with self.fresh_page(root) as (gui,projects,records,shell):
                self.prepare_valid(gui);state=gui._capture_form_state()
                for identity in state['fgd_identities'].values():
                    identity['facility_id']='saved-unknown-facility';identity.pop('facility_name',None)
                workspace=projects.new_workspace('旧设施名称未保存');projects.save(replace(workspace,units=(replace(workspace.units[0],form_state=state),)));project_id=workspace.project_id
            with self.fresh_page(root) as (gui,projects,records,shell):
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):self.assertTrue(gui.open_project_from_home(project_id))
                selector=gui._process_rows['fgd'][0]['facility_selector']
                self.assertEqual(selector.currentData(),'saved-unknown-facility')
                self.assertIn('待核实',selector.currentText());self.assertIn('对应关系已保留',selector.toolTip())
                self.assertEqual(len(records.list_all()),0)

    def test_f01_incomplete_batch_roundtrips_sqlite_and_still_blocks(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            with self.fresh_page(root) as (gui,projects,records,shell):
                self.prepare_valid(gui)
                second=gui._add_process_row('fgd',instance_id='unfinished-persisted')
                second['reagent_name'].setText('已填未完第二批');second['group_amount'].setText('20')
                gui.project_name.setText('未完输入保存');self.assertTrue(gui._save_project());project_id=gui._workspace.project_id
                self.assertEqual(len(records.list_all()),0)
            with self.fresh_page(root) as (gui,projects,records,shell):
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):self.assertTrue(gui.open_project_from_home(project_id))
                second=next(row for row in gui._process_rows['fgd'] if row['instance_id']=='unfinished-persisted')
                self.assertEqual(second['reagent_name'].text(),'已填未完第二批');self.assertEqual(second['group_amount'].text(),'20')
                gui._run_calculation();self.assertFalse(gui._calculation_has_result);self.assertEqual(len(records.list_all()),0)
                fgd_fixtures.FGDChainTests.select(gui,second['components'][0],'CaCO3')
                gui._run_calculation();self.assertTrue(gui._calculation_has_result);self.assertEqual(len(records.list_all()),1)
                self.assertEqual(records.list_all()[0].calculation_result.total_amount,Decimal('11.88'))
    def test_f02_two_independent_batches_same_facility_do_not_add_percentages(self):
        with tempfile.TemporaryDirectory() as directory,self.fresh_page(Path(directory)) as (gui,projects,records,shell):
            self.prepare_valid(gui);root=Path(directory)
            adapter=shell.pages[AppRoute.EXCEL_IMPORT].excel_adapter
            context=adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name='独立批次')
            workbook=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()));ws=workbook['B.7']
            for cell,value in {'A3':'第一批','B3':10,'J3':'缺省值','K3':'缺省值','L3':'缺省值','M3':'设备A',
                'A6':'第二批','B6':20,'C6':'未完第二批组分','I6':'自定义','M6':'设备A'}.items():ws[cell]=value
            path=root/'two-batches.xlsx';workbook.save(path);workbook.close()
            bad=adapter.import_preview(path,context).units[0]
            self.assertFalse(bad.can_calculate);self.assertEqual(len(records.list_all()),0)
            workbook=load_workbook(path);ws=workbook['B.7']
            for cell,value in {'C6':'CaCO3','I6':'目录','J6':'缺省值','K6':'缺省值','L6':'缺省值'}.items():ws[cell]=value
            workbook.save(path);workbook.close()
            complete=adapter.import_preview(path,context).units[0]
            self.assertTrue(complete.can_calculate,complete.errors)
            units=complete.input_value.fgd_units
            self.assertEqual(len(units),2);self.assertEqual(units[0].facility_id,units[1].facility_id)
            self.assertNotEqual(units[0].batch_id,units[1].batch_id)
            self.assertEqual([unit.components[0].carbonate_fraction.value for unit in units],[Decimal('.9'),Decimal('.9')])
            formal=gui.calculation_use_case.calculate(complete.input_value)
            self.assertTrue(formal.successful,formal.problems);self.assertEqual(formal.result.total_amount,Decimal('11.88'))
            self.assertEqual(len(records.list_all()),1)
    def test_f03_two_v2_batches_shared_facility_reopen_and_formal_total(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            with self.fresh_page(root) as (gui,projects,records,shell):
                first=self.prepare_valid(gui)
                second=gui._add_process_row('fgd',instance_id='shared-second')
                second['reagent_name'].setText('同设施第二批');second['group_amount'].setText('20')
                fgd_fixtures.FGDChainTests.select(gui,second['components'][0],'CaCO3')
                identifier=first['facility_id'];display=first['facility_selector'].currentText()
                second['facility_selector'].setCurrentIndex(second['facility_selector'].findData(identifier))
                fgd_fixtures.FGDChainTests.measured(second['components'][0],'tr','75')
                gui.project_name.setText('同设施两批');self.assertTrue(gui._save_project());project_id=gui._workspace.project_id
            with self.fresh_page(root) as (gui,projects,records,shell):
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):self.assertTrue(gui.open_project_from_home(project_id))
                units=gui._process_rows['fgd'];self.assertEqual(len(units),2)
                for unit in units:
                    selector=unit['facility_selector']
                    self.assertEqual(selector.currentData(),identifier);self.assertEqual(selector.currentText(),display)
                    self.assertEqual(selector.count(),1)
                component=units[1]['components'][0]
                component['tr_source'].setCurrentIndex(component['tr_source'].findData('STANDARD_DEFAULT'))
                component['tr_source'].setCurrentIndex(component['tr_source'].findData('MEASURED'))
                self.assertEqual(component['tr'].text(),'75')
                gui._run_calculation();self.assertTrue(gui._calculation_has_result);self.assertEqual(len(records.list_all()),1)
                self.assertEqual(records.list_all()[0].calculation_result.total_amount,Decimal('9.9'))
