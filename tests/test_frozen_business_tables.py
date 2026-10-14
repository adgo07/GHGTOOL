"""Frozen display evidence and one business table projection across outputs."""
from datetime import date
from decimal import Decimal
from dataclasses import replace
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from packages.application.carbon_accounting import CarbonAccountingUseCase,create_g06_parameter_resolver
from packages.application.catalog_queries import CatalogQueryService
from packages.application.reporting import build_saved_record_report,build_report_model
from packages.application.reporting.presentation import report_html,report_text
from packages.core.models import AccountingPeriod,ParameterType,PeriodType
from packages.core.parameter_resolution import ParameterResolutionContext
from packages.persistence.in_memory_records import InMemoryRecordRepository
from packages.persistence import SQLiteCatalogRepository,build_catalog_database
from packages.standards.carbon_material import CarbonMaterialCalculator,CarbonMaterialInput,FuelInput,FuelPath,FuelType,ParameterValue,ParameterSourceKind,EmissionSourceState,EmissionSourceStatus,SOURCE_FUEL,STANDARD_ID
from packages.standards.carbon_material_input_policy import InputOrigin
from packages.ui.record_experience import build_record_report_model


class FrozenBusinessTableTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory=tempfile.TemporaryDirectory()
        cls.catalog=CatalogQueryService(SQLiteCatalogRepository(build_catalog_database(output_path=Path(cls.directory.name)/"catalog.sqlite")))
        cls.resolver=create_g06_parameter_resolver(cls.catalog.repository)
        cls.period=AccountingPeriod(PeriodType.ANNUAL,date(2026,1,1),date(2026,12,31))

    @classmethod
    def tearDownClass(cls):cls.directory.cleanup()

    def calculate(self, *, catalog=False):
        if catalog:
            parameters = {item.parameter_type: item for item in self.catalog.repository.list_parameters()
                          if item.subject_id == "natural_gas"}
            def parameter(parameter_type):
                metadata = parameters[parameter_type]
                resolution = self.resolver.resolve(ParameterResolutionContext(
                    parameter_id=metadata.parameter_id, standard_id=STANDARD_ID,
                    accounting_period=self.period, subject_id=metadata.subject_id,
                    parameter_type=metadata.parameter_type,
                ))
                self.assertFalse(resolution.blocked, resolution.warnings)
                factor = resolution.recommended.factor
                return ParameterValue(
                    factor.parameter_id, factor.value, factor.unit, ParameterSourceKind.STANDARD_DEFAULT,
                    factor_id=factor.factor_id, source_id=factor.source_id,
                    source_version=factor.version, source_location=factor.source_location,
                    factor_year=factor.factor_year,
                )
            lhv=parameter(ParameterType.LOWER_HEATING_VALUE)
            carbon=parameter(ParameterType.CARBON_CONTENT_PER_HEAT)
            fox=parameter(ParameterType.OXIDATION_RATE)
        else:
            def parameter(identifier,value,unit):return ParameterValue(identifier,Decimal(value),unit,ParameterSourceKind.MEASURED,source_location="专项填写依据")
            lhv=parameter("custom.lhv","389.31","GJ/10⁴Nm³");carbon=parameter("custom.carbon","0.0153","tC/GJ");fox=parameter("custom.fox","0.99","ratio")
        fuel=FuelInput("frozen-fuel",FuelPath.VOLUME,"1",carbon,fox,lhv,
            fuel_type=FuelType.NATURAL_GAS if catalog else FuelType.OTHER,fuel_label="冻结燃料名称",
            origin=InputOrigin.CATALOG if catalog else InputOrigin.CUSTOM,catalog_subject_id="natural_gas" if catalog else None)
        input_value=CarbonMaterialInput("input.frozen-tables","enterprise.frozen","冻结表式验收",self.period,boundary_confirmed=True,
            fuel_inputs=(fuel,),source_states=(EmissionSourceState(SOURCE_FUEL,EmissionSourceStatus.INVOLVED),))
        repository=InMemoryRecordRepository()
        outcome=CarbonAccountingUseCase(CarbonMaterialCalculator(parameter_resolver=self.resolver),repository).calculate(input_value)
        self.assertTrue(outcome.successful,outcome.problems)
        return repository,outcome.record

    def test_derived_carbon_is_frozen_exactly_and_missing_history_is_not_recomputed(self):
        repository,record=self.calculate()
        trace=repository.get_trace_snapshot(record.record_id)
        prefix="CAR-FLD-F01-frozen-fuel"
        self.assertEqual(trace["derived_display_evidence"][prefix]["value"],"5.956443")
        raw=repository.get_raw_input_snapshot(record.record_id)
        with patch.object(CarbonMaterialCalculator,"calculate",side_effect=AssertionError("report must not calculate")),patch.object(type(self.resolver),"get_factor",side_effect=AssertionError("report must not read live factors")):
            model=build_saved_record_report(repository,record)
            table=next(s for s in model.sections if s.section_id=="b2").tables[0]
            self.assertEqual(table.rows[0].cells[3].value,"5.96")
            self.assertEqual(table.rows[0].cells[4].value,"计算值")
            self.assertIn("5.956443",table.rows[0].cells[3].source_detail)
            self.assertEqual(table.rows[0].cells[7].source,"实测值")
            self.assertIn("专项填写依据",table.rows[0].cells[7].source_detail)
            legacy_trace={key:value for key,value in trace.items() if key not in {"derived_display_evidence","display_evidence_schema_version","parameter_source_types"}}
            legacy=build_report_model(record,raw,legacy_trace,{}, {}, {})
            old_table=next(s for s in legacy.sections if s.section_id=="b2").tables[0]
            self.assertEqual(old_table.rows[0].cells[3].value,"未记录（采用单位热值含碳量）")
        self.assertEqual(len(repository.list_all()),1)
        self.assertEqual(repository.get_trace_snapshot(record.record_id),trace)

    def test_report_has_no_debug_ids_or_title_decoration_and_keeps_source_audit(self):
        from io import BytesIO
        from zipfile import ZipFile
        from docx import Document
        from packages.infrastructure.reporting import render_report_docx
        repository,record=self.calculate(catalog=True)
        model=build_saved_record_report(repository,record)
        self.assertTrue(all("CAR-FLD-" not in note and "GEN-" not in note for note in model.notices))
        table=next(s for s in model.sections if s.section_id=="b2").tables[0]
        self.assertTrue(any("PDF" in (cell.source_detail or "") for row in table.rows for cell in row.cells))
        self.assertTrue(all("PDF" not in note for note in table.footnotes))
        content=render_report_docx(model)
        with ZipFile(BytesIO(content)) as archive:
            self.assertNotIn(b"<w:pBdr",archive.read("word/styles.xml"))
        document=Document(BytesIO(content))
        self.assertTrue(all(not paragraph.text.startswith("适用版本：") for paragraph in document.paragraphs))

    def test_record_view_and_word_share_values_units_sources_and_missing_states(self):
        repository,record=self.calculate(catalog=True)
        raw=repository.get_raw_input_snapshot(record.record_id);trace=repository.get_trace_snapshot(record.record_id)
        reporting=repository.get_reporting_snapshot(record.record_id);provenance=repository.get_provenance_snapshot(record.record_id)
        qualification=repository.get_report_qualification(record.record_id)
        view=build_record_report_model(record,raw,trace,reporting,provenance=provenance,qualification=qualification)
        word=build_saved_record_report(repository,record)
        self.assertEqual(view.sections,word.sections)
        b2=next(s for s in view.sections if s.section_id=="b2").tables[0]
        self.assertEqual(b2.rows[0].cells[7].source,"GB/T32151.34—2024 附录C")
        self.assertIn("选择方式",b2.rows[0].cells[7].source_detail)
        self.assertIn("<table",report_html(view))
        self.assertIn("冻结燃料名称",report_text(view))
        self.assertEqual(len(repository.list_all()),1)
