from __future__ import annotations

from hashlib import sha256
from importlib.resources import files
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from packages.excel.templates import (
    APPENDIX_B_SHEETS,
    APPROVED_TEMPLATE_SHA256,
    APPROVED_TEMPLATE_V2_RESOURCE,
    APPROVED_TEMPLATE_V3_RESOURCE,
    APPROVED_TEMPLATE_V3_SHA256,
    APPROVED_TEMPLATE_V2_SHA256,
    ExcelTemplateService,
)


class AppendixBTemplateAcceptanceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.service = ExcelTemplateService.default()

    def test_controlled_resource_is_formula_free_and_keeps_key_layout(self) -> None:
        content = self.service.read_bytes()
        self.assertEqual(sha256(content).hexdigest(), APPROVED_TEMPLATE_SHA256)
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
        try:
            self.assertEqual(tuple(workbook.sheetnames), APPENDIX_B_SHEETS)
            self.assertTrue(all(workbook[name].sheet_state == "visible" for name in APPENDIX_B_SHEETS))
            self.assertEqual(workbook["B.2"]["N2"].value, "FCh来源")
            self.assertEqual(workbook["B.7"]["J2"].value, "因子来源")
            self.assertEqual(workbook["B.2"]["L2"].value, "其他来源FCh填写值\n（tC/GJ）")
            formulas = [
                f"{sheet.title}!{cell.coordinate}"
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.data_type == "f"
            ]
            self.assertEqual(formulas, [])
            self.service.validate_structure(workbook)

            self.assertIn("A3:A5", {str(area) for area in workbook["B.7"].merged_cells.ranges})
            self.assertIn("B3:B5", {str(area) for area in workbook["B.7"].merged_cells.ranges})
            self.assertEqual(workbook["B.9"]["G3"].value, 0.11)
            self.assertEqual(workbook["B.9"]["G4"].value, 0.11)
            self.assertIn("焓值可留空由软件解析", workbook["B.9"]["A6"].value)

            validations = {
                sheet.title: {
                    (str(item.sqref), item.type, item.formula1)
                    for item in sheet.data_validations.dataValidation
                }
                for sheet in workbook.worksheets
            }
            self.assertIn(("E4:E18", "list", '"计算值,实测值,直接计算值"'), validations["B.2"])
            self.assertIn(("G4:G18", "list", '"缺省值,实测值"'), validations["B.2"])
            self.assertIn(("J4:J18", "list", '"缺省值,实测值"'), validations["B.2"])
            self.assertIn(("N4:N18", "list", '"缺省值,实测值"'), validations["B.2"])
            self.assertIn(("J3:J8", "list", '"缺省值,实测值"'), validations["B.7"])
            b2_by_range = {str(item.sqref): item for item in workbook["B.2"].data_validations.dataValidation}
            self.assertEqual(b2_by_range["D4:D18"].formula1, 'OR($E4="实测值",$E4="直接计算值",D4="")')
            self.assertEqual(b2_by_range["F4:F18"].error, "仅在相应数据来源选择“实测值”后可填写 F 列。")
            self.assertEqual(b2_by_range["I4:I18"].error, "仅在相应数据来源选择“实测值”后可填写 I 列。")
            b7_by_range = {str(item.sqref): item for item in workbook["B.7"].data_validations.dataValidation}
            self.assertEqual(b7_by_range["J3:J8"].error, "请选择缺省值或实测值。")
            self.assertTrue(all(workbook[name].protection.sheet for name in APPENDIX_B_SHEETS))
            self.assertTrue(workbook["B.2"]["A20"].value.startswith("燃料输入规则 V5："))
            self.assertIn("多组分逐项填实测", workbook["B.7"]["A9"].value)
            self.assertIn(("B3:B4", "list", '"电网电力,非化石电力"'), validations["B.8"])
            self.assertIn(("B3:B4", "list", '"饱和蒸汽,过热蒸汽"'), validations["B.9"])
        finally:
            workbook.close()

    def test_activity_inputs_are_unlocked_and_catalog_adopted_values_are_locked(self) -> None:
        workbook=load_workbook(BytesIO(self.service.read_bytes()))
        try:
            for sheet,cells in {
                "B.2":("A4","B4","C4","D4","E4","F4","G4","I4","J4","L4","M4","N4"),
                "B.3":("B4","C4","D4","E4","F4","G4","H4"),"B.4":("B4","C4","D4","E4","F4","G4","H4"),
                "B.5":("B4","C4","D4","E4","F4","G4","H4"),"B.6":("A3","F3","A7","F7"),
                "B.7":("A3","B3","C3","D3","F3","H3","I3","J3","A6","B6"),
                "B.8":("A3","B3","C3","D3"),"B.9":("A3","B3","C3","D3","E3","F3","G3"),
            }.items():
                self.assertTrue(workbook[sheet].protection.sheet,sheet)
                for cell in cells:
                    self.assertFalse(workbook[sheet][cell].protection.locked,(sheet,cell))
            for sheet,cell in (("B.2","H4"),("B.2","K4"),("B.7","E3"),("B.7","G3"),("B.8","E3")):
                self.assertTrue(workbook[sheet][cell].protection.locked,(sheet,cell))
            for sheet in ("B.3", "B.4", "B.5"):
                for col in "IJKLM":
                    self.assertFalse(workbook[sheet][f"{col}4"].protection.locked, (sheet,col))
            self.assertEqual(self.service.validate_structure(workbook),"2.5.0")
        finally:
            workbook.close()

    def test_real_v21_template_keeps_its_original_identity_and_bytes(self) -> None:
        content = files("resources").joinpath(APPROVED_TEMPLATE_V3_RESOURCE).read_bytes()
        self.assertEqual(sha256(content).hexdigest(), APPROVED_TEMPLATE_V3_SHA256)
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
        try:
            self.assertEqual(self.service.validate_structure(workbook), "2.1.0")
        finally:
            workbook.close()

    def test_real_v2_template_remains_recognized_without_relabeling(self) -> None:
        content = files("resources").joinpath(APPROVED_TEMPLATE_V2_RESOURCE).read_bytes()
        self.assertEqual(sha256(content).hexdigest(), APPROVED_TEMPLATE_V2_SHA256)
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
        try:
            self.assertEqual(self.service.validate_structure(workbook), "2.0.0")
        finally:
            workbook.close()

    def test_download_copy_is_byte_identical_to_approved_resource(self) -> None:
        content = self.service.read_bytes()
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "appendix-b.xlsx"
            self.service.copy_to(destination)
            downloaded = destination.read_bytes()
        self.assertEqual(downloaded, content)
        self.assertEqual(sha256(downloaded).hexdigest(), APPROVED_TEMPLATE_SHA256)


if __name__ == "__main__":
    unittest.main()
