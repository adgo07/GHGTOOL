"""Fuel-only canonical/project/record/Word chain, with independent arithmetic."""
from dataclasses import asdict, replace
from datetime import date
from decimal import Decimal
from io import BytesIO
import tempfile
from pathlib import Path
import unittest

from docx import Document
from openpyxl import load_workbook
from packages.application.carbon_accounting import CarbonAccountingUseCase, CarbonAccountingPreviewUseCase, create_g06_parameter_resolver
from packages.application.catalog_queries import CatalogQueryService
from packages.application.canonical_input_codec import encode_canonical_input, decode_canonical_input
from packages.application.project_workspaces import ProjectWorkspaceService
from packages.application.reporting import build_saved_record_report
from packages.core.models import AccountingPeriod, PeriodType
from packages.excel.appendix_b import AppendixBImportContext, AppendixBWorkbookImporter
from packages.excel.templates import ExcelTemplateService
from packages.infrastructure.reporting import render_report_docx
from packages.persistence import SQLiteProjectWorkspaceRepository, InMemoryRecordRepository, SQLiteCatalogRepository, build_catalog_database
from packages.standards.carbon_material_input_policy import InputOrigin
from packages.standards.carbon_material import (
    CarbonMaterialInput, CarbonMaterialCalculator, FuelInput, FuelPath, FuelType,
    ParameterValue, ParameterSourceKind, EmissionSourceState, EmissionSourceStatus, SOURCE_FUEL,
)

PERIOD = AccountingPeriod(PeriodType.ANNUAL, date(2025, 1, 1), date(2025, 12, 31))


def parameter(name, value, unit, kind=ParameterSourceKind.MEASURED):
    return ParameterValue(name, value, unit, source_kind=kind,
                          source_id=None, source_version=None, source_location=None,
                          selection_reason="独立燃料闭环测试")


def fuel_input(*, path=FuelPath.MASS, direct=False, kind=ParameterSourceKind.MEASURED):
    carbon_unit = ("tC/t" if path is FuelPath.MASS else "tC/10^4Nm3") if direct else "tC/GJ"
    fuel = FuelInput("fuel.b1", path, "10", parameter("carbon", "0.3" if direct else "0.03", carbon_unit, kind),
                     parameter("oxidation", "0.98", "ratio"),
                     None if direct else parameter("lhv", "8", "GJ/t" if path is FuelPath.MASS else "GJ/10⁴Nm³"),
                     fuel_type=FuelType.OTHER, fuel_label="企业混合燃料", origin=InputOrigin.CUSTOM)
    return CarbonMaterialInput("input.b1", "enterprise.b1", "燃料闭环企业", PERIOD,
                               boundary_confirmed=True, fuel_inputs=(fuel,),
                               source_states=(EmissionSourceState(SOURCE_FUEL, EmissionSourceStatus.INVOLVED),))


class FuelLifecycleTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.records = InMemoryRecordRepository()
        self.calculator = CarbonMaterialCalculator()
        self.formal = CarbonAccountingUseCase(self.calculator, record_repository=self.records)

    def tearDown(self):
        self.directory.cleanup()

    def project_round_trip(self, value):
        service = ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(self.root / "projects.sqlite"))
        workspace = service.new_workspace("燃料项目")
        unit = replace(workspace.units[0], canonical_input=value)
        workspace = replace(workspace, units=(unit,))
        service.save(workspace)
        reopened = ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(self.root / "projects.sqlite")).get(workspace.project_id)
        self.assertEqual(reopened.units[0].canonical_input, value)
        self.assertEqual(self.records.list_all(), ())
        return reopened.units[0].canonical_input

    def test_measured_heat_mass_and_volume_project_preview_and_append(self):
        for path in (FuelPath.MASS, FuelPath.VOLUME):
            with self.subTest(path=path):
                value = fuel_input(path=path)
                self.assertEqual(decode_canonical_input(encode_canonical_input(value)), value)
                value = self.project_round_trip(value)
                preview = CarbonAccountingPreviewUseCase(self.calculator).calculate(value)
                self.assertFalse(preview.blocked, preview.problems)
                self.assertIsNone(preview.record)
                count = len(self.records.list_all())
                first = self.formal.calculate(value)
                self.assertTrue(first.successful, first.problems)
                self.assertEqual(first.result.total_amount, Decimal("8.624000000000000000000000000000000000001"))
                second = self.formal.calculate(value)
                self.assertEqual(len(self.records.list_all()), count + 2)
                self.assertNotEqual(first.record.record_id, second.record.record_id)
                # Each subcase starts its own project check with an empty store.
                self.records = InMemoryRecordRepository()
                self.formal = CarbonAccountingUseCase(self.calculator, record_repository=self.records)

    def test_direct_calculated_mass_and_volume_retain_legal_source(self):
        for path in (FuelPath.MASS, FuelPath.VOLUME):
            with self.subTest(path=path):
                value = fuel_input(path=path, direct=True, kind=ParameterSourceKind.CALCULATED)
                outcome = self.formal.calculate(value)
                self.assertTrue(outcome.successful, outcome.problems)
                self.assertEqual(outcome.result.total_amount, Decimal("10.78000000000000000000000000000000000000"))
                self.assertIs(outcome.input.fuel_inputs[0].carbon_content.source_kind, ParameterSourceKind.CALCULATED)

    def test_illegal_fch_and_missing_measured_input_write_no_record(self):
        for value in (fuel_input(kind=ParameterSourceKind.CALCULATED),
                      replace(fuel_input(), fuel_inputs=(replace(fuel_input().fuel_inputs[0], carbon_content=None),))):
            with self.subTest(value=value):
                self.records = InMemoryRecordRepository()
                self.formal = CarbonAccountingUseCase(self.calculator, record_repository=self.records)
                outcome = self.formal.calculate(value)
                self.assertTrue(outcome.blocked)
                self.assertIsNone(outcome.record)
                self.assertEqual(self.records.list_all(), ())

    def test_excel_preview_project_formal_and_frozen_word(self):
        repository = SQLiteCatalogRepository(build_catalog_database(output_path=self.root / "catalog.sqlite"))
        catalog = CatalogQueryService(repository)
        calculator = CarbonMaterialCalculator(parameter_resolver=create_g06_parameter_resolver(repository))
        importer = AppendixBWorkbookImporter(preview_use_case=CarbonAccountingPreviewUseCase(calculator), catalog_service=catalog)
        workbook = load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
        sheet = workbook["B.2"]
        for cell, value in {"A14": "天然气", "B14": 1.25, "C14": "10⁴Nm³", "E14": "计算值",
                            "G14": "缺省值", "J14": "缺省值", "N14": "实测值", "L14": 0.016}.items():
            sheet[cell] = value
        path = self.root / "燃料输入.xlsx"
        workbook.save(path)
        workbook.close()
        preview = importer.import_preview(path, context=AppendixBImportContext(PERIOD, True, "燃料闭环企业"))
        self.assertTrue(preview.units[0].can_calculate, preview.units[0].errors)
        self.assertIsNone(preview.units[0].calculation.record)
        self.assertEqual(self.records.list_all(), ())
        value = self.project_round_trip(preview.units[0].input_value)
        path.unlink()
        formal = CarbonAccountingUseCase(calculator, record_repository=self.records)
        outcome = formal.calculate(value)
        self.assertTrue(outcome.successful, outcome.problems)
        record = outcome.record
        self.assertIs(value.fuel_inputs[0].carbon_content.source_kind, ParameterSourceKind.MEASURED)
        self.assertEqual(next(snapshot for snapshot in record.parameter_snapshots
                              if snapshot.detail_id.endswith("-CARBON")).selection_method.value, "ENTERPRISE_MEASURED")
        before = asdict(record)
        model = build_saved_record_report(self.records, record)
        word = Document(BytesIO(render_report_docx(model)))
        fuel_row = next(section for section in model.sections if section.section_id == "b2").tables[0].rows[0]
        self.assertEqual(fuel_row.cells[7].value, "0.02")
        self.assertIn("天然气", "\n".join(cell.text for table in word.tables for row in table.rows for cell in row.cells))
        # A later input/result cannot rewrite the original Record or its report.
        changed = replace(value, fuel_inputs=(replace(value.fuel_inputs[0], activity="99"),))
        formal.calculate(changed)
        self.assertEqual(len(self.records.list_all()), 2)
        self.assertEqual(asdict(self.records.get(record.record_id)), before)
        self.assertEqual(build_saved_record_report(self.records, record), model)

    def test_v5_excel_direct_calculated_project_record_and_word(self):
        repository = SQLiteCatalogRepository(build_catalog_database(output_path=self.root / "catalog.sqlite"))
        calculator = CarbonMaterialCalculator(parameter_resolver=create_g06_parameter_resolver(repository))
        importer = AppendixBWorkbookImporter(preview_use_case=CarbonAccountingPreviewUseCase(calculator),
                                            catalog_service=CatalogQueryService(repository))
        workbook = load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
        sheet = workbook["B.2"]
        for cell, value in {"A14": "天然气", "B14": 1.25, "C14": "10⁴Nm³", "D14": 0.3,
                            "E14": "直接计算值", "I14": 98, "J14": "实测值"}.items():
            sheet[cell] = value
        path = self.root / "直接计算燃料.xlsx"
        workbook.save(path)
        workbook.close()
        preview = importer.import_preview(path, context=AppendixBImportContext(PERIOD, True, "燃料闭环企业"))
        self.assertTrue(preview.units[0].can_calculate, preview.units[0].errors)
        value = self.project_round_trip(preview.units[0].input_value)
        self.assertIs(value.fuel_inputs[0].carbon_content.source_kind, ParameterSourceKind.CALCULATED)
        self.assertEqual(value.fuel_inputs[0].carbon_content.unit, "tC/10^4Nm3")
        self.assertIsNone(value.fuel_inputs[0].lower_heating_value)
        path.unlink()
        outcome = CarbonAccountingUseCase(calculator, record_repository=self.records).calculate(value)
        self.assertTrue(outcome.successful, outcome.problems)
        self.assertEqual(outcome.result.total_amount, Decimal("1.347500000000000000000000000000000000000"))
        report = build_saved_record_report(self.records, outcome.record)
        fuel_row = next(section for section in report.sections if section.section_id == "b2").tables[0].rows[0]
        self.assertEqual(fuel_row.cells[3].value, "0.30")
        self.assertEqual(fuel_row.cells[7].value, "不适用")
        self.assertEqual(self.records.get_raw_input_snapshot(outcome.record.record_id)["fuel_inputs"][0]["carbon_content"]["source_kind"], "CALCULATED")
        word = Document(BytesIO(render_report_docx(report)))
        self.assertIn("天然气", "\n".join(cell.text for table in word.tables for row in table.rows for cell in row.cells))

    def test_old_v4_cannot_silently_adopt_new_direct_calculated_label(self):
        repository = SQLiteCatalogRepository(build_catalog_database(output_path=self.root / "catalog.sqlite"))
        calculator = CarbonMaterialCalculator(parameter_resolver=create_g06_parameter_resolver(repository))
        importer = AppendixBWorkbookImporter(preview_use_case=CarbonAccountingPreviewUseCase(calculator),
                                            catalog_service=CatalogQueryService(repository))
        old_path = Path(__file__).resolve().parents[1] / "resources/excel_templates/gb_t_32151_34_2024_appendix_b_v4.xlsx"
        workbook = load_workbook(old_path)
        sheet = workbook["B.2"]
        for cell, value in {"A14": "天然气", "B14": 1.25, "C14": "10⁴Nm³", "D14": 0.3,
                            "E14": "直接计算值", "I14": 98, "J14": "实测值"}.items():
            sheet[cell] = value
        path = self.root / "旧v4错误标签.xlsx"
        workbook.save(path)
        workbook.close()
        preview = importer.import_preview(path, context=AppendixBImportContext(PERIOD, True, "燃料闭环企业"))
        self.assertFalse(preview.units[0].can_calculate)
        self.assertTrue(preview.units[0].errors)
        self.assertEqual(self.records.list_all(), ())
