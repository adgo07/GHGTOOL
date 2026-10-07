from __future__ import annotations

from dataclasses import asdict, replace
from io import BytesIO
from pathlib import Path
import ast
import tempfile
import unittest

from openpyxl import load_workbook

from packages.application.reporting.model import (
    ReportCell,
    ReportModel,
    ReportRow,
    ReportSection,
    ReportTable,
)
from packages.excel.r2 import ExcelWorkbookImporter, WorkbookFatalError
from packages.infrastructure.reporting import REPORT_DISCLAIMER, render_report_xlsx


def _full_model() -> ReportModel:
    basic_information = (
        ("报告主体", "青舟炭素材料有限公司"),
        ("核算期间", "2025年度"),
        ("排放总量", "0"),
    )
    basic_table = ReportTable(
        "basic-information",
        "基本信息",
        ("项目", "内容"),
        tuple(ReportRow((ReportCell(label), ReportCell(value))) for label, value in basic_information),
    )
    basic_extra = ReportTable(
        "basic-supplement",
        "基本信息补充表",
        ("补充项目", "补充值", "附件来源"),
        (ReportRow((
            ReportCell("附件", "份", "检测报告"),
            ReportCell("1"),
            ReportCell("报告编号"),
        )),),
    )
    sections = [ReportSection("basic", "基本信息", (basic_table, basic_extra), ("基本信息补充说明",))]

    for index in range(1, 10):
        rows = (
            ReportRow((
                ReportCell("0", "tCO₂", "标准规定值"),
                ReportCell("", None, None),
                ReportCell(f"固定文本-B{index}", "单位原文", "来源原文"),
            ), note=f"行说明-B{index}"),
            ReportRow((
                ReportCell("0001.2300", "t", "企业台账"),
                ReportCell("=1+1", "=@unit", "+来源"),
                ReportCell("-cmd|' /C calc'!A0", "@SUM(1,1)", "@来源"),
            )),
            ReportRow((
                ReportCell("+来源"),
                ReportCell("@SUM(1,1)"),
                ReportCell("=末行文本"),
            ), note="=row-note"),
        )
        table = ReportTable(
            f"b{index}-table",
            f"B.{index} 明细标题",
            ("排放类别", "排放量", "审计值"),
            rows,
            landscape=True,
        )
        if index == 3:
            long_rows = tuple(
                ReportRow((ReportCell(f"动态行-{row}", "kg", "动态来源"),))
                for row in range(125)
            )
            extra = ReportTable("b3-many", "=B.3 第二张表", ("动态条目",), long_rows)
            sections.append(ReportSection("b3", f"B.3 生产过程 {index}", (table, extra), ("B.3 节说明",)))
        else:
            sections.append(ReportSection(f"b{index}", f"B.{index} 生产过程 {index}", (table,), (f"B.{index} 节说明",)))

    sections.append(ReportSection(
        "evidence",
        "数据来源及支撑信息",
        (ReportTable(
            "evidence-table",
            "支撑信息",
            ("信息类型", "描述"),
            (ReportRow((ReportCell("检测报告", "份", "记录来源"), ReportCell("@evidence"))),),
        ),),
        ("来源说明",),
    ))
    return ReportModel(
        schema_version="report-schema-v1",
        record_id="record-PRIVATE-123",
        standard_name="GB/T 32151.34—2024 炭素材料生产企业报告",
        standard_version="GB/T 32151.34—2024",
        period_text="2025年度",
        prepared_on="2026-10-07",
        basic_information=basic_information,
        sections=tuple(sections),
        notices=("提示文本", "=notice-formula"),
        supplementary_note="+补充说明",
    )


def _visible_cells(workbook):
    return [
        cell
        for sheet in workbook.worksheets if sheet.sheet_state == "visible"
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]


def _expected_visible_text(model: ReportModel) -> set[str]:
    expected = {
        model.standard_name,
        model.standard_version,
        model.period_text,
        model.prepared_on,
        REPORT_DISCLAIMER,
        *(value for pair in model.basic_information for value in pair),
        *model.notices,
    }
    if model.supplementary_note is not None:
        expected.add(model.supplementary_note)
    for section in model.sections:
        expected.add(section.title)
        expected.update(section.notes)
        for table in section.tables:
            expected.add(table.title)
            expected.update(table.columns)
            for row in table.rows:
                if row.note is not None:
                    expected.add(row.note)
                for cell in row.cells:
                    if cell.value:
                        expected.add(cell.value)
                    if cell.unit:
                        expected.add(cell.unit)
                    if cell.source:
                        expected.add(cell.source)
    return {value for value in expected if value}


class ReportExcelRendererTests(unittest.TestCase):
    def test_lossless_full_b1_to_b9_report_and_metadata_visibility(self) -> None:
        base = _full_model()
        model = replace(base, standard_name="青舟温室气体排放核算报告" * 8)
        before = asdict(model)
        workbook = load_workbook(BytesIO(render_report_xlsx(model)), data_only=False)

        self.assertEqual(workbook.sheetnames[0], "基本信息")
        self.assertEqual(
            [name for name in workbook.sheetnames if name not in {"基本信息", "数据来源", "__export_metadata__"}],
            [f"B.{index} 生产过程 {index}" for index in range(1, 10)],
        )
        self.assertEqual(workbook["__export_metadata__"].sheet_state, "hidden")
        self.assertNotIn("__metadata__", workbook.sheetnames)
        metadata_values = [cell.value for row in workbook["__export_metadata__"].iter_rows() for cell in row]
        self.assertIn(model.record_id, metadata_values)

        visible_cells = _visible_cells(workbook)
        visible_text = "\n".join(str(cell.value) for cell in visible_cells)
        self.assertNotIn(model.record_id, visible_text)
        for expected in _expected_visible_text(model):
            self.assertIn(expected, visible_text)

        basic = workbook["基本信息"]
        self.assertEqual(basic["B10"].value, "0")
        self.assertEqual(basic["B10"].data_type, "s")
        self.assertEqual(basic["C14"].value, "报告编号")
        self.assertIn("补充项目：单位：份；来源：检测报告", basic["D14"].value)
        self.assertIn(":$D$", basic.print_area)
        self.assertGreater(basic.row_dimensions[1].height, 28)

        b2 = workbook["B.2 生产过程 2"]
        self.assertEqual(b2.freeze_panes, "A7")
        self.assertEqual(b2.page_setup.orientation, "landscape")
        self.assertEqual(str(b2.page_setup.paperSize), str(b2.PAPERSIZE_A4))
        self.assertEqual(b2["A1"].font.name, "Microsoft YaHei")
        self.assertEqual(b2["A3"].value, REPORT_DISCLAIMER)
        self.assertEqual((b2["A7"].value, b2["B7"].value, b2["C7"].value), ("0", None, "固定文本-B2"))
        self.assertIn("排放类别：单位：tCO₂；来源：标准规定值", b2["D7"].value)
        self.assertIn("审计值：单位：单位原文；来源：来源原文", b2["D7"].value)
        self.assertNotIn("排放量：", b2["D7"].value)
        self.assertEqual(b2["E7"].value, "行说明-B2")
        self.assertLess(b2.max_column, 3 * len(("排放类别", "排放量", "审计值")))
        self.assertGreater(b2.row_dimensions[7].height, 15)

        b3 = workbook["B.3 生产过程 3"]
        self.assertEqual(b3["A137"].value, "动态行-124")
        self.assertEqual(b3["B137"].value, "动态条目：单位：kg；来源：动态来源")
        self.assertLess(b3.max_column, 3 * len(("排放类别", "排放量", "审计值")))
        self.assertEqual(asdict(model), before)
        self.assertTrue(all(cell.data_type == "s" for cell in visible_cells))
        self.assertTrue(all(
            cell.data_type != "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        ))

    def test_formula_prefixes_are_literal_and_destination_bytes_match(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / "saved-report.xlsx"
            content = render_report_xlsx(_full_model(), destination)
            self.assertEqual(destination.read_bytes(), content)

            workbook = load_workbook(BytesIO(content), data_only=False)
            visible_cells = _visible_cells(workbook)
            visible_text = "\n".join(str(cell.value) for cell in visible_cells)
            for text in (
                "=1+1",
                "+来源",
                "-cmd|' /C calc'!A0",
                "@SUM(1,1)",
                "=@unit",
                "@来源",
                "=row-note",
                "=notice-formula",
            ):
                self.assertIn(text, visible_text)
            injected_cells = [
                cell for cell in visible_cells
                if any(token in str(cell.value) for token in (
                    "=1+1", "+来源", "-cmd|' /C calc'!A0", "@SUM(1,1)",
                    "=@unit", "@来源", "=row-note", "=notice-formula",
                ))
            ]
            self.assertTrue(injected_cells)
            self.assertTrue(all(cell.data_type == "s" for cell in injected_cells))

    def test_empty_report_section_keeps_its_notes(self) -> None:
        model = _full_model()
        model = replace(model, sections=(*model.sections, ReportSection("appendix", "补充说明区", (), ("空表章节说明",))))
        workbook = load_workbook(BytesIO(render_report_xlsx(model)), data_only=False)
        sheet = workbook["补充说明区"]
        self.assertEqual(sheet["A1"].value, "补充说明区")
        self.assertIn("空表章节说明", [cell.value for row in sheet.iter_rows() for cell in row])
    def test_result_workbook_is_not_an_r2_input_template(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            result_path = Path(folder) / "report.xlsx"
            result_path.write_bytes(render_report_xlsx(_full_model()))
            with self.assertRaises(WorkbookFatalError) as raised:
                ExcelWorkbookImporter().import_preview(result_path)
        self.assertIn("缺少模板工作表", str(raised.exception))

    def test_renderer_has_no_record_repository_catalog_or_calculator_dependency(self) -> None:
        source_path = Path(__file__).parents[1] / "packages" / "infrastructure" / "reporting" / "excel_renderer.py"
        tree = ast.parse(source_path.read_text(encoding="utf-8"))
        imported = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported.extend(alias.name for alias in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                imported.append(node.module)
        forbidden = ("record", "repository", "catalog", "calculator", "sqlite")
        self.assertFalse(
            [module for module in imported if any(word in module.lower() for word in forbidden)],
            imported,
        )


if __name__ == "__main__":
    unittest.main()
