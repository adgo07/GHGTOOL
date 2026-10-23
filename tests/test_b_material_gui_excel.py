from packages.ui.dialogs import MessageBox as QMessageBox
"""真实离屏Qt、实际v6工作簿和冻结正式链（不冒充原生人工验证）。"""
import os
os.environ.setdefault("QT_QPA_PLATFORM","offscreen")
from dataclasses import replace
from decimal import Decimal, localcontext
from io import BytesIO
from pathlib import Path
import tempfile, unittest
from unittest.mock import patch
from PySide6.QtWidgets import QApplication
from openpyxl import load_workbook
from apps.carbon_accounting_desktop.app import create_main_window
from apps.carbon_accounting_desktop.config import AppConfig
from packages.application import ProjectWorkspaceService
from packages.persistence import SQLiteRecordRepository, SQLiteProjectWorkspaceRepository, build_catalog_database
from packages.reference_data import DEFAULT_SOURCE_PATH
from packages.ui.view_models import AppRoute
from packages.excel.templates import ExcelTemplateService
from packages.standards.carbon_material_normalization import MaterialRole as Role, MaterialDataSource as Source

class MaterialGuiExcelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])

    def test_gui_excel_dry_parity_drafts_identity_sqlite_and_upgrade_cancel(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            catalog=build_catalog_database(DEFAULT_SOURCE_PATH,root/'catalog.sqlite')
            records=SQLiteRecordRepository(root/'records.sqlite')
            projects=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/'projects.sqlite'))
            window=create_main_window(AppConfig(catalog_database=catalog,records_database=root/'records.sqlite',projects_database=root/'projects.sqlite'),record_repository=records,project_service=projects)
            shell=window.centralWidget();shell.navigate(AppRoute.NEW_ACCOUNTING);self.app.processEvents()
            gui=shell.pages[AppRoute.NEW_ACCOUNTING]
            try:
                gui.period_year.setValue(2025);gui.enterprise_name.setText('材料整链');gui.boundary_confirmed.setChecked(True)
                for status in gui._source_statuses.values():status.setCurrentIndex(status.findData('NOT_INVOLVED'))
                status=gui._source_statuses['CAR-SRC-CALCINATION-001'];status.setCurrentIndex(status.findData('INVOLVED'))
                process=gui._process_rows['calcination'][0]
                byrole={row['role'].currentData():row for row in process['materials']}
                feed=byrole[Role.CALCINATION_FEED]
                for key,text in [('name','同名'),('mass','80'),('fixed_carbon','90'),('volatile_matter','10')]:feed[key].setText(text)
                feed['fixed_carbon_source'].setCurrentIndex(feed['fixed_carbon_source'].findData(Source.CHEMICAL_CALCULATION))
                self.assertEqual(feed['fixed_carbon'].text(),'')
                feed['fixed_carbon'].setText('90')
                feed['fixed_carbon_source'].setCurrentIndex(feed['fixed_carbon_source'].findData(Source.MEASURED))
                self.assertEqual(feed['fixed_carbon'].text(),'90')
                feed['fixed_carbon'].setText('72')
                feed['fixed_carbon_source'].setCurrentIndex(feed['fixed_carbon_source'].findData(Source.CHEMICAL_CALCULATION))
                self.assertEqual(feed['fixed_carbon'].text(),'90')
                feed['volatile_matter_source'].setCurrentIndex(feed['volatile_matter_source'].findData(Source.CHEMICAL_CALCULATION));feed['volatile_matter'].setText('10')
                for key in ('mass_basis','composition_basis'):feed[key].setCurrentIndex(feed[key].findData('DRY'))
                feed['moisture_received_percent'].setText('20');feed['conversion_method'].setCurrentIndex(feed['conversion_method'].findData('DRY_TO_RECEIVED_MOISTURE'))
                for role,mass,carbon,volatile in [(Role.CALCINED_PRODUCT,'20','80','5'),(Role.UNDERBURN_RECOVERED,'30','40',''),(Role.CARBON_DUST,'10','25','')]:
                    row=byrole[role]
                    for key,text in [('name','同名'),('mass',mass),('fixed_carbon',carbon),('volatile_matter',volatile)]:row[key].setText(text)
                # Names are business labels; line identity owns source/method drafts.
                feed['name'].setText('更名草稿');feed['name'].editingFinished.emit()
                self.assertEqual(feed['fixed_carbon_source'].currentData(),Source.CHEMICAL_CALCULATION)
                self.assertEqual(feed['moisture_received_percent'].text(),'20')
                feed['name'].setText('同名')
                middle=gui._add_material_line(process,line_id='remove-middle')
                survivor=gui._add_material_line(process,line_id='keep-after-middle',role_value=Role.CARBON_DUST)
                survivor['name'].setText('同名');survivor['mass'].setText('0')
                survivor['fixed_carbon_source'].setCurrentIndex(survivor['fixed_carbon_source'].findData(Source.CHEMICAL_CALCULATION))
                survivor['fixed_carbon'].setText('90')
                survivor['composition_basis'].setCurrentIndex(survivor['composition_basis'].findData('DRY'))
                survivor['moisture_received_percent'].setText('20')
                survivor['conversion_method'].setCurrentIndex(survivor['conversion_method'].findData('DRY_TO_RECEIVED_MOISTURE'))
                gui._remove_material_line(process,middle['line_id'])
                remaining=next(row for row in process['materials'] if row['line_id']=='keep-after-middle')
                self.assertIs(remaining,survivor)
                self.assertEqual(remaining['fixed_carbon'].text(),'90')
                self.assertEqual(remaining['moisture_received_percent'].text(),'20')
                self.assertEqual(remaining['composition_basis'].currentData(),'DRY')
                gui._remove_material_line(process,survivor['line_id'])
                identity=feed['line_id'];state=gui._capture_form_state()
                workspace=projects.new_workspace('物料草稿');workspace=replace(workspace,units=(replace(workspace.units[0],form_state=state),));projects.save(workspace)
                restored=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/'projects.sqlite')).get(workspace.project_id)
                gui._restore_form_state(restored.units[0].form_state)
                process=gui._process_rows['calcination'][0];feed=next(row for row in process['materials'] if row['line_id']==identity)
                self.assertEqual(feed['moisture_received_percent'].text(),'20')
                self.assertEqual(feed['composition_basis'].currentData(),'DRY')
                before=gui._capture_form_state();gui._material_execution_version='MATERIAL_V1'
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):gui._upgrade_material_rules()
                self.assertEqual(gui._material_execution_version,'MATERIAL_V1')
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.Yes):gui._upgrade_material_rules()
                self.assertEqual(gui._material_execution_version,'MATERIAL_V2')
                self.assertEqual(gui._material_upgrade_original_state['material_execution_version'],'MATERIAL_V1')
                gui.quick_calculate_button.click();self.app.processEvents()
                self.assertEqual(len(records.list_all()),1)
                gui_record=records.get(gui._latest_record_id)
                w=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()));s=w['B.3']
                for row,mass,percentage in [(4,'80','90'),(5,'80','10'),(6,'20','80'),(7,'30','40'),(8,'10','25'),(9,'20','5')]:
                    s[f'C{row}']='同名';s[f'D{row}']=Decimal(mass);s[f'E{row}']=Decimal(percentage)
                    s[f'F{row}']='计算值' if row in (4,5) else '实测值'
                    if row in (4,5):
                        s[f'I{row}']='干基';s[f'J{row}']='干基';s[f'K{row}']=20;s[f'L{row}']='按收到基水分归一'
                workbook_path=root/'材料v6.xlsx';w.save(workbook_path);w.close()
                page=shell.pages[AppRoute.EXCEL_IMPORT]
                page.import_enterprise_name.setText('材料整链');page.import_period_type.setCurrentIndex(page.import_period_type.findData('ANNUAL'))
                page.import_period_start.setText('2025-01-01');page.import_period_end.setText('2025-12-31');page.import_boundary_confirmed.setChecked(True)
                with patch('packages.ui.pages.QMessageBox.warning'):
                    page._preview_workbook(workbook_path)
                self.assertEqual(len(records.list_all()),1)
                self.assertIsNotNone(page._last_preview)
                self.assertTrue(page._last_preview.units[0].can_calculate, page._last_preview.units[0].errors)
                page.project_name.setText('材料Excel项目');page._save_preview_as_project()
                saved_id=page._workspace.project_id
                reopened=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/'projects.sqlite')).get(saved_id)
                frozen_ingress=reopened.units[0].ingress_provenance['workbook']
                self.assertEqual(frozen_ingress['template_master_sha256'],ExcelTemplateService.default().get().sha256)
                page.open_project(saved_id);page._calculate_selected_unit()
                self.assertEqual(len(records.list_all()),2)
                latest=records.list_all()[0]
                self.assertEqual(latest.calculation_result.total_amount,gui_record.calculation_result.total_amount)
                imported_record=next(record for record in records.list_all() if record.record_id != gui_record.record_id)
                ingress=records.get_raw_input_snapshot(imported_record.record_id)['ingress_provenance']['workbook']
                self.assertEqual(ingress['template_master_sha256'],ExcelTemplateService.default().get().sha256)
                self.assertNotEqual(ingress['template_master_sha256'],ingress['sha256'])
                from packages.application.reporting import build_report_model
                frozen_model=build_report_model(imported_record,records.get_raw_input_snapshot(imported_record.record_id),
                    records.get_trace_snapshot(imported_record.record_id),records.get_provenance_snapshot(imported_record.record_id),
                    records.get_reporting_snapshot(imported_record.record_id),records.get_report_qualification(imported_record.record_id))
                self.assertEqual(frozen_model.input_template_version,'2.9.0')
                self.assertEqual(frozen_model.input_template_sha256,ingress['template_master_sha256'])
                material_table=next(section for section in frozen_model.sections if section.section_id=='b3').tables[0]
                self.assertEqual(material_table.rows[0].cells[5].value,'计算值')
                self.assertTrue(any('按水分质量守恒归一' in note for note in material_table.footnotes))
                from packages.infrastructure.reporting import render_report_docx
                from docx import Document
                document=Document(BytesIO(render_report_docx(frozen_model)))
                self.assertIn('计算值',' '.join(cell.text for table in document.tables for row in table.rows for cell in row.cells))
                self.assertIn('按水分质量守恒归一',' '.join(p.text for p in document.paragraphs))
                from importlib.resources import files
                from packages.excel.templates import APPROVED_TEMPLATE_V5_RESOURCE
                old_book=load_workbook(BytesIO(files('resources').joinpath(APPROVED_TEMPLATE_V5_RESOURCE).read_bytes()))
                old_book['B.8']['B3']='电网电力';old_book['B.8']['C3']=1;old_book['B.8']['D3']=Decimal('.5')
                old_path=root/'旧v5.xlsx';old_book.save(old_path);old_book.close()
                context=page.excel_adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name='历史v5')
                old_preview=page.excel_adapter.import_preview(old_path,context)
                self.assertEqual(old_preview.provenance.template_version,'2.3.0')
                self.assertIsNone(old_preview.provenance.template_master_sha256)
                self.assertEqual(len(records.list_all()),2)
                broken_book=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
                broken_book['B.4']['I3']='任意新增表头';broken_path=root/'非受控结构.xlsx';broken_book.save(broken_path);broken_book.close()
                from packages.excel.ingress import WorkbookFatalError
                with self.assertRaisesRegex(WorkbookFatalError,'B.4!I3'):
                    page.excel_adapter.import_preview(broken_path,context)
                self.assertEqual(len(records.list_all()),2)
                original_workspace=page._workspace
                legacy_units=tuple(replace(unit,canonical_input=replace(unit.canonical_input,material_execution_version='MATERIAL_V1')) for unit in original_workspace.units)
                legacy_workspace=replace(original_workspace,units=legacy_units)
                projects.save(legacy_workspace);page.open_project(legacy_workspace.project_id)
                count=len(projects.list_all())
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):page._upgrade_material_project()
                self.assertEqual(len(projects.list_all()),count)
                self.assertEqual(page._workspace.project_id,legacy_workspace.project_id)
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.Yes):page._upgrade_material_project()
                self.assertEqual(len(projects.list_all()),count+1)
                self.assertIsNotNone(page._selected_unit())
                self.assertEqual(page._selected_unit().canonical_input.material_execution_version,'MATERIAL_V2')
                self.assertEqual(projects.get(legacy_workspace.project_id).units[0].canonical_input.material_execution_version,'MATERIAL_V1')
                self.assertEqual(len(records.list_all()),2)
            finally:
                with patch.object(QMessageBox,'question',return_value=QMessageBox.StandardButton.No):window.close()
                window.deleteLater();self.app.processEvents()

if __name__=='__main__':unittest.main()
