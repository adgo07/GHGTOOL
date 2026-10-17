"""材料正式链：独立预期、非法零Record、SQLite和冻结Word。"""
from dataclasses import replace
from decimal import Decimal, localcontext
from pathlib import Path
import json, tempfile, unittest
from unittest.mock import patch
from tests.test_g06_carbon_material import _input, _measured_parameter
from tests.test_b_material_normalization import rows, line, Role, Source
from packages.standards.carbon_material import CarbonMaterialCalculator, CalcinationInput, ALGORITHM_VERSION
from packages.application.carbon_accounting import CarbonAccountingUseCase
from packages.application.canonical_input_codec import encode_canonical_input, decode_canonical_input
from packages.application.carbon_material_implementation import CarbonMaterialExcel
from packages.persistence import SQLiteRecordRepository, SQLiteProjectWorkspaceRepository
from packages.application.project_workspaces import ProjectWorkspaceService

class MaterialChainTests(unittest.TestCase):
    def payload(self, material_rows=None, **kwargs):
        return _input(calcination=CalcinationInput(material_rows=rows() if material_rows is None else material_rows,
            k1=_measured_parameter("CAR-PAR-K1", ".35", "ratio")), **kwargs)

    def test_formal_distinct_roles_and_zero_cc(self):
        calculator=CarbonMaterialCalculator()
        value=calculator.calculate(self.payload())
        self.assertTrue(value.successful, value.problems)
        self.assertEqual(ALGORITHM_VERSION,"CAR-SM01-2026-10-17-MATERIAL.1")
        self.assertEqual(value.algorithm_version,ALGORITHM_VERSION)
        self.assertEqual(value.result.algorithm_version,ALGORITHM_VERSION)
        self.assertEqual(value.result.total_amount, Decimal("226.8291666666666666666666666666666666667"))
        value=calculator.calculate(self.payload(rows(cc="0",ucc="5",du="0",carbon=(None,"80",None))))
        self.assertTrue(value.successful, value.problems)
        # Input feed remains 100t, .9/.1: independent role balance 90-4 plus methane10*.35.
        with localcontext() as c:
            c.prec=40
            expected=Decimal("86")*44/12+Decimal("10")*Decimal(".35")*44/16
        self.assertEqual(value.result.total_amount,expected)
        zero_cc=(line("feed",Role.CALCINATION_FEED,"10","90","0"),
                 line("cc",Role.CALCINED_PRODUCT,"0"),
                 line("ucc",Role.UNDERBURN_RECOVERED,"5","80"),
                 line("dust",Role.CARBON_DUST,"0"))
        actual=calculator.calculate(self.payload(zero_cc))
        self.assertTrue(actual.successful,actual.problems)
        self.assertEqual(actual.result.total_amount,Decimal("18.33333333333333333333333333333333333333"))

    def test_illegal_definition_no_record_and_valid_dry_freezes_both_values(self):
        dry=line("feed",Role.CALCINATION_FEED,"80","90","10",mass_basis="DRY",composition_basis="DRY",
            moisture_received_percent="20",conversion_method="DRY_TO_RECEIVED_MOISTURE",
            fixed_carbon_source=Source.CHEMICAL_CALCULATION,volatile_matter_source=Source.CHEMICAL_CALCULATION)
        received=self.payload((dry,*rows()[1:]))
        with tempfile.TemporaryDirectory() as directory:
            records=SQLiteRecordRepository(Path(directory)/"records.sqlite")
            usecase=CarbonAccountingUseCase(CarbonMaterialCalculator(),records)
            bad=replace(dry,conversion_method=None)
            rejected=usecase.calculate(self.payload((bad,*rows()[1:])))
            self.assertFalse(rejected.successful)
            self.assertEqual(len(records.list_all()),0)
            result=usecase.calculate(received)
            self.assertTrue(result.successful,result.problems)
            self.assertEqual(len(records.list_all()),1)
            record_id=result.record.record_id
            reopened=SQLiteRecordRepository(Path(directory)/"records.sqlite")
            original=reopened.get_raw_input_snapshot(record_id)
            self.assertIn('DRY_TO_RECEIVED_MOISTURE',json.dumps(original))
            trace=reopened.get_trace_snapshot(record_id)
            frozen=json.dumps(trace)
            self.assertIn('material.feed.adopted.mass_t',frozen)
            self.assertIn('100',frozen)
            # Subsequent catalog/rule access is forbidden when exporting the saved evidence.
            from packages.application.reporting.model import build_report_model
            model=build_report_model(reopened.get(record_id),raw_input_snapshot=original,
                trace_snapshot=trace, reporting_snapshot=reopened.get_reporting_snapshot(record_id),
                provenance_snapshot=reopened.get_provenance_snapshot(record_id),
                report_qualification=reopened.get_report_qualification(record_id))
            table=next(t for section in model.sections for t in section.tables if t.table_id.startswith('b3'))
            text=' '.join(cell.value for row in table.rows for cell in row.cells)
            self.assertIn('72',text)
            self.assertIn('收到基换算',text)
            self.assertTrue(any('原质量80' in note for note in table.footnotes))
            from packages.infrastructure.reporting import render_report_docx
            from docx import Document
            from io import BytesIO
            document=Document(BytesIO(render_report_docx(model)))
            word_text=' '.join(cell.text for t in document.tables for row in t.rows for cell in row.cells)
            self.assertIn('收到基换算',word_text)
            self.assertIn('72',word_text)
            self.assertIn('原质量80',' '.join(p.text for p in document.paragraphs))

    def test_numeric_guard_does_not_relabel_historical_record(self):
        from packages.standards._numeric_authority import N01C_ALGORITHM_VERSION, install_declared_decimal_context
        from packages.persistence.in_memory_records import InMemoryRecordRepository
        records=InMemoryRecordRepository()
        current=CarbonAccountingUseCase(CarbonMaterialCalculator(),records).calculate(self.payload())
        old=replace(current.record,record_id="old-fixture",algorithm_version=N01C_ALGORITHM_VERSION,
            calculation_result=replace(current.record.calculation_result,algorithm_version=N01C_ALGORITHM_VERSION))
        records.create(old)
        install_declared_decimal_context()
        self.assertEqual(records.get("old-fixture"),old)
        self.assertEqual(records.get("old-fixture").algorithm_version,N01C_ALGORITHM_VERSION)
        self.assertEqual(current.record.algorithm_version,ALGORITHM_VERSION)

    def test_codec_roundtrip_and_legacy_read_execution_identity(self):
        current=self.payload()
        self.assertEqual(decode_canonical_input(encode_canonical_input(current)), current)
        from packages.application import carbon_material_input_codec as codec
        document=json.loads(codec.encode_canonical_input(current))
        document['version']=1
        document['value']['fields'].pop('material_execution_version')
        def legacy(node):
            if isinstance(node,dict):
                if node.get('type')=='material_input_line':
                    for key in ('mass_basis','composition_basis','moisture_received_percent','conversion_method','fixed_carbon_unit','volatile_matter_unit'):
                        node['fields'].pop(key)
                for value in node.values():legacy(value)
            elif isinstance(node,list):
                for value in node:legacy(value)
        legacy(document)
        restored=codec.decode_canonical_input(json.dumps(document))
        self.assertEqual(restored.material_execution_version,'MATERIAL_V1')
        self.assertEqual(restored.calcinations[0].material_rows[0].mass_basis,'RECEIVED')
        unknown=replace(current,material_execution_version='UNKNOWN')
        self.assertEqual(decode_canonical_input(encode_canonical_input(unknown)),unknown)
        self.assertFalse(CarbonMaterialCalculator().calculate(unknown).successful)

    def test_explicit_upgrade_copies_sqlite_project_and_original_does_not_change(self):
        with tempfile.TemporaryDirectory() as directory:
            service=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(Path(directory)/'projects.sqlite'))
            workspace=service.new_workspace('原项目')
            old=replace(self.payload(),material_execution_version='MATERIAL_V1')
            workspace=replace(workspace,units=(replace(workspace.units[0],canonical_input=old),))
            service.save(workspace)
            upgraded=CarbonMaterialExcel.upgrade_workspace(workspace)
            self.assertNotEqual(upgraded.project_id,workspace.project_id)
            service.save(upgraded)
            self.assertEqual(service.get(workspace.project_id).units[0].canonical_input,old)
            self.assertEqual(service.get(upgraded.project_id).units[0].canonical_input.material_execution_version,'MATERIAL_V2')

if __name__=='__main__':unittest.main()
