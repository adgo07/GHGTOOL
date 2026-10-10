from __future__ import annotations

from contextlib import closing
from dataclasses import asdict
from hashlib import sha256
from io import BytesIO
import gc
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtWidgets import QApplication
from docx import Document
from openpyxl import load_workbook

from apps.carbon_accounting_desktop.app import create_main_window
from apps.carbon_accounting_desktop.config import AppConfig
from packages.application.project_workspaces import ProjectWorkspaceService
from packages.application.reporting import build_saved_record_report
from packages.application.reporting.model import frozen_totals
from packages.core.models import RecordStatus
from packages.excel.templates import APPROVED_SOURCE_SHA256, APPROVED_TEMPLATE_SHA256, ExcelTemplateService
from packages.persistence import SQLiteRecordRepository, build_catalog_database
from packages.persistence.projects_repository import SQLiteProjectWorkspaceRepository
from packages.reference_data import DEFAULT_SOURCE_PATH
from packages.ui.view_models import AppRoute


ROOT = Path(__file__).resolve().parents[1]


def _rendered_cell_text(cell) -> str:
    value = cell.value
    if cell.unit:
        value = f"{value} {cell.unit}".strip()
    if cell.source:
        value = f"{value}\n来源：{cell.source}" if value else f"来源：{cell.source}"
    return value


def _covered_by_merges(merges, row_offset: int = 0) -> set[tuple[int, int]]:
    covered = set()
    for row_start, column_start, row_end, column_end in merges:
        covered.update(
            (row + row_offset, column)
            for row in range(row_start, row_end + 1)
            for column in range(column_start, column_end + 1)
            if (row, column) != (row_start, column_start)
        )
    return covered


def _assert_word_tables_match_model(case: unittest.TestCase, document, report) -> None:
    report_tables = tuple(
        table
        for section in report.sections
        for table in section.tables
        if table.rows
    )
    case.assertEqual(len(document.tables), len(report_tables))
    for word_table, report_table in zip(document.tables, report_tables):
        headers = report_table.header_rows
        if not headers:
            from packages.application.reporting.model import ReportCell, ReportRow

            headers = (ReportRow(tuple(ReportCell(column) for column in report_table.columns)),)
        expected_rows = (*headers, *report_table.rows)
        case.assertEqual(len(word_table.rows), len(expected_rows), report_table.table_id)
        case.assertEqual(len(word_table.columns), len(report_table.columns), report_table.table_id)
        covered = _covered_by_merges(report_table.header_merges)
        covered.update(_covered_by_merges(report_table.body_merges, len(headers)))
        for row_index, row in enumerate(expected_rows):
            for column_index, cell in enumerate(row.cells):
                if (row_index, column_index) in covered:
                    continue
                case.assertEqual(
                    word_table.rows[row_index].cells[column_index].text,
                    _rendered_cell_text(cell),
                    f"{report_table.table_id}[{row_index},{column_index}]",
                )


class Exb01RPT02IntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.catalog = build_catalog_database(DEFAULT_SOURCE_PATH, self.root / "catalog.sqlite")
        self.records_path = self.root / "records.sqlite"
        self.projects_path = self.root / "projects.sqlite"
        self.repository = SQLiteRecordRepository(self.records_path)
        self.project_service = ProjectWorkspaceService(
            SQLiteProjectWorkspaceRepository(self.projects_path)
        )
        self.window = create_main_window(
            AppConfig(
                catalog_database=self.catalog,
                records_database=self.records_path,
                projects_database=self.projects_path,
            ),
            record_repository=self.repository,
            project_service=self.project_service,
        )
        self.window.show()
        self.shell = self.window.centralWidget()
        self.shell.navigate(AppRoute.EXCEL_IMPORT)
        self.app.processEvents()
        self.page = self.shell.pages[AppRoute.EXCEL_IMPORT]

    def tearDown(self) -> None:
        self.window.close()
        self.window.deleteLater()
        self.app.processEvents()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        gc.collect()
        self.directory.cleanup()

    def _write_real_appendix_b_workbook(self) -> tuple[Path, str]:
        template_bytes = ExcelTemplateService.default().read_bytes()
        self.assertEqual(sha256(template_bytes).hexdigest(), APPROVED_TEMPLATE_SHA256)
        workbook = load_workbook(BytesIO(template_bytes))
        fuel = workbook["B.2"]
        fuel["A14"], fuel["B14"] = "天然气", 1.25
        fuel["D14"], fuel["F14"], fuel["I14"] = "计算值", "缺省值", "缺省值"
        destination = self.root / "经批准附录B工作簿.xlsx"
        workbook.save(destination)
        workbook.close()
        return destination, sha256(destination.read_bytes()).hexdigest()

    def _set_import_context(self) -> None:
        self.page.import_enterprise_name.setText("EXB01 与 RPT02 联合验收企业")
        self.page.import_period_type.setCurrentIndex(self.page.import_period_type.findData("ANNUAL"))
        self.page.import_period_start.setText("2025-01-01")
        self.page.import_period_end.setText("2025-12-31")
        self.page.import_boundary_confirmed.setChecked(True)

    def test_approved_workbook_preview_project_record_and_frozen_word_report(self) -> None:
        workbook_path, workbook_digest = self._write_real_appendix_b_workbook()
        self._set_import_context()
        self.assertEqual(self.repository.list_all(), ())

        with patch(
            "packages.ui.pages.QFileDialog.getOpenFileName",
            return_value=(str(workbook_path), ""),
        ):
            self.page.selectFileButton.click()
            self.app.processEvents()

        preview = self.page._last_preview
        self.assertIsNotNone(preview)
        self.assertEqual(preview.provenance.workbook_sha256, workbook_digest)
        self.assertEqual(len(preview.units), 1)
        imported = preview.units[0]
        self.assertTrue(imported.can_calculate, imported.errors)
        self.assertIsNotNone(imported.calculation)
        self.assertIsNone(imported.calculation.record)
        self.assertEqual(self.repository.list_all(), ())

        expected_input = imported.input_value
        self.page.project_name.setText("附录B导入并冻结报告")
        self.page.save_project_button.click()
        self.app.processEvents()

        project_id = self.page._workspace.project_id
        projects = self.project_service
        saved = projects.get(project_id)
        self.assertIsNotNone(saved)
        self.assertEqual(saved.name, "附录B导入并冻结报告")
        self.assertEqual(len(saved.units), 1)
        saved_unit = saved.units[0]
        self.assertEqual(saved_unit.canonical_input, expected_input)
        self.assertEqual(saved_unit.record_ids, ())
        self.assertEqual(saved_unit.ingress_provenance["source"], "EXCEL_APPENDIX_B")
        self.assertEqual(saved_unit.ingress_provenance["workbook"]["sha256"], workbook_digest)
        activity = next(
            item
            for item in saved_unit.ingress_provenance["numeric_cell_evidence"]
            if item["sheet"] == "B.2" and item["cell"] == "B14"
        )
        self.assertEqual(activity["raw_numeric_lexical"], "1.25")
        self.assertEqual(activity["normalized_decimal_lexical"], "1.25")
        self.assertEqual(self.repository.list_all(), ())
        self.assertIn("尚未生成正式核算记录", self.page.status_label.text())

        self.assertTrue(self.page.calculate_button.isEnabled())
        self.page.calculate_button.click()
        self.app.processEvents()
        records = self.repository.list_all()
        self.assertEqual(len(records), 1)
        record = records[0]
        self.assertIn(record.status, {RecordStatus.COMPLETED, RecordStatus.COMPLETED_WITH_WARNINGS})
        self.assertEqual(saved.units[0].record_ids, ())
        linked = projects.get(project_id)
        self.assertEqual(linked.units[0].record_ids, (record.record_id,))
        raw_snapshot = self.repository.get_raw_input_snapshot(record.record_id)
        self.assertEqual(raw_snapshot["ingress_provenance"]["source"], "EXCEL_APPENDIX_B")
        self.assertEqual(raw_snapshot["ingress_provenance"]["workbook"]["sha256"], workbook_digest)
        self.assertEqual(
            next(
                item
                for item in raw_snapshot["ingress_provenance"]["numeric_cell_evidence"]
                if item["sheet"] == "B.2" and item["cell"] == "B14"
            )["raw_numeric_lexical"],
            "1.25",
        )

        frozen_record = asdict(record)
        report_supplementary = {"prepared_on": "2026-10-10"}
        report = build_saved_record_report(
            self.repository,
            record,
            supplementary_info=report_supplementary,
        )
        self.assertEqual(report.layout_id, "gbt-32151-34-2024-appendix-b")
        self.assertEqual(report.template_sha256, APPROVED_SOURCE_SHA256)
        b1 = next(section for section in report.sections if section.section_id == "b1").tables[0]
        totals = frozen_totals(record, self.repository.get_trace_snapshot(record.record_id))
        self.assertEqual((b1.rows[-2].cells[2].value, b1.rows[-1].cells[2].value), (totals["ES"], totals["ET"]))

        self.page.open_record_button.click()
        self.app.processEvents()
        self.assertEqual(self.shell.current_route, AppRoute.RECORDS)
        record_page = self.shell.pages[AppRoute.RECORDS]
        self.assertTrue(record_page.open_record(record.record_id))
        destination = self.root / "EXB01-正式记录-RPT02.docx"
        with (
            patch.object(record_page, "_report_supplementary_dialog", return_value=report_supplementary),
            patch("packages.ui.report_export.QFileDialog.getSaveFileName", return_value=(str(destination), "")),
            patch("packages.ui.report_export.QMessageBox.information") as success,
            patch("packages.ui.report_export.QMessageBox.critical") as failure,
        ):
            record_page.export_word_button.click()
            self.app.processEvents()

        failure.assert_not_called()
        success.assert_called_once()
        self.assertTrue(destination.is_file())
        document = Document(destination)
        headings = {paragraph.text for paragraph in document.paragraphs}
        appendix_sections = tuple(
            section for section in report.sections
            if section.section_id in {f"b{number}" for number in range(1, 10)}
        )
        self.assertEqual(tuple(section.section_id for section in appendix_sections), tuple(f"b{number}" for number in range(1, 10)))
        for section in appendix_sections:
            self.assertIn(section.title, headings)
        _assert_word_tables_match_model(self, document, report)
        self.assertEqual(document.core_properties.comments, f"布局：{report.layout_id}；批准模板 SHA256：{report.template_sha256}")
        self.assertEqual(asdict(self.repository.get(record.record_id)), frozen_record)
        self.assertEqual(len(self.repository.list_all()), 1)
        self.assertEqual(
            self.repository.get_raw_input_snapshot(record.record_id)["ingress_provenance"]["source"],
            "EXCEL_APPENDIX_B",
        )
        with closing(sqlite3.connect(self.records_path)) as connection:
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM report_export_history WHERE record_id = ?", (record.record_id,)).fetchone()[0],
                1,
            )


if __name__ == "__main__":
    unittest.main()
