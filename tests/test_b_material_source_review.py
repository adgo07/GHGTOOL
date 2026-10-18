"""F02 controlled template, real importer and source/method separation."""
from io import BytesIO
from importlib.resources import files
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import unittest
from decimal import Decimal
from PySide6.QtWidgets import QApplication
from openpyxl import load_workbook
from tests import test_b_material_draft_scope as draft_helpers
from packages.excel.templates import ExcelTemplateService,APPROVED_TEMPLATE_V6_RESOURCE,APPROVED_TEMPLATE_V6_SHA256
from packages.ui.view_models import AppRoute
from packages.standards.carbon_material_normalization import MaterialDataSource as Source
from scripts.project_appendix_b_v7 import generated_bytes,TARGET,SOURCE as MASTER_SOURCE

class MaterialSourceReviewTests(unittest.TestCase):
 page=draft_helpers.MaterialDraftScopeTests.page
 @classmethod
 def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
 @staticmethod
 def workbook(content,label='计算值',missing=None):
  w=load_workbook(BytesIO(content));s=w['B.3']
  for row,mass,percentage in [(4,'80','90'),(5,'80','10'),(6,'20','80'),(7,'30','40'),(8,'10','25'),(9,'20','5')]:
   s[f'C{row}']='合成材料';s[f'D{row}']=Decimal(mass);s[f'E{row}']=Decimal(percentage)
   s[f'F{row}']=label if row in (4,5) else '实测值'
   if row in (4,5):
    s[f'I{row}']='干基';s[f'J{row}']='干基';s[f'K{row}']=20;s[f'L{row}']='按收到基水分归一'
    if missing:s[f'{missing}{row}']=None
  return w
 def test_controlled_v7_is_mechanical_and_preserves_v6_and_other_parts(self):
  self.assertEqual(sha256(MASTER_SOURCE.read_bytes()).hexdigest(),APPROVED_TEMPLATE_V6_SHA256)
  self.assertEqual(TARGET.read_bytes(),generated_bytes())
  with ZipFile(MASTER_SOURCE) as old,ZipFile(TARGET) as new:
   self.assertEqual(old.namelist(),new.namelist())
   self.assertEqual([n for n in old.namelist() if old.read(n)!=new.read(n)],['xl/worksheets/sheet3.xml','xl/worksheets/sheet4.xml','xl/worksheets/sheet5.xml'])
  w=load_workbook(TARGET)
  try:
   self.assertEqual(ExcelTemplateService.default().validate_structure(w),'2.5.0')
   for name in ('B.3','B.4','B.5'):
    sources=[v.formula1 for v in w[name].data_validations.dataValidation if str(v.sqref).startswith('F4:')]
    self.assertEqual(sources,['"实测值,计算值"'])
  finally:w.close()
 def test_actual_import_calculated_missing_data_and_scoped_old_alias(self):
  with self.page() as (gui,projects,root,shell,records):
   row=gui._process_rows['calcination'][0]['materials'][0];combo=row['fixed_carbon_source']
   self.assertEqual(combo.itemText(combo.findData(Source.CHEMICAL_CALCULATION)),'计算值')
   gui.period_year.setValue(2025);page=shell.pages[AppRoute.EXCEL_IMPORT]
   context=page.excel_adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name='材料来源验证')
   cases=[('new',ExcelTemplateService.default().read_bytes(),'计算值',None,True,'2.7.0'),
          ('missing-method',ExcelTemplateService.default().read_bytes(),'计算值','L',False,'2.7.0'),
          ('missing-water',ExcelTemplateService.default().read_bytes(),'计算值','K',False,'2.7.0'),
          ('new-reject-old-word',ExcelTemplateService.default().read_bytes(),'收到基换算',None,False,'2.7.0'),
          ('old-alias',files('resources').joinpath(APPROVED_TEMPLATE_V6_RESOURCE).read_bytes(),'收到基换算',None,True,'2.4.0')]
   for name,content,label,missing,valid,version in cases:
    with self.subTest(name=name):
     w=self.workbook(content,label,missing);path=root/(name+'.xlsx');w.save(path);w.close()
     preview=page.excel_adapter.import_preview(path,context)
     self.assertEqual(preview.provenance.template_version,version)
     self.assertEqual(preview.units[0].can_calculate,valid,preview.units[0].errors)
     if valid:
      material=preview.units[0].input_value.calcinations[0].material_rows[0]
      self.assertEqual(material.fixed_carbon_source,Source.CHEMICAL_CALCULATION)
      self.assertEqual(material.conversion_method,'DRY_TO_RECEIVED_MOISTURE')
     if version=='2.4.0':self.assertIsNone(preview.provenance.template_master_sha256)
     else:self.assertEqual(preview.provenance.template_master_sha256,ExcelTemplateService.default().get().sha256)
   self.assertEqual(len(records.list_all()),0)

 def test_mixed_v6_v7_markers_are_rejected(self):
  w=load_workbook(TARGET)
  try:
   old=load_workbook(MASTER_SOURCE)
   try:w['B.4']['A14']=old['B.4']['A14'].value
   finally:old.close()
   with self.assertRaisesRegex(ValueError,'冲突模板版本'):ExcelTemplateService.default().validate_structure(w)
  finally:w.close()
