"""真实v2—v7母版及受控v1结构fixture的脱硫版本兼容边界。"""
from contextlib import closing
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
import shutil,sqlite3,tempfile,unittest
from openpyxl import load_workbook
from packages.application import CarbonAccountingPreviewUseCase
from packages.application.carbon_accounting import create_g06_parameter_resolver
from packages.application.catalog_queries import CatalogQueryService
from packages.excel.appendix_b import AppendixBWorkbookImporter
from packages.persistence import build_catalog_database,SQLiteCatalogRepository
from packages.reference_data import DEFAULT_SOURCE_PATH
from packages.standards.carbon_material import CarbonMaterialCalculator,ParameterSourceKind
from tests.exb01_fixture_helper import appendix_b_context
ROOT=Path(__file__).resolve().parents[1]

class FGDOldTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory=tempfile.TemporaryDirectory(prefix="fgd-old-template-")
        cls.root=Path(cls.directory.name)
        cls.catalog=build_catalog_database(DEFAULT_SOURCE_PATH,cls.root/"catalog.sqlite")
        cls.repository=SQLiteCatalogRepository(cls.catalog)
        cls.resolver=create_g06_parameter_resolver(cls.repository)
    @classmethod
    def tearDownClass(cls):cls.directory.cleanup()
    def setUp(self):
        self.case=tempfile.TemporaryDirectory(dir=self.root)
        self.addCleanup(self.case.cleanup)
        self.output=Path(self.case.name)
    def master(self,version):
        if version==1:
            path=ROOT/"tests/fixtures/exb01/appendix_b_v1.xlsx"
            self.assertEqual(sha256(path.read_bytes()).hexdigest(),"e6a070bf28adb47939e24713f676136b5a023017d6f695e0031c68d51c389c3d")
            return path  # Controlled historical structure fixture, not original mother.
        return ROOT/f"resources/excel_templates/gb_t_32151_34_2024_appendix_b_v{version}.xlsx"
    def preview(self,version,changes=None,resolver=None):
        master=self.master(version);before=master.read_bytes()
        wb=load_workbook(master);ws=wb["B.7"]
        assignments={"A3":"旧单组分","B3":10,"E3":0.440}
        if version>=3:assignments["J3"]="缺省值"
        assignments.update(changes or {})
        for cell,value in assignments.items():ws[cell]=value
        path=self.output/f"v{version}.xlsx";wb.save(path);wb.close()
        actual=resolver or self.resolver
        importer=AppendixBWorkbookImporter(preview_use_case=CarbonAccountingPreviewUseCase(CarbonMaterialCalculator(parameter_resolver=actual)),catalog_service=CatalogQueryService(self.repository))
        preview=importer.import_preview(path,context=appendix_b_context())
        self.assertEqual(master.read_bytes(),before)
        if preview.units[0].calculation is not None:self.assertIsNone(preview.units[0].calculation.record)
        return preview.units[0]
    def test_old_real_masters_and_v1_fixture_keep_filled_ratios(self):
        for version in range(1,8):
            with self.subTest(version=version):
                unit=self.preview(version,{"D3":90,"F3":100})
                self.assertTrue(unit.can_calculate,unit.errors)
                self.assertEqual(unit.calculation.result.total_amount,Decimal("3.96"))
                component=unit.input_value.fgd_units[0].components[0]
                self.assertEqual(component.carbonate_fraction.source_kind,ParameterSourceKind.MEASURED)
                self.assertEqual(component.conversion_rate.source_kind,ParameterSourceKind.MEASURED)
                self.assertEqual(component.emission_factor.source_kind,ParameterSourceKind.STANDARD_SPECIFIED)
    def test_old_real_masters_and_v1_fixture_keep_blank_qualified_defaults(self):
        for version in range(1,8):
            with self.subTest(version=version):
                unit=self.preview(version)
                self.assertTrue(unit.can_calculate,unit.errors)
                self.assertEqual(unit.calculation.result.total_amount,Decimal("3.96"))
                component=unit.input_value.fgd_units[0].components[0]
                for value,expected in ((component.carbonate_fraction,Decimal(".9")),(component.conversion_rate,Decimal("1"))):
                    self.assertEqual(value.value,expected)
                    self.assertEqual(value.source_kind,ParameterSourceKind.STANDARD_DEFAULT)
                    self.assertTrue(value.factor_id)
                self.assertEqual(sum(w.code=="EXB01_STANDARD_DEFAULT_APPLIED" for w in unit.warnings),2)
    def test_old_present_invalid_ratios_never_become_defaults(self):
        for version in range(1,8):
            for column in ("D","F"):
                for invalid in ("bad"," ",True):
                    with self.subTest(version=version,column=column,invalid=invalid):
                        unit=self.preview(version,{"D3":90,"F3":100,column+"3":invalid})
                        self.assertFalse(unit.can_calculate)
                        if unit.input_value is not None:
                            component=unit.input_value.fgd_units[0].components[0]
                            self.assertIsNone(component.carbonate_fraction if column=="D" else component.conversion_rate)
    def test_old_zero_is_measured_and_not_default(self):
        for version in range(1,8):
            with self.subTest(version=version):
                unit=self.preview(version,{"D3":0,"F3":0})
                self.assertTrue(unit.can_calculate,unit.errors)
                self.assertEqual(unit.calculation.result.total_amount,Decimal("0"))
                component=unit.input_value.fgd_units[0].components[0]
                self.assertEqual(component.carbonate_fraction.source_kind,ParameterSourceKind.MEASURED)
                self.assertEqual(component.conversion_rate.source_kind,ParameterSourceKind.MEASURED)
    def test_old_multi_component_blank_i_is_rejected(self):
        for version in range(1,8):
            with self.subTest(version=version):
                changes={"E4":.522}
                if version>=3:changes["J4"]="缺省值"
                unit=self.preview(version,changes)
                self.assertFalse(unit.can_calculate)
                self.assertIn("EXB01_MULTI_CARBONATE_CONTENT_REQUIRED",{e.code for e in unit.errors})
                if unit.input_value is not None:
                    self.assertTrue(all(c.carbonate_fraction is None for c in unit.input_value.fgd_units[0].components))
    def test_old_blank_requires_real_generic_binding(self):
        for fault in ("missing","wrong"):
            path=self.output/(fault+".sqlite");shutil.copyfile(self.catalog,path)
            with closing(sqlite3.connect(path)) as connection:
                with connection:
                    query=("DELETE FROM reference_data_bindings WHERE factor_id IN (SELECT factor_id FROM factor_values WHERE lower(parameter_id) IN ('car-par-p04b-i','car-par-p04b-tr'))" if fault=="missing" else
                           "UPDATE reference_data_bindings SET asset_id='asset-wrong-fgd' WHERE factor_id IN (SELECT factor_id FROM factor_values WHERE lower(parameter_id) IN ('car-par-p04b-i','car-par-p04b-tr'))")
                    self.assertGreaterEqual(connection.execute(query).rowcount,2)
            resolver=create_g06_parameter_resolver(SQLiteCatalogRepository(path))
            for version in range(1,8):
                with self.subTest(fault=fault,version=version):
                    unit=self.preview(version,resolver=resolver)
                    self.assertFalse(unit.can_calculate)
                    if unit.input_value is not None:
                        component=unit.input_value.fgd_units[0].components[0]
                        self.assertIsNone(component.carbonate_fraction)
                        self.assertIsNone(component.conversion_rate)
    def test_current_explicit_measured_blank_still_rejects_and_zero_stays_measured(self):
        for column in ("D","F"):
            with self.subTest(column=column):
                unit=self.preview(8,{"D3":90,"F3":100,"K3":"实测值","L3":"实测值",column+"3":None})
                self.assertFalse(unit.can_calculate)
        for column in ("K","L"):
            with self.subTest(missing_source=column):
                unit=self.preview(8,{"D3":90,"F3":100,"K3":"实测值","L3":"实测值",column+"3":None})
                self.assertFalse(unit.can_calculate)
        unit=self.preview(8,{"D3":0,"F3":0,"K3":"实测值","L3":"实测值"})
        self.assertTrue(unit.can_calculate,unit.errors)
        self.assertEqual(unit.calculation.result.total_amount,Decimal("0"))
    def test_old_catalog_factor_mismatch_is_not_relabelled(self):
        for version in range(1,8):
            with self.subTest(version=version):
                unit=self.preview(version,{"D3":90,"F3":100,"E3":.5})
                self.assertFalse(unit.can_calculate)
                if unit.input_value is not None:self.assertIsNone(unit.input_value.fgd_units[0].components[0].emission_factor)

    def test_existing_factor_source_selectors_remain_required(self):
        for version in range(3,9):
            with self.subTest(version=version):
                changes={"D3":90,"F3":100,"J3":None}
                if version==8:changes.update(K3="实测值",L3="实测值")
                unit=self.preview(version,changes)
                self.assertFalse(unit.can_calculate)
                self.assertIn("EXB01_CARBONATE_FACTOR_SOURCE_REQUIRED",{e.code for e in unit.errors})
    def test_undeclared_legacy_custom_factor_does_not_become_measured(self):
        for version in (1,2):
            with self.subTest(version=version):
                changes={"C3":"未登记碳酸盐","D3":90,"F3":100,"E3":.1234}
                if version==2:changes.update(I3="自定义",E3=None,H3=.1234)
                unit=self.preview(version,changes)
                self.assertFalse(unit.can_calculate)
                self.assertIn("EXB01_CARBONATE_FACTOR_SOURCE_REQUIRED",{e.code for e in unit.errors})
