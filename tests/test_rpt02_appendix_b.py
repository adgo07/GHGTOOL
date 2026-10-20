from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
from datetime import date
from decimal import Decimal
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace
import unittest
from io import BytesIO
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

from packages.application.reporting import build_saved_record_report
from packages.application.reporting.model import build_report_model, frozen_totals
from packages.application.reporting.appendix_b import _energy_parameter, _frozen_step, _short_source
from packages.standards.carbon_material import SOURCE_PURCHASED_HEAT, SOURCE_EXPORTED_HEAT, SOURCE_PURCHASED_ELECTRICITY, SOURCE_EXPORTED_ELECTRICITY
from packages.core.models import CalculationLine
from packages.infrastructure.reporting import render_report_docx
from packages.persistence.catalog_repository import SQLiteCatalogRepository
from packages.standards.carbon_material import CarbonMaterialCalculator, SOURCE_BAKING
from packages.ui.record_experience import build_report_view
from scripts.rpt02_acceptance_samples import build_acceptance_case, build_acceptance_model


ROOT = Path(__file__).resolve().parents[1]
LAYOUT_PATH = ROOT / "packages" / "application" / "reporting" / "appendix_b_layout.json"
TEMPLATE_PATH = ROOT / "docs" / "GB_T32151_34_2024_附录B填报模板_脚注与边框修订版.xlsx"


def _section(model, section_id):
    return next(section for section in model.sections if section.section_id == section_id)


def _header_merges_from_sheet(sheet, header_rows):
    merges = []
    last_header_row = 1 + header_rows
    for merged in sheet.merged_cells.ranges:
        if merged.min_row < 2 or merged.max_row > last_header_row:
            continue
        merges.append((merged.min_row - 2, merged.min_col - 1, merged.max_row - 2, merged.max_col - 1))
    return tuple(sorted(merges))


class RPT02AppendixBTests(unittest.TestCase):
    """The manual records here are synthetic mapping fixtures, never formal business acceptance samples."""

    def test_layout_metadata_matches_formula_free_approved_workbook(self):
        layout = json.loads(LAYOUT_PATH.read_text(encoding="utf-8"))
        self.assertEqual(hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest(), layout["formula_free_sha256"])
        from packages.excel.templates import ExcelTemplateService
        workbook = load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()), data_only=False)
        try:
            self.assertEqual(workbook.sheetnames, [*[f"B.{number}" for number in range(1, 10)], "软件补充"])
            for key, table in layout["tables"].items():
                sheet = workbook[f"B.{key[1:]}"]
                headers = table["headers"]
                actual = tuple(
                    tuple(sheet.cell(row=2 + ri, column=ci + 1).value or "" for ci in range(len(headers[0])))
                    for ri in range(len(headers))
                )
                if key == "b6":
                    # v9 input asks for an explicit per-segment unit. Historical
                    # Word headers stay frozen; monitored records override T only.
                    self.assertEqual(actual[0][5], "监测段时长\n（单位见M列）")
                    self.assertEqual(headers[0][5], "报告期的时长\nd")
                    actual = tuple(tuple(headers[ri][ci] if (ri,ci)==(0,5) else value
                                         for ci,value in enumerate(row)) for ri,row in enumerate(actual))
                self.assertEqual(actual, tuple(tuple(row) for row in headers), sheet.title)
                self.assertEqual(
                    tuple(item for item in _header_merges_from_sheet(sheet,len(headers)) if item[3]<len(headers[0])),
                    tuple(sorted(tuple(item) for item in table["header_merges"])),
                    sheet.title,
                )
                self.assertEqual(len(table["column_weights"]), len(headers[0]), sheet.title)

            model = build_acceptance_model("GUI")
            self.assertEqual(
                [row.cells[0].value for row in _section(model, "b1").tables[0].rows[:10]],
                [workbook["B.1"].cell(row=row, column=1).value for row in range(3, 13)],
            )
            workbook.close()
            workbook=load_workbook(TEMPLATE_PATH,data_only=False)
            for cell_ref in layout["removed_formula_cells"]:
                sheet_name, address = cell_ref.split("!")
                cell = workbook[sheet_name][address]
                self.assertIsNone(cell.value, cell_ref)
                self.assertNotEqual(cell.data_type, "f", cell_ref)
                self.assertGreater(cell.style_id, 0, cell_ref)
        finally:
            workbook.close()

    def test_synthetic_frozen_record_expands_all_nine_appendix_tables(self):
        model = build_acceptance_model("GUI")
        self.assertEqual([section.section_id for section in model.sections], ["basic", *(f"b{i}" for i in range(1, 10))])
        self.assertEqual(model.layout_id, "gbt-32151-34-2024-appendix-b")
        b1 = _section(model, "b1").tables[0]
        self.assertEqual(len(b1.rows), 13)
        self.assertEqual((b1.rows[-2].cells[2].value, b1.rows[-1].cells[2].value), ("7.30", "9.70"))
        self.assertEqual(len(b1.footnotes), 1)

        expected_columns = {"b2": 11, "b3": 6, "b4": 6, "b5": 6, "b6": 7, "b7": 7, "b8": 5, "b9": 7}
        for section_id, count in expected_columns.items():
            for table in _section(model, section_id).tables:
                self.assertEqual(len(table.columns), count, table.table_id)
                self.assertEqual(len(table.header_rows[0].cells), count, table.table_id)
                self.assertEqual(len(table.column_weights), count, table.table_id)

        b2 = _section(model, "b2").tables[0]
        self.assertEqual(len(b2.rows), 2)
        self.assertEqual(b2.rows[0].cells[10].value, "1.40")
        self.assertEqual(tuple(sorted(b2.header_merges)), ((0, 0, 1, 0), (0, 1, 0, 2), (0, 3, 0, 4), (0, 5, 0, 6), (0, 7, 1, 7), (0, 8, 0, 9), (0, 10, 1, 10)))
        self.assertEqual(len(_section(model, "b3").tables), 2)
        self.assertGreater(len(_section(model, "b3").tables[0].rows), 8)
        self.assertEqual(len(_section(model, "b6").tables[0].rows), 2)
        self.assertEqual(len(_section(model, "b7").tables[0].rows), 3)
        b8 = _section(model, "b8").tables[0]
        self.assertEqual(len(b8.rows), 4)
        self.assertTrue(all(Decimal(row.cells[4].value) < 0 for row in b8.rows if row.cells[0].value == "输出"))
        b9 = _section(model, "b9")
        self.assertEqual(len(b9.tables[0].rows), 3)
        self.assertEqual(b9.notes, ())
        self.assertEqual(b9.tables[0].footnotes, tuple(json.loads(LAYOUT_PATH.read_text(encoding="utf-8"))["tables"]["b9"]["footnotes"]))
        self.assertEqual(b9.tables[0].rows[0].cells[5].source, "标准查表")

    def test_b7_uses_frozen_reagent_name_and_marks_legacy_name_missing(self):
        record, raw, trace, provenance, reporting, qualification = build_acceptance_case("GUI")
        raw = deepcopy(raw)
        raw["fgd_units"][0]["reagent_name"] = "石灰石甲"
        named = build_report_model(record, raw, trace, provenance, reporting, qualification, today=date(2026, 10, 11))
        named_rows = _section(named, "b7").tables[0].rows
        self.assertEqual(named_rows[0].cells[0].value, "石灰石甲")

        raw["fgd_units"][0].pop("reagent_name")
        legacy = build_report_model(record, raw, trace, provenance, reporting, qualification, today=date(2026, 10, 11))
        legacy_rows = _section(legacy, "b7").tables[0].rows
        self.assertEqual(legacy_rows[0].cells[0].value, "历史未记录")

    def test_zero_empty_disabled_and_legacy_missing_fields_stay_distinct(self):
        record, raw, trace, provenance, reporting, qualification = build_acceptance_case("GUI")
        raw = deepcopy(raw)
        raw["fuel_inputs"][1]["oxidation_rate"] = {"value": "0", "unit": "ratio"}
        raw["fume_incinerations"][1]["fch"] = {}
        raw["source_states"] = [
            {**state, "status": "NOT_INVOLVED"} if state["source_id"] == SOURCE_BAKING else state
            for state in raw["source_states"]
        ]
        raw["bakings"] = []
        model = build_report_model(record, raw, trace, provenance, reporting, qualification, today=date(2026, 10, 9))

        self.assertEqual(_section(model, "b2").tables[0].rows[1].cells[8].value, "0.00")
        self.assertEqual(_section(model, "b6").tables[0].rows[1].cells[3].value, "未填写")
        self.assertNotIn("{", _section(model, "b6").tables[0].rows[1].cells[3].value)
        self.assertEqual(_section(model, "b4").tables[0].rows[0].cells[0].value, "未启用")
        self.assertIn("未启用", _section(model, "b1").tables[0].rows[2].cells[2].value)

        legacy = build_report_model(record, {}, {}, {}, {}, {}, snapshot_schema_version=0, today=date(2026, 10, 9))
        fuel_row = _section(legacy, "b2").tables[0].rows[0]
        self.assertTrue(fuel_row.cells[0].value.startswith("历史燃料结果"))
        self.assertEqual(fuel_row.cells[1].value, "历史未记录")
        self.assertEqual(fuel_row.cells[-1].value, "1.40")
        self.assertEqual(_section(legacy, "b4").tables[0].rows[0].cells[0].value, "历史未记录此项明细")
        self.assertTrue(any("未保存完整业务输入快照" in notice for notice in legacy.notices))

    def test_saved_report_uses_frozen_record_without_catalog_or_calculator(self):
        record, raw, trace, provenance, reporting, qualification = build_acceptance_case("EXCEL_R2")
        before = deepcopy((record, raw, trace, provenance, reporting, qualification))

        class FrozenRepository:
            def _check(self, record_id):
                if record_id != record.record_id:
                    raise AssertionError("report requested another Record")

            def get_raw_input_snapshot(self, record_id):
                self._check(record_id)
                return deepcopy(raw)

            def get_trace_snapshot(self, record_id):
                self._check(record_id)
                return deepcopy(trace)

            def get_provenance_snapshot(self, record_id):
                self._check(record_id)
                return deepcopy(provenance)

            def get_reporting_snapshot(self, record_id):
                self._check(record_id)
                return deepcopy(reporting)

            def get_report_qualification(self, record_id):
                self._check(record_id)
                return deepcopy(qualification)

            def get_snapshot_schema_version(self, record_id):
                self._check(record_id)
                return 1

        with (
            patch.object(CarbonMaterialCalculator, "calculate", side_effect=AssertionError("report recalculated")) as calculate,
            patch.object(SQLiteCatalogRepository, "list_parameters", side_effect=AssertionError("report queried Catalog")) as parameters,
            patch.object(SQLiteCatalogRepository, "list_factors", side_effect=AssertionError("report queried Catalog")) as factors,
        ):
            model = build_saved_record_report(FrozenRepository(), record)
        self.assertEqual((record, raw, trace, provenance, reporting, qualification), before)
        self.assertEqual((calculate.call_count, parameters.call_count, factors.call_count), (0, 0, 0))
        self.assertEqual(model.record_id, record.record_id)
        self.assertEqual(model.layout_id, "gbt-32151-34-2024-appendix-b")

    def test_native_word_tables_are_editable_merged_and_repeat_headers(self):
        document = Document(BytesIO(render_report_docx(build_acceptance_model("GUI"))))
        self.assertEqual(len(document.tables), 11)
        b1, b2 = document.tables[1], document.tables[2]
        self.assertEqual(len(b1.rows), 14)
        self.assertIs(b1.rows[0].cells[0]._tc, b1.rows[0].cells[1]._tc)
        self.assertIs(b2.rows[0].cells[1]._tc, b2.rows[0].cells[2]._tc)
        for row in (*b1.rows[:1], *b2.rows[:2]):
            self.assertIsNotNone(row._tr.trPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"))
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        self.assertIn("标准查表",text)
        self.assertNotIn("选择方式：",text)
        self.assertNotIn("来源定位：",text)
        self.assertNotIn("因子年度：",text)
        b1_heading = next(p for p in document.paragraphs if p.text == "B.1 温室气体排放量汇总")
        self.assertTrue(b1_heading.paragraph_format.page_break_before)

    def test_energy_snapshot_bare_id_and_same_id_different_units(self):
        record, *_ = build_acceptance_case("GUI")
        seed = record.parameter_snapshots[0]
        heat = replace(seed, detail_id="shared", value_used="0.11", unit_used="tCO2/GJ", source_location="热力冻结来源")
        power = replace(seed, snapshot_id="snapshot.power.collision", detail_id="shared", value_used="0.7", unit_used="tCO2/MWh", source_location="电力冻结来源")
        record = replace(record, parameter_snapshots=(heat, power))
        cell = _energy_parameter(record, {}, "shared", "CAR-FLD-HEAT-shared-EF3", "tCO2/GJ", "EF3", None)
        self.assertEqual(cell.value, "0.11")
        self.assertIn("热力冻结来源", cell.source_detail)
        self.assertNotIn("电力冻结来源", cell.source_detail)

    def test_purchase_and_export_same_id_keep_separate_factor_sources(self):
        record, raw, trace, *_ = build_acceptance_case("GUI")
        seed = record.parameter_snapshots[0]
        purchase = replace(seed, snapshot_id="snapshot.purchase", detail_id="shared", value_used="0.7", unit_used="tCO2/MWh", source_location="购电合同来源")
        export = replace(purchase, snapshot_id="snapshot.export", detail_id="CAR-FLD-POWER-EXPORTED-EF.shared", source_location="输电合同来源")
        record = replace(record, parameter_snapshots=(purchase, export))
        raw = deepcopy(raw)
        raw["electricity_details"] = [{"detail_id":"shared", "electricity_amount":"1", "electricity_unit":"MWh"}]
        raw["exported_electricity"] = [{"line_id":"shared", "amount":"1", "unit":"MWh"}]
        model = build_report_model(record, raw, trace, {}, {}, {})
        rows = _section(model, "b8").tables[0].rows
        self.assertIn("购电合同来源", rows[0].cells[3].source_detail)
        self.assertNotIn("输电合同来源", rows[0].cells[3].source_detail)
        self.assertIn("输电合同来源", rows[1].cells[3].source_detail)

    def test_incomplete_colliding_power_snapshots_do_not_invent_direction_or_value(self):
        record, raw, trace, *_ = build_acceptance_case("GUI")
        seed = record.parameter_snapshots[0]
        purchase = replace(seed, detail_id="shared", value_used="0.7", unit_used="tCO2/MWh", source_location="购电合同来源")
        record = replace(record, parameter_snapshots=(purchase,))
        raw = deepcopy(raw)
        raw["electricity_details"] = [{"detail_id": "shared", "electricity_amount": "1", "electricity_unit": "MWh"}]
        raw["exported_electricity"] = [{"line_id": "shared", "amount": "1", "unit": "MWh"}]
        trace = {"formula_steps": [
            {"emission_source_id": source, "process_instance_id": "shared", "intermediate_result": "0.7",
             "input_variables": [{"name": "EF2", "value": "0.7"}]}
            for source in (SOURCE_PURCHASED_ELECTRICITY, SOURCE_EXPORTED_ELECTRICITY)
        ]}
        rows = _section(build_report_model(record, raw, trace, {}, {}, {}), "b8").tables[0].rows
        for row in rows:
            self.assertEqual(row.cells[3].value, "0.7")
            self.assertEqual(row.cells[3].source, "历史来源快照无法唯一关联")
        # A non-purchased self-consumed fossil row cannot make output provenance ambiguous.
        raw["electricity_details"][0].update(acquisition_mode="SELF_CONSUMED", attribute="FOSSIL")
        output = replace(purchase, source_location="自动输出冻结来源")
        record = replace(record, parameter_snapshots=(output,))
        rows = _section(build_report_model(record, raw, trace, {}, {}, {}), "b8").tables[0].rows
        self.assertEqual(len(rows), 1)
        self.assertIn("自动输出冻结来源", rows[0].cells[3].source_detail)
        # A single prefixed candidate also must agree with the frozen adopted value.
        record = replace(record, parameter_snapshots=(replace(purchase, detail_id="CAR-FLD-POWER-EXPORTED-EF.shared"),))
        step = {"input_variables": [{"name": "EF2", "value": "0.5"}]}
        cell = _energy_parameter(record, {}, "shared", "CAR-FLD-POWER-EXPORTED-EF.shared", "tCO2/MWh", "EF2", step)
        self.assertEqual(cell.value, "0.5")
        self.assertEqual(cell.source, "历史来源快照与冻结采用值不一致")

    def test_manual_heat_keeps_frozen_reference_and_auto_heat_does_not_repeat_it(self):
        record,raw,trace,*_=build_acceptance_case("GUI")
        trace=deepcopy(trace)
        step=next(item for item in trace["formula_steps"] if item.get("process_instance_id")=="heat-in-1")
        step["calculation_provenance"]["enthalpy_source"]="用户手动填写"
        step["calculation_provenance"]["enthalpy_used_kj_per_kg"]="2800"
        model=build_report_model(record,raw,trace,{}, {}, {})
        b9=_section(model,"b9")
        self.assertEqual(b9.tables[0].rows[0].cells[5].value,"2800.00")
        self.assertEqual(b9.tables[0].rows[0].cells[5].source,"填写值")
        self.assertFalse(any("自动参考焓值" in note for note in b9.notes))
        self.assertIn("冻结自动参考焓值：2716 kJ/kg",b9.tables[0].rows[0].cells[5].source_detail)

    def test_b_table_notes_keep_business_exceptions_and_audit_stays_on_cells(self):
        model = build_acceptance_model("GUI")
        notes = []
        for section in model.sections:
            if not section.section_id.startswith("b"):
                continue
            notes.extend(section.notes)
            for table in section.tables:
                notes.extend(table.footnotes)
                notes.extend(row.note for row in table.rows if row.note)
        self.assertFalse(any("行第" in note and "列" in note for note in notes))
        for note in notes:
            self.assertNotIn("选择方式：", note)
            self.assertNotIn("来源定位：", note)
            self.assertNotIn("来源版本：", note)
            self.assertNotIn("因子年度：", note)
        factor = _section(model, "b9").tables[0].rows[0].cells[6]
        self.assertIn("来源定位：", factor.source_detail)
        self.assertIn("来源版本：", factor.source_detail)

    def test_report_source_uses_clause_without_dropping_frozen_pdf_audit(self):
        record, raw, trace, provenance, reporting, qualification = build_acceptance_case("GUI")
        snapshot = next(
            item for item in record.parameter_snapshots
            if item.detail_id == "CAR-FLD-P04B-fgd-1-0-I"
        )
        frozen_location = (
            "GB/T 32151.34—2024 第5.2.6.1条；PDF第30页；印刷页22"
        )
        snapshots = tuple(
            replace(item, source_location=frozen_location)
            if item.snapshot_id == snapshot.snapshot_id else item
            for item in record.parameter_snapshots
        )
        record = replace(record, parameter_snapshots=snapshots)

        model = build_report_model(record, raw, trace, provenance, reporting, qualification)
        cell = _section(model, "b7").tables[0].rows[0].cells[3]

        self.assertEqual(cell.source, "GB/T32151.34—2024 第5.2.6.1条")
        self.assertNotIn("PDF", cell.source)
        self.assertNotIn("印刷页", cell.source)
        self.assertIn(f"来源定位：{frozen_location}", cell.source_detail)

    def test_legacy_user_defined_source_is_not_relabelled_as_measured(self):
        self.assertEqual(
            _short_source({"source_kind": "USER_DEFINED"}, None, "用户提供值"),
            "填写值",
        )

    def test_b9_missing_history_explains_missing_input_without_repeating_results(self):
        record, raw, trace, provenance, reporting, qualification = build_acceptance_case("GUI")
        raw = deepcopy(raw)
        raw["purchased_heat"] = []
        raw["exported_heat"] = []
        history_lines = (
            CalculationLine("legacy.purchased-heat", SOURCE_PURCHASED_HEAT, "GEN-GAS-CO2", "0.50", "tCO2"),
            CalculationLine("legacy.exported-heat", SOURCE_EXPORTED_HEAT, "GEN-GAS-CO2", "0.20", "tCO2"),
        )
        record = replace(record, calculation_result=replace(
            record.calculation_result,
            lines=(*record.calculation_result.lines, *history_lines),
        ))
        section = _section(build_report_model(record, raw, trace, provenance, reporting, qualification), "b9")
        note = "；".join(section.notes)
        self.assertIn("历史记录未保存购入热力、输出热力活动输入", note)
        self.assertIn("历史未记录", note)
        self.assertNotIn("0.50", note)
        self.assertNotIn("0.30", note)
        self.assertNotIn("0.20", note)
        self.assertFalse(any("第" in item for item in section.notes))

    def test_b9_suppresses_obvious_manual_entry_labels_but_keeps_audit(self):
        record, raw, trace, provenance, reporting, qualification = build_acceptance_case("GUI")
        raw = deepcopy(raw)
        raw["purchased_heat"][0]["steam_amount_t"] = {
            "value": "4", "unit": "t", "source_type": "MANUAL", "source_reference": "手工录入",
        }
        section = _section(build_report_model(record, raw, trace, provenance, reporting, qualification), "b9")
        quantity = section.tables[0].rows[0].cells[2]
        self.assertIsNone(quantity.source)
        self.assertIn("手工录入", quantity.source_detail)

    def test_missing_trace_uses_exact_frozen_line_without_sharing_aggregate(self):
        record, *_ = build_acceptance_case("GUI")
        line = CalculationLine("CAR-FLD-HEAT-PURCHASED-RESULT.h1", SOURCE_PURCHASED_HEAT, "GEN-GAS-CO2", "12.345", "tCO2")
        record = replace(record, calculation_result=replace(record.calculation_result, lines=(line,)))
        self.assertEqual(_frozen_step(record, {}, SOURCE_PURCHASED_HEAT, "h1")["intermediate_result"], "12.345")
        self.assertIsNone(_frozen_step(record, {}, SOURCE_PURCHASED_HEAT, "h2"))
        aggregate = replace(line, line_id="CAR-FLD-HEAT-PURCHASED-RESULT")
        record = replace(record, calculation_result=replace(record.calculation_result, lines=(aggregate,)))
        self.assertIsNone(_frozen_step(record, {}, SOURCE_PURCHASED_HEAT, "h1"))
        self.assertEqual(_frozen_step(record, {}, SOURCE_PURCHASED_HEAT, "h1", single_instance=True)["intermediate_result"], "12.345")

    def test_record_report_b1_preserves_frozen_es_ei_et_precision(self):
        record, raw, trace, *_ = build_acceptance_case("GUI")
        values = frozen_totals(record, trace)
        self.assertEqual(values, {"ES": "7.30", "EI": "2.40", "ET": "9.70"})
        ui_text = build_report_view(record,raw,trace,{})
        report = build_report_model(record, raw, trace, {}, {}, {}, today=date(2026, 10, 9))
        b1 = _section(report, "b1").tables[0]
        self.assertIn("不包括购入和输出的电力、热力所产生的 CO₂",ui_text)
        self.assertIn("7.30",ui_text)
        self.assertEqual(
            {"ES": b1.rows[-2].cells[2].value, "ET": b1.rows[-1].cells[2].value},
            {"ES": values["ES"], "ET": values["ET"]},
        )
        empty_record = SimpleNamespace(calculation_result=SimpleNamespace(lines=(), total_amount=None))
        self.assertEqual(frozen_totals(empty_record, {}), {"ES": "", "EI": "", "ET": ""})


if __name__ == "__main__":
    unittest.main()
