"""V5 template resource and compatibility protections; no native Office claim."""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from zipfile import ZipFile
from openpyxl import load_workbook
from packages.excel.templates import (ExcelTemplateService, APPENDIX_B_SHEETS, APPENDIX_B_TEMPLATE,
 APPROVED_TEMPLATE_V5_SHA256 as APPROVED_TEMPLATE_SHA256, APPROVED_TEMPLATE_V5_RESOURCE as APPROVED_TEMPLATE_RESOURCE, APPROVED_TEMPLATE_V4_RESOURCE,
 APPROVED_TEMPLATE_V4_SHA256, APPROVED_TEMPLATE_V3_RESOURCE, APPROVED_TEMPLATE_V3_SHA256,
 APPROVED_TEMPLATE_V2_RESOURCE, APPROVED_TEMPLATE_V2_SHA256)

class FuelPowerTemplateV5Tests(TestCase):
 def test_download_copies_exact_registered_bytes_and_recognizes_v5(self):
  # V5 is now a preserved historical resource; current default download is V6.
  service=ExcelTemplateService((replace(APPENDIX_B_TEMPLATE,template_version="2.3.0",
      resource_path=APPROVED_TEMPLATE_RESOURCE,sha256=APPROVED_TEMPLATE_SHA256),))
  content=service.read_bytes()
  self.assertEqual(sha256(content).hexdigest(),APPROVED_TEMPLATE_SHA256)
  with TemporaryDirectory() as tmp:
   output=service.copy_to(Path(tmp)/"template.xlsx")
   self.assertEqual(output.read_bytes(),content)
  workbook=load_workbook(BytesIO(content))
  try:
   self.assertEqual(tuple(workbook.sheetnames),APPENDIX_B_SHEETS)
   self.assertEqual(service.validate_structure(workbook),"2.3.0")
   self.assertTrue(workbook["B.2"]["A20"].value.startswith("燃料输入规则 V5："))
   formulas={str(d.sqref):d.formula1 for d in workbook["B.2"].data_validations.dataValidation}
   self.assertIn("直接计算值",formulas["E4:E18"])
   self.assertEqual(formulas["G4:G18"],'"缺省值,实测值"')
   self.assertEqual(formulas["N4:N18"],'"缺省值,实测值"')
  finally: workbook.close()
 def test_all_other_ooxml_parts_are_identical_to_v4(self):
  with ZipFile(Path("resources")/APPROVED_TEMPLATE_RESOURCE) as new, ZipFile(Path("resources")/APPROVED_TEMPLATE_V4_RESOURCE) as old:
   self.assertEqual(new.namelist(),old.namelist())
   changed=[name for name in old.namelist() if old.read(name)!=new.read(name)]
   self.assertEqual(changed,["xl/worksheets/sheet2.xml"])
 def test_prior_template_bytes_and_version_identity_remain_intact(self):
  service=ExcelTemplateService.default()
  for resource,digest,version in [(APPROVED_TEMPLATE_V2_RESOURCE,APPROVED_TEMPLATE_V2_SHA256,"2.0.0"),(APPROVED_TEMPLATE_V3_RESOURCE,APPROVED_TEMPLATE_V3_SHA256,"2.1.0"),(APPROVED_TEMPLATE_V4_RESOURCE,APPROVED_TEMPLATE_V4_SHA256,"2.2.0")]:
   with self.subTest(version=version):
    content=(Path("resources")/resource).read_bytes()
    self.assertEqual(sha256(content).hexdigest(),digest)
    workbook=load_workbook(BytesIO(content))
    try: self.assertEqual(service.validate_structure(workbook),version)
    finally: workbook.close()

 def test_inserted_activity_rows_do_not_downgrade_template_identity(self):
  service=ExcelTemplateService.default()
  for resource,version in [(APPROVED_TEMPLATE_RESOURCE,"2.3.0"),(APPROVED_TEMPLATE_V4_RESOURCE,"2.2.0")]:
   with self.subTest(version=version):
    workbook=load_workbook(Path("resources")/resource)
    try:
     workbook["B.2"].insert_rows(19,3)
     self.assertEqual(service.validate_structure(workbook),version)
    finally: workbook.close()
 def test_conflicting_version_footers_are_rejected(self):
  service=ExcelTemplateService.default()
  workbook=load_workbook(BytesIO(service.read_bytes()))
  try:
   workbook["B.2"]["A21"]="来源栏用于确认参数来源；缺省值须有当前期间适用的 Resolver 绑定。"
   with self.assertRaisesRegex(ValueError,"冲突"): service.validate_structure(workbook)
  finally: workbook.close()

 def test_row_boolean_lexemes_and_exact_height_are_valid(self):
  from xml.etree import ElementTree as ET
  ns={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
  with ZipFile(Path("resources")/APPROVED_TEMPLATE_RESOURCE) as archive:
   for name in archive.namelist():
    if name.startswith("xl/worksheets/") and name.endswith(".xml"):
     root=ET.fromstring(archive.read(name))
     for row in root.findall(".//s:row",ns):
      for attribute in ("customHeight","hidden","collapsed","customFormat","thickTop","thickBot","ph"):
       if attribute in row.attrib:
        self.assertIn(row.get(attribute),{"0","1","true","false"},(name,row.get("r"),attribute))
   row=ET.fromstring(archive.read("xl/worksheets/sheet2.xml")).find(".//s:row[@r='20']",ns)
   self.assertEqual(row.get("customHeight"),"1")
   self.assertEqual(row.get("ht"),"45")
 def test_narrow_repair_does_not_match_custom_height_as_height(self):
  from scripts.repair_appendix_b_v5_xml import repair_row_height
  xml=b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="20" ht="34" customHeight="45" s="42"/></sheetData></worksheet>'
  repaired=repair_row_height(xml)
  self.assertIn(b'ht="45" customHeight="1"',repaired)
  self.assertEqual(repair_row_height(repaired),repaired)

 def test_delivery_audit_rejects_invalid_boolean_even_when_xml_parses(self):
  from scripts.inspect_release import excel_row_boolean_issues
  xml=b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="20" ht="34" customHeight="45"/></sheetData></worksheet>'
  with TemporaryDirectory() as directory:
   bad=Path(directory)/"bad.xlsx"
   with ZipFile(bad,"w") as archive: archive.writestr("xl/worksheets/sheet2.xml",xml)
   issues=excel_row_boolean_issues(bad)
   self.assertTrue(any("customHeight=45" in issue for issue in issues))
  self.assertEqual(excel_row_boolean_issues(Path("resources")/APPROVED_TEMPLATE_RESOURCE),[])
