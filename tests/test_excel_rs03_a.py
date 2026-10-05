from __future__ import annotations

from dataclasses import fields, is_dataclass, replace
from datetime import date
from decimal import Decimal
from enum import Enum
from io import BytesIO
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from packages.application.carbon_accounting import create_g06_parameter_resolver
from packages.application.catalog_queries import CatalogQueryService
from packages.core.models import (
    ElectricityAcquisitionMode,
    ElectricityAttribute,
    ElectricityProofStatus,
    ElectricityProofType,
)
from packages.excel.gbt32151_34_v1 import (
    BODY_ROWS,
    CARBONATE_LABELS,
    ELECTRICITY_HEADERS,
    FUEL_C1_SUBJECT_IDS,
    FUEL_HEADERS,
    FUEL_LABEL_BY_TYPE,
    HEAT_HEADERS,
    MATERIAL_HEADERS,
    SECTION_HEADERS,
    SOURCE_LABELS,
    SOURCE_IDS,
    TEMPLATE_ID,
    TEMPLATE_VERSION,
    UNIT_HEADERS,
    VISIBLE_SHEETS,
    ExcelWorkbookImporter,
    WorkbookFatalError,
    _CellReader,
    _UnitContext,
    create_template_bytes,
    significant_digit_count,
    write_template,
)
from packages.persistence import SQLiteCatalogRepository, build_catalog_database
from packages.reference_data import DEFAULT_SOURCE_PATH
from packages.standards.carbon_material import (
    CarbonMaterialCalculator,
    EmissionSourceStatus,
    FuelPath,
    FuelType,
    InMemoryRecordRepository,
    SOURCE_CALCINATION,
    SOURCE_EXPORTED_ELECTRICITY,
    SOURCE_EXPORTED_HEAT,
    SOURCE_FUEL,
    SOURCE_PURCHASED_ELECTRICITY,
    SOURCE_PURCHASED_HEAT,
    SteamKind,
    saturated_steam_enthalpy,
    superheated_steam_enthalpy,
)
from packages.ui.carbon_material_page import CarbonMaterialAccountingPage
from packages.ui.pages import ExcelImportPage


SHEET_FOR_MARKER = {
    "核算单元清单": "核算单元",
    "B.2 化石燃料": "燃料与能源",
    "B.8 电力": "燃料与能源",
    "B.9 热力": "燃料与能源",
    "B.3 煅烧": "过程排放",
    "B.4 焙烧/炭化": "过程排放",
    "B.5 石墨化": "过程排放",
    "B.6 烟气焚烧": "过程排放",
    "B.7 烟气脱硫": "过程排放",
}


def section_bounds(sheet, marker: str, headers: tuple[str, ...]) -> tuple[int, int, int]:
    marker_rows = [row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == marker]
    if len(marker_rows) != 1:
        raise AssertionError(f"template section missing or repeated: {marker}")
    marker_row = marker_rows[0]
    header_row = marker_row + 2
    assert tuple(sheet.cell(header_row, col).value for col in range(1, len(headers) + 1)) == headers
    later = [
        row for row in range(header_row + 1, sheet.max_row + 1)
        if sheet.cell(row, 1).value in SECTION_HEADERS
    ]
    first = header_row + 1
    last = min(later) - 1 if later else sheet.max_row
    return header_row, first, max(first - 1, last)


def _source_reference(source_location: str | None) -> str | None:
    if source_location is None:
        return None
    if "：" in source_location:
        return source_location.rsplit("：", 1)[-1]
    return source_location


def _business_semantics(value):
    """Normalize generated identity while retaining business values and provenance."""
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (tuple, list)):
        return tuple(_business_semantics(item) for item in value)
    if isinstance(value, dict):
        return tuple(sorted((key, _business_semantics(item)) for key, item in value.items()))
    if not is_dataclass(value) or isinstance(value, type):
        return value

    kind = type(value).__name__
    if kind == "InputValue":
        # UI and Excel use different activity-source enums for manual cells, but
        # their saved numeric value, unit, and any user source reference remain
        # the business semantics to compare here.
        return (value.value, value.unit, value.source_reference, _business_semantics(value.source_level))
    if kind == "ParameterValue":
        source_kind = _business_semantics(value.source_kind)
        measured = source_kind in {"MEASURED", "USER_DEFINED", "PROJECT_SPECIFIED"}
        return (
            value.parameter_id,
            value.value,
            value.unit,
            source_kind,
            None if measured else value.source_id,
            value.source_version,
            _source_reference(value.source_location) if measured else value.source_location,
            value.factor_id,
            value.factor_year,
        )
    if kind == "ParameterSnapshot":
        selection_method = _business_semantics(value.selection_method)
        measured = selection_method in {"ENTERPRISE_MEASURED", "MANUAL_OVERRIDE"}
        return (
            value.parameter_id,
            value.value_used,
            value.unit_used,
            None if measured else value.source_id,
            value.source_version,
            selection_method,
            value.standard_id,
            value.factor_version,
            _source_reference(value.source_location) if measured else value.source_location,
            value.factor_id,
            value.factor_year,
        )

    ignored = {
        "input_id", "enterprise_id", "boundary_component_ids", "reporting_data",
        "fuel_id", "electricity_detail_id", "detail_id", "line_id", "instance_id",
        "evidence_ref_ids", "evidence_id", "snapshot_id", "result_id", "record_id",
        "calculated_at", "selection_reason",
    }
    normalized = []
    for field in fields(value):
        if field.name in ignored:
            continue
        field_value = getattr(value, field.name)
        if field.name == "source_states":
            projected = tuple(sorted((_business_semantics(item) for item in field_value), key=repr))
        elif field.name in {"parameter_snapshots", "traces"}:
            projected = tuple(sorted((_business_semantics(item) for item in field_value), key=repr))
        else:
            projected = _business_semantics(field_value)
        normalized.append((field.name, projected))
    return tuple(normalized)


def _calculation_semantics(outcome):
    result = outcome.result
    result_semantics = None
    es_ei_et = None
    if result is not None:
        lines = tuple(sorted(
            (
                (
                    line.emission_source_id,
                    line.greenhouse_gas_id,
                    line.amount,
                    line.unit,
                )
                for line in result.lines
            ),
            key=repr,
        ))
        result_semantics = (
            result.standard_id,
            result.algorithm_version,
            lines,
            result.total_amount,
            result.total_unit,
            tuple(sorted((problem.code, problem.level.value, problem.message) for problem in result.problems)),
        )
        by_id = {line.line_id: line.amount for line in result.lines}
        es_ei_et = (
            by_id["CAR-FLD-DIRECT-RESULT"],
            by_id["CAR-FLD-INDIRECT-RESULT"],
            by_id["CAR-FLD-TOTAL-RESULT"],
        )
    problems = tuple(sorted((problem.code, problem.level.value, problem.message) for problem in outcome.problems))
    snapshots = tuple(sorted((_business_semantics(item) for item in outcome.parameter_snapshots), key=repr))
    return {
        "successful": outcome.successful,
        "blocked": outcome.blocked,
        "problems": problems,
        "parameter_snapshots": snapshots,
        "result": result_semantics,
        "ES_EI_ET": es_ei_et,
    }


class ExcelRS03ATests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.application = QApplication.instance() or QApplication([])
        cls.temp_root = tempfile.TemporaryDirectory()
        cls.catalog_path = Path(cls.temp_root.name) / "catalog.sqlite"
        build_catalog_database(DEFAULT_SOURCE_PATH, cls.catalog_path)
        cls.catalog_service = CatalogQueryService(
            SQLiteCatalogRepository(cls.catalog_path), as_of=date(2026, 9, 12)
        )
        cls.resolver = create_g06_parameter_resolver(cls.catalog_service.repository)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temp_root.cleanup()

    def new_book(self):
        return load_workbook(BytesIO(create_template_bytes()), data_only=False)

    @staticmethod
    def set_project(workbook, enterprise: str = "示例企业", *, boundary: str = "是") -> None:
        sheet = workbook["核算单元"]
        sheet["B3"] = enterprise
        sheet["B4"] = "年度"
        sheet["B5"] = date(2025, 1, 1)
        sheet["B6"] = date(2025, 12, 31)
        sheet["B7"] = boundary
        sheet["B8"] = "炭素材料生产边界"

    @staticmethod
    def append_section(workbook, marker: str, values: dict[str, object]) -> int:
        headers = SECTION_HEADERS[marker]
        sheet = workbook[SHEET_FOR_MARKER[marker]]
        header_row, first, last = section_bounds(sheet, marker, headers)
        row = next(
            (
                candidate for candidate in range(first, last + 1)
                if all(sheet.cell(candidate, col).value is None for col in range(1, len(headers) + 1))
            ),
            last + 1,
        )
        for header, value in values.items():
            if header not in headers:
                raise AssertionError(f"unknown {marker} column: {header}")
            sheet.cell(row, headers.index(header) + 1).value = value
        return row

    def add_unit(self, workbook, name: str, unit_type: str = "全厂", *, enabled: str = "是", boundary_description: str | None = None) -> int:
        if workbook["核算单元"]["B3"].value in (None, ""):
            self.set_project(workbook)
        return self.append_section(workbook, "核算单元清单", {
            UNIT_HEADERS[0]: name,
            UNIT_HEADERS[1]: unit_type,
            UNIT_HEADERS[2]: enabled,
            UNIT_HEADERS[3]: boundary_description,
        })

    def add_fuel(
        self,
        workbook,
        unit_name: str,
        amount,
        *,
        fuel: FuelType = FuelType.NATURAL_GAS,
        path: str = "体积",
        carbon=None,
        oxidation=None,
        lhv=None,
        parameter_kind: str | None = None,
        parameter_ref: str | None = None,
        activity_kind: str | None = "计量/仪表记录",
        activity_ref: str | None = "活动数据-1",
    ) -> int:
        return self.append_section(workbook, "B.2 化石燃料", {
            FUEL_HEADERS[0]: unit_name,
            FUEL_HEADERS[1]: FUEL_LABEL_BY_TYPE[fuel],
            FUEL_HEADERS[2]: path,
            FUEL_HEADERS[3]: amount,
            FUEL_HEADERS[4]: lhv,
            FUEL_HEADERS[5]: carbon,
            FUEL_HEADERS[6]: oxidation,
            FUEL_HEADERS[7]: activity_kind,
            FUEL_HEADERS[8]: activity_ref,
            FUEL_HEADERS[9]: parameter_kind,
            FUEL_HEADERS[10]: parameter_ref,
        })

    def save_book(self, workbook) -> Path:
        path = Path(self.temp_root.name) / f"book-{len(list(Path(self.temp_root.name).glob('book-*.xlsx')))}.xlsx"
        workbook.save(path)
        workbook.close()
        return path

    def import_unit(self, workbook, *, index: int = 0):
        path = self.save_book(workbook)
        preview = ExcelWorkbookImporter(self.resolver).import_preview(path)
        return preview, preview.units[index]

    def test_template_has_four_user_sheets_and_hidden_metadata_only(self) -> None:
        content = create_template_bytes()
        self.assertGreater(len(content), 1000)
        generated = Path(self.temp_root.name) / "runtime.xlsx"
        write_template(generated)
        workbook = load_workbook(generated, data_only=False)
        self.assertEqual(tuple(name for name in workbook.sheetnames if workbook[name].sheet_state == "visible"), VISIBLE_SHEETS)
        self.assertEqual(tuple(name for name in workbook.sheetnames if name not in VISIBLE_SHEETS), ("__metadata__",))
        metadata = workbook["__metadata__"]
        self.assertEqual(metadata["B2"].value, TEMPLATE_ID)
        self.assertEqual(metadata["B3"].value, TEMPLATE_VERSION)
        self.assertEqual(metadata.sheet_state, "hidden")
        self.assertNotIn("排放源", workbook.sheetnames)
        self.assertNotIn("报告信息", workbook.sheetnames)
        self.assertNotIn("证据来源", workbook.sheetnames)
        self.assertFalse(any("回收" in str(cell.value) for sheet in workbook for row in sheet.iter_rows() for cell in row))
        self.assertEqual(len(workbook["燃料与能源"].tables), 3)
        self.assertEqual(len(workbook["过程排放"].tables), 5)
        self.assertEqual(len(next(iter(workbook["燃料与能源"].tables.values())).ref.split(":")), 2)
        self.assertEqual(BODY_ROWS, 15)
        self.assertTrue(any(item.formula1 == "=UnitNames" for sheet in workbook for item in sheet.data_validations.dataValidation))
        self.assertFalse(any(cell.data_type == "f" for sheet in workbook for row in sheet.iter_rows() for cell in row))
        workbook.close()

    def test_single_unit_fuel_standard_defaults_paths_and_provenance(self) -> None:
        workbook = self.new_book()
        self.add_unit(workbook, "天然气单元")
        self.add_fuel(workbook, "天然气单元", Decimal("1.5"), fuel=FuelType.NATURAL_GAS)
        self.add_fuel(workbook, "天然气单元", Decimal("2"), fuel=FuelType.ANTHRACITE, path="质量")
        self.add_fuel(workbook, "天然气单元", Decimal("3"), fuel=FuelType.NATURAL_GAS, path="热量")
        preview, unit = self.import_unit(workbook)
        self.assertTrue(unit.can_calculate, tuple(issue.message for issue in unit.errors))
        self.assertEqual(len(unit.input_value.fuel_inputs), 3)
        fuel_inputs = unit.input_value.fuel_inputs
        natural_gas_volume = next(item for item in fuel_inputs if item.fuel_type is FuelType.NATURAL_GAS and item.path is FuelPath.VOLUME)
        natural_gas_heat = next(item for item in fuel_inputs if item.fuel_type is FuelType.NATURAL_GAS and item.path is FuelPath.HEAT)
        anthracite = next(item for item in fuel_inputs if item.fuel_type is FuelType.ANTHRACITE)
        self.assertEqual(natural_gas_volume.path, FuelPath.VOLUME)
        self.assertEqual(anthracite.path, FuelPath.MASS)
        self.assertEqual(natural_gas_volume.carbon_content.source_kind.value, "STANDARD_DEFAULT")
        self.assertEqual(natural_gas_volume.carbon_content.unit, "tC/GJ")
        self.assertIn("C.1", natural_gas_volume.carbon_content.source_location)
        self.assertEqual(natural_gas_volume.activity.source_type.value, "METER")
        self.assertIsNone(natural_gas_heat.lower_heating_value)
        self.assertEqual(natural_gas_heat.carbon_content.unit, "tC/GJ")
        self.assertEqual(len(unit.input_value.reporting_data.activity_evidence), 3)
        self.assertIsNone(unit.calculation.record)
        self.assertEqual(len(preview.provenance.workbook_sha256), 64)
        self.assertTrue(any(item.sheet == "燃料与能源" and item.raw_cell_type == "n" for item in preview.numeric_evidence))

        measured_book = self.new_book()
        self.add_unit(measured_book, "实测参数")
        self.add_fuel(measured_book, "实测参数", 1, lhv=Decimal("400"), carbon=Decimal("0.014"), oxidation=Decimal("0.98"), parameter_kind="实测值", parameter_ref="化验报告-燃料-1")
        _preview, measured_unit = self.import_unit(measured_book)
        self.assertTrue(measured_unit.can_calculate, tuple(item.message for item in measured_unit.errors))
        measured_fuel = measured_unit.input_value.fuel_inputs[0]
        self.assertEqual(measured_fuel.carbon_content.source_kind.value, "MEASURED")
        self.assertIn("化验报告-燃料-1", measured_fuel.carbon_content.source_location)
        self.assertTrue(measured_fuel.carbon_content.evidence_ref_ids)

    def test_c1_resolver_has_verified_parameters_for_every_specific_selectable_fuel(self) -> None:
        for fuel_type, subject in FUEL_C1_SUBJECT_IDS.items():
            with self.subTest(fuel=fuel_type.value):
                for suffix in ("carbon_content", "oxidation_rate", "lhv"):
                    parameter_id = f"{subject}_{suffix}"
                    factors = self.resolver.repository.list_factors(parameter_id)
                    self.assertTrue(any(
                        factor.review_status.value == "VERIFIED"
                        and "gbt_32151_34_2024" in factor.applicable_standard_ids
                        for factor in factors
                    ), parameter_id)

    def test_numeric_ingress_rejects_text_formula_date_and_more_than_15_digits_but_accepts_zero(self) -> None:
        cases = (
            ("123.45", "EXCEL-NUMBER-CELL-REQUIRED"),
            ("1,234 t", "EXCEL-NUMBER-CELL-REQUIRED"),
            ("=1+2", "EXCEL-FORMULA-REJECTED"),
            (date(2025, 1, 2), "EXCEL-NUMBER-CELL-REQUIRED"),
            (Decimal("1234567890123456"), "EXCEL-NUMBER-SIGNIFICANT-DIGITS"),
        )
        for value, expected in cases:
            with self.subTest(value=value):
                workbook = self.new_book()
                self.add_unit(workbook, "坏数据")
                self.add_fuel(workbook, "坏数据", value)
                _preview, unit = self.import_unit(workbook)
                self.assertFalse(unit.can_calculate)
                self.assertIn(expected, {item.code for item in unit.errors})

        workbook = self.new_book()
        self.add_unit(workbook, "零活动量")
        self.add_fuel(workbook, "零活动量", Decimal("0"))
        preview, unit = self.import_unit(workbook)
        self.assertTrue(unit.can_calculate, tuple(item.message for item in unit.errors))
        self.assertEqual(unit.input_value.fuel_inputs[0].activity.value, Decimal(0))
        self.assertIn(Decimal(0), [item.normalized_decimal for item in preview.numeric_evidence])
        self.assertEqual(significant_digit_count(Decimal("0.0012300")), 5)

    def test_custom_factor_needs_traceable_source_and_fraction_is_validated_by_calculator(self) -> None:
        workbook = self.new_book()
        self.add_unit(workbook, "缺来源")
        self.add_fuel(workbook, "缺来源", 1, oxidation=Decimal("0.98"))
        _preview, unit = self.import_unit(workbook)
        self.assertFalse(unit.can_calculate)
        self.assertIn("EXCEL-PARAMETER-SOURCE-REQUIRED", {item.code for item in unit.errors})

        for value in (Decimal("-0.1"), Decimal("1.1")):
            with self.subTest(ratio=value):
                workbook = self.new_book()
                self.add_unit(workbook, "比例错误")
                self.add_fuel(workbook, "比例错误", 1, oxidation=value, parameter_kind="实测值", parameter_ref="检测报告-氧化率")
                _preview, unit = self.import_unit(workbook)
                self.assertFalse(unit.can_calculate)
                self.assertTrue(any("0 到 1" in item.message for item in unit.errors))

    def test_unit_isolation_derived_source_states_disabled_units_and_unassociated_rows(self) -> None:
        workbook = self.new_book()
        self.add_unit(workbook, "有效单元")
        self.add_unit(workbook, "错误单元", unit_type="工序")
        self.add_unit(workbook, "停用单元", unit_type="其他", enabled="否")
        self.add_fuel(workbook, "有效单元", 1)
        self.add_fuel(workbook, "错误单元", 1, fuel=FuelType.COAL, path="质量")
        self.add_fuel(workbook, "停用单元", 1)
        self.add_fuel(workbook, "不存在的单元", 1)
        preview = ExcelWorkbookImporter(self.resolver).import_preview(self.save_book(workbook))
        good, bad = preview.units
        self.assertTrue(good.can_calculate)
        self.assertFalse(bad.can_calculate)
        self.assertTrue(bad.errors)
        states = {item.source_id: item.status for item in good.input_value.source_states}
        self.assertEqual(states[SOURCE_FUEL], EmissionSourceStatus.INVOLVED)
        self.assertTrue(all(states[key] is EmissionSourceStatus.NOT_INVOLVED for key in SOURCE_IDS if key != SOURCE_FUEL))
        bad_states = {item.source_id: item.status for item in bad.input_value.source_states}
        self.assertEqual(bad_states[SOURCE_FUEL], EmissionSourceStatus.INVOLVED)
        self.assertTrue(any(item.code == "EXCEL-DISABLED-UNIT-DATA-IGNORED" for item in preview.warnings))
        self.assertTrue(any(item.code == "EXCEL-UNKNOWN-UNIT-DATA-IGNORED" for item in preview.warnings))
        self.assertIsNone(good.calculation.record)

    def _add_material(self, workbook, unit: str, marker: str, instance: str, category: str, name: str, mass, fixed, volatile=None, *, source="实测值", reference="化验报告-1"):
        return self.append_section(workbook, marker, {
            MATERIAL_HEADERS[0]: unit,
            MATERIAL_HEADERS[1]: instance,
            MATERIAL_HEADERS[2]: category,
            MATERIAL_HEADERS[3]: name,
            MATERIAL_HEADERS[4]: mass,
            MATERIAL_HEADERS[5]: fixed,
            MATERIAL_HEADERS[6]: volatile,
            MATERIAL_HEADERS[7]: source,
            MATERIAL_HEADERS[8]: reference,
        })

    def test_multi_material_rows_normalize_to_existing_calcination_baking_and_graphitization_inputs(self) -> None:
        workbook = self.new_book()
        self.add_unit(workbook, "多物料")
        self.add_fuel(workbook, "多物料", 0)
        for category, name, mass, fixed, volatile in (
            ("待煅烧原料", "原料一", 6, Decimal("0.7"), Decimal("0.1")),
            ("待煅烧原料", "原料二", 4, Decimal("0.9"), Decimal("0.2")),
            ("煅后料", "煅后焦", 8, Decimal("0.9"), Decimal("0.05")),
            ("欠烧煅料", "欠烧料", 1, Decimal("0.5"), None),
            ("炭粉尘", "粉尘", 1, Decimal("0.2"), None),
        ):
            self._add_material(workbook, "多物料", "B.3 煅烧", "煅烧一", category, name, mass, fixed, volatile)
        for category, name, mass, fixed, volatile in (
            ("填充料", "填充料", 2, Decimal("0.2"), Decimal("0.3")),
            ("待焙烧/炭化品", "生坯", 10, Decimal("0.8"), Decimal("0.1")),
            ("粉尘/碎屑/副产品", "副产品", 1, Decimal("0.1"), None),
            ("焙烧/炭化品", "焙烧品", 9, Decimal("0.85"), None),
        ):
            self._add_material(workbook, "多物料", "B.4 焙烧/炭化", "焙烧一", category, name, mass, fixed, volatile)
        for category, name, mass, fixed, volatile in (
            ("保温料/电阻料", "保温料", 2, Decimal("0.2"), Decimal("0.3")),
            ("待石墨化品", "生料", 10, Decimal("0.8"), None),
            ("粉尘/碎屑/残块/副产品", "残块", 1, Decimal("0.1"), None),
            ("石墨化产品", "石墨", 9, Decimal("0.85"), None),
        ):
            self._add_material(workbook, "多物料", "B.5 石墨化", "石墨化一", category, name, mass, fixed, volatile)
        _preview, unit = self.import_unit(workbook)
        self.assertTrue(unit.can_calculate, tuple(item.message for item in unit.errors))
        self.assertEqual(len(unit.input_value.calcinations), 1)
        self.assertEqual(len(unit.input_value.bakings), 1)
        self.assertEqual(len(unit.input_value.graphitizations), 1)
        calcination = unit.input_value.calcinations[0]
        self.assertEqual(calcination.gc.value, Decimal(10))
        self.assertEqual(calcination.wfc.value, Decimal("0.78"))
        self.assertEqual(calcination.wfc_c.value, Decimal("0.79"))
        self.assertEqual(calcination.k1.source_kind.value, "STANDARD_DEFAULT")
        self.assertIn("第5.2.2条", calcination.k1.source_location)
        self.assertEqual(len(unit.input_value.reporting_data.activity_evidence), 14)

    def test_complete_c2_factors_and_missing_carbonate_kind_is_blocked(self) -> None:
        workbook = self.new_book()
        self.add_unit(workbook, "C2完整表")
        for index, kind in enumerate(CARBONATE_LABELS, 1):
            self.append_section(workbook, "B.7 烟气脱硫", {
                "核算单元（必填）": "C2完整表",
                "设施/批次（必填）": f"设施{index}",
                "碳酸盐种类（必填）": kind,
                "脱硫剂消耗量（t，必填）": 1,
            })
        preview, unit = self.import_unit(workbook)
        self.assertTrue(unit.can_calculate, tuple(item.message for item in unit.errors))
        components = [component for facility in unit.input_value.fgd_units for component in facility.components]
        self.assertEqual(len(components), 11)
        self.assertEqual(len({item.emission_factor.parameter_id for item in components}), 11)
        for item in components:
            self.assertEqual(item.emission_factor.source_kind.value, "STANDARD_SPECIFIED")
            self.assertEqual(item.emission_factor.source_id, "SRC-32151-34-2024")
            self.assertTrue(item.emission_factor.source_version)
            self.assertIn("C.2", item.emission_factor.source_location)
            self.assertEqual(item.carbonate_fraction.value, Decimal("0.90"))
            self.assertEqual(item.conversion_rate.value, Decimal("1"))
        for kind, expected in (("CaCO₃", "0.440"), ("MgCO₃", "0.522"), ("Na₂CO₃", "0.415")):
            factor = next(item.emission_factor for item in components if item.carbonate_type == kind)
            self.assertEqual(factor.value, Decimal(expected))
        self.assertTrue(any(item.sheet == "过程排放" for item in preview.numeric_evidence))

        workbook = self.new_book()
        self.add_unit(workbook, "未知碳酸盐")
        self.append_section(workbook, "B.7 烟气脱硫", {
            "核算单元（必填）": "未知碳酸盐",
            "设施/批次（必填）": "脱硫设施一",
            "脱硫剂消耗量（t，必填）": 1,
        })
        _preview, unknown = self.import_unit(workbook)
        self.assertFalse(unknown.can_calculate)
        self.assertTrue(any("选择碳酸盐种类" in item.message and "CaCO₃" in item.message for item in unknown.errors))
        self.assertEqual(unknown.input_value.fgd_units, ())

    def test_multiple_electricity_heat_rows_and_c4_c5_calculator_anchors(self) -> None:
        workbook = self.new_book()
        self.add_unit(workbook, "多能源")
        for direction, amount, factor, reference in (
            ("输出", 2, Decimal("0.5"), "输出电力-1"),
            ("输出", 3, Decimal("0.4"), "输出电力-2"),
            ("购入", 100, None, None),
            ("购入", 200, None, None),
        ):
            values = {
                ELECTRICITY_HEADERS[0]: "多能源",
                ELECTRICITY_HEADERS[1]: direction,
                ELECTRICITY_HEADERS[2]: amount,
                ELECTRICITY_HEADERS[3]: factor,
                ELECTRICITY_HEADERS[6]: "实测值" if factor is not None else None,
                ELECTRICITY_HEADERS[7]: reference,
                ELECTRICITY_HEADERS[8]: "购入" if direction == "购入" else None,
                ELECTRICITY_HEADERS[9]: "常规" if direction == "购入" else None,
                ELECTRICITY_HEADERS[10]: "无" if direction == "购入" else None,
                ELECTRICITY_HEADERS[11]: "未提供" if direction == "购入" else None,
            }
            self.append_section(workbook, "B.8 电力", values)
        self.append_section(workbook, "B.9 热力", {
            HEAT_HEADERS[0]: "多能源", HEAT_HEADERS[1]: "购入", HEAT_HEADERS[2]: 1000,
            HEAT_HEADERS[3]: "饱和蒸汽", HEAT_HEADERS[4]: 2675,
        })
        self.append_section(workbook, "B.9 热力", {
            HEAT_HEADERS[0]: "多能源", HEAT_HEADERS[1]: "输出", HEAT_HEADERS[2]: 500,
            HEAT_HEADERS[3]: "饱和蒸汽", HEAT_HEADERS[4]: 2675,
            HEAT_HEADERS[7]: Decimal("0.12"), HEAT_HEADERS[10]: "实测值", HEAT_HEADERS[11]: "热力因子报告-1",
        })
        _preview, unit = self.import_unit(workbook)
        self.assertTrue(unit.can_calculate, tuple(item.message for item in unit.errors))
        self.assertEqual(len(unit.input_value.exported_electricity), 2)
        self.assertEqual([item.factor.value for item in unit.input_value.exported_electricity], [Decimal("0.5"), Decimal("0.4")])
        self.assertEqual(len(unit.input_value.electricity_details), 2)
        self.assertEqual(len(unit.input_value.purchased_heat), 1)
        self.assertEqual(len(unit.input_value.exported_heat), 1)
        self.assertTrue(any(item.parameter_id == "heat_emission_factor_default" for item in unit.calculation.parameter_snapshots))
        self.assertEqual(unit.input_value.exported_electricity[0].factor.source_kind.value, "MEASURED")
        self.assertEqual(len(unit.input_value.reporting_data.measured_factor_evidence), 3)
        self.assertEqual(unit.source_breakdown[next(key for key in SOURCE_IDS if SOURCE_LABELS[key] == "输出电力")], Decimal("2.2"))
        for pressure in (Decimal("1.7"), Decimal("1.8")):
            enthalpy, interpolated, _endpoints = saturated_steam_enthalpy(pressure)
            self.assertIsInstance(enthalpy, Decimal)
            self.assertFalse(interpolated)
        self.assertEqual(saturated_steam_enthalpy("1.70")[0], Decimal("2793.8"))
        superheated_anchor, interpolated, endpoints = superheated_steam_enthalpy("1.5", "325")
        self.assertGreater(superheated_anchor, Decimal("3000"))
        self.assertTrue(interpolated)
        self.assertEqual(endpoints, (Decimal("1"), Decimal("3")))

    def test_gui_and_excel_semantic_parity_without_aligning_generated_ids(self) -> None:
        workbook = self.new_book()
        self.set_project(workbook, "示例企业")
        self.add_unit(workbook, "全厂")
        fuel_row = self.add_fuel(
            workbook,
            "全厂",
            Decimal("1.25"),
            fuel=FuelType.NATURAL_GAS,
            path="体积",
            activity_kind=None,
            activity_ref=None,
        )
        for category, name, mass, fixed, volatile in (
            ("待煅烧原料", "原料一", 6, Decimal("0.7"), Decimal("0.1")),
            ("待煅烧原料", "原料二", 4, Decimal("0.9"), Decimal("0.2")),
            ("煅后料", "煅后焦", 8, Decimal("0.9"), Decimal("0.05")),
            ("欠烧煅料", "欠烧料", 1, Decimal("0.5"), None),
            ("炭粉尘", "粉尘", 1, Decimal("0.2"), None),
        ):
            self._add_material(workbook, "全厂", "B.3 煅烧", "煅烧一", category, name, mass, fixed, volatile)

        self.append_section(workbook, "B.8 电力", {
            ELECTRICITY_HEADERS[0]: "全厂",
            ELECTRICITY_HEADERS[1]: "购入",
            ELECTRICITY_HEADERS[2]: 100,
            ELECTRICITY_HEADERS[8]: "购入",
            ELECTRICITY_HEADERS[9]: "常规",
            ELECTRICITY_HEADERS[10]: "无",
            ELECTRICITY_HEADERS[11]: "未提供",
        })
        power_rows = ((2, Decimal("0.50"), "输出电力报告-A"), (3, Decimal("0.40"), "输出电力报告-B"))
        for amount, factor, reference in power_rows:
            self.append_section(workbook, "B.8 电力", {
                ELECTRICITY_HEADERS[0]: "全厂",
                ELECTRICITY_HEADERS[1]: "输出",
                ELECTRICITY_HEADERS[2]: amount,
                ELECTRICITY_HEADERS[3]: factor,
                ELECTRICITY_HEADERS[6]: "实测值",
                ELECTRICITY_HEADERS[7]: reference,
            })
        heat_rows = (
            ("购入", 1000, Decimal("0.11"), "购入热力报告-A"),
            ("购入", 500, Decimal("0.10"), "购入热力报告-B"),
            ("输出", 250, Decimal("0.12"), "输出热力报告-A"),
        )
        for direction, amount, factor, reference in heat_rows:
            self.append_section(workbook, "B.9 热力", {
                HEAT_HEADERS[0]: "全厂",
                HEAT_HEADERS[1]: direction,
                HEAT_HEADERS[2]: amount,
                HEAT_HEADERS[3]: "饱和蒸汽",
                HEAT_HEADERS[4]: 2675,
                HEAT_HEADERS[7]: factor,
                HEAT_HEADERS[10]: "实测值",
                HEAT_HEADERS[11]: reference,
            })

        book_path = self.save_book(workbook)
        preview = ExcelWorkbookImporter(self.resolver).import_preview(book_path)
        excel = preview.units[0]
        self.assertTrue(excel.can_calculate)

        page = CarbonMaterialAccountingPage(
            catalog_service=self.catalog_service,
            record_repository=InMemoryRecordRepository(),
        )
        page.enterprise_name.setText("示例企业")
        page.period_year.setValue(2025)
        page.boundary_confirmed.setChecked(True)
        active_sources = {
            SOURCE_FUEL,
            SOURCE_CALCINATION,
            SOURCE_PURCHASED_ELECTRICITY,
            SOURCE_EXPORTED_ELECTRICITY,
            SOURCE_PURCHASED_HEAT,
            SOURCE_EXPORTED_HEAT,
        }
        for source_id, combo in page._source_statuses.items():
            status = EmissionSourceStatus.INVOLVED if source_id in active_sources else EmissionSourceStatus.NOT_INVOLVED
            combo.setCurrentIndex(combo.findData(status))

        fuel = page._fuel_rows[0]
        fuel.fuel_type.setCurrentIndex(fuel.fuel_type.findData(FuelType.NATURAL_GAS))
        fuel.path.setCurrentIndex(fuel.path.findData(FuelPath.VOLUME))
        page._refresh_fuel_defaults()
        fuel.activity.setText("1.25")

        process = page._process_rows["calcination"][0]
        for field, value in {
            "gc": "10",
            "wfc": "78",
            "cc": "8",
            "ucc": "1",
            "du": "1",
            "wfc_c": "79",
            "wvar": "14",
            "wvar_c": "5",
        }.items():
            process["fields"][field].setText(value)

        purchased_power = page._electricity_rows[0]
        purchased_power.amount.setText("100")
        purchased_power.acquisition.setCurrentIndex(
            purchased_power.acquisition.findData(ElectricityAcquisitionMode.PURCHASED)
        )
        purchased_power.attribute.setCurrentIndex(
            purchased_power.attribute.findData(ElectricityAttribute.ORDINARY)
        )
        purchased_power.proof_type.setCurrentIndex(
            purchased_power.proof_type.findData(ElectricityProofType.NONE)
        )
        purchased_power.proof_status.setCurrentIndex(
            purchased_power.proof_status.findData(ElectricityProofStatus.NOT_PROVIDED)
        )
        for index, (amount, factor, reference) in enumerate(power_rows):
            row = page._output_electricity_rows[0] if index == 0 else page._add_output_electricity_row()
            row["amount"].setText(str(amount))
            row["measured"].setText(str(factor))
            row["source"].setText(reference)
        for prefix, entries in (("heat", heat_rows[:2]), ("exported_heat", heat_rows[2:])):
            for index, (_direction, amount, factor, reference) in enumerate(entries):
                row = page._heat_rows[prefix][0] if index == 0 else page._add_heat_row(prefix)
                row["amount"].setText(str(amount))
                row["enthalpy"].setText("2675")
                row["steam"].setCurrentIndex(row["steam"].findData(SteamKind.SATURATED))
                row["measured"].setText(str(factor))
                row["source"].setText(reference)

        gui_input = page._input(increment=False, render_electricity=False)
        gui_records = InMemoryRecordRepository()
        gui_result = CarbonMaterialCalculator(
            parameter_resolver=self.resolver,
            record_repository=gui_records,
        ).calculate(gui_input)
        self.assertTrue(gui_result.successful, gui_result.problems)
        self.assertEqual(_business_semantics(gui_input), _business_semantics(excel.input_value))
        self.assertEqual(
            _calculation_semantics(gui_result),
            _calculation_semantics(excel.calculation),
        )
        self.assertEqual(gui_records.list_all()[0].status.value, "COMPLETED")
        self.assertIsNone(excel.calculation.record)
        self.assertTrue(excel.can_calculate)
        self.assertEqual(excel.errors, ())
        self.assertEqual(excel.warnings, ())
        self.assertEqual(preview.warnings, ())
        self.assertEqual(len(gui_input.calcinations), 1)
        calcination = gui_input.calcinations[0]
        excel_calcination = excel.input_value.calcinations[0]
        self.assertEqual(
            tuple(getattr(calcination, field).value for field in ("gc", "wfc", "cc", "ucc", "du", "wfc_c", "wvar", "wvar_c")),
            tuple(getattr(excel_calcination, field).value for field in ("gc", "wfc", "cc", "ucc", "du", "wfc_c", "wvar", "wvar_c")),
        )
        self.assertEqual(tuple(getattr(excel_calcination, field).value for field in ("gc", "wfc", "wfc_c", "wvar", "wvar_c")),
                         (Decimal("10"), Decimal("0.78"), Decimal("0.79"), Decimal("0.14"), Decimal("0.05")))
        self.assertEqual(len(gui_input.exported_electricity), 2)
        self.assertEqual(len(excel.input_value.exported_electricity), 2)
        self.assertEqual(len(gui_input.purchased_heat), 2)
        self.assertEqual(len(excel.input_value.purchased_heat), 2)
        self.assertGreater(len(preview.numeric_evidence), 0)
        self.assertEqual(
            {item.source_reference for item in excel.input_value.reporting_data.measured_factor_evidence},
            {"输出电力报告-A", "输出电力报告-B", "购入热力报告-A", "购入热力报告-B", "输出热力报告-A"},
        )

        bad_workbook = load_workbook(book_path)
        bad_workbook["燃料与能源"].cell(fuel_row, FUEL_HEADERS.index(FUEL_HEADERS[3]) + 1).value = Decimal("-1")
        bad_path = Path(self.temp_root.name) / "parity-negative.xlsx"
        bad_workbook.save(bad_path)
        bad_workbook.close()
        bad_preview = ExcelWorkbookImporter(self.resolver).import_preview(bad_path)
        bad_excel = bad_preview.units[0]
        self.assertFalse(bad_excel.can_calculate)
        self.assertIsNotNone(bad_excel.calculation)
        assert bad_excel.calculation is not None
        self.assertFalse(bad_excel.calculation.successful)
        self.assertIsNone(bad_excel.calculation.result)
        self.assertTrue(any(item.code == "CAR-VAL-NONNEGATIVE" for item in bad_excel.calculation.problems))

        # Change exactly one Domain activity value in the GUI-generated payload.
        # Its generated identities remain untouched; this is solely the negative
        # case used to prove that the same illegal business value is blocked.
        bad_gui_input = replace(
            gui_input,
            fuel_inputs=(replace(gui_input.fuel_inputs[0], activity=Decimal("-1")),),
        )
        bad_gui_records = InMemoryRecordRepository()
        gui_bad = CarbonMaterialCalculator(
            parameter_resolver=self.resolver,
            record_repository=bad_gui_records,
        ).calculate(bad_gui_input)
        page.close()
        self.assertNotEqual(_business_semantics(gui_input), _business_semantics(bad_gui_input))
        self.assertEqual(_business_semantics(bad_gui_input), _business_semantics(bad_excel.input_value))
        self.assertFalse(gui_bad.successful)
        self.assertIsNone(gui_bad.result)
        self.assertTrue(any(item.code == "CAR-VAL-NONNEGATIVE" for item in gui_bad.problems))
        self.assertEqual(bad_gui_records.list_all(), ())
        self.assertEqual(
            tuple((problem.code, problem.level.value) for problem in gui_bad.problems),
            tuple((problem.code, problem.level.value) for problem in bad_excel.calculation.problems),
        )

    def test_preview_shows_business_breakdown_and_never_persists_record(self) -> None:
        workbook = self.new_book()
        self.add_unit(workbook, "预览单元")
        self.add_fuel(workbook, "预览单元", 1, fuel=FuelType.DIESEL, path="质量")
        path = self.save_book(workbook)
        page = ExcelImportPage(catalog_service=self.catalog_service)
        page.file_input.setText(str(path))
        with patch("packages.ui.pages.QMessageBox.information") as information:
            page._preview_workbook()
        displayed = information.call_args.args[2]
        for label in ("化石燃料：", "生产过程：", "烟气治理：", "购入电力：", "购入热力：", "输出电力抵扣：", "输出热力抵扣：", "直接排放：", "净间接排放：", "排放总量："):
            self.assertIn(label, displayed)
        self.assertNotIn("CAR-", displayed)
        self.assertIn("此预览不会保存项目或生成正式核算记录。", displayed)
        issue = type("Issue", (), {"location": "燃料与能源!D13", "message": "请输入有效数值。"})()
        self.assertEqual(page._format_excel_issue(issue, "· "), "· 燃料与能源!D13：请输入有效数值。")
        internal_issue = type("Issue", (), {
            "location": "燃料与能源!A13", "message": "排放源 CAR-SRC-FUEL-001 已标记涉及但缺少输入。"
        })()
        formatted = page._format_excel_issue(internal_issue, "· ")
        self.assertNotIn("CAR-", formatted)
        page.close()

    def test_malformed_metadata_or_missing_core_sheet_is_fatal_and_adapter_has_no_qt_dependency(self) -> None:
        workbook = self.new_book()
        workbook["__metadata__"]["B2"] = "unexpected"
        path = self.save_book(workbook)
        with self.assertRaises(WorkbookFatalError):
            ExcelWorkbookImporter(self.resolver).import_preview(path)

        workbook = self.new_book()
        workbook.remove(workbook["过程排放"])
        path = self.save_book(workbook)
        with self.assertRaises(WorkbookFatalError):
            ExcelWorkbookImporter(self.resolver).import_preview(path)

        root = str(Path(__file__).parents[1])
        script = "import sys; import packages.excel.gbt32151_34_v1; print('PySide6' in sys.modules)"
        result = subprocess.run([sys.executable, "-c", script], cwd=root, capture_output=True, text=True, check=True)
        self.assertEqual(result.stdout.strip(), "False")

    def test_numeric_helpers_and_nonfinite_serialized_lexemes(self) -> None:
        self.assertEqual(significant_digit_count(Decimal("0")), 1)
        self.assertEqual(significant_digit_count(Decimal("123456789012345")), 15)
        self.assertEqual(significant_digit_count(Decimal("1.234567890123456")), 16)
        self.assertEqual(significant_digit_count(Decimal("1E-6")), 1)
        from openpyxl import Workbook

        cell = Workbook().active["A1"]
        ctx = _UnitContext("单元", "u", self._whole_site(), "e", "企业", None, False, None)
        reader = _CellReader({("测试", "A1"): "NaN"})
        self.assertIsNone(reader.number(ctx, "测试", 1, cell))
        self.assertEqual(ctx.errors[-1].code, "EXCEL-NUMBER-NONFINITE")
        for nonfinite in ("Infinity", "-Infinity"):
            ctx = _UnitContext("单元", "u", self._whole_site(), "e", "企业", None, False, None)
            reader = _CellReader({("测试", "A1"): nonfinite})
            self.assertIsNone(reader.number(ctx, "测试", 1, cell))
            self.assertEqual(ctx.errors[-1].code, "EXCEL-NUMBER-NONFINITE")

    def test_nonfinite_numeric_lexemes_are_rejected_through_the_real_import_path(self) -> None:
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        for lexeme in ("NaN", "Infinity", "-Infinity"):
            with self.subTest(lexeme=lexeme):
                workbook = self.new_book()
                self.add_unit(workbook, "非有限数")
                fuel_row = self.add_fuel(workbook, "非有限数", Decimal("1.25"))
                source_path = self.save_book(workbook)
                edited_path = Path(self.temp_root.name) / f"nonfinite-{lexeme.replace('-', 'minus')}.xlsx"
                members: dict[str, bytes] = {}
                with ZipFile(source_path, "r") as archive:
                    members = {name: archive.read(name) for name in archive.namelist()}
                root = ElementTree.fromstring(members["xl/worksheets/sheet3.xml"])
                cell = root.find(f".//{{{ns}}}c[@r='D{fuel_row}']")
                self.assertIsNotNone(cell)
                value_node = cell.find(f"{{{ns}}}v")
                self.assertIsNotNone(value_node)
                value_node.text = lexeme
                members["xl/worksheets/sheet3.xml"] = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
                with ZipFile(edited_path, "w", ZIP_DEFLATED) as archive:
                    for name, content in members.items():
                        archive.writestr(name, content)
                with self.assertRaises(WorkbookFatalError):
                    ExcelWorkbookImporter(self.resolver).import_preview(edited_path)

    @staticmethod
    def _whole_site():
        from packages.application.project_workspaces import AccountingUnitType
        return AccountingUnitType.WHOLE_SITE


if __name__ == "__main__":
    unittest.main()
