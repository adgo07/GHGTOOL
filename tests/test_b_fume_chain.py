"""Issue020真实正式入口/GUI/B.6/SQLite/冻结Word整链；非人工CU。"""
import os
os.environ.setdefault("QT_QPA_PLATFORM","offscreen")
from dataclasses import replace
from datetime import date,datetime,timedelta
from decimal import Decimal,localcontext,ROUND_DOWN
from io import BytesIO
from pathlib import Path
import json,tempfile,unittest
from unittest.mock import patch
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication,QMessageBox
import tests.test_b_fgd_acceptance_fixes as qt_fixtures
import tests.test_b_fume_segments as fume_fixtures
from packages.standards.carbon_material import CarbonMaterialInput,CarbonMaterialCalculator,EmissionSourceState,EmissionSourceStatus,SOURCE_FUME,FumeIncinerationInput,InputValue
from packages.core.models import AccountingPeriod,PeriodType
from packages.application.carbon_accounting import CarbonAccountingUseCase
from packages.application.canonical_input_codec import encode_canonical_input,decode_canonical_input
from packages.application import ProjectWorkspaceService
from packages.persistence import SQLiteRecordRepository,SQLiteProjectWorkspaceRepository
from packages.excel.templates import ExcelTemplateService
from packages.ui.view_models import AppRoute
from packages.application.reporting.model import build_report_model
from packages.infrastructure.reporting import render_report_docx as render_docx

PERIOD=AccountingPeriod(PeriodType.ANNUAL,date(2025,1,1),date(2025,12,31))
def payload(*segments,period=PERIOD,**kwargs):
    return CarbonMaterialInput("fume-chain","enterprise.fume","烟气测试",period,boundary_confirmed=True,
        source_states=(EmissionSourceState(SOURCE_FUME,EmissionSourceStatus.INVOLVED),),
        fume_segments=tuple(segments),fume_execution_version=kwargs.pop("fume_execution_version","FUME_V2"),**kwargs)

def report(repository,record_id):
    return build_report_model(repository.get(record_id),raw_input_snapshot=repository.get_raw_input_snapshot(record_id),
        trace_snapshot=repository.get_trace_snapshot(record_id),reporting_snapshot=repository.get_reporting_snapshot(record_id),
        provenance_snapshot=repository.get_provenance_snapshot(record_id),report_qualification=repository.get_report_qualification(record_id))

class FumeFormalTests(unittest.TestCase):
    def test_mapping_partial_formal_snapshot_and_word_frozen(self):
        with tempfile.TemporaryDirectory() as directory:
            records=SQLiteRecordRepository(Path(directory)/"records.sqlite")
            calculator=CarbonMaterialCalculator();usecase=CarbonAccountingUseCase(calculator,records)
            value=payload(fume_fixtures.segment(),fume_fixtures.segment(1,3,sid="s2",q="500",qvar="2000"))
            result=usecase.calculate(value)
            self.assertTrue(result.successful,result.problems)
            self.assertEqual(result.record.calculation_result.total_amount,Decimal(".2112"))
            self.assertTrue(any(p.code=="CAR-WARN-FUME-COVERAGE" for p in result.problems))
            trace=records.get_trace_snapshot(result.record.record_id)
            self.assertEqual(trace["fume_monitoring"]["coverage"][0]["covered_duration_days"],"3")
            self.assertEqual(len(trace["fume_monitoring"]["coverage"][0]["uncovered_intervals"]),1)
            self.assertEqual(len(result.parameter_snapshots),2)
            with patch.object(calculator,"calculate",side_effect=AssertionError("历史不得重算")):
                model=report(records,result.record.record_id)
                section=next(s for s in model.sections if s.section_id=="b6")
                self.assertEqual(len(section.tables[0].rows),2)
                self.assertIn("不补零",str(section))
                self.assertIn("2025-01-04",str(section))
                self.assertIn("监测段",str(section.tables[0].rows[0].note))
                data=render_docx(model)
                self.assertGreater(len(data),30000)
            self.assertEqual(len(records.list_all()),1)
    def test_full_coverage_hours_and_per_segment_parameters(self):
        for duration in ("1",InputValue("24","h")):
            result=CarbonMaterialCalculator().calculate(payload(fume_fixtures.segment(duration=duration),period=AccountingPeriod(PeriodType.CUSTOM,date(2025,1,1),date(2025,1,1))))
            self.assertTrue(result.successful,result.problems)
            fume_fixtures.FumeSegmentArithmeticTests().assertMappingValue(result.result.total_amount,".0704")
            self.assertFalse(any(p.code=="CAR-WARN-FUME-COVERAGE" for p in result.problems))
        value=payload(fume_fixtures.segment(),fume_fixtures.segment(1,2,sid="s2",hm="20",fch=".04",fox=".5"))
        fume_fixtures.FumeSegmentArithmeticTests().assertMappingValue(CarbonMaterialCalculator().calculate(value).result.total_amount,".1056")
    def test_overlap_bad_values_duplicate_duration_mismatch_no_record(self):
        with tempfile.TemporaryDirectory() as directory:
            records=SQLiteRecordRepository(Path(directory)/"records.sqlite")
            usecase=CarbonAccountingUseCase(CarbonMaterialCalculator(),records)
            first=fume_fixtures.segment(0,2)
            bads=(fume_fixtures.segment(1,3,sid="s2"),fume_fixtures.segment(2,3,sid="bad",q=None),
                fume_fixtures.segment(2,3,sid="bad",fch=None),fume_fixtures.segment(2,3,sid="bad",duration="365"),
                fume_fixtures.segment(2,3,sid="segment-1"),fume_fixtures.segment(2,3,sid="bad",fox="1.1"))
            for bad in bads:
                with self.subTest(bad=bad):
                    result=usecase.calculate(payload(first,bad));self.assertFalse(result.successful,result.problems)
                    self.assertEqual(len(records.list_all()),0)
            self.assertFalse(usecase.calculate(payload()).successful)
            self.assertEqual(len(records.list_all()),0)
    def test_cross_instance_and_boundary_not_mutual_coverage_and_ambient(self):
        value=payload(fume_fixtures.segment(0,180),fume_fixtures.segment(180,365,source="source-B"))
        results=[]
        for precision in (6,17,60):
            with localcontext() as ctx:
                ctx.prec=precision;ctx.rounding=ROUND_DOWN
                result=CarbonMaterialCalculator().calculate(value)
                self.assertTrue(result.successful,result.problems)
                results.append(result.result.total_amount)
                self.assertEqual(sum(p.code=="CAR-WARN-FUME-COVERAGE" for p in result.problems),2)
        self.assertEqual(results,[Decimal("25.696")]*3)
    def test_zero_valid_legacy_read_save_and_old_codec_no_fake_intervals(self):
        value=payload(fume_fixtures.segment(q="0"))
        self.assertEqual(CarbonMaterialCalculator().calculate(value).result.total_amount,0)
        self.assertEqual(decode_canonical_input(encode_canonical_input(value)),value)
        old=payload(fume_execution_version="FUME_V1",fume_incinerations=(FumeIncinerationInput(q="1000",qvar="1000",hm="40",fch=".02",fox="1",duration="1"),))
        encoded=json.loads(encode_canonical_input(old));legacy=encoded["payload"];legacy["version"]=3
        legacy["value"]["fields"].pop("fume_segments");legacy["value"]["fields"].pop("fume_execution_version")
        for item in legacy["value"]["fields"]["fume_incinerations"]["items"]:
            item["fields"].pop("source_name");item["fields"].pop("boundary_name")
        encoded["input_version"]=3
        restored=decode_canonical_input(json.dumps(encoded))
        self.assertEqual(restored.fume_segments,())
        outcome=CarbonMaterialCalculator().calculate(restored)
        self.assertTrue(outcome.successful,outcome.problems)
        self.assertTrue(any(p.code=="CAR-WARN-FUME-LEGACY-COVERAGE" for p in outcome.problems))
        self.assertFalse(CarbonMaterialCalculator().calculate(replace(restored,fume_execution_version="FUME_V2")).successful)

    def test_full_annual_and_gap_freeze_separate_coverage_qualification(self):
        with tempfile.TemporaryDirectory() as directory:
            records=SQLiteRecordRepository(Path(directory)/"records.sqlite")
            usecase=CarbonAccountingUseCase(CarbonMaterialCalculator(),records)
            for start,complete,expected in ((180,True,"25.696"),(185,False,"25.344")):
                result=usecase.calculate(payload(fume_fixtures.segment(0,180),fume_fixtures.segment(start,365,sid="s2",q="500",qvar="2000")))
                self.assertTrue(result.successful,result.problems)
                fume_fixtures.FumeSegmentArithmeticTests().assertMappingValue(result.result.total_amount,expected)
                saved=records.get_trace_snapshot(result.record.record_id)["fume_monitoring"]
                self.assertEqual(saved["complete"],complete);self.assertEqual(saved["coverage"][0]["complete"],complete)
                self.assertTrue(records.get_report_qualification(result.record.record_id)["eligible"])
                section=next(x for x in report(records,result.record.record_id).sections if x.section_id=="b6")
                self.assertEqual("覆盖不完整" in str(section),not complete)
                self.assertIn("监测段时长",section.tables[0].columns[5])
    def test_old_canonical_upgrade_preserves_original_sqlite_and_does_not_invent_segments(self):
        from packages.application.carbon_material_implementation import CarbonMaterialExcel
        old=payload(fume_execution_version="FUME_V1",fume_incinerations=(FumeIncinerationInput(q="1",qvar="1",hm="40",fch=".02",fox="1",duration="365"),))
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);projects=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/"projects.sqlite"))
            records=SQLiteRecordRepository(root/"records.sqlite")
            workspace=projects.new_workspace("旧烟气")
            workspace=replace(workspace,units=(replace(workspace.units[0],canonical_input=old),));projects.save(workspace)
            self.assertTrue(CarbonMaterialExcel.can_upgrade_workspace(workspace))
            upgraded=CarbonMaterialExcel.upgrade_workspace(workspace);projects.save(upgraded)
            reopened=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/"projects.sqlite"))
            self.assertEqual(reopened.get(workspace.project_id).units[0].canonical_input,old)
            value=reopened.get(upgraded.project_id).units[0].canonical_input
            self.assertNotEqual(upgraded.project_id,workspace.project_id)
            self.assertEqual(value.fume_incinerations,old.fume_incinerations);self.assertEqual(value.fume_segments,())
            self.assertEqual(value.fume_execution_version,"FUME_V2")
            self.assertFalse(CarbonAccountingUseCase(CarbonMaterialCalculator(),records).calculate(value).successful)
            self.assertEqual(len(records.list_all()),0)
    def test_same_segment_id_across_real_boundaries_has_unique_result_and_snapshot(self):
        with tempfile.TemporaryDirectory() as directory:
            records=SQLiteRecordRepository(Path(directory)/"records.sqlite")
            result=CarbonAccountingUseCase(CarbonMaterialCalculator(),records).calculate(payload(fume_fixtures.segment(),fume_fixtures.segment(boundary="other-stack")))
            self.assertTrue(result.successful,result.problems)
            lines=[line for line in result.result.lines if line.emission_source_id==SOURCE_FUME]
            self.assertEqual(len({line.line_id for line in lines}),2)
            self.assertEqual(len({item.detail_id for item in result.parameter_snapshots}),2)
            section=next(x for x in report(records,result.record.record_id).sections if x.section_id=="b6")
            self.assertEqual(len(section.tables),2)

    def test_declared_custom_profiles_reach_formal_segments_without_fallback(self):
        from packages.core.decimal_policy import DecimalPolicy
        from decimal import ROUND_HALF_EVEN,ROUND_HALF_UP
        value=payload(fume_fixtures.segment(q="1234.56789123456789123456789123456789",qvar="987.654321987654321"))
        for precision,rounding in ((28,ROUND_HALF_EVEN),(40,ROUND_HALF_UP),(50,ROUND_DOWN)):
            policy=DecimalPolicy(precision=precision,rounding=rounding)
            calculator=CarbonMaterialCalculator(policy=policy)
            expected=fume_fixtures.evaluate(value.fume_segments,policy=policy,numeric_profile_id=calculator.numeric_profile_id).covered_amount_tco2
            for ambient in (6,17,60):
                with localcontext() as context:
                    context.prec=ambient;context.rounding=ROUND_DOWN
                    result=calculator.calculate(value)
                    self.assertTrue(result.successful,result.problems)
                    self.assertEqual(result.result.total_amount,expected)
                    self.assertEqual(json.loads(result.evidence.trace_snapshot_json)["fume_monitoring"]["numeric_profile_id"],calculator.numeric_profile_id)
                    self.assertEqual(json.loads(result.evidence.provenance_snapshot_json)["numeric_provenance"]["profile_id"],calculator.numeric_profile_id)
                    self.assertEqual(context.prec,ambient)

class FumeQtExcelTests(unittest.TestCase):
    fresh_page=qt_fixtures.FGDAcceptanceFixTests.fresh_page
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    @staticmethod
    def prepare(gui):
        gui.enterprise_name.setText("烟气整链");gui.period_year.setValue(2025);gui.boundary_confirmed.setChecked(True)
        for status in gui._source_statuses.values():status.setCurrentIndex(status.findData("NOT_INVOLVED"))
        status=gui._source_statuses[SOURCE_FUME];status.setCurrentIndex(status.findData("INVOLVED"))
    @staticmethod
    def fill(row,start="2025-01-01 00:00:00",end="2025-01-02 00:00:00",duration="1",q="1000",qvar="1000"):
        for field,value in dict(q=q,qvar=qvar,hm="40",fch=".02",fox="100",duration=duration).items():row["fields"][field].setText(value)
        row["segment_start"].setText(start);row["segment_end"].setText(end)
    @staticmethod
    def workbook(root,rows):
        book=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
        for i,values in enumerate(rows,3):
            for col,value in zip("ABCDEFHIJKLM",values):book["B.6"][f"{col}{i}"]=value
        path=root/"烟气分段.xlsx";book.save(path);book.close();return path
    @staticmethod
    def preview(gui,shell,path):
        from packages.excel.appendix_b import AppendixBImportContext
        return shell.pages[AppRoute.EXCEL_IMPORT].excel_adapter.import_preview(path,context=AppendixBImportContext(gui._period(),boundary_confirmed=True,enterprise_name="烟气整链"))
    def test_gui_excel_canonical_project_reopen_formal_word_parity(self):
        with tempfile.TemporaryDirectory() as directory,self.fresh_page(Path(directory)) as (gui,projects,records,shell):
            root=Path(directory);self.prepare(gui);first=gui._process_rows["fume"][0];self.fill(first)
            second=gui._add_process_row("fume",instance_id="second-segment")
            self.fill(second,"2025-01-02 00:00:00","2025-01-04 00:00:00","2","500","2000")
            for key in ("source","boundary"):
                combo=second[key+"_selector"];combo.setCurrentIndex(combo.findData(first[key+"_selector"].currentData()))
            value=gui._input(increment=False);expected=gui.calculator.calculate(value)
            self.assertTrue(expected.successful,expected.problems);self.assertEqual(expected.result.total_amount,Decimal(".2112"))
            path=self.workbook(root,[(1000,1000,40,.02,100,1,"A","炉口","1","2025-01-01 00:00:00","2025-01-02 00:00:00","d"),
                (500,2000,40,.02,100,48,"A","炉口","2","2025-01-02 00:00:00","2025-01-04 00:00:00","h")])
            preview=self.preview(gui,shell,path);unit=preview.units[0]
            self.assertTrue(unit.can_calculate,unit.errors);canonical=unit.input_value
            self.assertEqual(len(canonical.fume_segments),2)
            self.assertEqual(canonical.fume_segments[0].source_instance_id,canonical.fume_segments[1].source_instance_id)
            self.assertEqual(gui.calculator.calculate(canonical).result.total_amount,expected.result.total_amount)
            self.assertEqual(len(records.list_all()),0)
            workspace=projects.new_workspace("烟气表格")
            provenance={"source":"EXCEL_APPENDIX_B","workbook":{"template_version":preview.provenance.template_version,"template_master_sha256":preview.provenance.template_master_sha256}}
            projects.save(replace(workspace,units=(replace(workspace.units[0],canonical_input=canonical,ingress_provenance=provenance),)))
            reopened=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/"projects.sqlite")).get(workspace.project_id)
            usecase=CarbonAccountingUseCase(gui.calculator,records)
            result=usecase.calculate(reopened.units[0].canonical_input,ingress_provenance=reopened.units[0].ingress_provenance)
            self.assertTrue(result.successful,result.problems);self.assertEqual(len(records.list_all()),1)
            with patch.object(gui.calculator,"calculate",side_effect=AssertionError("不得重算")),patch.object(gui._parameter_resolver,"resolve",side_effect=AssertionError("不得回查")):
                model=report(records,result.record.record_id);self.assertEqual(model.input_template_version,"2.7.0")
                self.assertIn("不补零",str(model));self.assertGreater(len(render_docx(model)),30000)
                from packages.ui.record_tables import ReportTableCard
                from PySide6.QtWidgets import QLabel
                from packages.application.reporting.presentation import report_text,report_html
                table=next(section for section in model.sections if section.section_id=="b6").tables[0]
                card=ReportTableCard(table)
                labels="\n".join(label.text() for label in card.findChildren(QLabel))
                self.assertIn("监测段：2025-01-01",labels);self.assertIn("监测段：2025-01-02",labels)
                self.assertIn("不补零",labels)
                self.assertIn("监测段：2025-01-02",report_text(model));self.assertIn("监测段：2025-01-02",report_html(model))
                card.deleteLater()
    def test_gui_invalid_second_row_no_record_and_hour_day_equivalence(self):
        with tempfile.TemporaryDirectory() as directory,self.fresh_page(Path(directory)) as (gui,projects,records,shell):
            self.prepare(gui);first=gui._process_rows["fume"][0];self.fill(first)
            second=gui._add_process_row("fume",instance_id="bad-segment");self.fill(second)
            for key in ("source","boundary"):
                combo=second[key+"_selector"];combo.setCurrentIndex(combo.findData(first[key+"_selector"].currentData()))
            gui._run_calculation();self.assertEqual(len(records.list_all()),0)
            second["segment_start"].setText("2025-01-02 00:00:00");second["segment_end"].setText("2025-01-03 00:00:00")
            second["duration_unit"].setCurrentIndex(second["duration_unit"].findData("h"));second["fields"]["duration"].setText("24")
            gui._run_calculation();self.assertEqual(len(records.list_all()),1)
            self.assertEqual(records.list_all()[0].calculation_result.total_amount,Decimal(".1408"))
    def test_excel_bad_identity_interval_and_blank_numeric_never_partial_canonical(self):
        with tempfile.TemporaryDirectory() as directory,self.fresh_page(Path(directory)) as (gui,projects,records,shell):
            root=Path(directory);self.prepare(gui)
            good=[1000,1000,40,.02,100,1,"A","炉口","1","2025-01-01 00:00:00","2025-01-02 00:00:00","d"]
            for position,bad in ((0,None),(6,None),(9,""),(10,"2024-01-01 00:00:00"),(11,"月")):
                row=list(good);row[8]="2";row[position]=bad
                preview=self.preview(gui,shell,self.workbook(root,[good,row]));unit=preview.units[0]
                self.assertFalse(unit.can_calculate);self.assertIsNone(unit.input_value)
                self.assertEqual(len(records.list_all()),0)
    def test_gui_sqlite_real_new_window_stable_ids_drafts_and_upgrade_cancel(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            with self.fresh_page(root) as (gui,projects,records,shell):
                self.prepare(gui);first=gui._process_rows["fume"][0];self.fill(first)
                first["source_selector"].setEditText("同名源");first["boundary_selector"].setEditText("排口计量")
                value=gui._input(increment=False);identity=value.fume_segments[0]
                gui.project_name.setText("烟气恢复");self.assertTrue(gui._save_project());project_id=gui._workspace.project_id
            with self.fresh_page(root) as (gui,projects,records,shell):
                with patch.object(QMessageBox,"question",return_value=QMessageBox.StandardButton.No):self.assertTrue(gui.open_project_from_home(project_id))
                actual=gui._input(increment=False).fume_segments[0];self.assertEqual(actual,identity)
                self.assertEqual(len(records.list_all()),0)
                state=gui._capture_form_state();state.pop("fume_execution_version");state.pop("fume_identities")
                gui._restore_form_state(state);before=gui._capture_form_state()
                with patch.object(QMessageBox,"question",return_value=QMessageBox.StandardButton.No):gui._upgrade_fume_rules()
                self.assertEqual(before,gui._capture_form_state())
                with patch.object(QMessageBox,"question",return_value=QMessageBox.StandardButton.Yes):gui._upgrade_fume_rules()
                original=dict(before);original.pop("fume_upgrade_original_state",None)
                self.assertEqual(gui._fume_upgrade_original_state,original)
                self.assertEqual(gui._fume_execution_version,"FUME_V2")
                self.assertEqual(len(records.list_all()),0)

    def test_identity_only_gui_segment_is_not_ignored_and_deleted_owner_reopens(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            with self.fresh_page(root) as (gui,projects,records,shell):
                self.prepare(gui);self.fill(gui._process_rows["fume"][0])
                owner=gui._add_process_row("fume",instance_id="owner")
                owner["source_selector"].setEditText("真实第二源")
                gui._run_calculation();self.assertEqual(len(records.list_all()),0)
                survivor=gui._add_process_row("fume",instance_id="survivor")
                self.fill(survivor)
                for key in ("source","boundary"):
                    combo=survivor[key+"_selector"];combo.setCurrentIndex(combo.findData(owner[key+"_selector"].currentData()))
                gui._remove_process_row("fume","owner")
                value=gui._input(increment=False);expected=value.fume_segments[1]
                gui.project_name.setText("删除源拥有行后恢复");self.assertTrue(gui._save_project());project_id=gui._workspace.project_id
            with self.fresh_page(root) as (gui,projects,records,shell):
                with patch.object(QMessageBox,"question",return_value=QMessageBox.StandardButton.No):self.assertTrue(gui.open_project_from_home(project_id))
                self.assertEqual(gui._input(increment=False).fume_segments[1],expected)
                gui._run_calculation();self.assertEqual(len(records.list_all()),1)
                fume_fixtures.FumeSegmentArithmeticTests().assertMappingValue(records.list_all()[0].calculation_result.total_amount,".1408")
    def test_v9_only_b6_changes_and_real_v8_retains_legacy_identity(self):
        from scripts.project_appendix_b_v9 import SOURCE,TARGET,generated_bytes
        from zipfile import ZipFile
        self.assertEqual(TARGET.read_bytes(),generated_bytes())
        with ZipFile(SOURCE) as old,ZipFile(TARGET) as new:
            self.assertEqual(old.namelist(),new.namelist())
            self.assertEqual([name for name in old.namelist() if old.read(name)!=new.read(name)],["xl/worksheets/sheet6.xml"])
        with tempfile.TemporaryDirectory() as directory,self.fresh_page(Path(directory)) as (gui,projects,records,shell):
            root=Path(directory);self.prepare(gui)
            book=load_workbook(SOURCE)
            for col,value in zip("ABCDEF",(1000,1000,40,.02,100,1)):book["B.6"][f"{col}3"]=value
            path=root/"旧v8.xlsx";book.save(path);book.close()
            preview=self.preview(gui,shell,path);unit=preview.units[0]
            self.assertTrue(unit.can_calculate,unit.errors);self.assertEqual(preview.provenance.template_version,"2.6.0")
            self.assertEqual(unit.input_value.fume_execution_version,"FUME_V1");self.assertEqual(unit.input_value.fume_segments,())
            self.assertEqual(len(unit.input_value.fume_incinerations),1)
            self.assertTrue(any(x.code=="CAR-WARN-FUME-LEGACY-COVERAGE" for x in unit.warnings))
            path=self.workbook(root,[(1000,1000,40,.02,100,1,"A","炉口","1","2025-01-01 00:00:00","2025-01-02 00:00:00","d")])
            book=load_workbook(path);book["B.6"]["H9"]="误放表后真实源";book.save(path);book.close()
            preview=self.preview(gui,shell,path);self.assertFalse(preview.units[0].can_calculate);self.assertIsNone(preview.units[0].input_value)
