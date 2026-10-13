"""High-value integration checks for Appendix B ingress semantics."""
from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP, localcontext
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from packages.application import CarbonAccountingPreviewUseCase
from packages.application.carbon_accounting import create_g06_parameter_resolver
from packages.application.catalog_queries import CatalogQueryService
from packages.core.models import ElectricityAttribute
from packages.excel.appendix_b import AppendixBImportContext, AppendixBWorkbookImporter
from packages.persistence.catalog_builder import build_catalog_database
from packages.persistence.catalog_repository import SQLiteCatalogRepository
from packages.reference_data import DEFAULT_SOURCE_PATH
from packages.standards.carbon_material import (
    CarbonMaterialCalculator,
    FuelPath,
    FuelType,
    HeatFactorMode,
    ParameterSourceKind,
    SOURCE_EXPORTED_ELECTRICITY,
    SOURCE_FGD,
    SOURCE_FUEL,
    SOURCE_PURCHASED_ELECTRICITY,
    SOURCE_PURCHASED_HEAT,
    SteamKind,
)
from packages.standards.carbon_material_normalization import MaterialRole
from tests.exb01_fixture_helper import appendix_b_context, write_valid_appendix_b_workbook


class AppendixBIngressIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls._catalog_directory = tempfile.TemporaryDirectory()
        catalog_path = build_catalog_database(
            DEFAULT_SOURCE_PATH,
            Path(cls._catalog_directory.name) / "catalog.sqlite",
        )
        repository = SQLiteCatalogRepository(catalog_path)
        cls.catalog_service = CatalogQueryService(repository)
        cls.resolver = create_g06_parameter_resolver(repository)
        cls.preview_use_case = CarbonAccountingPreviewUseCase(
            CarbonMaterialCalculator(parameter_resolver=cls.resolver)
        )

    @classmethod
    def tearDownClass(cls) -> None:
        cls._catalog_directory.cleanup()

    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.importer = AppendixBWorkbookImporter(
            preview_use_case=self.preview_use_case,
            catalog_service=self.catalog_service,
        )

    def tearDown(self) -> None:
        self.directory.cleanup()

    @staticmethod
    def _context(*, fuel_path: FuelPath | None = None) -> AppendixBImportContext:
        overrides = {"B.2!B4": fuel_path} if fuel_path is not None else {}
        return AppendixBImportContext(
            period=appendix_b_context().period,
            boundary_confirmed=True,
            enterprise_name="附录B专项验收企业",
            fuel_path_overrides=overrides,
        )

    @staticmethod
    def _set_cells(path: Path, changes: dict[str, dict[str, object]]) -> None:
        workbook = load_workbook(path, data_only=False)
        try:
            for sheet_name, assignments in changes.items():
                for cell, value in assignments.items():
                    workbook[sheet_name][cell] = value
            workbook.save(path)
        finally:
            workbook.close()

    def _workbook(self, name: str, *, all_sources: bool = False) -> Path:
        return write_valid_appendix_b_workbook(self.root / name, all_sources=all_sources)

    def _custom_fuel(self, name: str, *, route: str = "实测值") -> Path:
        path = self._workbook(name)
        self._set_cells(path, {
            "B.2": {
                "A4": "自定义能源品种",
                "B4": 10,
                "C4": 0.4 if route == "实测值" else 999,
                "D4": route,
                "E4": 8,
                "F4": "实测值",
                "G4": None,
                "K4": 0.03,
                "L4": "自定义",
                "H4": 98,
                "I4": "实测值",
            },
        })
        return path

    def test_b6_additional_rows_are_consumed_once_with_independent_formula_total(self) -> None:
        from packages.standards.carbon_material import SOURCE_FUME
        path = self._workbook("two-fume-facilities.xlsx", all_sources=True)
        workbook = load_workbook(path)
        sheet = workbook["B.6"]
        footer_merges = [area.bounds for area in sheet.merged_cells.ranges if area.min_row >= 4]
        for left, top, right, bottom in footer_merges:
            sheet.unmerge_cells(start_row=top, start_column=left, end_row=bottom, end_column=right)
        sheet.insert_rows(4)
        for left, top, right, bottom in footer_merges:
            sheet.merge_cells(start_row=top + 1, start_column=left, end_row=bottom + 1, end_column=right)
        for column, value in zip("ABCDEF", (200, 10, 30, 0.02, 98, 365)):
            sheet[f"{column}4"] = value
        workbook.save(path)
        workbook.close()
        preview = self.importer.import_preview(path, context=self._context(fuel_path=FuelPath.VOLUME))
        unit = preview.units[0]
        self.assertTrue(unit.can_calculate, unit.errors)
        self.assertEqual([item.q.value for item in unit.input_value.fume_incinerations], [Decimal("100"), Decimal("200")])
        amounts = [line.amount for line in unit.result.lines if line.emission_source_id == SOURCE_FUME]
        self.assertEqual(sum(amounts), Decimal("0.05665968"))
        self.assertIsNone(unit.calculation.record)
        self.assertIn("B.6!A4", {item.q.source_reference for item in unit.input_value.fume_incinerations})

    def test_b6_invalid_second_row_blocks_at_its_actual_cell(self) -> None:
        path = self._workbook("invalid-second-fume.xlsx", all_sources=True)
        workbook = load_workbook(path)
        sheet = workbook["B.6"]
        footer_merges = [area.bounds for area in sheet.merged_cells.ranges if area.min_row >= 4]
        for left, top, right, bottom in footer_merges:
            sheet.unmerge_cells(start_row=top, start_column=left, end_row=bottom, end_column=right)
        sheet.insert_rows(4)
        for left, top, right, bottom in footer_merges:
            sheet.merge_cells(start_row=top + 1, start_column=left, end_row=bottom + 1, end_column=right)
        workbook["B.6"]["A4"] = "不是数值"
        workbook.save(path)
        workbook.close()
        preview = self.importer.import_preview(path, context=self._context(fuel_path=FuelPath.VOLUME))
        unit = preview.units[0]
        self.assertFalse(unit.can_calculate)
        self.assertIn("B.6!A4", {issue.location for issue in unit.errors})
        self.assertIsNone(unit.calculation.record if unit.calculation else None)

    @staticmethod
    def _result_lines(unit) -> tuple[tuple[str, Decimal], ...]:
        return tuple(
            (line.emission_source_id, Decimal(str(line.amount)))
            for line in unit.calculation.result.lines
        )

    def test_custom_carbonate_does_not_inherit_default_fraction_or_conversion(self) -> None:
        path = self._workbook("custom-carbonate-missing-parameters.xlsx")
        self._set_cells(path, {"B.7":{"A3":"自定义脱硫剂","B3":2,"C3":"自定义碳酸盐","H3":0.5,"I3":"自定义"}})
        unit = self.importer.import_preview(path,context=self._context()).units[0]
        self.assertFalse(unit.can_calculate)
        self.assertIn(("EXB01_CUSTOM_CARBONATE_CONTENT_REQUIRED","B.7!D3"),{(e.code,e.location) for e in unit.errors})
        self.assertIn(("EXB01_CUSTOM_CARBONATE_CONVERSION_REQUIRED","B.7!F3"),{(e.code,e.location) for e in unit.errors})

    def test_custom_carbonate_filled_parameters_keep_user_defined_source(self) -> None:
        from packages.standards.carbon_material_input_policy import InputOrigin
        path = self._workbook("custom-carbonate-filled-parameters.xlsx")
        self._set_cells(path, {"B.7":{"A3":"自定义脱硫剂","B3":2,"C3":"自定义碳酸盐",
                                     "D3":80,"H3":0.5,"F3":90,"I3":"自定义"}})
        unit = self.importer.import_preview(path,context=self._context()).units[0]
        self.assertTrue(unit.can_calculate,unit.errors)
        component=unit.input_value.fgd_units[0].components[0]
        self.assertEqual(component.origin,InputOrigin.CUSTOM)
        for parameter in (component.carbonate_fraction,component.emission_factor,component.conversion_rate):
            self.assertEqual(parameter.source_kind,ParameterSourceKind.USER_DEFINED)
            self.assertIsNone(parameter.factor_id)
        self.assertEqual(unit.result.total_amount,Decimal("0.720"))
        self.assertIsNone(unit.calculation.record)

    def test_catalog_fch_tampering_is_blocked_at_the_fch_cell(self) -> None:
        path = self._workbook("tampered-catalog-fch.xlsx")
        self._set_cells(path, {"B.2": {"A4":"天然气","B4":1,"D4":"计算值","F4":"缺省值",
                                     "G4":0.03,"I4":"缺省值"}})
        unit = self.importer.import_preview(path,context=self._context(fuel_path=FuelPath.VOLUME)).units[0]
        self.assertFalse(unit.can_calculate)
        self.assertIn(("EXB01_CATALOG_FCH_READONLY","B.2!G4"),{(e.code,e.location) for e in unit.errors})

    def test_registered_coke_keeps_its_catalog_type_instead_of_becoming_custom(self) -> None:
        path = self._workbook("registered-coke.xlsx")
        self._set_cells(path, {"B.2": {"A4":"焦炭","B4":1,"C4":0.4,"D4":"实测值",
                                     "H4":98,"I4":"填写值"}})
        unit = self.importer.import_preview(path,context=self._context(fuel_path=FuelPath.MASS)).units[0]
        self.assertTrue(unit.can_calculate,unit.errors)
        self.assertEqual(unit.input_value.fuel_inputs[0].fuel_type,FuelType.COKE)

    def test_custom_fuel_direct_carbon_uses_explicit_mass_and_ignores_lhv_columns(self) -> None:
        path = self._custom_fuel("custom-mass-direct.xlsx")
        preview = self.importer.import_preview(path, context=self._context(fuel_path=FuelPath.MASS))
        unit = preview.units[0]

        self.assertTrue(unit.can_calculate, tuple((item.code, item.location) for item in unit.errors))
        fuel = unit.input_value.fuel_inputs[0]
        self.assertIs(fuel.fuel_type, FuelType.OTHER)
        self.assertEqual(fuel.fuel_label, "自定义能源品种")
        self.assertIs(fuel.path, FuelPath.MASS)
        self.assertEqual(fuel.activity.unit, "t")
        self.assertEqual(fuel.carbon_content.value, Decimal("0.4"))
        self.assertEqual(fuel.carbon_content.unit, "tC/t")
        self.assertIs(fuel.carbon_content.source_kind, ParameterSourceKind.MEASURED)
        self.assertIsNone(fuel.lower_heating_value)
        with localcontext() as context:
            context.prec = 40
            context.rounding = ROUND_HALF_UP
            expected = Decimal("10") * Decimal("0.4") * Decimal("0.98") * Decimal("44") / Decimal("12")
        self.assertEqual(unit.source_breakdown[SOURCE_FUEL], expected)

    def test_custom_fuel_computed_carbon_uses_explicit_volume_energy_path(self) -> None:
        path = self._custom_fuel("custom-volume-calculated.xlsx", route="计算值")
        self._set_cells(path, {"B.2": {"B4": 2}})
        preview = self.importer.import_preview(path, context=self._context(fuel_path=FuelPath.VOLUME))
        unit = preview.units[0]

        self.assertTrue(unit.can_calculate, tuple((item.code, item.location) for item in unit.errors))
        fuel = unit.input_value.fuel_inputs[0]
        self.assertIs(fuel.path, FuelPath.VOLUME)
        self.assertEqual(fuel.activity.unit, "ten_thousand_Nm3")
        self.assertEqual(fuel.lower_heating_value.value, Decimal("8"))
        self.assertEqual(fuel.lower_heating_value.unit, "GJ/10⁴Nm³")
        self.assertEqual(fuel.carbon_content.value, Decimal("0.03"))
        self.assertEqual(fuel.carbon_content.unit, "tC/GJ")
        with localcontext() as context:
            context.prec = 40
            context.rounding = ROUND_HALF_UP
            expected = Decimal("2") * Decimal("8") * Decimal("0.03") * Decimal("0.98") * Decimal("44") / Decimal("12")
        self.assertEqual(unit.source_breakdown[SOURCE_FUEL], expected)

    def test_custom_fuel_without_catalog_unit_requires_an_explicit_path(self) -> None:
        path = self._custom_fuel("custom-unit-required.xlsx")
        preview = self.importer.import_preview(path, context=self._context())
        unit = preview.units[0]

        self.assertFalse(unit.can_calculate)
        self.assertIsNone(unit.calculation)
        self.assertIn(
            ("EXB01_FUEL_PATH_REQUIRED", "B.2!B4"),
            {(item.code, item.location) for item in unit.errors},
        )
        self.assertEqual(unit.input_value.fuel_inputs, ())

    def test_process_role_slots_keep_custom_names_and_pair_each_mass_once(self) -> None:
        path = self._workbook("custom-process-materials.xlsx", all_sources=True)
        self._set_cells(path, {
            "B.3": {
                "B4": "自配煅烧料甲", "B5": "自配煅烧料甲",
                "B6": "煅后产品甲", "B7": "欠烧回收料甲",
                "B8": "煅烧粉尘甲", "B9": "煅后产品甲",
            },
            "B.4": {
                "B4": "焙烧填充料甲", "B5": "待焙烧生坯甲",
                "B6": "焙烧填充料甲", "B7": "待焙烧生坯甲",
                "B8": "焙烧副产料甲", "B9": "焙烧成品甲",
            },
            "B.5": {
                "B4": "石墨化保温料甲", "B5": "待石墨化品甲",
                "B6": "石墨化保温料甲", "B7": "石墨化副产物甲",
                "B8": "石墨化成品甲",
            },
        })
        preview = self.importer.import_preview(path, context=self._context(fuel_path=FuelPath.VOLUME))
        unit = preview.units[0]

        self.assertTrue(unit.can_calculate, tuple((item.code, item.location) for item in unit.errors))
        cases = (
            (unit.input_value.calcinations[0].material_rows, MaterialRole.CALCINATION_FEED, "自配煅烧料甲", Decimal("100"), Decimal("90"), Decimal("5")),
            (unit.input_value.bakings[0].material_rows, MaterialRole.BAKING_FILLER, "焙烧填充料甲", Decimal("40"), Decimal("90"), Decimal("5")),
            (unit.input_value.graphitizations[0].material_rows, MaterialRole.GRAPHITIZATION_PACKING, "石墨化保温料甲", Decimal("40"), Decimal("80"), Decimal("4")),
        )
        for rows, role, name, mass, fixed, volatile in cases:
            with self.subTest(name=name):
                matches = [item for item in rows if item.role is role and item.name == name]
                self.assertEqual(len(matches), 1)
                self.assertEqual(matches[0].mass_t, mass)
                self.assertEqual(matches[0].fixed_carbon_percent, fixed)
                self.assertEqual(matches[0].volatile_matter_percent, volatile)

    def test_paired_mass_mismatch_is_a_located_block_for_each_process_sheet(self) -> None:
        cases = (
            ("B.3", {"B4": "新煅烧料", "B5": "新煅烧料"}, "C5"),
            ("B.4", {"B4": "新填充料", "B6": "新填充料"}, "C6"),
            ("B.5", {"B4": "新保温料", "B6": "新保温料"}, "C6"),
        )
        for sheet, names, mismatch_cell in cases:
            with self.subTest(sheet=sheet):
                path = self._workbook(f"mass-mismatch-{sheet.replace('.', '-')}.xlsx", all_sources=True)
                updates = dict(names)
                updates[mismatch_cell] = 999
                self._set_cells(path, {sheet: updates})
                preview = self.importer.import_preview(path, context=self._context(fuel_path=FuelPath.VOLUME))
                unit = preview.units[0]
                self.assertFalse(unit.can_calculate)
                self.assertIsNone(unit.calculation)
                self.assertIn(
                    ("EXB01_PROCESS_PAIRED_MASS_MISMATCH", f"{sheet}!{mismatch_cell}"),
                    {(item.code, item.location) for item in unit.errors},
                )

    def test_fgd_batch_applies_one_amount_to_each_of_two_components(self) -> None:
        path = self._workbook("two-carbonates.xlsx", all_sources=True)
        preview = self.importer.import_preview(path, context=self._context(fuel_path=FuelPath.VOLUME))
        unit = preview.units[0]

        self.assertTrue(unit.can_calculate, tuple((item.code, item.location) for item in unit.errors))
        self.assertEqual(len(unit.input_value.fgd_units), 1)
        fgd_unit = unit.input_value.fgd_units[0]
        self.assertEqual(fgd_unit.reagent_name, "脱硫批次A")
        components = fgd_unit.components
        self.assertEqual(len(components), 2)
        self.assertEqual([item.amount.value for item in components], [Decimal("2"), Decimal("2")])
        with localcontext() as context:
            context.prec = 40
            context.rounding = ROUND_HALF_UP
            expected = (
                Decimal("2") * Decimal("0.90") * Decimal("0.440") * Decimal("0.95")
                + Decimal("2") * Decimal("0.10") * Decimal("0.522") * Decimal("0.90")
            )
        self.assertEqual(unit.source_breakdown[SOURCE_FGD], expected)
        self.assertEqual(
            sum((line.amount for line in unit.result.lines if line.emission_source_id == SOURCE_FGD), Decimal("0")),
            expected,
        )

    def test_fgd_batch_name_is_not_footer_and_trailing_activity_is_located_error(self) -> None:
        path = self._workbook("fgd-a-prefix-batch.xlsx")
        self._set_cells(path, {"B.7": {
            "A3":"A Sample Batch","B3":2,"D3":None,"E3":0.440,"F3":None,
            "A10":"Activity after footer","B10":99,"C10":"CaCO₃",
            "D10":100,"E10":0.440,"F10":100,
        }})
        unit=self.importer.import_preview(path,context=self._context(fuel_path=FuelPath.VOLUME)).units[0]
        self.assertFalse(unit.can_calculate)
        self.assertIn(("EXB01_TRAILING_ACTIVITY","B.7!B10"),{(e.code,e.location) for e in unit.errors})
        self.assertEqual(len(unit.input_value.fgd_units),1)
        self.assertEqual(unit.input_value.fgd_units[0].components[0].amount.value,Decimal("2"))
        self.assertIsNone(unit.calculation)

    def test_legacy_template_version_and_custom_parameters_remain_readable(self) -> None:
        path=self.root/"legacy-v1-filled.xlsx"
        path.write_bytes((Path(__file__).parent/"fixtures/exb01/appendix_b_v1.xlsx").read_bytes())
        self._set_cells(path,{"B.2":{"A4":"旧模板自定义燃料","B4":10,"D4":"计算值",
                                    "E4":8,"F4":"填写值","G4":0.03,"H4":98,"I4":"填写值"}})
        preview=self.importer.import_preview(path,context=self._context(fuel_path=FuelPath.MASS))
        self.assertEqual(preview.provenance.template_version,"1.0.0")
        self.assertTrue(preview.units[0].can_calculate,preview.units[0].errors)
        self.assertEqual(preview.units[0].input_value.fuel_inputs[0].carbon_content.source_kind,ParameterSourceKind.USER_DEFINED)
        self.assertIsNone(preview.units[0].calculation.record)

    def test_fuel_name_starting_with_usage_prompt_is_not_a_footer(self) -> None:
        path = self._custom_fuel("fuel-usage-prompt-prefix.xlsx")
        self._set_cells(path, {"B.2": {"A4": "使用提示型煤"}})
        unit = self.importer.import_preview(
            path, context=self._context(fuel_path=FuelPath.MASS)
        ).units[0]

        self.assertTrue(unit.can_calculate, tuple((item.code, item.location) for item in unit.errors))
        self.assertEqual(len(unit.input_value.fuel_inputs), 1)
        self.assertEqual(unit.input_value.fuel_inputs[0].fuel_label, "使用提示型煤")

    def test_nonfossil_purchase_is_zero_and_export_row_uses_its_applicable_factor(self) -> None:
        path = self._workbook("nonfossil-power.xlsx")
        self._set_cells(path, {
            "B.8": {
                "B3": "非化石电力", "C3": 5,
                # PR37 treats the purchased-electricity attribute as not
                # applicable to output; retain the row and use its stated EF.
                "B4": "非化石电力", "C4": 2, "D4": 0.5,
            },
        })
        preview = self.importer.import_preview(path, context=self._context())
        unit = preview.units[0]

        self.assertEqual(len(unit.input_value.electricity_details), 1)
        self.assertIs(unit.input_value.electricity_details[0].attribute, ElectricityAttribute.NONFOSSIL)
        self.assertEqual(unit.input_value.electricity_details[0].electricity_amount, Decimal("5"))
        self.assertEqual(len(unit.input_value.exported_electricity), 1)
        exported = unit.input_value.exported_electricity[0]
        self.assertEqual(exported.amount.value, Decimal("2"))
        self.assertEqual(exported.factor.value, Decimal("0.5"))
        self.assertEqual(exported.factor.unit, "tCO2/MWh")
        self.assertTrue(unit.can_calculate, tuple((item.code, item.location) for item in unit.errors))
        electricity_lines = [
            line for line in unit.result.lines
            if line.emission_source_id in {SOURCE_PURCHASED_ELECTRICITY, SOURCE_EXPORTED_ELECTRICITY}
        ]
        by_source = {line.emission_source_id: line.amount for line in electricity_lines}
        self.assertEqual(set(by_source), {SOURCE_PURCHASED_ELECTRICITY, SOURCE_EXPORTED_ELECTRICITY})
        self.assertEqual(by_source[SOURCE_PURCHASED_ELECTRICITY], Decimal("0"))
        self.assertEqual(by_source[SOURCE_EXPORTED_ELECTRICITY], Decimal("1.0"))

    def test_output_without_attribute_uses_resolved_default_factor_and_keeps_the_row(self) -> None:
        path = self._workbook("output-default-factor.xlsx")
        self._set_cells(path, {"B.8": {"B4": None, "C4": 2, "D4": None}})
        preview = self.importer.import_preview(path, context=self._context())
        unit = preview.units[0]

        self.assertTrue(unit.can_calculate, tuple((item.code, item.location) for item in unit.errors))
        self.assertEqual(len(unit.input_value.exported_electricity), 1)
        exported = unit.input_value.exported_electricity[0]
        self.assertEqual(exported.amount.value, Decimal("2"))
        self.assertEqual(exported.factor.parameter_id, "electricity_emission_factor_national")
        self.assertEqual(exported.factor.value, Decimal("0.5306"))
        snapshots = [
            item for item in unit.calculation.parameter_snapshots
            if item.parameter_id == "electricity_emission_factor_national"
        ]
        self.assertTrue(snapshots)
        self.assertEqual(snapshots[0].value_used, Decimal("0.5306"))
        self.assertEqual(unit.source_breakdown[SOURCE_EXPORTED_ELECTRICITY], Decimal("1.0612"))

    def test_process_activity_without_first_group_label_is_a_located_fatal(self) -> None:
        path = self._workbook("missing-first-process-group.xlsx", all_sources=True)
        self._set_cells(path, {"B.3": {"A4": None}})
        preview = self.importer.import_preview(path, context=self._context(fuel_path=FuelPath.VOLUME))
        unit = preview.units[0]

        self.assertFalse(unit.can_calculate)
        self.assertIsNone(unit.calculation)
        self.assertIn(
            ("EXB01_PROCESS_GROUP_REQUIRED", "B.3!A4"),
            {(item.code, item.location) for item in unit.errors},
        )

    def test_fgd_blank_resolver_defaults_are_snapshotted_and_excess_defaults_block(self) -> None:
        path = self._workbook("fgd-resolved-defaults.xlsx", all_sources=True)
        self._set_cells(path, {"B.7": {
            "D3": None, "E3": 0.440, "F3": None,
            "D4": None, "E4": None, "F4": None,
        }})
        unit = self.importer.import_preview(
            path, context=self._context(fuel_path=FuelPath.VOLUME)
        ).units[0]

        self.assertTrue(unit.can_calculate, tuple((item.code, item.location) for item in unit.errors))
        components = unit.input_value.fgd_units[0].components
        self.assertEqual(len(components), 1)
        self.assertEqual(components[0].carbonate_fraction.value, Decimal("0.90"))
        self.assertEqual(components[0].conversion_rate.value, Decimal("1.00"))
        snapshots = {item.parameter_id: item for item in unit.calculation.parameter_snapshots}
        self.assertEqual(snapshots["car-par-p04b-i"].value_used, Decimal("0.90"))
        self.assertEqual(snapshots["car-par-p04b-tr"].value_used, Decimal("1.00"))
        # 2 t × 90% × 0.440 tCO₂/t × 100% conversion.
        self.assertEqual(unit.source_breakdown[SOURCE_FGD], Decimal("0.79200"))

        over_path = self._workbook("fgd-excess-default-fractions.xlsx", all_sources=True)
        self._set_cells(over_path, {"B.7": {
            "D3": None, "E3": 0.440, "F3": None,
            "D4": None, "E4": 0.522, "F4": None,
        }})
        blocked = self.importer.import_preview(
            over_path, context=self._context(fuel_path=FuelPath.VOLUME)
        ).units[0]

        self.assertFalse(blocked.can_calculate)
        self.assertIsNone(blocked.calculation)
        self.assertIn(
            "EXB01_MULTI_CARBONATE_CONTENT_REQUIRED",
            {item.code for item in blocked.errors},
        )
    def test_saturated_steam_default_is_resolved_and_uses_the_reference_enthalpy(self) -> None:
        path = self._workbook("default-steam.xlsx")
        self._set_cells(path, {"B.9": {"B3": "饱和蒸汽", "C3": 1000, "D3": 1}})
        preview = self.importer.import_preview(path, context=self._context())
        unit = preview.units[0]

        self.assertTrue(unit.can_calculate, tuple((item.code, item.location) for item in unit.errors))
        heat = unit.input_value.purchased_heat[0]
        self.assertIs(heat.steam_kind, SteamKind.SATURATED)
        self.assertIs(heat.factor_mode, HeatFactorMode.STANDARD_DEFAULT)
        self.assertIsNone(heat.factor)
        snapshots = [item for item in unit.calculation.parameter_snapshots if item.parameter_id == "heat_emission_factor_default"]
        self.assertTrue(snapshots)
        self.assertEqual(snapshots[0].value_used, Decimal("0.11"))
        # 1,000 t × 1,000 kg/t × 2,777 kJ/kg × 0.11 tCO2/GJ ÷ 10^6.
        self.assertEqual(unit.source_breakdown[SOURCE_PURCHASED_HEAT], Decimal("305.47"))

    def test_hot_water_is_rejected_instead_of_treated_as_steam(self) -> None:
        path = self._workbook("unsupported-hot-water.xlsx")
        self._set_cells(path, {"B.9": {"B3": "热水", "C3": 10, "D3": 1}})
        preview = self.importer.import_preview(path, context=self._context())
        unit = preview.units[0]

        self.assertFalse(unit.can_calculate)
        self.assertIsNone(unit.calculation)
        self.assertIn(("EXB01_HEAT_KIND_UNSUPPORTED", "B.9!B3"), {(item.code, item.location) for item in unit.errors})

    def test_domain_fatal_validation_is_copied_into_import_errors(self) -> None:
        path = self._workbook("domain-fatal.xlsx", all_sources=True)
        self._set_cells(path, {"B.3": {"D4": 101}})
        preview = self.importer.import_preview(path, context=self._context(fuel_path=FuelPath.VOLUME))
        unit = preview.units[0]

        self.assertFalse(unit.can_calculate)
        self.assertIsNotNone(unit.calculation)
        self.assertIsNone(unit.calculation.result)
        self.assertIn("CAR-VAL-PERCENT-RANGE", {item.code for item in unit.errors})

    def test_ambient_low_precision_does_not_change_canonical_input_or_result(self) -> None:
        path = self._workbook("ambient-independent.xlsx", all_sources=True)
        context = self._context(fuel_path=FuelPath.VOLUME)
        with localcontext() as decimal_context:
            decimal_context.prec = 40
            decimal_context.rounding = ROUND_HALF_UP
            normal = self.importer.import_preview(path, context=context).units[0]
        with localcontext() as decimal_context:
            decimal_context.prec = 6
            decimal_context.rounding = "ROUND_DOWN"
            low_precision = self.importer.import_preview(path, context=context).units[0]

        self.assertTrue(normal.can_calculate, tuple((item.code, item.location) for item in normal.errors))
        self.assertTrue(low_precision.can_calculate, tuple((item.code, item.location) for item in low_precision.errors))
        self.assertEqual(low_precision.input_value, normal.input_value)
        self.assertEqual(self._result_lines(low_precision), self._result_lines(normal))

    def test_result_formula_is_ignored_but_input_formula_blocks(self) -> None:
        ordinary_path = self._custom_fuel("formula-free-result.xlsx")
        formula_result_path = self._custom_fuel("display-result-formula.xlsx")
        self._set_cells(formula_result_path, {"B.2": {"J4": "=999"}})
        ordinary = self.importer.import_preview(ordinary_path, context=self._context(fuel_path=FuelPath.MASS)).units[0]
        with_result_formula = self.importer.import_preview(formula_result_path, context=self._context(fuel_path=FuelPath.MASS)).units[0]

        self.assertTrue(ordinary.can_calculate, ordinary.errors)
        self.assertTrue(with_result_formula.can_calculate, with_result_formula.errors)
        self.assertEqual(
            tuple(amount for _, amount in self._result_lines(with_result_formula)),
            tuple(amount for _, amount in self._result_lines(ordinary)),
        )
        self.assertNotIn("EXB01_FORMULA_REJECTED", {item.code for item in with_result_formula.errors})

        input_formula_path = self._custom_fuel("input-formula.xlsx")
        self._set_cells(input_formula_path, {"B.2": {"B4": "=10"}})
        rejected = self.importer.import_preview(input_formula_path, context=self._context(fuel_path=FuelPath.MASS)).units[0]
        self.assertFalse(rejected.can_calculate)
        self.assertIsNone(rejected.calculation)
        self.assertIn(("EXB01_FORMULA_REJECTED", "B.2!B4"), {(item.code, item.location) for item in rejected.errors})


if __name__ == "__main__":
    unittest.main()
