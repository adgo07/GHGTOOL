"""B-EXTERNAL表格预览/项目/正式记录链，旧九表逐部件保护。"""
from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile
import unittest
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication
from packages.excel.templates import ExcelTemplateService
from packages.application.carbon_accounting import CarbonAccountingUseCase
from packages.application.project_workspaces import ProjectWorkspaceService
from packages.persistence import SQLiteProjectWorkspaceRepository
from packages.ui.view_models import AppRoute
from tests import test_b_material_draft_scope as drafts
from tests import test_b_heat_chain as heat

class ExternalExcelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    page=drafts.MaterialDraftScopeTests.page

    def test_projection_original_nine_parts_and_old_v10_identity(self):
        from scripts.project_appendix_b_v11 import SOURCE,TARGET,generated_bytes
        self.assertEqual(TARGET.read_bytes(),generated_bytes())
        with ZipFile(SOURCE) as old,ZipFile(TARGET) as new:
            for i in range(1,10):self.assertEqual(old.read(f'xl/worksheets/sheet{i}.xml'),new.read(f'xl/worksheets/sheet{i}.xml'))
        for path,version in ((SOURCE,'2.8.0'),(TARGET,'2.9.0')):
            wb=load_workbook(path);self.assertEqual(ExcelTemplateService.default().validate_structure(wb),version);wb.close()

    def test_state_matrix_preview_zero_records_and_project_reopen(self):
        with self.page() as (gui,projects,root,shell,records):
            heat.HeatQtTests.prepare(gui)
            page=shell.pages[AppRoute.EXCEL_IMPORT]
            context=page.excel_adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name='外部补充整链')
            for include,amount,status,valid in [('不纳入','NaN','NOT_INCLUDED',True),('纳入',None,'PENDING',True),('纳入',0,'COMPLETED_ZERO',True),('纳入',20,'COMPLETED_POSITIVE',True),('纳入',-1,None,False),('纳入','20',None,False),('纳入',1234567890123456,None,False),('不纳入','=1+1',None,False),('纳入','NaN',None,False)]:
                with self.subTest(include=include,amount=amount):
                    wb=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
                    for c,v in {'B3':'饱和蒸汽','C3':1,'F3':2800,'H3':'已知焓','I3':'缺省值','J3':'边界外'}.items():wb['B.9'][c]=v
                    ws=wb['软件补充'];ws['B4']=include;ws['B5']=amount;ws['B6']='其他核算';ws['B10']='有';ws['B11']='有'
                    path=root/'external.xlsx';wb.save(path);wb.close()
                    preview=page.excel_adapter.import_preview(path,context)
                    self.assertEqual(len(records.list_all()),0)
                    self.assertEqual(preview.units[0].can_calculate,valid,preview.units[0].errors)
                    if not valid:continue
                    value=preview.units[0].input_value
                    self.assertEqual(value.external_co2.amount_lexeme,None if amount is None else str(amount))
                    self.assertTrue(value.other_activity_present);self.assertTrue(value.transport_present)
                    workspace=projects.new_workspace('外部补充')
                    projects.save(replace(workspace,units=(replace(workspace.units[0],canonical_input=value),)))
                    recovered=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/'projects.sqlite')).get(workspace.project_id)
                    self.assertEqual(recovered.units[0].canonical_input,value)
                    result=gui.calculator.calculate(value)
                    self.assertTrue(result.successful,result.problems)
                    import json
                    node=json.loads(result.evidence.trace_snapshot_json)['software_supplement']
                    self.assertEqual(node['status'],status)
                    self.assertEqual(result.result.total_amount,Decimal('.308'))
                    if status=='COMPLETED_POSITIVE':self.assertEqual(Decimal(node['supplementary_total_tco2']),Decimal('20.308'))
            # Save exactly twice after preview: immutable records append and sums stay20.308.
            value=replace(value,external_co2=replace(value.external_co2,amount_lexeme='20'))
            for count in (1,2):
                result=CarbonAccountingUseCase(gui.calculator,records).calculate(value)
                self.assertTrue(result.successful,result.problems);self.assertEqual(len(records.list_all()),count)

if __name__=='__main__':unittest.main()
