from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
import sqlite3
import tempfile
import unittest

from docx import Document

from packages.application.reporting import build_saved_record_report
from packages.application.carbon_accounting import CarbonAccountingUseCase
from packages.application.reporting.model import build_report_model
from packages.core.models import (
    ActivityDataSource,
    AccountingInput,
    AccountingPeriod,
    AccountingRecord,
    CalculationLine,
    CalculationResult,
    PeriodType,
    ParameterSelectionMethod,
    ParameterSnapshot,
    RecordStatus,
)
from packages.core.errors import IssueLevel, ValidationProblem
from packages.infrastructure.reporting import render_report_docx
from packages.persistence.records_repository import SQLiteRecordRepository
from packages.persistence.in_memory_records import InMemoryRecordRepository
from packages.standards.carbon_material import (
    ActivityDataEvidence,
    CarbonMaterialCalculator,
    CarbonMaterialInput,
    CarbonReportingData,
    EmissionSourceState,
    EmissionSourceStatus,
    FuelInput,
    FuelPath,
    FuelType,
    InputValue,
    ParameterSourceKind,
    ParameterValue,
    CalcinationInput,
    BakingInput,
    GraphitizationInput,
    SOURCE_CALCINATION,
    SOURCE_BAKING,
    SOURCE_GRAPHITIZATION,
    SOURCE_PURCHASED_ELECTRICITY,
    SOURCE_EXPORTED_ELECTRICITY,
    SOURCE_PURCHASED_HEAT,
    SOURCE_EXPORTED_HEAT,
    SOURCE_FUEL,
    STANDARD_ID,
)
from packages.standards.carbon_material_normalization import MaterialDataSource, MaterialInputLine, MaterialRole


NOW = datetime(2026, 10, 6, 12, 0, tzinfo=timezone.utc)
PERIOD = AccountingPeriod(PeriodType.ANNUAL, date(2025, 1, 1), date(2025, 12, 31))


def _record(
    record_id: str = "record.rpt01",
    *,
    period=PERIOD,
    parameter_snapshots=(),
    problems=(),
) -> AccountingRecord:
    accounting_input = AccountingInput(
        input_id=f"input.{record_id}",
        standard_id=STANDARD_ID,
        period=period,
        enterprise_name="报告测试企业",
        boundary_component_ids=("boundary.main",),
    )
    result = CalculationResult(
        result_id=f"result.{record_id}",
        standard_id=STANDARD_ID,
        algorithm_version="CAR-RPT01-TEST",
        lines=(
            CalculationLine("CAR-FLD-DIRECT-RESULT", "CAR-RULE-TOTAL-001", "GEN-GAS-CO2", "3.25", "tCO2"),
            CalculationLine("CAR-FLD-INDIRECT-RESULT", "CAR-RULE-TOTAL-001", "GEN-GAS-CO2", "1.75", "tCO2"),
            CalculationLine("CAR-FLD-TOTAL-RESULT", "CAR-RULE-TOTAL-001", "GEN-GAS-CO2", "5.00", "tCO2"),
        ),
        total_amount="5.00",
        total_unit="tCO2",
        calculated_at=NOW,
    )
    return AccountingRecord(
        record_id=record_id,
        standard_id=STANDARD_ID,
        algorithm_version="CAR-RPT01-TEST",
        created_at=NOW,
        input_snapshot=accounting_input,
        calculation_result=result,
        status=RecordStatus.COMPLETED_WITH_WARNINGS if problems else RecordStatus.COMPLETED,
        parameter_snapshots=tuple(parameter_snapshots),
        problems=tuple(problems),
        standard_version="2024",
    )


def _material(line_id: str, role: MaterialRole, name: str, mass: str, fixed: str, volatile: str | None) -> MaterialInputLine:
    return MaterialInputLine(
        line_id, role, name, mass, fixed, MaterialDataSource.MEASURED,
        volatile, MaterialDataSource.MEASURED,
    )


class RPT01ReportAndExcelTests(unittest.TestCase):
    def test_word_report_uses_saved_snapshots_and_export_history_is_additive(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            repository = SQLiteRecordRepository(Path(directory) / "records.sqlite")
            record = _record()
            repository.create_with_details(
                record,
                raw_input={"enterprise_name": "快照企业", "fuel_inputs": [], "calcinations": [], "bakings": [], "graphitizations": [], "fume_incinerations": [], "fgd_units": [], "electricity_details": [], "exported_electricity": [], "purchased_heat": [], "exported_heat": []},
                trace_snapshot={"aggregations": {"ES": "3.25", "EI": "1.75", "ET": "5.00"}, "source_subtotals": {"fuel": "3.25"}, "formula_steps": []},
                provenance_snapshot={"mapping_version": "SM01-2026-09-13-R6"},
                reporting_snapshot={"boundary_description": "厂区边界"},
                report_qualification={"eligible": True},
            )
            original = repository.get(record.record_id)
            model = build_saved_record_report(
                repository,
                original,
                supplementary_info={"enterprise_name": "导出时企业名称", "prepared_on": "2026-10-06"},
            )
            docx = render_report_docx(model)
            parsed = Document(BytesIO(docx))
            visible_text = "\n".join(
                [paragraph.text for paragraph in parsed.paragraphs]
                + [cell.text for table in parsed.tables for row in table.rows for cell in row.cells]
            )
            self.assertIn("导出时企业名称", visible_text)
            self.assertIn("B.1 温室气体排放量汇总", visible_text)
            self.assertIn("B.9 购入和输出热力", visible_text)
            self.assertIn("5.00 tCO₂", visible_text)
            self.assertEqual(repository.get(record.record_id), original)

            repository.record_report_export(
                record.record_id,
                export_id="report-export.one",
                format="DOCX",
                template_version="1.0.0",
                document_filename="report-one.docx",
                document_sha256="1" * 64,
                supplementary_info={"enterprise_name": "导出时企业名称", "prepared_on": "2026-10-06"},
            )
            repository.record_report_export(
                record.record_id,
                export_id="report-export.two",
                format="DOCX",
                template_version="1.0.0",
                document_filename="report-two.docx",
                document_sha256="2" * 64,
                supplementary_info={"enterprise_name": "第二次导出企业", "prepared_on": "2026-10-07"},
            )
            self.assertEqual(repository.get_latest_report_export_supplementary(record.record_id)["enterprise_name"], "第二次导出企业")
            self.assertEqual(repository.get(record.record_id), original)
            self.assertEqual([item.action for item in repository.list_audit(record.record_id)], ["CREATE", "REPORT_EXPORT", "REPORT_EXPORT"])
            connection = sqlite3.connect(repository.path)
            try:
                exports = connection.execute(
                    "SELECT format, template_version, document_sha256 FROM report_export_history ORDER BY rowid"
                ).fetchall()
                self.assertEqual(exports, [("DOCX", "1.0.0", "1" * 64), ("DOCX", "1.0.0", "2" * 64)])
                with self.assertRaises(sqlite3.IntegrityError):
                    connection.execute("UPDATE report_export_history SET document_filename='changed.docx' WHERE export_id='report-export.one'")
                connection.rollback()
                with self.assertRaises(sqlite3.IntegrityError):
                    connection.execute("DELETE FROM report_export_history WHERE export_id='report-export.one'")
                connection.rollback()
            finally:
                connection.close()

    def test_real_record_report_reads_fuel_parameters_and_evidence_snapshots(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            repository = SQLiteRecordRepository(Path(directory) / "records.sqlite")
            measured = lambda pid, value, unit: ParameterValue(
                pid, value, unit, source_kind=ParameterSourceKind.MEASURED,
                source_id="company-measurement", source_version="2025",
                source_location="燃料化验单第2页", selection_reason="企业实测参数。",
            )
            fuel = FuelInput(
                "natural-gas", FuelPath.VOLUME,
                activity=InputValue("10", "ten_thousand_Nm3", source_type=ActivityDataSource.METER, source_reference="年度燃气计量台账"),
                lower_heating_value=measured("natural-gas_lhv", "35.5", "GJ/10⁴Nm³"),
                carbon_content=measured("natural-gas_carbon", "0.0153", "tC/GJ"),
                oxidation_rate=measured("natural-gas_oxidation", "0.99", "ratio"),
                fuel_type=FuelType.NATURAL_GAS,
            )
            input_value = CarbonMaterialInput(
                "input.rpt01-real", "enterprise.rpt01", "报告测试企业", PERIOD,
                boundary_confirmed=True, boundary_component_ids=("boundary.main",),
                source_states=(EmissionSourceState(SOURCE_FUEL, EmissionSourceStatus.INVOLVED),),
                fuel_inputs=(fuel,),
                reporting_data=CarbonReportingData(
                    boundary_description="厂区生产边界",
                    activity_evidence=(ActivityDataEvidence(
                        "evidence.fuel.activity", "天然气活动量", (SOURCE_FUEL,),
                        source_reference="年度燃气计量台账", monitoring_location="燃气总表",
                    ),),
                ),
            )
            outcome = CarbonAccountingUseCase(
                CarbonMaterialCalculator(), repository
            ).calculate(input_value, calculated_at=NOW)
            self.assertTrue(outcome.successful)
            record = repository.get(outcome.record.record_id)
            report = build_saved_record_report(repository, record, supplementary_info={"prepared_on": "2026-10-06"})
            b2 = next(section for section in report.sections if section.section_id == "b2").tables[0]
            self.assertEqual(b2.rows[0].cells[0].value, "天然气")
            # Presentation rounds saved emissions; the original Record remains exact.
            exact_amount = record.calculation_result.lines[0].amount
            self.assertEqual(b2.rows[0].cells[6].value, "19.72")
            self.assertEqual(repository.get(record.record_id).calculation_result.lines[0].amount, exact_amount)
            self.assertEqual(b2.rows[0].cells[4].value, "0.0153")
            self.assertEqual(b2.rows[0].cells[3].source.split("；", 1)[0], "企业实测值")
            evidence = next(section for section in report.sections if section.section_id == "evidence")
            self.assertIn("年度燃气计量台账", evidence.tables[0].rows[0].cells[2].value)

    def test_report_model_keeps_parameter_sources_multi_instance_rows_and_warnings(self) -> None:
        snapshots = tuple(
            ParameterSnapshot(
                snapshot_id=f"snapshot.{suffix}",
                parameter_id=f"parameter.{suffix}",
                factor_id=None,
                value_used=value,
                unit_used=unit,
                source_id=f"source.{suffix}",
                source_version="2024",
                selection_method=method,
                selection_reason="来自本条记录保存的参数快照。",
                standard_id=STANDARD_ID,
                snapshot_at=NOW,
                source_location=location,
                detail_id=detail,
            )
            for suffix, value, unit, method, location, detail in (
                ("lhv", "35.5", "GJ/t", ParameterSelectionMethod.ENTERPRISE_MEASURED, "燃料检测报告", "CAR-FLD-F01-coke-LHV"),
                ("carbon", "0.028", "tC/GJ", ParameterSelectionMethod.STANDARD_REQUIRED, "附录C.1", "CAR-FLD-F01-coke-CARBON"),
                ("oxidation", "0.98", "ratio", ParameterSelectionMethod.SYSTEM_RECOMMENDED, "附录C.1", "CAR-FLD-F01-coke-FOX"),
            )
        )
        warning = ValidationProblem("CAR-VAL-RPT01-WARN", IssueLevel.WARNING, "部分来源说明未填写；本次核算结果可用。")
        record = _record("record.rpt01.multiple", parameter_snapshots=snapshots, problems=(warning,))
        raw = {
            "fuel_inputs": [{
                "fuel_id": "coke", "fuel_type": "COKE", "path": "MASS",
                "activity": {"value": "4", "unit": "t", "source_type": "METER"},
                "lower_heating_value": {"value": "35.5", "unit": "GJ/t", "source_kind": "MEASURED"},
                "carbon_content": {"value": "0.028", "unit": "tC/GJ", "source_kind": "STANDARD_SPECIFIED"},
                "oxidation_rate": {"value": "0.98", "unit": "ratio", "source_kind": "SYSTEM_RECOMMENDED"},
            }],
            "calcinations": [{
                "instance_id": f"calc-{index}",
                "material_rows": [
                    {"role": "calcination_feed", "name": f"原料{index}甲", "mass_t": "100", "fixed_carbon_percent": "95", "fixed_carbon_source": "MEASURED", "volatile_matter_percent": "10", "volatile_matter_source": "CHEMICAL_CALCULATION"},
                    {"role": "calcination_feed", "name": f"原料{index}乙", "mass_t": "20", "fixed_carbon_percent": "80", "fixed_carbon_source": "MEASURED", "volatile_matter_percent": "5", "volatile_matter_source": "MEASURED"},
                    {"role": "calcined_product", "name": f"煅后料{index}", "mass_t": "90", "fixed_carbon_percent": "98", "fixed_carbon_source": "MEASURED", "volatile_matter_percent": "1", "volatile_matter_source": "MEASURED"},
                ],
            } for index in range(1, 4)],
            "bakings": [{"instance_id": "bake-1", "material_rows": [
                {"role": "green_baking_product", "name": "待焙烧品甲", "mass_t": "80", "fixed_carbon_percent": "80", "fixed_carbon_source": "MEASURED", "volatile_matter_percent": "2", "volatile_matter_source": "MEASURED"},
                {"role": "green_baking_product", "name": "待焙烧品乙", "mass_t": "20", "fixed_carbon_percent": "60", "fixed_carbon_source": "MEASURED", "volatile_matter_percent": "4", "volatile_matter_source": "MEASURED"},
            ]}],
            "graphitizations": [{"instance_id": "graph-1", "material_rows": [
                {"role": "green_graphitization_product", "name": "待石墨化品甲", "mass_t": "70", "fixed_carbon_percent": "80", "fixed_carbon_source": "MEASURED", "volatile_matter_percent": "1", "volatile_matter_source": "MEASURED"},
                {"role": "green_graphitization_product", "name": "待石墨化品乙", "mass_t": "30", "fixed_carbon_percent": "60", "fixed_carbon_source": "MEASURED", "volatile_matter_percent": "2", "volatile_matter_source": "MEASURED"},
            ]}],
            "electricity_details": [
                {"detail_id": "power-a", "attribute": "ORDINARY", "electricity_amount": {"value": "12", "unit": "MWh", "source_type": "METER"}, "factor": {"value": "0.5", "unit": "tCO2/MWh", "source_kind": "STANDARD_DEFAULT"}},
                {"detail_id": "power-b", "attribute": "NONFOSSIL", "electricity_amount": {"value": "3", "unit": "MWh", "source_type": "METER"}, "factor": {"value": "0.2", "unit": "tCO2/MWh", "source_kind": "MEASURED"}},
            ],
            "exported_electricity": [{"line_id": "power-out", "electricity_amount": "2", "electricity_unit": "MWh", "attribute": "ORDINARY", "factor": {"value": "0.6", "unit": "tCO2/MWh", "source_kind": "MEASURED"}}],
            "purchased_heat": [{"line_id": "heat-in", "steam_amount_t": "5", "amount": {"value": "5", "unit": "t"}, "steam_kind": "SATURATED", "enthalpy": {"value": "2777", "unit": "kJ/kg"}, "pressure_mpa": {"value": "1.0", "unit": "MPa"}}],
            "exported_heat": [{"line_id": "heat-out", "steam_amount_t": "1", "amount": {"value": "1", "unit": "t"}, "steam_kind": "SUPERHEATED", "enthalpy": {"value": "2810", "unit": "kJ/kg"}, "pressure_mpa": {"value": "1.0", "unit": "MPa"}, "temperature_c": {"value": "250", "unit": "C"}}],
        }
        trace = {
            "aggregations": {"ES": "3.25", "EI": "1.75", "ET": "5.00"},
            "source_subtotals": {"fuel": "1", "calcination": "0.7", "baking": "0.3", "graphitization": "0.2", "gas_control": "0.1", "purchased_electricity": "0.6", "exported_electricity": "0.1", "purchased_heat": "1", "exported_heat": "0.05"},
            "formula_steps": [
                {"emission_source_id": SOURCE_CALCINATION, "process_instance_id": f"calc-{index}", "intermediate_result": str(index / 10), "input_variables": [{"name": "gc", "value": "100", "unit": "t"}]}
                for index in range(1, 4)
            ] + [
                {"emission_source_id": SOURCE_BAKING, "process_instance_id": "bake-1", "intermediate_result": "0.3", "input_variables": [{"name": "bpm", "value": "20", "unit": "t"}]},
                {"emission_source_id": SOURCE_GRAPHITIZATION, "process_instance_id": "graph-1", "intermediate_result": "0.2", "input_variables": [{"name": "gta", "value": "100", "unit": "t"}]},
                {"emission_source_id": SOURCE_PURCHASED_ELECTRICITY, "process_instance_id": "power-a", "intermediate_result": "0.6"},
                {"emission_source_id": SOURCE_PURCHASED_ELECTRICITY, "process_instance_id": "power-b", "intermediate_result": "0.2"},
                {"emission_source_id": SOURCE_EXPORTED_ELECTRICITY, "process_instance_id": "power-out", "intermediate_result": "0.1"},
                {"emission_source_id": SOURCE_PURCHASED_HEAT, "process_instance_id": "heat-in", "intermediate_result": "1"},
                {"emission_source_id": SOURCE_EXPORTED_HEAT, "process_instance_id": "heat-out", "intermediate_result": "0.05"},
            ],
        }
        report = build_report_model(record, raw, trace, {"mapping_version": "SM01-2026-09-13-R6"}, {}, {"eligible": False, "message": "月度期间仅支持核算，不符合年度报告资格。"}, today=date(2026, 10, 6))
        sections = {section.section_id: section for section in report.sections}
        fuel_rows = sections["b2"].tables[0].rows
        self.assertEqual(len(fuel_rows), 1)
        self.assertEqual([fuel_rows[0].cells[index].source.split("；", 1)[0] for index in (3, 4, 5)], ["企业实测值", "标准规定值", "软件按标准推荐"])
        self.assertEqual(len(sections["b3"].tables), 6)
        self.assertEqual(len(sections["b3"].tables[0].rows), 3)
        self.assertEqual(len(sections["b4"].tables[0].rows), 2)
        self.assertEqual(len(sections["b5"].tables[0].rows), 2)
        self.assertEqual(len(sections["b8"].tables[0].rows), 3)
        self.assertEqual(len(sections["b9"].tables[0].rows), 2)
        self.assertTrue(any("来源说明未填写" in notice for notice in report.notices))
        self.assertTrue(any("月度期间" in notice for notice in report.notices))
        report_text = " ".join(cell.value for section in report.sections for table in section.tables for row in table.rows for cell in row.cells)
        self.assertNotIn("CAR-", report_text)
        self.assertNotIn("record.rpt01.multiple", report_text)

    def test_monthly_custom_and_old_record_qualification_are_explained(self) -> None:
        monthly = AccountingPeriod(PeriodType.MONTHLY, date(2025, 1, 1), date(2025, 1, 31))
        custom = AccountingPeriod(PeriodType.CUSTOM, date(2025, 1, 1), date(2025, 6, 30))
        for period in (monthly, custom):
            report = build_report_model(_record(f"record.period.{period.period_type.value.lower()}" , period=period), {}, {}, {}, {}, {"eligible": False, "message": "该期间不满足年度报告资格。"}, snapshot_schema_version=0, today=date(2026, 10, 6))
            self.assertIn("至", report.period_text)
            self.assertTrue(any("年度报告资格" in item for item in report.notices))
            self.assertTrue(any("历史记录" in item for item in report.notices))



if __name__ == "__main__":
    unittest.main()
