"""Integration checks for the RS03 Excel preview and frozen-record export entrypoints."""

from __future__ import annotations

from contextlib import closing
from dataclasses import asdict
from datetime import date, datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication, QMessageBox, QPushButton
from docx import Document
from openpyxl import load_workbook

from apps.carbon_accounting_desktop.app import create_main_window
from apps.carbon_accounting_desktop.config import AppConfig
from packages.application.carbon_accounting import (
    CarbonAccountingPreviewUseCase,
    create_g06_parameter_resolver,
)
from packages.application.reporting import build_saved_record_report
from packages.application.project_workspaces import ProjectWorkspaceService
from packages.excel.r2 import METADATA_SHEET, VISIBLE_SHEETS, ExcelWorkbookImporter, create_template_bytes
from packages.persistence.projects_repository import ProjectRecordAssociationError, SQLiteProjectWorkspaceRepository
from packages.persistence import (
    SQLiteCatalogRepository,
    SQLiteRecordRepository,
    build_catalog_database,
)
from packages.reference_data import DEFAULT_SOURCE_PATH
from packages.ui.view_models import AppRoute


FIXED_CALCULATED_AT = datetime(2024, 12, 31, 12, 34, 56, tzinfo=timezone.utc)
SNAPSHOT_GETTERS = {
    "raw_input": "get_raw_input_snapshot",
    "effective_rule_set": "get_effective_rule_set",
    "trace": "get_trace_snapshot",
    "provenance": "get_provenance_snapshot",
    "reporting": "get_reporting_snapshot",
    "qualification": "get_report_qualification",
}


class _FixedDateTime(datetime):
    @classmethod
    def now(cls, tz=None):
        return FIXED_CALCULATED_AT if tz is not None else FIXED_CALCULATED_AT.replace(tzinfo=None)


class Rs03ExcelEntrypointTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.catalog = build_catalog_database(DEFAULT_SOURCE_PATH, self.root / "catalog.sqlite")
        self.records_path = self.root / "records.sqlite"
        self.repository = SQLiteRecordRepository(self.records_path)
        self.window = create_main_window(
            AppConfig(
                catalog_database=self.catalog,
                records_database=self.records_path,
                projects_database=self.root / "projects.sqlite",
            ),
            record_repository=self.repository,
        )
        self.window.show()
        self.shell = self.window.centralWidget()
        self.shell.navigate(AppRoute.NEW_ACCOUNTING)
        self.app.processEvents()
        self.page = self.shell.pages[AppRoute.NEW_ACCOUNTING]
        self.page.period_year.setValue(2025)
        self.page.enterprise_name.setText("RS03 集成验收企业")
        self.page.boundary_confirmed.setChecked(True)
        status = self.page._source_statuses["CAR-SRC-FUEL-001"]
        status.setCurrentIndex(status.findData("INVOLVED"))
        row = self.page._fuel_rows[0]
        row.fuel_type.setCurrentIndex(row.fuel_type.findData("NATURAL_GAS"))
        row.path.setCurrentIndex(row.path.findData("VOLUME"))
        row.activity.setText("1.25")

    def tearDown(self) -> None:
        self.window.close()
        self.window.deleteLater()
        self.app.processEvents()
        self.directory.cleanup()

    def _workbook(self, year: int = 2025) -> Path:
        workbook = load_workbook(BytesIO(create_template_bytes()))
        basic = workbook["基本信息"]
        basic["B3"] = "RS03 集成验收企业"
        basic["B4"] = "年度"
        basic["B5"] = f"{year}-01-01"
        basic["B6"] = f"{year}-12-31"
        basic["B7"] = "是"
        basic["A13"], basic["B13"], basic["C13"] = "全厂", "全厂", "是"
        fuel = workbook["B.2 化石燃料"]
        fuel["A6"], fuel["B6"], fuel["C6"], fuel["D6"] = "全厂", "天然气", "体积", 1.25
        destination = self.root / f"输入-{year}.xlsx"
        workbook.save(destination)
        workbook.close()
        return destination

    def _record_page(self, record_id: str):
        self.shell.navigate(AppRoute.RECORDS)
        self.app.processEvents()
        page = self.shell.pages[AppRoute.RECORDS]
        self.assertTrue(page.open_record(record_id))
        return page

    def _create_record(self):
        self.page.quick_calculate_button.click()
        self.app.processEvents()
        records = self.repository.list_all()
        self.assertEqual(len(records), 1)
        return records[0]

    def _frozen_state(self, record_id: str) -> dict[str, object]:
        record = self.repository.get(record_id)
        return {
            "record": asdict(record),
            **{
                key: getattr(self.repository, getter)(record_id)
                for key, getter in SNAPSHOT_GETTERS.items()
            },
        }

    def test_gui_preview_matches_formal_calculator_metadata_and_never_persists(self) -> None:
        workbook = self._workbook(year=2024)
        self.assertEqual(self.repository.list_all(), ())
        self.shell.navigate(AppRoute.EXCEL_IMPORT)
        self.app.processEvents()
        page = self.shell.pages[AppRoute.EXCEL_IMPORT]
        self.assertIs(page.calculation_use_case, self.page.calculation_use_case)
        self.assertIs(page.calculation_use_case.calculator, self.page.calculation_use_case.calculator)

        with (
            patch("packages.excel.r2.datetime", _FixedDateTime),
            patch("packages.ui.pages.QFileDialog.getOpenFileName", return_value=(str(workbook), "")),
        ):
            page.selectFileButton.click()
            self.app.processEvents()

        preview = page._last_preview
        self.assertIsNotNone(preview)
        self.assertEqual(len(preview.units), 1)
        unit = preview.units[0]
        self.assertTrue(unit.can_calculate, unit.errors)
        self.assertIsNotNone(unit.calculation)
        actual = unit.calculation
        self.assertIsNone(actual.record)
        self.assertEqual(preview.provenance.imported_at, FIXED_CALCULATED_AT)
        self.assertEqual(actual.result.calculated_at, FIXED_CALCULATED_AT)

        calculator = self.page.calculation_use_case.calculator
        expected = calculator.calculate(unit.input_value, calculated_at=FIXED_CALCULATED_AT)
        self.assertEqual(asdict(actual), asdict(expected))
        self.assertEqual(actual.evidence.standard_version, calculator.standard_version)
        self.assertTrue(actual.parameter_snapshots)
        warning_codes = {
            problem.code for problem in actual.problems
            if getattr(problem.level, "value", problem.level) == "WARNING"
        }
        self.assertIn("CAR-VAL-STANDARD-IMPLEMENTATION-PERIOD", warning_codes)
        self.assertIn("CAR-VAL-STANDARD-IMPLEMENTATION-PERIOD", {message.code for message in unit.warnings})
        provenance = json.loads(actual.evidence.provenance_snapshot_json)
        expected_identity = calculator.reference_data_identity_provider()
        self.assertEqual(provenance["reference_data"], {"status": "CAPTURED", **expected_identity})
        self.assertEqual(self.repository.list_all(), ())

    def test_r2_importer_uses_preview_calculators_resolver_and_rejects_a_different_one(self) -> None:
        calculator = self.page.calculation_use_case.calculator
        preview_use_case = CarbonAccountingPreviewUseCase(calculator)
        resolver = calculator.parameter_resolver
        importer = ExcelWorkbookImporter(resolver, preview_use_case=preview_use_case)
        self.assertIs(importer.preview_use_case, preview_use_case)
        self.assertIs(importer.parameter_resolver, resolver)
        self.assertIs(importer.preview_use_case.calculator.parameter_resolver, resolver)

        preview = importer.import_preview(self._workbook())
        self.assertTrue(preview.units[0].can_calculate, preview.units[0].errors)
        self.assertTrue(preview.units[0].calculation.parameter_snapshots)
        self.assertEqual(self.repository.list_all(), ())

        different_resolver = create_g06_parameter_resolver(SQLiteCatalogRepository(self.catalog))
        self.assertIsNot(different_resolver, resolver)
        with self.assertRaisesRegex(ValueError, "必须使用同一参数选择服务"):
            ExcelWorkbookImporter(different_resolver, preview_use_case=preview_use_case)

    def test_word_export_reads_only_frozen_record_after_catalog_is_unavailable(self) -> None:
        record = self._create_record()
        record_before = self._frozen_state(record.record_id)
        supplementary = {
            "prepared_on": date(2026, 10, 7).isoformat(),
            "supplementary_note": "RS03 历史核算报告导出验证",
        }
        report_before = build_saved_record_report(
            self.repository,
            record,
            supplementary_info=supplementary,
        )
        self._record_page(record.record_id)
        page = self.shell.pages[AppRoute.RECORDS]
        destination = self.root / "历史核算报告.docx"
        archived_catalog = self.root / "catalog-unavailable.sqlite"
        self.catalog.rename(archived_catalog)

        with (
            patch.object(page, "_report_supplementary_dialog", return_value=supplementary),
            patch("packages.ui.pages.QFileDialog.getSaveFileName", return_value=(str(destination), "")),
            patch("packages.ui.pages.QMessageBox.information") as success,
            patch("packages.ui.pages.QMessageBox.critical") as failure,
            patch("packages.ui.pages.QMessageBox.warning") as warning,
        ):
            page.export_word_button.click()
            self.app.processEvents()

        failure.assert_not_called()
        warning.assert_not_called()
        success.assert_called_once()
        self.assertTrue(destination.is_file())
        self.assertFalse(self.catalog.exists())
        self.assertTrue(archived_catalog.exists())
        parsed = Document(BytesIO(destination.read_bytes()))
        visible_text = "\n".join(
            [paragraph.text for paragraph in parsed.paragraphs]
            + [cell.text for table in parsed.tables for row in table.rows for cell in row.cells]
        )
        self.assertIn("B.1 温室气体排放量汇总", visible_text)
        self.assertIn("B.9 购入和输出热力", visible_text)
        self.assertIn("27.03 tCO₂", visible_text)
        report_after = build_saved_record_report(
            self.repository,
            self.repository.get(record.record_id),
            supplementary_info=supplementary,
        )
        self.assertEqual(asdict(report_after), asdict(report_before))

        with closing(sqlite3.connect(self.records_path)) as connection:
            history = connection.execute(
                "SELECT format, document_sha256 FROM report_export_history WHERE record_id = ?",
                (record.record_id,),
            ).fetchall()
            audit = connection.execute(
                "SELECT action, details_json FROM audit_log WHERE record_id = ? AND action = 'REPORT_EXPORT'",
                (record.record_id,),
            ).fetchall()
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0][0], "DOCX")
        digest = sha256(destination.read_bytes()).hexdigest().upper()
        self.assertEqual(history[0][1], digest)
        self.assertEqual(len(audit), 1)
        self.assertEqual(audit[0][0], "REPORT_EXPORT")
        self.assertEqual(json.loads(audit[0][1])["document_sha256"], digest)
        self.assertEqual(self._frozen_state(record.record_id), record_before)

    def test_word_generation_failure_preserves_existing_file_and_record(self) -> None:
        record = self._create_record()
        before = self._frozen_state(record.record_id)
        page = self._record_page(record.record_id)
        destination = self.root / "existing-report.docx"
        original = b"existing report content"
        destination.write_bytes(original)

        with (
            patch.object(page, "_report_supplementary_dialog", return_value={"prepared_on": "2026-10-07"}),
            patch("packages.ui.pages.QFileDialog.getSaveFileName", return_value=(str(destination), "")),
            patch("packages.ui.pages.render_report_docx", side_effect=RuntimeError("private renderer detail")),
            patch("packages.ui.pages.QMessageBox.critical") as failure,
            patch("packages.ui.pages.QMessageBox.information") as success,
            patch("packages.ui.pages.QMessageBox.warning") as warning,
        ):
            page.export_word_button.click()
            self.app.processEvents()

        failure.assert_called_once()
        success.assert_not_called()
        warning.assert_not_called()
        self.assertIn(str(destination), failure.call_args.args[2])
        self.assertNotIn("private renderer detail", failure.call_args.args[2])
        self.assertIn("文件未保存", failure.call_args.args[2])
        self.assertEqual(destination.read_bytes(), original)
        self.assertEqual(self._frozen_state(record.record_id), before)
        with closing(sqlite3.connect(self.records_path)) as connection:
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM report_export_history WHERE record_id = ?", (record.record_id,)).fetchone()[0],
                0,
            )
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) FROM audit_log WHERE record_id = ? AND action = 'REPORT_EXPORT'",
                    (record.record_id,),
                ).fetchone()[0],
                0,
            )

    def test_word_write_failure_preserves_existing_file_and_record(self) -> None:
        record = self._create_record()
        before = self._frozen_state(record.record_id)
        page = self._record_page(record.record_id)
        destination = self.root / "existing-report.docx"
        original = b"existing report content"
        destination.write_bytes(original)

        with (
            patch.object(page, "_report_supplementary_dialog", return_value={"prepared_on": "2026-10-07"}),
            patch("packages.ui.pages.QFileDialog.getSaveFileName", return_value=(str(destination), "")),
            patch("packages.ui.pages.os.replace", side_effect=OSError("private filesystem detail")),
            patch("packages.ui.pages.QMessageBox.critical") as failure,
            patch("packages.ui.pages.QMessageBox.information") as success,
            patch("packages.ui.pages.QMessageBox.warning") as warning,
        ):
            page.export_word_button.click()
            self.app.processEvents()

        failure.assert_called_once()
        success.assert_not_called()
        warning.assert_not_called()
        self.assertIn(str(destination), failure.call_args.args[2])
        self.assertNotIn("private filesystem detail", failure.call_args.args[2])
        self.assertIn("文件未保存", failure.call_args.args[2])
        self.assertEqual(destination.read_bytes(), original)
        self.assertEqual(tuple(self.root.glob(".ghg-report-*.tmp")), ())
        self.assertEqual(self._frozen_state(record.record_id), before)
        with closing(sqlite3.connect(self.records_path)) as connection:
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM report_export_history WHERE record_id = ?", (record.record_id,)).fetchone()[0],
                0,
            )
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) FROM audit_log WHERE record_id = ? AND action = 'REPORT_EXPORT'",
                    (record.record_id,),
                ).fetchone()[0],
                0,
            )

    def test_word_audit_failure_reports_saved_path_without_claiming_completion(self) -> None:
        record = self._create_record()
        before = self._frozen_state(record.record_id)
        page = self._record_page(record.record_id)
        destination = self.root / "saved-without-audit.docx"

        with (
            patch.object(page, "_report_supplementary_dialog", return_value={"prepared_on": "2026-10-07"}),
            patch("packages.ui.pages.QFileDialog.getSaveFileName", return_value=(str(destination), "")),
            patch.object(self.repository, "record_report_export", side_effect=RuntimeError("private audit detail")),
            patch("packages.ui.pages.QMessageBox.warning") as warning,
            patch("packages.ui.pages.QMessageBox.information") as success,
            patch("packages.ui.pages.QMessageBox.critical") as failure,
        ):
            page.export_word_button.click()
            self.app.processEvents()

        warning.assert_called_once()
        success.assert_not_called()
        failure.assert_not_called()
        message = warning.call_args.args[2]
        self.assertIn(str(destination), message)
        self.assertIn("已保存", message)
        self.assertIn("导出审计未完成", message)
        self.assertIn("正式记录未改", message)
        self.assertIn("保留该文件", message)
        self.assertNotIn("稍后重试", message)
        self.assertNotIn("private audit detail", message)
        self.assertTrue(destination.is_file())
        parsed = Document(BytesIO(destination.read_bytes()))
        self.assertTrue(parsed.paragraphs)
        self.assertEqual(self._frozen_state(record.record_id), before)
        with closing(sqlite3.connect(self.records_path)) as connection:
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM report_export_history WHERE record_id = ?", (record.record_id,)).fetchone()[0],
                0,
            )
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) FROM audit_log WHERE record_id = ? AND action = 'REPORT_EXPORT'",
                    (record.record_id,),
                ).fetchone()[0],
                0,
            )

    def test_adjusted_word_suffix_collision_requires_confirmation(self) -> None:
        record = self._create_record()
        page = self._record_page(record.record_id)
        selected = self.root / "report.xlsx"
        final_path = self.root / "report.docx"
        original = b"keep this existing file"
        final_path.write_bytes(original)

        with (
            patch.object(page, "_report_supplementary_dialog", return_value={"prepared_on": "2026-10-07"}),
            patch("packages.ui.pages.QFileDialog.getSaveFileName", return_value=(str(selected), "")),
            patch(
                "packages.ui.pages.QMessageBox.question",
                return_value=QMessageBox.StandardButton.No,
            ) as question,
            patch("packages.ui.pages.render_report_docx") as renderer,
        ):
            page.export_word_button.click()
            self.app.processEvents()

        question.assert_called_once()
        self.assertIn(str(final_path), question.call_args.args[2])
        renderer.assert_not_called()
        self.assertEqual(final_path.read_bytes(), original)

    def test_record_page_has_no_excel_result_export_and_r2_template_still_exports(self) -> None:
        record = self._create_record()
        record_page = self._record_page(record.record_id)
        self.assertTrue(record_page.export_word_button.isEnabled())
        self.assertFalse(hasattr(record_page, "export_excel_button"))
        self.assertIsNone(record_page.findChild(QPushButton, "exportExcelReportButton"))

        self.shell.navigate(AppRoute.EXCEL_IMPORT)
        self.app.processEvents()
        page = self.shell.pages[AppRoute.EXCEL_IMPORT]
        destination = self.root / "exported-R2-template.xlsx"
        with (
            patch("packages.ui.pages.QFileDialog.getSaveFileName", return_value=(str(destination), "")),
            patch("packages.ui.pages.QMessageBox.information") as success,
        ):
            page.templateButton.click()
            self.app.processEvents()
        success.assert_called_once()
        self.assertTrue(destination.is_file())
        workbook = load_workbook(destination, read_only=False)
        self.assertEqual(tuple(name for name in workbook.sheetnames if name != METADATA_SHEET), VISIBLE_SHEETS)
        self.assertEqual(workbook[METADATA_SHEET].sheet_state, "hidden")
        workbook.close()


    def test_project_association_failures_explain_recovery_state_without_technical_details(self) -> None:
        self.shell.navigate(AppRoute.EXCEL_IMPORT)
        self.app.processEvents()
        page = self.shell.pages[AppRoute.EXCEL_IMPORT]
        page.project_service = ProjectWorkspaceService(
            SQLiteProjectWorkspaceRepository(self.root / "projects.sqlite")
        )
        page._refresh_saved_projects()
        page._preview_workbook(self._workbook())
        page.project_name.setText("关联恢复状态测试")
        page.save_project_button.click()
        self.app.processEvents()
        self.assertIsNotNone(page._workspace)
        self.assertTrue(page.calculate_button.isEnabled())

        expected_states = (
            (False, False, "恢复信息均未保存"),
            (True, False, "恢复信息已保留"),
            (True, True, "项目关联也已保存"),
        )
        for recovery_pending, association_saved, expected in expected_states:
            error = ProjectRecordAssociationError(
                "private repository detail",
                recovery_pending=recovery_pending,
                association_saved=association_saved,
            )
            with (
                patch.object(page.project_service, "save_after_record", side_effect=error),
                patch("packages.ui.pages.QMessageBox.warning") as warning,
            ):
                page.calculate_button.click()
                self.app.processEvents()
            message = page.status_label.text()
            self.assertIn(expected, message)
            self.assertNotIn("private repository detail", message)
            warning.assert_called_once()
            self.assertIn(expected, warning.call_args.args[2])
            self.assertEqual(len(self.repository.list_all()), expected_states.index((recovery_pending, association_saved, expected)) + 1)

        with (
            patch.object(page.project_service, "save_after_record", side_effect=OSError("unknown private failure")),
            patch("packages.ui.pages.QMessageBox.warning") as warning,
        ):
            page.calculate_button.click()
            self.app.processEvents()
        message = page.status_label.text()
        self.assertIn("项目关联状态无法确认", message)
        self.assertNotIn("恢复信息已保留", message)
        self.assertNotIn("unknown private failure", message)
        warning.assert_called_once()
        self.assertIn("状态无法确认", warning.call_args.args[1])
        self.assertEqual(len(self.repository.list_all()), 4)


if __name__ == "__main__":
    unittest.main()
