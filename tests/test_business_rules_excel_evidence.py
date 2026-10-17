"""Real OOXML cache evidence for A; does not claim native Excel/WPS validation."""
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook

from packages.application import CarbonAccountingPreviewUseCase
from packages.application.carbon_accounting import create_g06_parameter_resolver
from packages.application.catalog_queries import CatalogQueryService
from packages.excel.appendix_b import AppendixBWorkbookImporter
from packages.persistence.catalog_builder import build_catalog_database
from packages.persistence.catalog_repository import SQLiteCatalogRepository
from packages.reference_data import DEFAULT_SOURCE_PATH
from packages.standards.carbon_material import CarbonMaterialCalculator
from tests.exb01_fixture_helper import appendix_b_context, write_valid_appendix_b_workbook

NS="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG="http://schemas.openxmlformats.org/package/2006/relationships"


def replace_with_formula(path, sheet_name, reference, cached_value):
    """Patch actual ZIP XML: openpyxl save would discard the cache being tested."""
    out=BytesIO()
    with ZipFile(path) as source:
        book=ET.fromstring(source.read("xl/workbook.xml"))
        rels=ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
        targets={r.attrib["Id"]:r.attrib["Target"] for r in rels.findall(f"{{{PKG}}}Relationship")}
        sheet=next(s for s in book.findall(f"{{{NS}}}sheets/{{{NS}}}sheet") if s.attrib["name"]==sheet_name)
        target=targets[sheet.attrib[f"{{{REL}}}id"]].lstrip("/")
        if not target.startswith("xl/"): target="xl/"+target
        with ZipFile(out,"w",ZIP_DEFLATED) as updated:
            for info in source.infolist():
                data=source.read(info.filename)
                if info.filename==target:
                    root=ET.fromstring(data)
                    cell=root.find(f".//{{{NS}}}c[@r='{reference}']")
                    if cell is None: raise AssertionError(f"missing fixture cell {sheet_name}!{reference}")
                    for child in list(cell): cell.remove(child)
                    cell.attrib.pop("t",None)
                    ET.SubElement(cell,f"{{{NS}}}f").text="1+1"
                    if cached_value is not None:
                        ET.SubElement(cell,f"{{{NS}}}v").text=cached_value
                    data=ET.tostring(root,encoding="utf-8",xml_declaration=True)
                updated.writestr(info,data)
    path.write_bytes(out.getvalue())


class BusinessRulesExcelCacheEvidenceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=TemporaryDirectory()
        cls.root=Path(cls.temp.name)
        repository=SQLiteCatalogRepository(build_catalog_database(DEFAULT_SOURCE_PATH,cls.root/"catalog.sqlite"))
        cls.importer=AppendixBWorkbookImporter(
            preview_use_case=CarbonAccountingPreviewUseCase(CarbonMaterialCalculator(parameter_resolver=create_g06_parameter_resolver(repository))),
            catalog_service=CatalogQueryService(repository))

    @classmethod
    def tearDownClass(cls): cls.temp.cleanup()

    def test_input_formula_with_absent_or_stale_cache_never_becomes_activity(self):
        for cache in (None,"987654321.123456789"):
            with self.subTest(cache=cache):
                path=write_valid_appendix_b_workbook(self.root/("missing.xlsx" if cache is None else "stale.xlsx"),all_sources=True)
                self.assertTrue(self.importer.import_preview(path,context=appendix_b_context()).units[0].can_calculate)
                replace_with_formula(path,"B.2","B4",cache)
                unit=self.importer.import_preview(path,context=appendix_b_context()).units[0]
                self.assertFalse(unit.can_calculate)
                self.assertIsNone(unit.calculation)
                self.assertIn(("EXB01_FORMULA_REJECTED","B.2!B4"),{(p.code,p.location) for p in unit.errors})

    def test_result_formula_cache_never_changes_full_precision_preview(self):
        baseline=write_valid_appendix_b_workbook(self.root/"baseline.xlsx",all_sources=True)
        ordinary=self.importer.import_preview(baseline,context=appendix_b_context()).units[0]
        self.assertTrue(ordinary.can_calculate,ordinary.errors)
        expected=tuple((line.emission_source_id,line.amount) for line in ordinary.calculation.result.lines)
        for cache in (None,"987654321.123456789"):
            with self.subTest(cache=cache):
                path=write_valid_appendix_b_workbook(self.root/("result-missing.xlsx" if cache is None else "result-stale.xlsx"),all_sources=True)
                workbook=load_workbook(path,read_only=True,data_only=False)
                try:
                    result_headers=[cell for row in workbook["B.2"].iter_rows(min_row=2,max_row=3)
                                    for cell in row if isinstance(cell.value,str)
                                    and cell.value.startswith("排放量") and "tCO₂" in cell.value]
                    self.assertEqual(len(result_headers),1,"正式模板应有唯一的排放量结果列")
                    result_cell=f"{result_headers[0].column_letter}4"
                finally:
                    workbook.close()
                replace_with_formula(path,"B.2",result_cell,cache)
                unit=self.importer.import_preview(path,context=appendix_b_context()).units[0]
                self.assertTrue(unit.can_calculate,unit.errors)
                self.assertEqual(tuple((line.emission_source_id,line.amount) for line in unit.calculation.result.lines),expected)
