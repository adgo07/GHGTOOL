"""Integration checks for the approved Appendix B workbook ingress."""

from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP, localcontext
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from packages.application import CarbonAccountingPreviewUseCase
from packages.application.carbon_accounting import create_g06_parameter_resolver
from packages.application.catalog_queries import CatalogQueryService
from packages.excel.appendix_b import AppendixBWorkbookImporter
from packages.excel.ingress import WorkbookFatalError
from packages.persistence.catalog_builder import build_catalog_database
from packages.persistence.catalog_repository import SQLiteCatalogRepository
from packages.reference_data import DEFAULT_SOURCE_PATH
from packages.standards.carbon_material import CarbonMaterialCalculator, EmissionSourceStatus, FuelPath
from tests.exb01_fixture_helper import (
    appendix_b_context,
    rewrite_numeric_lexeme,
    write_valid_appendix_b_workbook,
)


class MainIntegrationExcelTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        catalog_path = build_catalog_database(DEFAULT_SOURCE_PATH, self.root / "catalog.sqlite")
        repository = SQLiteCatalogRepository(catalog_path)
        self.catalog_service = CatalogQueryService(repository)
        self.resolver = create_g06_parameter_resolver(repository)
        self.calculator = CarbonMaterialCalculator(parameter_resolver=self.resolver)
        self.preview_use_case = CarbonAccountingPreviewUseCase(self.calculator)
        self.importer = AppendixBWorkbookImporter(
            preview_use_case=self.preview_use_case,
            catalog_service=self.catalog_service,
        )

    def tearDown(self) -> None:
        self.directory.cleanup()

    def _context(self, *, fuel_row: int = 4):
        return appendix_b_context(fuel_path_overrides={f"B.2!B{fuel_row}": FuelPath.VOLUME})

    def test_b2_through_b9_form_one_canonical_input_and_use_the_shared_calculator(self) -> None:
        path = write_valid_appendix_b_workbook(self.root / "all-sources.xlsx", all_sources=True)
        workbook = load_workbook(path)
        process = workbook["B.3"]
        # Remove data-region merges before inserting rows. openpyxl does not
        # update merged-range internals when insert_rows() shifts the sheet.
        for area in list(process.merged_cells.ranges):
            if area.min_row >= 4:
                process.unmerge_cells(str(area))
        process.insert_rows(5)
        process["B5"], process["C5"], process["D5"] = "新增待煅烧原料", 20, 88
        process.insert_rows(7)
        process["B7"], process["C7"], process["D7"] = "新增待煅烧原料", 20, 2
        # Make group boundaries explicit so role inference remains driven by
        # the approved labels after the test-only row insertions.
        group_labels = {
            4: "进入煅烧炉的碳", 5: "进入煅烧炉的碳",
            6: "进入煅烧炉的挥发分", 7: "进入煅烧炉的挥发分",
            8: "输出煅烧炉的碳", 9: "输出煅烧炉的碳", 10: "输出煅烧炉的碳",
            11: "输出煅烧炉的挥发分",
        }
        for row, label in group_labels.items():
            process[f"A{row}"] = label
        workbook["B.9"]["G3"], workbook["B.9"]["G4"] = 0.10, 0.12
        workbook.save(path)
        workbook.close()
        preview = self.importer.import_preview(path, context=self._context())

        self.assertEqual(len(preview.units), 1)
        unit = preview.units[0]
        self.assertTrue(unit.can_calculate, tuple((item.code, item.location) for item in unit.errors))
        value = unit.input_value
        self.assertEqual(len(value.fuel_inputs), 1)
        self.assertEqual(len(value.calcinations), 1)
        added_materials = {row.name: row for row in value.calcinations[0].material_rows}
        added = added_materials["新增待煅烧原料"]
        self.assertEqual(added.mass_t, Decimal("20"))
        self.assertEqual(added.fixed_carbon_percent, Decimal("88"))
        self.assertEqual(added.volatile_matter_percent, Decimal("2"))
        feed = [row for row in value.calcinations[0].material_rows if row.role.value == "calcination_feed"]
        self.assertEqual(sum((row.mass_t for row in feed), Decimal("0")), Decimal("120"))
        self.assertEqual(len(value.bakings), 1)
        self.assertEqual(len(value.graphitizations), 1)
        self.assertEqual(len(value.fume_incinerations), 1)
        self.assertEqual(len(value.fgd_units), 1)
        self.assertEqual(len(value.electricity_details), 1)
        self.assertEqual(len(value.exported_electricity), 1)
        self.assertEqual(len(value.purchased_heat), 1)
        self.assertEqual(len(value.exported_heat), 1)
        self.assertEqual(
            {item.source_id for item in value.source_states if item.status is EmissionSourceStatus.INVOLVED},
            {item.source_id for item in value.source_states},
        )

        preview_outcome = unit.calculation
        # Independent formula check from the filled template values. B.3's
        # added 20 t material is paired across carbon / volatile rows and
        # contributes mass once. B.7 has two carbonate components; B.8/B.9
        # indirect emissions are checked as purchased minus exported amounts.
        with localcontext() as decimal_context:
            decimal_context.prec = 40
            decimal_context.rounding = ROUND_HALF_UP
            d = Decimal
            expected_sources = {
                "fuel": d("1") * d("0.5") * d("0.98") * d("44") / d("12"),
                "calcination": (
                    (d("100") * d("0.90") + d("20") * d("0.88") - d("80") * d("0.95"))
                    * d("44") / d("12")
                    + (d("100") * d("0.05") + d("20") * d("0.02") - d("80") * d("0.01"))
                    * d("0.35") * d("44") / d("16")
                ),
                "baking": (
                    (d("40") * d("0.90") + d("50") * d("0.80") - d("45") * d("0.92"))
                    * d("44") / d("12")
                    + (d("40") * d("0.05") + d("50") * d("0.03"))
                    * d("0.35") * d("44") / d("16")
                ),
                "graphitization": (
                    (d("40") * d("0.80") + d("35") * d("0.90") - d("32") * d("0.95"))
                    * d("44") / d("12")
                    + d("40") * d("0.04") * d("0.35") * d("44") / d("16")
                ),
                "gas_control": (
                    d("100") * d("10") * d("30") * d("0.02") * d("0.98")
                    * d("365") * d("24") * d("44") / d("12") * d("1e-9")
                    + d("2") * d("0.90") * d("0.440") * d("0.95")
                    + d("2") * d("0.10") * d("0.522") * d("0.90")
                ),
                "purchased_electricity": d("100") * d("0.5306"),
                "purchased_heat": d("1000000") * d("2777") * d("0.10") / d("1000000"),
                "exported_electricity": d("2") * d("0.5"),
                "exported_heat": d("100000") * d("2777") * d("0.12") / d("1000000"),
            }
            expected_es = sum(
                (expected_sources[key] for key in ("fuel", "calcination", "baking", "graphitization", "gas_control")),
                d("0"),
            )
            expected_ei = (
                expected_sources["purchased_electricity"] + expected_sources["purchased_heat"]
                - expected_sources["exported_electricity"] - expected_sources["exported_heat"]
            )
            expected_et = expected_es + expected_ei

        trace_snapshot = json.loads(preview_outcome.evidence.trace_snapshot_json)
        actual_sources = {
            key: Decimal(amount)
            for key, amount in trace_snapshot["source_subtotals"].items()
        }
        self.assertEqual(actual_sources, expected_sources)
        self.assertEqual(
            {key: Decimal(amount) for key, amount in trace_snapshot["aggregations"].items()},
            {"ES": expected_es, "EI": expected_ei, "ET": expected_et},
        )
        self.assertEqual(preview_outcome.result.total_amount, expected_et)

        expected = self.calculator.calculate(value, calculated_at=preview_outcome.result.calculated_at)
        self.assertEqual(preview_outcome.result, expected.result)
        self.assertEqual(preview_outcome.traces, expected.traces)
        self.assertIsNone(preview_outcome.record)

    def test_invalid_cell_blocks_the_entire_workbook_even_with_other_valid_sources(self) -> None:
        for case, invalid_value, code in (
            ("text-number", "1.25", "EXB01_NUMERIC_CELL_REQUIRED"),
            ("formula", "=1+1", "EXB01_FORMULA_REJECTED"),
        ):
            with self.subTest(case=case):
                path = write_valid_appendix_b_workbook(self.root / f"{case}.xlsx")
                workbook = load_workbook(path)
                workbook["B.2"]["A14"] = "天然气"
                workbook["B.2"]["B14"] = 1.25
                workbook["B.8"]["B3"] = "电网电力"
                workbook["B.8"]["C3"] = invalid_value
                workbook.save(path)
                workbook.close()

                preview = self.importer.import_preview(path, context=self._context(fuel_row=14))
                self.assertEqual(len(preview.units), 1)
                unit = preview.units[0]
                self.assertFalse(unit.can_calculate)
                self.assertIsNone(unit.calculation)
                self.assertIn(code, {item.code for item in unit.errors})
                self.assertTrue(unit.input_value.fuel_inputs)
                self.assertFalse(unit.input_value.electricity_details)

    def test_nan_and_infinity_are_rejected_by_the_real_import_path(self) -> None:
        for index, lexeme in enumerate(("NaN", "Infinity", "-Infinity")):
            with self.subTest(lexeme=lexeme):
                path = write_valid_appendix_b_workbook(self.root / f"nonfinite-{index}.xlsx", all_sources=True)
                rewrite_numeric_lexeme(path, "B.2", "B4", lexeme)
                try:
                    preview = self.importer.import_preview(path, context=self._context())
                except WorkbookFatalError as exc:
                    # openpyxl may reject the malformed numeric token before cell parsing.
                    self.assertIn("Excel", str(exc))
                else:
                    unit = preview.units[0]
                    self.assertFalse(unit.can_calculate)
                    self.assertIsNone(unit.calculation)
                    self.assertIn("EXB01_NUMBER_NONFINITE", {item.code for item in unit.errors})

    def test_saved_decimal_lexical_and_scientific_15_digit_values_are_preserved(self) -> None:
        path = write_valid_appendix_b_workbook(self.root / "exact-decimals.xlsx")
        workbook = load_workbook(path)
        fuel = workbook["B.2"]
        fuel["A14"] = "天然气"
        fuel["B14"] = 1
        fuel["C14"] = Decimal("0.01")
        fuel["D14"] = "实测值"
        workbook.save(path)
        workbook.close()
        raw_activity = "1.23000000000000"
        raw_carbon = "1.23456789012345E-2"
        rewrite_numeric_lexeme(path, "B.2", "B14", raw_activity)
        rewrite_numeric_lexeme(path, "B.2", "C14", raw_carbon)

        preview = self.importer.import_preview(path, context=self._context(fuel_row=14))
        unit = preview.units[0]
        self.assertTrue(unit.can_calculate, tuple((item.code, item.location) for item in unit.errors))
        fuel_input = unit.input_value.fuel_inputs[0]
        self.assertEqual(fuel_input.activity.value, Decimal(raw_activity))
        self.assertEqual(fuel_input.carbon_content.value, Decimal(raw_carbon))
        evidence = {(item.sheet, item.cell): item for item in preview.numeric_evidence}
        self.assertEqual(evidence[("B.2", "B14")].serialized_numeric_text, raw_activity)
        self.assertEqual(evidence[("B.2", "C14")].serialized_numeric_text, raw_carbon)
        self.assertEqual(evidence[("B.2", "B14")].normalized_decimal, Decimal(raw_activity))
        self.assertEqual(evidence[("B.2", "C14")].normalized_decimal, Decimal(raw_carbon))

    def test_more_than_15_significant_digits_blocks_workbook(self) -> None:
        path = write_valid_appendix_b_workbook(self.root / "overprecision.xlsx", all_sources=True)
        rewrite_numeric_lexeme(path, "B.2", "B4", "1.234567890123456E+2")
        preview = self.importer.import_preview(path, context=self._context())
        unit = preview.units[0]
        self.assertFalse(unit.can_calculate)
        self.assertIsNone(unit.calculation)
        self.assertIn("EXB01_NUMBER_SIGNIFICANT_DIGITS", {item.code for item in unit.errors})

    def test_missing_required_sheet_is_fatal(self) -> None:
        path = write_valid_appendix_b_workbook(self.root / "missing-sheet.xlsx")
        workbook = load_workbook(path)
        del workbook["B.9"]
        workbook.save(path)
        workbook.close()

        with self.assertRaises(WorkbookFatalError):
            self.importer.import_preview(path, context=self._context())


if __name__ == "__main__":
    unittest.main()
