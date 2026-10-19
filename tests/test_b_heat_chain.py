"""真实offscreen Qt/SQLite项目/正式Record/冻结Word热力链。"""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from dataclasses import replace
from decimal import Decimal
import unittest
from unittest.mock import patch
from PySide6.QtWidgets import QApplication
from packages.application.carbon_accounting import CarbonAccountingUseCase
from packages.application.project_workspaces import ProjectWorkspaceService
from packages.persistence import SQLiteProjectWorkspaceRepository, SQLiteRecordRepository
from packages.standards.carbon_material import SOURCE_PURCHASED_HEAT, SOURCE_EXPORTED_HEAT, HeatFactorMode
from tests.test_b_material_draft_scope import MaterialDraftScopeTests


class HeatQtTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])
    page = MaterialDraftScopeTests.page

    @staticmethod
    def prepare(gui):
        gui.enterprise_name.setText("热力整链")
        gui.period_year.setValue(2026)
        gui.boundary_confirmed.setChecked(True)
        for status in gui._source_statuses.values():
            status.setCurrentIndex(status.findData("NOT_INVOLVED"))
        status = gui._source_statuses[SOURCE_PURCHASED_HEAT]
        status.setCurrentIndex(status.findData("INVOLVED"))
        row = gui._energy_family_rows["heat"][0]
        row.heat_boundary.setCurrentIndex(row.heat_boundary.findData("OUTSIDE_ENTERPRISE"))
        row.enthalpy_mode.setCurrentIndex(row.enthalpy_mode.findData("MANUAL"))
        row.amount.setText("1")
        row.enthalpy.setText("2800")
        return row

    def test_known_gui_project_reopen_and_frozen_record_word(self):
        with self.page() as (gui, projects, root, shell, records):
            row = self.prepare(gui)
            # Auxiliary invalid draft stays in Presentation State only.
            row.pressure.setText("999")
            state = gui._capture_form_state()
            workspace = projects.new_workspace("热力草稿")
            projects.save(replace(workspace, units=(replace(workspace.units[0], form_state=state),)))
            recovered = ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/"projects.sqlite")).get(workspace.project_id)
            gui._restore_form_state(recovered.units[0].form_state)
            row = gui._energy_family_rows["heat"][0]
            self.assertEqual(row.pressure.text(), "999")
            self.assertEqual(row.heat_boundary.currentData(), "OUTSIDE_ENTERPRISE")
            value = gui._input(increment=False)
            self.assertIsNone(value.purchased_heat[0].pressure_mpa)
            self.assertEqual(len(records.list_all()), 0)
            result = CarbonAccountingUseCase(gui.calculator, records).calculate(value)
            self.assertTrue(result.successful, result.problems)
            self.assertEqual(result.result.total_amount, Decimal("0.308"))
            again = CarbonAccountingUseCase(gui.calculator, records).calculate(value)
            self.assertTrue(again.successful)
            self.assertEqual(len(records.list_all()), 2)
            self.assertNotEqual(result.record.record_id, again.record.record_id)
            from packages.application.reporting.model import build_report_model
            saved = SQLiteRecordRepository(root/"records.sqlite")
            record_id = result.record.record_id
            with patch.object(type(gui.calculator), "calculate", side_effect=AssertionError("frozen")), patch.object(gui._parameter_resolver, "resolve", side_effect=AssertionError("frozen")):
                model = build_report_model(saved.get(record_id), raw_input_snapshot=saved.get_raw_input_snapshot(record_id),
                    trace_snapshot=saved.get_trace_snapshot(record_id), reporting_snapshot=saved.get_reporting_snapshot(record_id),
                    provenance_snapshot=saved.get_provenance_snapshot(record_id), report_qualification=saved.get_report_qualification(record_id))
            text = str(model)
            self.assertIn("2800", text)
            self.assertNotIn("C.4", text)
            self.assertNotIn("C.5", text)
            from packages.infrastructure.reporting import render_report_docx
            from io import BytesIO
            from zipfile import ZipFile
            with patch.object(type(gui.calculator), "calculate", side_effect=AssertionError("frozen")), patch.object(gui._parameter_resolver, "resolve", side_effect=AssertionError("frozen")):
                word = render_report_docx(model)
                with ZipFile(BytesIO(word)) as archive:
                    xml = archive.read("word/document.xml").decode()
            self.assertIn("2800", xml)
            self.assertIn("0.11", xml)
            self.assertNotIn("附录C.4", xml)

    def test_internal_no_unrelated_values_and_boundary_conflict(self):
        with self.page() as (gui, projects, root, shell, records):
            row = self.prepare(gui)
            row.heat_boundary.setCurrentIndex(row.heat_boundary.findData("INSIDE_ENTERPRISE"))
            row.heat_fuel_accounted.setCurrentIndex(row.heat_fuel_accounted.findData(True))
            row.amount.clear()
            row.enthalpy.clear()
            row.heat_factor_mode.setCurrentIndex(row.heat_factor_mode.findData(HeatFactorMode.MEASURED))
            value = gui._input(increment=False)
            self.assertEqual(len(value.purchased_heat), 1)
            from packages.ui.source_cards import SourceCardPresentationState
            self.assertEqual(gui._derive_source_card_state(SOURCE_PURCHASED_HEAT)[0],SourceCardPresentationState.COMPLETED)
            result = CarbonAccountingUseCase(gui.calculator, records).calculate(value)
            self.assertTrue(result.successful, result.problems)
            from packages.application.reporting.model import build_report_model
            rid=result.record.record_id
            model=build_report_model(records.get(rid),raw_input_snapshot=records.get_raw_input_snapshot(rid),trace_snapshot=records.get_trace_snapshot(rid),
                provenance_snapshot=records.get_provenance_snapshot(rid),reporting_snapshot=records.get_reporting_snapshot(rid),report_qualification=records.get_report_qualification(rid))
            b9=next(section for section in model.sections if section.section_id=="b9")
            self.assertIn("不重复",str(b9))
            self.assertNotIn("未保存明细",str(b9))
            row.heat_boundary.setCurrentIndex(row.heat_boundary.findData("OUTSIDE_ENTERPRISE"))
            row.amount.setText("1")
            row.enthalpy.setText("2800")
            row.heat_factor_mode.setCurrentIndex(row.heat_factor_mode.findData(HeatFactorMode.STANDARD_DEFAULT))
            bad = CarbonAccountingUseCase(gui.calculator, records).calculate(gui._input(increment=False))
            self.assertFalse(bad.successful)
            self.assertEqual(len(records.list_all()), 1)

    def test_measured_blank_partial_row_and_negative_known_warning(self):
        from packages.core.errors import DomainValidationError
        with self.page() as (gui, projects, root, shell, records):
            row = self.prepare(gui)
            row.heat_factor_mode.setCurrentIndex(row.heat_factor_mode.findData(HeatFactorMode.MEASURED))
            with self.assertRaises(DomainValidationError): gui._input(increment=False)
            row.measured_heat_factor.setText("0")
            result = CarbonAccountingUseCase(gui.calculator, records).calculate(gui._input(increment=False))
            self.assertTrue(result.successful, result.problems)
            row.enthalpy.setText("-1")
            self.assertTrue(row.enthalpy.hasAcceptableInput())
            result = CarbonAccountingUseCase(gui.calculator, records).calculate(gui._input(increment=False))
            self.assertTrue(result.successful, result.problems)
            self.assertTrue(any(p.code == "CAR-VAL-STEAM-MANUAL-RANGE" for p in result.problems))
            row.amount.clear()
            with self.assertRaises(DomainValidationError): gui._input(increment=False)
            self.assertEqual(len(records.list_all()), 2)


class HeatExcelTests(HeatQtTests):
    # Reuse fixtures without collecting the inherited Qt cases a second time.
    test_known_gui_project_reopen_and_frozen_record_word = None
    test_internal_no_unrelated_values_and_boundary_conflict = None
    test_measured_blank_partial_row_and_negative_known_warning = None

    def test_v10_only_b9_changed_and_real_v9_remains_recognized(self):
        from scripts.project_appendix_b_v10 import SOURCE,TARGET,generated_bytes
        from packages.excel.templates import ExcelTemplateService
        from zipfile import ZipFile
        from io import BytesIO
        from openpyxl import load_workbook
        self.assertEqual(TARGET.read_bytes(), generated_bytes())
        with ZipFile(SOURCE) as old, ZipFile(TARGET) as new:
            self.assertEqual(old.namelist(),new.namelist())
            self.assertEqual([name for name in old.namelist() if old.read(name)!=new.read(name)], ["xl/worksheets/sheet9.xml"])
        for path,version in ((SOURCE,"2.7.0"),(TARGET,"2.8.0")):
            wb=load_workbook(BytesIO(path.read_bytes()))
            try:self.assertEqual(ExcelTemplateService.default().validate_structure(wb),version)
            finally:wb.close()

    def test_explicit_known_and_factor_sources_preview_no_records(self):
        from packages.excel.templates import ExcelTemplateService
        from openpyxl import load_workbook
        from io import BytesIO
        from packages.ui.view_models import AppRoute
        with self.page() as (gui,projects,root,shell,records):
            self.prepare(gui)
            page=shell.pages[AppRoute.EXCEL_IMPORT]
            context=page.excel_adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name="热力整链")
            cases=(
                ("default",{"B3":"饱和蒸汽","C3":1,"F3":2800,"H3":"已知焓","I3":"缺省值","J3":"边界外","D3":"unused invalid draft"},True,Decimal("0.308")),
                ("measured-blank",{"B3":"饱和蒸汽","C3":1,"F3":2800,"H3":"已知焓","I3":"实测值","J3":"边界外"},False,None),
                ("measured-zero",{"B3":"饱和蒸汽","C3":1,"F3":2800,"G3":0,"H3":"已知焓","I3":"实测值","J3":"边界外"},True,Decimal("0")),
                ("measured-equal-default",{"B3":"饱和蒸汽","C3":1,"F3":2800,"G3":0.11,"H3":"已知焓","I3":"实测值","J3":"边界外"},True,Decimal("0.308")),
                ("unknown",{"B3":"饱和蒸汽","C3":1,"F3":2800,"H3":"已知焓","I3":"缺省值","J3":"待核实"},False,None),
                ("internal",{"J3":"边界内","K3":"已计入","G3":"unused invalid draft"},True,Decimal("0")),
                ("internal-unknown",{"J3":"边界内","K3":"待核实"},False,None),
                ("known-missing",{"B3":"饱和蒸汽","C3":1,"H3":"已知焓","I3":"缺省值","J3":"边界外","D3":1},False,None),
                ("auto-no-fallback",{"B3":"饱和蒸汽","C3":1,"F3":2800,"D3":1000,"H3":"自动物性","I3":"缺省值","J3":"边界外"},False,None),
            )
            for name,cells,valid,expected in cases:
                with self.subTest(name=name):
                    wb=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
                    for cell,value in cells.items():wb["B.9"][cell]=value
                    path=root/(name+".xlsx");wb.save(path);wb.close()
                    preview=page.excel_adapter.import_preview(path,context)
                    self.assertEqual(preview.units[0].can_calculate,valid,preview.units[0].errors)
                    self.assertEqual(len(records.list_all()),0)
                    if valid:
                        value=preview.units[0].input_value
                        result=gui.calculator.calculate(value)
                        self.assertTrue(result.successful,result.problems)
                        self.assertEqual(result.result.total_amount,expected)
                        if name=="measured-equal-default":
                            self.assertEqual(value.purchased_heat[0].factor_mode,HeatFactorMode.MEASURED)
                            self.assertIsNotNone(value.purchased_heat[0].factor)


    def test_excel_project_reopen_purchase_export_and_frozen_word(self):
        from packages.excel.templates import ExcelTemplateService
        from openpyxl import load_workbook
        from io import BytesIO
        from packages.ui.view_models import AppRoute
        from packages.application.canonical_input_codec import encode_canonical_input,decode_canonical_input
        from packages.application.reporting.model import build_report_model
        from packages.infrastructure.reporting import render_report_docx
        from zipfile import ZipFile
        with self.page() as (gui,projects,root,shell,records):
            self.prepare(gui)
            page=shell.pages[AppRoute.EXCEL_IMPORT]
            context=page.excel_adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name="热力整链")
            wb=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
            for cell,value in {"B3":"饱和蒸汽","C3":1,"F3":2800,"H3":"已知焓","I3":"缺省值","J3":"边界外",
                               "B4":"过热蒸汽","C4":0.5,"F4":2800,"H4":"已知焓","I4":"缺省值"}.items():wb["B.9"][cell]=value
            path=root/"purchase-export.xlsx";wb.save(path);wb.close()
            preview=page.excel_adapter.import_preview(path,context)
            self.assertTrue(preview.units[0].can_calculate,preview.units[0].errors)
            canonical=preview.units[0].input_value
            self.assertEqual(decode_canonical_input(encode_canonical_input(canonical)),canonical)
            workspace=projects.new_workspace("热力Canonical项目")
            provenance={"source":"EXCEL_APPENDIX_B","workbook":{"template_version":preview.provenance.template_version,
                        "template_master_sha256":preview.provenance.template_master_sha256,"sha256":preview.provenance.workbook_sha256}}
            projects.save(replace(workspace,units=(replace(workspace.units[0],canonical_input=canonical,ingress_provenance=provenance),)))
            reopened=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/"projects.sqlite")).get(workspace.project_id)
            self.assertEqual(reopened.units[0].canonical_input,canonical)
            self.assertEqual(len(records.list_all()),0)
            outcome=CarbonAccountingUseCase(gui.calculator,records).calculate(reopened.units[0].canonical_input,
                ingress_provenance=reopened.units[0].ingress_provenance)
            self.assertTrue(outcome.successful,outcome.problems)
            self.assertEqual(outcome.result.total_amount,Decimal("0.154")) #.308 purchase-.154 export once
            self.assertEqual({x.emission_source_id:x.amount for x in outcome.result.lines if x.emission_source_id in {SOURCE_PURCHASED_HEAT,SOURCE_EXPORTED_HEAT}},
                             {SOURCE_PURCHASED_HEAT:Decimal("0.308"),SOURCE_EXPORTED_HEAT:Decimal("0.154")})
            rid=outcome.record.record_id
            saved=SQLiteRecordRepository(root/"records.sqlite")
            with patch.object(type(gui.calculator),"calculate",side_effect=AssertionError("frozen")),patch.object(gui._parameter_resolver,"resolve",side_effect=AssertionError("frozen")):
                model=build_report_model(saved.get(rid),raw_input_snapshot=saved.get_raw_input_snapshot(rid),trace_snapshot=saved.get_trace_snapshot(rid),
                    provenance_snapshot=saved.get_provenance_snapshot(rid),reporting_snapshot=saved.get_reporting_snapshot(rid),report_qualification=saved.get_report_qualification(rid))
                self.assertEqual(model.input_template_version,"2.8.0")
                b9=next(table for section in model.sections for table in section.tables if table.table_id=="b9")
                self.assertEqual(len(b9.rows),2)
                word=render_report_docx(model)
                with ZipFile(BytesIO(word)) as archive:xml=archive.read("word/document.xml").decode()
                self.assertIn("2800",xml)
                self.assertIn("0.11",xml)
                self.assertNotIn("附录C.4",xml)
            self.assertEqual(len(records.list_all()),1)


    def test_auto_gui_excel_project_record_word_same_if97_and_no_invalid_record(self):
        from packages.standards.carbon_material import SteamKind
        from packages.excel.templates import ExcelTemplateService
        from packages.ui.view_models import AppRoute
        from packages.application.reporting.model import build_report_model
        from packages.infrastructure.reporting import render_report_docx
        from openpyxl import load_workbook
        from io import BytesIO
        from zipfile import ZipFile
        import json
        with self.page() as (gui,projects,root,shell,records):
            row=self.prepare(gui)
            row.enthalpy_mode.setCurrentIndex(row.enthalpy_mode.findData("AUTO"))
            row.heat_kind.setCurrentIndex(row.heat_kind.findData(SteamKind.SUPERHEATED))
            row.pressure.setText(".0035");row.temperature.setText("26.85")
            row.enthalpy.setText("9999") # unselected known draft must not replace automatic h
            gui_value=gui._input(increment=False)
            self.assertIsNone(gui_value.purchased_heat[0].enthalpy)
            expected=gui.calculator.calculate(gui_value)
            self.assertTrue(expected.successful,expected.problems)
            wb=load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
            for cell,value in {"B3":"过热蒸汽","C3":1,"D3":.0035,"E3":26.85,"F3":9999,"H3":"自动物性","I3":"缺省值","J3":"边界外"}.items():wb["B.9"][cell]=value
            path=root/"if97.xlsx";wb.save(path);wb.close()
            page=shell.pages[AppRoute.EXCEL_IMPORT]
            preview=page.excel_adapter.import_preview(path,page.excel_adapter.context(period=gui._period(),boundary_confirmed=True,enterprise_name="热力整链"))
            self.assertTrue(preview.units[0].can_calculate,preview.units[0].errors)
            canonical=preview.units[0].input_value
            self.assertIsNone(canonical.purchased_heat[0].enthalpy)
            workspace=projects.new_workspace("IF97项目")
            projects.save(replace(workspace,units=(replace(workspace.units[0],canonical_input=canonical),)))
            reopened=ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(root/"projects.sqlite")).get(workspace.project_id)
            self.assertEqual(reopened.units[0].canonical_input,canonical)
            self.assertEqual(len(records.list_all()),0)
            result=CarbonAccountingUseCase(gui.calculator,records).calculate(reopened.units[0].canonical_input)
            self.assertTrue(result.successful,result.problems)
            self.assertEqual(result.result.total_amount,expected.result.total_amount)
            self.assertEqual(result.record.algorithm_version,"CAR-SM01-2026-10-18-HEAT.1")
            rid=result.record.record_id
            trace=json.dumps(records.get_trace_snapshot(rid),ensure_ascii=False)
            for key in ("property_library_output_j_per_kg_lexeme","property_working_pressure_pa_lexeme","property_working_temperature_k_lexeme","actual_phase","GHGTOOL_COOLPROP_IF97_PROPERTY_V1"):
                self.assertIn(key,trace)
            with patch.object(type(gui.calculator),"calculate",side_effect=AssertionError("frozen")),patch.object(gui._parameter_resolver,"resolve",side_effect=AssertionError("frozen")),patch("packages.standards.carbon_material_steam_if97._runtime",side_effect=AssertionError("frozen")):
                model=build_report_model(records.get(rid),raw_input_snapshot=records.get_raw_input_snapshot(rid),trace_snapshot=records.get_trace_snapshot(rid),
                    provenance_snapshot=records.get_provenance_snapshot(rid),reporting_snapshot=records.get_reporting_snapshot(rid),report_qualification=records.get_report_qualification(rid))
                with ZipFile(BytesIO(render_report_docx(model))) as archive:xml=archive.read("word/document.xml").decode()
                self.assertIn("IF97",xml);self.assertIn("8.0.0",xml)
            row.pressure.setText(".1");row.temperature.setText("90")
            rejected=CarbonAccountingUseCase(gui.calculator,records).calculate(gui._input(increment=False))
            self.assertFalse(rejected.successful)
            self.assertEqual(len(records.list_all()),1)
