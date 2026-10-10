"""Integration checks for the real GUI ingress, SQLite records and Appendix B preview."""

from __future__ import annotations

from contextlib import closing
from dataclasses import asdict
from datetime import date
from io import BytesIO
import gc
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtWidgets import QApplication, QLabel
from PySide6.QtTest import QSignalSpy
from docx import Document
from openpyxl import load_workbook

from apps.carbon_accounting_desktop.app import create_main_window
from apps.carbon_accounting_desktop.config import AppConfig
from packages.application.carbon_accounting import CarbonAccountingUseCase, RecordRepositoryConfigurationError
from packages.application.reporting import build_saved_record_report
from packages.core.models import AccountingPeriod, PeriodType
from packages.excel.appendix_b import AppendixBImportContext, AppendixBWorkbookImporter
from packages.excel.templates import ExcelTemplateService
from packages.persistence import SQLiteRecordRepository, build_catalog_database
from packages.persistence.in_memory_records import InMemoryRecordRepository
from packages.reference_data import DEFAULT_SOURCE_PATH
from packages.standards.carbon_material import EmissionSourceStatus, FuelPath, FuelType
from packages.ui.view_models import AppRoute


class FailingDetailedRecordRepository(InMemoryRecordRepository):
    def __init__(self) -> None:
        super().__init__()
        self.persistence_attempts = 0

    def create_with_details(self, record, **kwargs) -> None:
        self.persistence_attempts += 1
        raise OSError("simulated record-store failure")


class MainIntegrationUiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.catalog = build_catalog_database(DEFAULT_SOURCE_PATH, self.root / "catalog.sqlite")
        self.records_path = self.root / "records.sqlite"
        self.repository = SQLiteRecordRepository(self.records_path)
        self.window = create_main_window(
            AppConfig(catalog_database=self.catalog), record_repository=self.repository,
        )
        self.window.show()
        self.shell = self.window.centralWidget()
        self.shell.navigate(AppRoute.NEW_ACCOUNTING)
        self.app.processEvents()
        self.page = self.shell.pages[AppRoute.NEW_ACCOUNTING]
        self.page.period_year.setValue(2025)
        self.page.enterprise_name.setText("集成验收企业")
        self.page.boundary_confirmed.setChecked(True)
        status = self.page._source_statuses["CAR-SRC-FUEL-001"]
        status.setCurrentIndex(status.findData(EmissionSourceStatus.INVOLVED))
        row = self.page._fuel_rows[0]
        row.fuel_type.setCurrentIndex(row.fuel_type.findData(FuelType.NATURAL_GAS))
        row.path.setCurrentIndex(row.path.findData(FuelPath.VOLUME))
        row.activity.setText("1.25")

    def tearDown(self) -> None:
        self.window.close()
        self.window.deleteLater()
        self.app.processEvents()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        # Release deferred Qt/Python resources before Windows removes SQLite files.
        gc.collect()
        self.directory.cleanup()

    def test_formal_calculation_and_browsing_reject_different_repositories(self) -> None:
        from apps.carbon_accounting_desktop.product import create_shell
        from packages.ui.carbon_material_page import CarbonMaterialAccountingPage
        from packages.ui.shell import AppShell
        use_case = self.page.calculation_use_case
        other = InMemoryRecordRepository()
        config = AppConfig(catalog_database=self.catalog)
        factories = (
            lambda: create_main_window(config, record_repository=other, calculation_use_case=use_case),
            lambda: create_shell(config, record_repository=other, calculation_use_case=use_case),
            lambda: AppShell(self.shell.view_model, self.shell._logo_path, self.shell._icon_directory,
                             record_repository=other, calculation_use_case=use_case),
            lambda: CarbonMaterialAccountingPage(record_repository=other, calculation_use_case=use_case),
        )
        for factory in factories:
            with self.subTest(factory=factory):
                with self.assertRaisesRegex(RecordRepositoryConfigurationError, "同一记录仓库"):
                    factory()
        derived = create_main_window(config, calculation_use_case=use_case)
        self.assertIs(derived.centralWidget().record_repository, self.repository)
        derived.close()
        derived.deleteLater()
        self.app.processEvents()

    def _set_import_context(self, page) -> None:
        page.import_enterprise_name.setText("集成验收企业")
        page.import_period_type.setCurrentIndex(page.import_period_type.findData("ANNUAL"))
        page.import_period_start.setText("2025-01-01")
        page.import_period_end.setText("2025-12-31")
        page.import_boundary_confirmed.setChecked(True)

    def _workbook(self) -> Path:
        workbook = load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
        fuel = workbook["B.2"]
        fuel["A14"], fuel["B14"] = "天然气", 1.25
        fuel["D14"], fuel["F14"], fuel["I14"] = "计算值", "缺省值", "缺省值"
        destination = self.root / "输入.xlsx"
        workbook.save(destination)
        workbook.close()
        return destination

    def test_gui_and_excel_have_exact_results_and_preview_does_not_save(self) -> None:
        self.page.quick_calculate_button.click()
        self.app.processEvents()
        record = self.repository.list_all()[0]
        excel_page = self.shell.pages[AppRoute.EXCEL_IMPORT]
        preview = AppendixBWorkbookImporter(
            preview_use_case=excel_page.preview_use_case,
            catalog_service=excel_page.catalog_service,
        ).import_preview(
            self._workbook(),
            context=AppendixBImportContext(
                period=AccountingPeriod(PeriodType.ANNUAL, date(2025, 1, 1), date(2025, 12, 31)),
                enterprise_name="集成验收企业",
                boundary_confirmed=True,
            ),
        )
        unit = preview.units[0]
        self.assertTrue(unit.can_calculate, unit.errors)
        self.assertEqual(record.calculation_result.total_amount, unit.result.total_amount)
        self.assertEqual(
            tuple((line.amount, line.unit) for line in record.calculation_result.lines),
            tuple((line.amount, line.unit) for line in unit.result.lines),
        )
        self.assertIsNone(unit.calculation.record)
        self.assertEqual(len(self.repository.list_all()), 1)
        excel_page = self.shell.pages[AppRoute.EXCEL_IMPORT]
        self._set_import_context(excel_page)
        with patch("packages.ui.pages.QFileDialog.getOpenFileName", return_value=(str(self._workbook()), "")):
            excel_page.selectFileButton.click()
        text = excel_page.preview_text.toPlainText()
        self.assertIn("预览排放总量：", text)
        self.assertIn("tCO₂", text)
        self.assertNotIn(str(unit.result.total_amount), text)
        self.assertEqual(len(self.repository.list_all()), 1)

    def test_repeated_success_appends_and_failed_input_preserves_saved_records(self) -> None:
        self.page.quick_calculate_button.click()
        self.page.quick_calculate_button.click()
        records = self.repository.list_all()
        self.assertEqual(len(records), 2)
        self.assertNotEqual(records[0].record_id, records[1].record_id)
        self.shell.navigate(AppRoute.HOME)
        self.app.processEvents()
        home = self.shell.pages[AppRoute.HOME]
        self.assertNotIn("暂无核算记录", home.findChild(QLabel, "statusSummary").text())
        self.shell.navigate(AppRoute.NEW_ACCOUNTING)
        originals = tuple(asdict(record) for record in records)
        self.page._fuel_rows[0].activity.setText("-1")
        self.page.quick_calculate_button.click()
        self.assertIn("未完成", self.page.calculation_status_hint.text())
        reopened = SQLiteRecordRepository(self.records_path)
        self.assertEqual(tuple(asdict(record) for record in reopened.list_all()), originals)

    def test_record_persistence_failure_does_not_show_or_link_a_result(self) -> None:
        failing_repository = FailingDetailedRecordRepository()
        self.page.calculation_use_case = CarbonAccountingUseCase(
            self.page.calculator,
            failing_repository,
        )
        self.page.record_repository = failing_repository
        created = QSignalSpy(self.page.record_created)
        requested = QSignalSpy(self.page.record_requested)

        self.page.quick_calculate_button.click()
        self.app.processEvents()

        self.assertEqual(failing_repository.persistence_attempts, 1)
        self.assertEqual(failing_repository.list_all(), ())
        self.assertEqual(self.repository.list_all(), ())
        self.assertFalse(self.page.result_card.isVisible())
        self.assertIn("未能保存核算记录", self.page.result_status.text())
        self.assertIn("未能保存核算记录", self.page.validation_list.topLevelItem(0).text(0))
        self.assertEqual(self.page._unit().record_ids, ())
        self.assertIsNone(self.page._unit().result_snapshot)
        self.assertIn("未建立新的项目关联", self.page.project_save_status.text())
        self.assertEqual(created.count(), 0)
        self.assertEqual(requested.count(), 0)

    def test_both_word_entrypoints_export_identical_business_content_for_same_record(self) -> None:
        self.page.quick_calculate_button.click()
        record = self.repository.list_all()[0]
        before = asdict(record)
        first, second = self.root / "新建页报告.docx", self.root / "记录页报告.docx"
        supplementary = {"prepared_on": "2026-10-10"}
        with (
            patch("packages.ui.report_export.report_supplementary_dialog", return_value=supplementary),
            patch("packages.ui.report_export.QFileDialog.getSaveFileName", return_value=(str(first), "")),
            patch("packages.ui.report_export.QMessageBox.information"),
            patch("packages.ui.report_export.QMessageBox.critical") as failure,
        ):
            self.page.export_report_button.click()
            failure.assert_not_called()
        self.shell.navigate(AppRoute.RECORDS)
        record_page = self.shell.pages[AppRoute.RECORDS]
        self.assertTrue(record_page.open_record(record.record_id))
        with (
            patch.object(record_page, "_report_supplementary_dialog", return_value=supplementary),
            patch("packages.ui.report_export.QFileDialog.getSaveFileName", return_value=(str(second), "")),
            patch("packages.ui.report_export.QMessageBox.information"),
            patch("packages.ui.report_export.QMessageBox.critical") as failure,
        ):
            record_page.export_word_button.click()
            failure.assert_not_called()
        self.assertTrue(first.is_file())
        self.assertTrue(second.is_file())
        first_document, second_document = Document(first), Document(second)
        self.assertEqual(first_document._element.body.xml, second_document._element.body.xml)
        report = build_saved_record_report(
            self.repository,
            self.repository.get(record.record_id),
            supplementary_info=supplementary,
        )
        report_headings = {paragraph.text for paragraph in first_document.paragraphs}
        appendix_sections = tuple(
            section for section in report.sections
            if section.section_id in {f"b{number}" for number in range(1, 10)}
        )
        self.assertEqual(tuple(section.section_id for section in appendix_sections), tuple(f"b{number}" for number in range(1, 10)))
        for section in appendix_sections:
            self.assertIn(section.title, report_headings)
        b1 = next(section for section in report.sections if section.section_id == "b1").tables[0]
        self.assertEqual(first_document.tables[1].rows[0].cells[2].text, "排放量ᵇ\ntCO₂")
        self.assertEqual(first_document.tables[1].rows[-2].cells[2].text, b1.rows[-2].cells[2].value)
        self.assertEqual(first_document.tables[1].rows[-1].cells[2].text, b1.rows[-1].cells[2].value)
        self.assertEqual(len(self.repository.list_all()), 1)
        self.assertEqual(asdict(self.repository.get(record.record_id)), before)
        with closing(sqlite3.connect(self.records_path)) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM report_export_history").fetchone()[0], 2)

    def test_word_export_after_catalog_change_uses_frozen_record(self) -> None:
        self.page.quick_calculate_button.click()
        record = self.repository.list_all()[0]
        original = asdict(record)
        model = build_saved_record_report(self.repository, record)
        self.shell.navigate(AppRoute.RECORDS)
        record_page = self.shell.pages[AppRoute.RECORDS]
        self.assertTrue(record_page.open_record(record.record_id))
        destination = self.root / "核算报告.docx"
        self.catalog.rename(self.root / "catalog-unavailable.sqlite")
        with (
            patch.object(record_page, "_report_supplementary_dialog", return_value={"prepared_on": date(2026, 10, 7).isoformat()}),
            patch("packages.ui.report_export.QFileDialog.getSaveFileName", return_value=(str(destination), "")),
            patch("packages.ui.report_export.QMessageBox.information") as success,
            patch("packages.ui.report_export.QMessageBox.critical") as failure,
        ):
            record_page.export_word_button.click()
        failure.assert_not_called()
        success.assert_called_once()
        self.assertTrue(destination.is_file())
        document = Document(destination)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("B.1", text)
        self.assertIn("B.9", text)
        self.assertEqual(asdict(self.repository.get(record.record_id)), original)
        self.assertEqual(asdict(build_saved_record_report(self.repository, record)), asdict(model))
        with closing(sqlite3.connect(self.records_path)) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM report_export_history").fetchone()[0], 1)


if __name__ == "__main__":
    unittest.main()
