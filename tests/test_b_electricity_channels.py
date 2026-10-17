"""电力五表跨入口采用值、方向与来源的独立回归。仅使用离屏Qt。"""
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from packages.application.carbon_accounting import CarbonAccountingPreviewUseCase, CarbonAccountingUseCase, create_g06_parameter_resolver
from packages.application.catalog_queries import CatalogQueryService
from packages.core.parameter_resolution import UserProvidedParameterValue
from packages.excel.appendix_b import AppendixBImportContext, AppendixBWorkbookImporter
from packages.excel.templates import ExcelTemplateService
from packages.persistence import SQLiteCatalogRepository, SQLiteRecordRepository, build_catalog_database
from packages.standards.carbon_material import CarbonMaterialCalculator, EmissionSourceStatus, SOURCE_EXPORTED_ELECTRICITY, SOURCE_PURCHASED_ELECTRICITY
from packages.ui.carbon_material_page import CarbonMaterialAccountingPage
from tests.test_pr46_baseline_integration import CASES, NOW, PERIOD


class ElectricityChannelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])
        cls.directory = tempfile.TemporaryDirectory()
        cls.root = Path(cls.directory.name)
        cls.repository = SQLiteCatalogRepository(build_catalog_database(output_path=cls.root / "catalog.sqlite"))
        cls.catalog = CatalogQueryService(cls.repository)
        cls.resolver = create_g06_parameter_resolver(cls.repository)
        cls.calculator = CarbonMaterialCalculator(parameter_resolver=cls.resolver)
        cls.preview = CarbonAccountingPreviewUseCase(cls.calculator)
        cls.importer = AppendixBWorkbookImporter(preview_use_case=cls.preview, catalog_service=cls.catalog)

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def page(self):
        page = CarbonMaterialAccountingPage(catalog_service=self.catalog, calculator=self.calculator)
        page.period_year.setValue(2026)
        page.enterprise_name.setText("电力跨入口验证")
        page.boundary_confirmed.setChecked(True)
        for combo in page._source_statuses.values():
            combo.setCurrentIndex(combo.findData(EmissionSourceStatus.NOT_INVOLVED))
        page._toggle_source_group("electricity")
        return page

    def workbook_input(self, case, output):
        workbook = load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
        try:
            workbook["B.8"]["A3"] = "输出" if output else "购入"
            workbook["B.8"]["B3"] = None if output else "电网电力"
            workbook["B.8"]["C3"] = Decimal("100")
            workbook["B.8"]["D3"] = Decimal(case[5])
            path = self.root / f"{case[2]}-{output}.xlsx"
            workbook.save(path)
        finally:
            workbook.close()
        preview = self.importer.import_preview(path, context=AppendixBImportContext(period=PERIOD, boundary_confirmed=True, enterprise_name="电力跨入口验证", region=case[1]))
        self.assertEqual(len(preview.units), 1)
        unit = preview.units[0]
        self.assertTrue(unit.can_calculate, unit.errors)
        self.assertIsNone(unit.calculation.record)
        return unit.input_value

    def test_gui_excel_formal_paths_keep_five_table_values_and_each_direction(self):
        records = SQLiteRecordRepository(self.root / "channels-records.sqlite")
        formal = CarbonAccountingUseCase(self.calculator, records)
        for output in (False, True):
            page = self.page()
            try:
                row = page._energy_family_rows["electricity"][0]
                row.kind.setCurrentIndex(row.kind.findData("exported_electricity" if output else "purchased_electricity"))
                for case in CASES:
                    with self.subTest(output=output, table=case[2]):
                        row.geographic_level.setCurrentIndex(row.geographic_level.findData(case[0]))
                        row.region.setCurrentIndex(row.region.findData(case[1]))
                        index = next(i for i in range(row.factor_selector.count()) if getattr(row.factor_selector.itemData(i), "table_number", None) == case[2])
                        row.factor_selector.setCurrentIndex(index)
                        row.amount.setText("100")
                        gui_input = page._input(increment=False, render_electricity=False)
                        excel_input = self.workbook_input(case, output)
                        expected_source = SOURCE_EXPORTED_ELECTRICITY if output else SOURCE_PURCHASED_ELECTRICITY
                        before = len(records.list_all())
                        for channel, input_value in (("GUI", gui_input), ("Excel", excel_input)):
                            with self.subTest(channel=channel):
                                result = formal.calculate(input_value, calculated_at=NOW)
                                self.assertTrue(result.successful, result.problems)
                                record = records.get(result.record.record_id)
                                lines = [line.amount for line in record.calculation_result.lines if line.emission_source_id == expected_source]
                                self.assertEqual(lines, [Decimal(case[6])])
                                self.assertEqual(record.calculation_result.total_amount, Decimal(case[6]) * (-1 if output else 1))
                                snapshots = [snapshot for snapshot in record.parameter_snapshots if snapshot.factor_id == case[4]]
                                self.assertEqual(len(snapshots), 1)
                                snapshot = snapshots[0]
                                self.assertEqual((snapshot.parameter_id, snapshot.value_used, snapshot.factor_year, snapshot.source_id), (case[3], Decimal(case[5]), 2023, "SRC-ELEC-2023-47"))
                                self.assertIn(case[2], snapshot.source_location)
                        self.assertEqual(len(records.list_all()), before + 2)
            finally:
                page._project_dirty = False
                page.close()
                page.deleteLater()
                self.app.processEvents()

    def test_word_projection_uses_frozen_electricity_when_live_catalog_and_calculator_are_unavailable(self):
        from packages.application.reporting import build_saved_record_report
        records = SQLiteRecordRepository(self.root / "frozen-word-records.sqlite")
        formal = CarbonAccountingUseCase(self.calculator, records)
        frozen = []
        for case in CASES:
            input_value = self.workbook_input(case, False)
            outcome = formal.calculate(input_value, calculated_at=NOW)
            self.assertTrue(outcome.successful, outcome.problems)
            record = records.get(outcome.record.record_id)
            model = build_saved_record_report(records, record, supplementary_info={"prepared_on": PERIOD.end.isoformat()})
            table = next(section for section in model.sections if section.section_id == "b8").tables[0]
            self.assertEqual(table.rows[0].cells[-1].value, case[6])
            frozen.append((record.record_id, model))
        with patch.object(SQLiteCatalogRepository, "list_factors", side_effect=AssertionError("报告不得查询实时因子")), patch.object(CarbonMaterialCalculator, "calculate", side_effect=AssertionError("报告不得重新计算")):
            for record_id, model in frozen:
                self.assertEqual(build_saved_record_report(records, records.get(record_id), supplementary_info={"prepared_on": PERIOD.end.isoformat()}), model)

    def test_other_source_description_is_optional_without_a_blank_specific_warning(self):
        from packages.core.models import ElectricityAcquisitionMode, ElectricityAttribute
        from packages.core.parameter_resolution import ElectricityConsumptionDetail
        for description in (None, "企业选用资料"):
            with self.subTest(description=description):
                detail = ElectricityConsumptionDetail("other-source", "enterprise.power", "gbt_32151_34_2024", PERIOD, "100", "MWh", ElectricityAcquisitionMode.PURCHASED, ElectricityAttribute.ORDINARY, factor_override=UserProvidedParameterValue("0.41", "tCO₂/MWh", source_reference=description))
                resolved = self.resolver.resolve_electricity_details((detail,), snapshot_at=NOW)[0]
                self.assertFalse(resolved.blocked)
                self.assertEqual(resolved.snapshot.value_used, Decimal("0.41"))
                self.assertIsNone(resolved.snapshot.factor_id)
                self.assertIsNone(resolved.snapshot.source_id)
                self.assertNotIn("GEN-VAL-USER-OVERRIDE-SOURCE-MISSING", {problem.code for problem in resolved.result.warnings})

    def test_switching_direction_keeps_independent_other_source_drafts(self):
        page = self.page()
        try:
            row = page._energy_family_rows["electricity"][0]
            row.factor_mode.setCurrentIndex(row.factor_mode.findData("MANUAL"))
            row.manual_factor.setText("0.41")
            row.electricity_source.setText("购入资料A")
            row.kind.setCurrentIndex(row.kind.findData("exported_electricity"))
            row.factor_mode.setCurrentIndex(row.factor_mode.findData("MANUAL"))
            self.assertEqual(row.manual_factor.text(), "")
            self.assertEqual(row.electricity_source.text(), "")
            row.manual_factor.setText("0.72")
            row.electricity_source.setText("输出资料B")
            row.kind.setCurrentIndex(row.kind.findData("purchased_electricity"))
            self.assertEqual((row.manual_factor.text(), row.electricity_source.text()), ("0.41", "购入资料A"))
            saved = page._energy_row_state(row)
            page._restore_unified_energy_rows([saved])
            row = page._energy_family_rows["electricity"][0]
            row.kind.setCurrentIndex(row.kind.findData("exported_electricity"))
            self.assertEqual((row.manual_factor.text(), row.electricity_source.text()), ("0.72", "输出资料B"))
        finally:
            page._project_dirty = False
            page.close()
            page.deleteLater()
            self.app.processEvents()
