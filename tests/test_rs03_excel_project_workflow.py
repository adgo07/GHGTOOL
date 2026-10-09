"""End-to-end Excel project workflow: preview, save, reopen, calculate and browse."""

from __future__ import annotations

from contextlib import closing
from dataclasses import replace
from decimal import Decimal
from io import BytesIO
import hashlib
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication
from openpyxl import load_workbook

from apps.carbon_accounting_desktop.app import create_main_window
from apps.carbon_accounting_desktop.config import AppConfig
from packages.application import ProjectWorkspaceService
from packages.application.carbon_accounting import CarbonAccountingUseCase
from packages.application.canonical_input_codec import encode_canonical_input
from packages.core.models import RecordStatus
from packages.excel.templates import ExcelTemplateService
from packages.persistence import (
    InMemoryRecordRepository,
    SQLiteProjectWorkspaceRepository,
    SQLiteRecordRepository,
    build_catalog_database,
)
from packages.reference_data import DEFAULT_SOURCE_PATH
from packages.ui.pages import ExcelImportPage
from packages.ui.view_models import AppRoute


class _FailingRecordRepository(InMemoryRecordRepository):
    def create_with_details(self, record, **kwargs) -> None:
        raise OSError("测试记录库写入失败")


class Rs03ExcelProjectWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.catalog_path = build_catalog_database(DEFAULT_SOURCE_PATH, self.root / "catalog.sqlite")
        self.records_path = self.root / "records.sqlite"
        self.projects_path = self.root / "projects.sqlite"
        self.records = SQLiteRecordRepository(self.records_path)
        self.projects = ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(self.projects_path))
        self.window = create_main_window(
            AppConfig(
                catalog_database=self.catalog_path,
                records_database=self.records_path,
                projects_database=self.projects_path,
            ),
            record_repository=self.records,
            project_service=self.projects,
        )
        self.window.show()
        self.shell = self.window.centralWidget()
        self.page = self.shell.pages[AppRoute.EXCEL_IMPORT]
        self._set_import_context()
        self.workbook_path = self._appendix_b_workbook()

    def tearDown(self) -> None:
        self.window.close()
        self.window.deleteLater()
        self.app.processEvents()
        self.directory.cleanup()

    def _set_import_context(self, *, enterprise_name: str = "RS03 项目工作流企业") -> None:
        self.page.import_enterprise_name.setText(enterprise_name)
        self.page.import_period_type.setCurrentIndex(self.page.import_period_type.findData("ANNUAL"))
        self.page.import_period_start.setText("2025-01-01")
        self.page.import_period_end.setText("2025-12-31")
        self.page.import_boundary_confirmed.setChecked(True)

    def _appendix_b_workbook(self) -> Path:
        workbook = load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
        fuel = workbook["B.2"]
        fuel["A14"], fuel["B14"] = "天然气", 1.25
        fuel["D14"], fuel["F14"], fuel["I14"] = "计算值", "缺省值", "缺省值"
        destination = self.root / "导入源.xlsx"
        workbook.save(destination)
        workbook.close()
        return destination

    def test_single_unit_keeps_its_original_cell_lexical_after_source_changes(self) -> None:
        original_bytes = self.workbook_path.read_bytes()
        original_digest = hashlib.sha256(original_bytes).hexdigest()
        self.page._preview_workbook(self.workbook_path)
        preview = self.page._last_preview
        self.assertEqual(preview.provenance.workbook_sha256, original_digest)
        self.assertEqual(len(preview.units), 1)
        self.assertTrue(preview.units[0].can_calculate, preview.units[0].errors)

        imported_unit_id = preview.units[0].unit_id
        imported_input = preview.units[0].input_value
        self.assertEqual(imported_input.fuel_inputs[0].activity.value, Decimal("1.25"))
        changed = load_workbook(self.workbook_path)
        changed["B.2"]["B14"] = 99
        changed.save(self.workbook_path)
        changed.close()
        self.assertNotEqual(hashlib.sha256(self.workbook_path.read_bytes()).hexdigest(), original_digest)
        self.workbook_path.unlink()

        self.page.project_name.setText("附录B导入证据")
        self.page.save_project_button.click()
        self.app.processEvents()
        saved = self.projects.list_all()[0]
        self.assertEqual(len(saved.units), 1)
        unit = saved.units[0]
        self.assertEqual(unit.canonical_input, imported_input)
        self.assertEqual(self.records.list_all(), ())
        provenance = unit.ingress_provenance
        self.assertEqual(provenance["source"], "EXCEL_APPENDIX_B")
        self.assertEqual(provenance["workbook"]["sha256"], original_digest)
        self.assertEqual(provenance["accounting_unit"]["importer_unit_id"], imported_unit_id)
        evidence = provenance["numeric_cell_evidence"]
        activity = next(item for item in evidence if item["sheet"] == "B.2" and item["cell"] == "B14")
        self.assertEqual(activity["raw_numeric_lexical"], "1.25")
        self.assertEqual(activity["normalized_decimal_lexical"], "1.25")

    def _preview_and_save(self) -> tuple[str, object]:
        with patch("packages.ui.pages.QMessageBox.warning"):
            self.page._preview_workbook(self.workbook_path)
        self.assertIsNotNone(self.page._last_preview)
        imported_input = self.page._last_preview.units[0].input_value
        self.imported_unit_id = self.page._last_preview.units[0].unit_id
        self.assertEqual(self.records.list_all(), ())
        self.assertEqual(len(self.page._last_preview.units), 1)
        self.assertTrue(self.page._last_preview.units[0].can_calculate)

        self.page.project_name.setText("RS03 有效核算单元项目")
        self.page.save_project_button.click()
        self.app.processEvents()
        projects = self.projects.list_all()
        self.assertEqual(len(projects), 1)
        self.assertGreaterEqual(self.page.saved_projects.findData(projects[0].project_id), 0)
        self.assertEqual(self.records.list_all(), ())
        return projects[0].project_id, imported_input

    def test_preview_saves_only_valid_units_and_reopens_without_source_workbook(self) -> None:
        project_id, imported_input = self._preview_and_save()
        saved = self.projects.get(project_id)
        self.assertIsNotNone(saved)
        self.assertEqual(saved.name, "RS03 有效核算单元项目")
        self.assertEqual(len(saved.units), 1)
        unit = saved.units[0]
        self.assertNotEqual(unit.unit_id, self.imported_unit_id)
        self.assertEqual(unit.canonical_input, imported_input)
        self.assertEqual(
            unit.input_fingerprint,
            hashlib.sha256(encode_canonical_input(imported_input).encode("utf-8")).hexdigest(),
        )
        self.assertEqual(unit.record_ids, ())
        self.assertEqual(self.records.list_all(), ())

        provenance = unit.ingress_provenance
        self.assertIsNotNone(provenance)
        self.assertEqual(provenance["source"], "EXCEL_APPENDIX_B")
        numeric = provenance["numeric_cell_evidence"]
        activity_cell = next(item for item in numeric if item["sheet"] == "B.2" and item["cell"] == "B14")
        self.assertEqual(activity_cell["raw_numeric_lexical"], "1.25")
        self.assertEqual(activity_cell["normalized_decimal_lexical"], "1.25")
        encoded_provenance = json.dumps(provenance, ensure_ascii=False)
        self.assertIn("EXCEL_APPENDIX_B", encoded_provenance)
        self.assertIn("template_id", encoded_provenance)

        self.workbook_path.unlink()
        restarted_projects = ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(self.projects_path))
        reopened = ExcelImportPage(
            catalog_service=self.shell.catalog_service,
            calculation_use_case=self.shell.calculation_use_case,
            project_service=restarted_projects,
            record_repository=self.records,
        )
        self.assertTrue(reopened.open_project(project_id))
        self.assertIsNone(reopened._last_workbook_path)
        self.assertEqual(reopened._selected_unit().canonical_input, imported_input)
        self.assertIn("预览排放总量", reopened.preview_text.toPlainText())
        self.assertIs(reopened.preview_use_case.calculator, self.shell.calculation_use_case.calculator)
        self.assertEqual(self.records.list_all(), ())

    def test_repeated_formal_calculation_appends_records_with_frozen_import_evidence(self) -> None:
        project_id, _ = self._preview_and_save()
        self.assertTrue(self.page.open_project(project_id))
        self.assertTrue(self.page.calculate_button.isEnabled())

        created_record_ids = []
        self.page.record_created.connect(created_record_ids.append)
        self.page.calculate_button.click()
        self.app.processEvents()
        first = self.records.list_all()
        self.assertEqual(len(first), 1)
        self.assertIn(first[0].status, {RecordStatus.COMPLETED, RecordStatus.COMPLETED_WITH_WARNINGS})

        self.page.calculate_button.click()
        self.app.processEvents()
        records = self.records.list_all()
        self.assertEqual(set(created_record_ids), {records[0].record_id, records[1].record_id})
        self.assertEqual(len(records), 2)
        self.assertNotEqual(records[0].record_id, records[1].record_id)
        saved = self.projects.get(project_id)
        self.assertEqual(len(saved.units[0].record_ids), 2)
        self.assertEqual(set(saved.units[0].record_ids), {item.record_id for item in records})

        requested_record_ids = []
        self.page.record_requested.connect(requested_record_ids.append)
        selected_record_index = self.page.unit_record_selector.findData(records[1].record_id)
        self.assertGreaterEqual(selected_record_index, 0)
        self.page.unit_record_selector.setCurrentIndex(selected_record_index)
        self.page.open_record_button.click()
        self.assertEqual(requested_record_ids, [records[1].record_id])

        for record in records:
            raw = self.records.get_raw_input_snapshot(record.record_id)
            self.assertIn("ingress_provenance", raw)
            self.assertEqual(raw["ingress_provenance"]["source"], "EXCEL_APPENDIX_B")
            self.assertEqual(
                raw["ingress_provenance"]["numeric_cell_evidence"][0]["raw_numeric_lexical"],
                "1.25",
            )

    def test_blocked_calculation_and_record_write_failure_do_not_claim_completion(self) -> None:
        project_id, _ = self._preview_and_save()
        workspace = self.projects.get(project_id)
        blocked_unit = replace(
            workspace.units[0],
            canonical_input=replace(workspace.units[0].canonical_input, boundary_confirmed=False),
        )
        self.projects.save(replace(workspace, units=(blocked_unit,)))
        self.assertTrue(self.page.open_project(project_id))
        self.page.calculate_button.click()
        self.app.processEvents()
        self.assertEqual(self.records.list_all(), ())
        self.assertIn("未通过校验", self.page.status_label.text())

        valid_unit = replace(
            blocked_unit,
            canonical_input=replace(blocked_unit.canonical_input, boundary_confirmed=True),
        )
        self.projects.save(replace(self.projects.get(project_id), units=(valid_unit,)))
        failing_store = _FailingRecordRepository()
        failing_use_case = CarbonAccountingUseCase(self.shell.calculation_use_case.calculator, failing_store)
        failing_page = ExcelImportPage(
            catalog_service=self.shell.catalog_service,
            calculation_use_case=failing_use_case,
            project_service=self.projects,
            record_repository=failing_store,
        )
        with patch("packages.ui.pages.QMessageBox.warning"):
            self.assertTrue(failing_page.open_project(project_id))
            failing_page.calculate_button.click()
            self.app.processEvents()
        self.assertEqual(failing_store.list_all(), ())
        self.assertIn("正式记录保存失败", failing_page.status_label.text())
        self.assertNotIn("已完成", failing_page.status_label.text())

    def test_record_exists_when_project_link_save_fails_and_no_automatic_recalculation(self) -> None:
        project_id, _ = self._preview_and_save()
        self.assertTrue(self.page.open_project(project_id))
        with closing(sqlite3.connect(self.projects_path)) as connection:
            connection.execute(
                "CREATE TRIGGER fail_project_update BEFORE UPDATE ON projects "
                "BEGIN SELECT RAISE(ABORT, 'test project write failure'); END"
            )
            connection.commit()
        with patch("packages.ui.pages.QMessageBox.warning"):
            self.page.calculate_button.click()
            self.app.processEvents()

        records = self.records.list_all()
        self.assertEqual(len(records), 1)
        record = records[0]
        self.assertIn("记录已保存，但项目关联失败", self.page.status_label.text())
        self.app.processEvents()
        self.assertEqual(len(self.records.list_all()), 1)
        before_recovery = self.projects.get(project_id)
        self.assertEqual(before_recovery.units[0].record_ids, ())
        self.assertIsNone(before_recovery.units[0].result_snapshot)

        with closing(sqlite3.connect(self.projects_path)) as connection:
            connection.execute("DROP TRIGGER fail_project_update")
            connection.commit()
        recovered_service = ProjectWorkspaceService(SQLiteProjectWorkspaceRepository(self.projects_path))
        recovered = recovered_service.get(project_id)
        self.assertEqual(recovered.units[0].record_ids, (record.record_id,))
        self.assertEqual(recovered.units[0].result_snapshot["record_id"], record.record_id)
        self.assertEqual(
            recovered.units[0].input_fingerprint,
            hashlib.sha256(encode_canonical_input(recovered.units[0].canonical_input).encode("utf-8")).hexdigest(),
        )
        with closing(sqlite3.connect(self.projects_path)) as connection:
            pending_count = connection.execute("SELECT COUNT(*) FROM pending_record_links").fetchone()[0]
        self.assertEqual(pending_count, 0)
        self.assertEqual(len(self.records.list_all()), 1)

    def test_opening_canonical_project_from_manual_accounting_routes_to_excel(self) -> None:
        project_id, _ = self._preview_and_save()
        self.shell.navigate(AppRoute.NEW_ACCOUNTING)
        accounting_page = self.shell.pages[AppRoute.NEW_ACCOUNTING]
        accounting_page._refresh_saved_projects()
        index = accounting_page.saved_projects.findData(project_id)
        self.assertGreaterEqual(index, 0)
        accounting_page.saved_projects.setCurrentIndex(index)
        accounting_page._open_selected_project()
        self.app.processEvents()

        self.assertEqual(self.shell.current_route, AppRoute.EXCEL_IMPORT)
        self.assertEqual(self.page._workspace.project_id, project_id)
        self.assertIsNotNone(self.page._selected_unit().canonical_input)
        self.assertEqual(self.page._selected_unit().record_ids, ())

    def test_failed_new_import_disables_the_previously_open_project_calculation(self) -> None:
        project_id, _ = self._preview_and_save()
        self.assertTrue(self.page.open_project(project_id))
        self.assertTrue(self.page.calculate_button.isEnabled())
        self.assertEqual(self.records.list_all(), ())

        bad_path = self.root / "坏工作簿.xlsx"
        workbook = load_workbook(BytesIO(ExcelTemplateService.default().read_bytes()))
        workbook["B.2"]["B14"] = "=1+1"
        workbook.save(bad_path)
        workbook.close()
        self.page._preview_workbook(bad_path)

        self.assertIsNone(self.page._workspace)
        self.assertIsNotNone(self.page._last_preview)
        self.assertFalse(self.page._last_preview.units[0].can_calculate)
        self.assertFalse(self.page.calculate_button.isEnabled())
        self.assertFalse(self.page.save_project_button.isEnabled())
        self.assertEqual(self.records.list_all(), ())
        self.assertIsNotNone(self.projects.get(project_id))
        self.assertTrue(self.page.open_project(project_id))
        self.assertTrue(self.page.calculate_button.isEnabled())
        self.assertEqual(self.records.list_all(), ())

    def test_formal_calculation_is_disabled_when_dependencies_are_missing(self) -> None:
        unconfigured = ExcelImportPage()
        self.assertFalse(unconfigured.calculate_button.isEnabled())
        self.assertFalse(unconfigured.save_project_button.isEnabled())
        self.assertIn("正式核算尚未配置", unconfigured.status_label.text())


if __name__ == "__main__":
    unittest.main()
